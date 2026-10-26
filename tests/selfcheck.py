#!/usr/bin/env python3
"""
The framework checked against itself.

    python tests/selfcheck.py

The validator skips `templates/` and `schemas/`, for good reasons written down in the
registry, and the cost of that is real: nothing was checking the templates that everything
else is validated against, which is how `RLM.yaml` went without `status` and `owners`
without anybody noticing. This file is what closes that hole. It runs in CI.

No pytest: one dependency fewer, and the output is meant to be read by a person who has
just broken something.
"""

from __future__ import annotations

import argparse
import ast
import atexit
import collections
import contextlib
import importlib.util
import io
import json
import os
import re
import shutil
import subprocess
import sys
import tempfile
import types
import zipfile
from datetime import date, datetime
from pathlib import Path

import yaml
from jsonschema import Draft202012Validator

ROOT = Path(__file__).resolve().parents[1]
REGISTRY = yaml.safe_load((ROOT / "schemas" / "artifact-types.yaml").read_text())
CHECKS = yaml.safe_load((ROOT / "skills" / "audit" / "checks.yaml").read_text())
VALIDATE = ROOT / "skills" / "audit" / "scripts" / "validate.py"
EXTRACT = ROOT / "skills" / "start" / "scripts" / "extract.py"
GENERATE = ROOT / "schemas" / "generate.py"

SECTION_MARK = re.compile(r"<!--\s*section:\s*([a-z0-9-]+)\s*-->")
PLACEHOLDERS = set(REGISTRY["placeholders"]["enforced"])


def _load(path: Path, name: str):
    """Import the validator so its constants are read, never restated here."""
    spec = importlib.util.spec_from_file_location(name, path)
    mod = importlib.util.module_from_spec(spec)
    # Registered before it runs: @dataclass resolves annotations through sys.modules, and
    # without this the import dies on the first decorated class.
    sys.modules[name] = mod
    spec.loader.exec_module(mod)
    return mod


GENERATED_MARK = _load(VALIDATE, "validate").GENERATED_MARK

# HOW A CHECK GETS THE FIXTURES IT NEEDS, AND WHY IT DOES NOT JUST READ THEM OFF DISK.
# `evals/fixtures/build/` is generated output and is gitignored, so in a fresh checkout it does
# not exist -- and a fresh checkout is exactly what CI is. Four checks written against that
# directory measured nothing there and said so in their own words, which is the worse of the
# two available failures: a suite that reports "the fixtures are not built" reads as a
# configuration problem rather than as the assertion having stopped running.
#
# So a check builds what it needs, once per run, into a temporary tree. The map of which
# generator writes where is read out of `make.py` rather than restated, for the reason
# everything else here is read rather than restated.
_MAKE = _load(ROOT / "evals" / "fixtures" / "make.py", "make")
_BUILT: dict[str, Path] = {}
_BUILD_DIR: list[Path] = []


def built(name: str) -> Path | None:
    """One generated fixture, built into a temporary tree that lives as long as the run."""
    if name in _BUILT:
        return _BUILT[name]
    spec = _MAKE.GENERATED.get(name)
    if spec is None:
        return None
    if not _BUILD_DIR:
        d = Path(tempfile.mkdtemp(prefix="selfcheck-fixtures-"))
        _BUILD_DIR.append(d)
        atexit.register(shutil.rmtree, str(d), True)
    script, dest = spec
    target = _BUILD_DIR[0] / (dest or name)
    target.parent.mkdir(parents=True, exist_ok=True)
    r = subprocess.run([sys.executable,
                        str(ROOT / "evals" / "fixtures" / "generators" / script), str(target)],
                       capture_output=True, text=True)
    if r.returncode != 0:
        return None
    _BUILT[name] = target
    return target


def corpus() -> list[Path]:
    """Every fixture repository there is: the generated ones, built, and the static ones."""
    roots: list[Path] = []
    for name in _MAKE.GENERATED:
        target = built(name)
        if target:
            roots += [p.parent for p in target.rglob("AGENTS.md")]
    roots += [p.parent for p in (ROOT / "evals" / "fixtures" / "static").rglob("AGENTS.md")]
    return sorted(roots)

failures: list[str] = []


def check(name: str):
    """Decorator that runs a check and records what it said."""
    def wrap(fn):
        problems = fn() or []
        mark = "ok  " if not problems else "FAIL"
        print(f"{mark} {name}")
        for p in problems:
            print(f"       {p}")
            failures.append(f"{name}: {p}")
        return fn
    return wrap


def jsonify(o):
    if isinstance(o, (date, datetime)):
        return o.isoformat()
    if isinstance(o, dict):
        return {k: jsonify(v) for k, v in o.items()}
    if isinstance(o, list):
        return [jsonify(v) for v in o]
    return o


def front_matter(text: str) -> dict:
    if text.startswith("---\n"):
        return yaml.safe_load(text[4:text.find("\n---", 4)])
    return yaml.safe_load(text)


def templates() -> dict[str, Path]:
    out = {}
    for p in sorted((ROOT / "templates").iterdir()):
        if p.suffix not in (".md", ".yaml") or p.name == "README.md":
            continue
        m = re.search(r"^artifact_type: (\S+)", p.read_text(), re.M)
        if m:
            out[m.group(1)] = p
    return out


TPL = templates()


# ─────────────────────────────────────────────────────────────────────────────

@check("the generated schemas match the registry")
def _schemas_fresh():
    r = subprocess.run([sys.executable, str(GENERATE), "--check"],
                       capture_output=True, text=True)
    return [] if r.returncode == 0 else [l for l in r.stdout.splitlines() if l.strip()]


@check("every type has a template and every template has a type")
def _types_match():
    reg, tpl = set(REGISTRY["types"]), set(TPL)
    return ([f"{t}: in the registry, no template" for t in sorted(reg - tpl)]
            + [f"{t}: has a template, not in the registry" for t in sorted(tpl - reg)])


@check("every template validates against its own schema")
def _templates_validate():
    problems = []
    for t, p in TPL.items():
        schema = json.loads((ROOT / "schemas" / "framework" / t / "v1.json").read_text())
        for e in Draft202012Validator(schema).iter_errors(jsonify(front_matter(p.read_text()))):
            # A template is allowed to fail on exactly one thing: carrying its own
            # placeholder in a field that must be filled in. That is what a template is.
            # Anything else means the template contradicts the registry.
            if any(ph in str(e.message) for ph in PLACEHOLDERS):
                continue
            where = "/".join(map(str, e.path)) or "front matter"
            problems.append(f"{p.name}: {where}: {e.message}")
    return problems


@check("every mandatory section has exactly one marker, above a heading")
def _markers():
    problems = []
    for t, spec in REGISTRY["types"].items():
        wanted = spec.get("sections") or []
        body = TPL[t].read_text() if wanted else ""
        found = SECTION_MARK.findall(TPL[t].read_text())
        if not wanted and found:
            problems.append(f"{TPL[t].name}: has markers, the registry declares no sections")
            continue
        for sid in wanted:
            n = found.count(sid)
            if n != 1:
                problems.append(f"{TPL[t].name}: marker {sid!r} appears {n} times")
                continue
            following = body.split(f"<!-- section: {sid} -->\n", 1)[1].splitlines()[0]
            if not following.startswith("#"):
                problems.append(f"{TPL[t].name}: {sid!r} does not sit above a heading")
        for extra in set(found) - set(wanted):
            problems.append(f"{TPL[t].name}: marker {extra!r} is not in the registry")
    return problems


@check("every check the validator emits is in the catalog, and the reverse")
def _catalog_matches_code():
    sources = [VALIDATE, *sorted((ROOT / "src/framework_data_ai").glob("*.py"))]
    emitted = set(re.findall(r'report\.add\(\s*"([A-Z]{2,3}\d{3})"',
                             "\n".join(path.read_text() for path in sources)))
    catalogued = set(CHECKS["checks"])
    off = {c for c, s in CHECKS["checks"].items()
           if s.get("level") in ("off", False)}
    # A RETIRED CODE IS THE THIRD STATE, AND IT IS NOT "OFF". Off is a check that cannot run
    # yet and says what has to exist first; retired is a check that ran, stopped, and keeps
    # its number so that no other failure is ever reported under it -- an annotation joins
    # on the code. So it must not be emitted anywhere, and it must say when and why.
    retired = {c for c, s in CHECKS["checks"].items() if s.get("level") == "retired"}

    problems = [f"{c}: emitted by the validator, absent from checks.yaml"
                for c in sorted(emitted - catalogued)]
    for c in sorted(emitted & retired):
        problems.append(f"{c}: retired in {CHECKS['checks'][c].get('retired_in')} and still "
                        "emitted somewhere. A retired number means one thing forever")
    for c in sorted(retired):
        if not (CHECKS["checks"][c].get("retired_in") and
                CHECKS["checks"][c].get("retired_because")):
            problems.append(f"{c}: retired without `retired_in` and `retired_because`")
    # A catalogued check that nothing emits is a promise the validator does not keep,
    # unless it is switched off and says why, or retired.
    for c in sorted(catalogued - emitted - off - retired):
        problems.append(f"{c}: in checks.yaml, never emitted, and not switched off")
    for c in sorted((catalogued - emitted) & off):
        if not CHECKS["checks"][c].get("blocked_by"):
            problems.append(f"{c}: off and unimplemented, with no `blocked_by` saying why")
    return problems


@check("every artifact is listed by the phase FRAMEWORK.md §5 says produces it")
def _phases_agree():
    # §5 tells the lifecycle as a narrative, §7 tabulates `born_in` per artifact, and the
    # mermaid draws it. Three tellings of one fact, and they had already drifted: `RSK`,
    # `RMP` and `LOG` were each born in a phase whose row did not mention them.
    #
    # The assertion runs one way only, §5 must contain what §7 says is born there, and
    # never the reverse: `WF` legitimately appears in two rows because its sections are
    # written at different times, and `EVR` appears under F6 although it is born at the
    # release gate. A diagram composes, a normative table enumerates, so the mermaid is
    # not asserted here either.
    fw = (ROOT / "FRAMEWORK.md").read_text()
    rows = dict(re.findall(r"\| \*\*(F\d) · [^|]*\*\* \| ([^|]+) \|", fw))
    problems = []
    for phase, cell in rows.items():
        listed = set(re.findall(r"`([A-Za-z.]+)(?:\s+§\w+)?`", cell))
        for spec in REGISTRY["types"].values():
            if spec.get("born_in") == phase and spec["id"] not in listed:
                problems.append(f"§5 {phase} does not list {spec['id']}, which §7 says is "
                                f"born there")
    if not rows:
        problems.append("no phase rows found in §5: the check is not running")
    return problems


@check("every skill only names checks, flags and files that exist")
def _skill_references():
    # Runs over every skill, not just the one that owns the validator. A skill is prose
    # that an agent executes, and prose naming a flag the validator does not accept fails
    # at the moment somebody is trusting it. The three things checked here are the ones
    # that go stale silently: the name that has to match the directory or the skill never
    # resolves, a check code that was renamed, and a flag that never existed.
    # Every script the skills ship, not only the validator. `audit` carries two now, and a
    # union across them is the honest shape of the guard: a skill that names a flag has to
    # be naming one that something it ships accepts. What the union gives up is telling a
    # `migrate.py` flag written against `validate.py` from a correct one, which is a typo
    # this cannot see and a reader can.
    # The operational CLI ships at the root, with its argparse definitions in the core --
    # when it ships at all. From 3.7.0 the documentary retrieval engine is not part of the
    # framework, so the file is included only if it exists rather than being assumed: a
    # hard-coded path here would fail as a missing file instead of reporting a stale flag,
    # which is the failure this check exists to make legible.
    scripts = [*sorted((ROOT / "skills").rglob("scripts/*.py"))]
    operational_cli = ROOT / "src/framework_data_ai/memory/cli.py"
    if operational_cli.exists():
        scripts.append(operational_cli)
    known_flags = {f for script in scripts
                   for f in re.findall(r'"(--[a-z-]+)"', script.read_text())}
    problems = []

    skills = sorted(p for p in (ROOT / "skills").iterdir() if (p / "SKILL.md").exists())
    if not skills:
        return ["no skills found under skills/: the check is not running"]

    for d in skills:
        skill = d / "SKILL.md"
        text = skill.read_text()
        head = (yaml.safe_load(text[4:text.find("\n---", 4)])
                if text.startswith("---\n") else {})
        if head.get("name") != d.name:
            problems.append(f"{d.name}: name is {head.get('name')!r}, the directory is "
                            f"{d.name!r}: the skill will not resolve")
        if not head.get("description"):
            problems.append(f"{d.name}: no description, so it never triggers")

        body = text[text.find("\n---", 4):]
        for code in sorted(set(re.findall(r"`([A-Z]{2,3}\d{3})`", body))):
            if code not in CHECKS["checks"]:
                problems.append(f"{d.name}: names {code}, which is not in the catalog")
        for flag in sorted(set(re.findall(r"`(--[a-z-]+)`", body))):
            if flag not in known_flags:
                problems.append(f"{d.name}: names {flag}, which the validator does not "
                                "accept")
        # A skill that points at a reference or a script it does not carry is a skill that
        # stops halfway through, at the point where the agent goes looking for the file.
        for ref in sorted(set(re.findall(r"`((?:references|scripts)/[\w./-]+)`", body))):
            if not (d / ref).exists() and not (ROOT / ref).exists():
                problems.append(f"{d.name}: points at {ref}, which does not exist")
    return problems


@check("the validator reports nothing on a minimal correct repository")
def _clean_repo():
    fm = lambda **kw: "---\n" + "\n".join(f"{k}: {v}" for k, v in kw.items()) + "\n---\n\n"
    files = {
        # `entries:` is what REG002 and REG003 read. Without it the register is reported by
        # REG004, which is the point: a register they cannot read used to report clean.
        "OPEN.md": fm(schema="framework/open-register/v1", artifact_type="open-register",
                      lifecycle="living", status="active", owners="[owner]",
                      created="2026-01-01", last_review="2026-01-01 09:00",
                      entries="\n  OD-001:\n    status: open\n"
                              "    cost_to_reverse: low\n"
                              # At the root, who an entry binds has no other answer, so
                              # `REG011` asks for one. `[all]` in a repository of one
                              # product is not the same claim as naming that product: it
                              # says the entry binds whatever this repository grows.
                              "    products: [all]\n"
                              "    default_in_force: nothing is scheduled\n")
        # THE BODY CARRIES THE ID TOO, AND `REG015` IS WHY. The map holds the fields a check
        # reads and the body holds the reasoning; an entry present in one half and absent
        # from the other is half a row, and this minimal repository has to be a correct one
        # in both directions.
        + "# Open\n\n## Cost to reverse LOW\n\n### OD-001 - the one open question\n",
        "decisions/DEC-001-slug.md": fm(
            schema="framework/decision-record/v1", artifact_type="decision-record",
            id="DEC-001", lifecycle="immutable", status="accepted", scope="architecture",
            products="[alpha]", owners="[owner]", created="2026-01-01",
            leaves_open="[]") + "# DEC-001\n",
        "products/alpha/PBR.md": fm(
            schema="framework/product-brief/v1", artifact_type="product-brief",
            lifecycle="living", status="active", products="[alpha]", owners="[owner]",
            created="2026-01-01", last_review="2026-01-01 09:00") + "# Brief\n",
        # A correct repository says which framework it was written against. Read from the
        # registry rather than written as 1, so this stays true through a bump instead of
        # becoming the first thing that breaks on the day one happens.
        "framework.yaml": f"framework_version: {REGISTRY['version']}\n",
    }
    with tempfile.TemporaryDirectory() as tmp:
        for rel, text in files.items():
            p = Path(tmp) / rel
            p.parent.mkdir(parents=True, exist_ok=True)
            p.write_text(text)
        r = subprocess.run(
            [sys.executable, str(VALIDATE), "--root", tmp, "--json", "--stale-days", "36500"],
            capture_output=True, text=True)
        if r.returncode not in (0, 1):
            return [f"the validator crashed: {r.stderr.strip().splitlines()[-1]}"]
        out = json.loads(r.stdout)
        return [f"{f['code']} {f['path']}: {f['message']}" for f in out["findings"]]


@check("every open entry reaches the generated view, in every shape it is allowed to declare")
def _every_entry_reaches_the_view():
    # THE GENERAL FORM OF A BUG THAT SHIPPED. `REG011` made one state illegal -- an entry at
    # the root naming no products -- and the generator's last section was selecting on
    # exactly that state, so the more a repository obeyed the check, the emptier the view
    # became. Nothing looked wrong: the register was right, the check was satisfied, and
    # `§5` said there was nothing above the products.
    #
    # So the invariant, rather than the instance: an entry that is open is in the view,
    # whatever legal shape it declares and wherever it is filed. A state added later that
    # the generator does not know about fails here, on the day it is added, instead of on
    # the day somebody finishes migrating to it.
    fm = lambda body: ("---\nschema: framework/open-register/v1\n"
                       "artifact_type: open-register\nlifecycle: living\nstatus: active\n"
                       "owners: [o]\ncreated: 2026-01-01 09:00\n"
                       "last_review: 2026-08-01 09:00\n" + body + "---\n\n# Open\n")
    man = lambda p, m: ("schema: framework/product-manifest/v1\n"
                        "artifact_type: product-manifest\nlifecycle: living\n"
                        "status: active\n"
                        f"products: [{p}]\nname: {p}\none_liner: A thing.\n"
                        "owners: [o]\ncreated: 2026-01-01 09:00\n"
                        f"last_review: 2026-01-01 09:{m}\nstage:\n  phase: F5\n  block: A\n")
    entry = lambda od, extra="": (f"  {od}:\n    status: open\n    cost_to_reverse: low\n"
                                  f"    default_in_force: nothing\n{extra}")

    # One entry per legal shape, and the shape is the point of each.
    shapes = {
        "OD-001": "declaring `[all]` at the root",
        "OD-002": "naming one product at the root",
        "OD-003": "naming no products at the root, which is legal and reported",
        "OD-004": "filed in the substrate's register",
        "OD-005": "filed in a product's own register",
    }
    files = {
        "framework.yaml": f"framework_version: {REGISTRY['version']}\n",
        "products/alpha/product.yaml": man("alpha", "10"),
        "products/beta/product.yaml": man("beta", "20"),
        "OPEN.md": fm("entries:\n"
                      + entry("OD-001", "    products: [all]\n")
                      + entry("OD-002", "    products: [beta]\n")
                      + entry("OD-003"))
        + "\n<!-- generated: open-union -->\nx\n<!-- /generated -->\n",
        "platform/OPEN.md": fm("entries:\n" + entry("OD-004")),
        "products/alpha/OPEN.md": fm("entries:\n" + entry("OD-005")),
        "products/beta/OPEN.md": fm("entries: {}\n"),
        # A decision that left something open whose entry nobody has written. It has no
        # register row by definition, so the view is the only place it can surface.
        "decisions/DEC-001-x.md": (
            "---\nschema: framework/decision-record/v1\n"
            "artifact_type: decision-record\nid: DEC-001\nlifecycle: immutable\n"
            "status: accepted\nscope: architecture\nproducts: [alpha]\nowners: [o]\n"
            "created: 2026-01-01 09:00\nleaves_open: [unregistered]\n---\n\n# DEC-001\n"),
    }

    problems = []
    with tempfile.TemporaryDirectory() as tmp:
        root = Path(tmp)
        for rel, text in files.items():
            f = root / rel
            f.parent.mkdir(parents=True, exist_ok=True)
            f.write_text(text)
        r = subprocess.run([sys.executable, str(VALIDATE), "--root", str(root),
                            "--emit-index"], capture_output=True, text=True)
        if r.returncode not in (0, 1):
            return [f"--emit-index failed: {r.stderr.strip() or r.stdout.strip()}"]

        union = (root / "OPEN.md").read_text()
        region = union.split("<!-- generated: open-union -->", 1)[-1]
        for od, why in sorted(shapes.items()):
            if od not in region:
                problems.append(f"{od}, {why}, does not appear anywhere in the generated "
                                "union. An entry that is open and invisible in the view is "
                                "the failure this view exists to prevent, and it reads as "
                                "nothing being open")

        # And the two the union has to keep apart, or the section stops meaning anything.
        tail = region.split("## Bound to no single product", 1)
        if len(tail) != 2:
            problems.append("the union has no section for what binds no single product")
        else:
            for od, want in (("OD-001", True), ("OD-003", True), ("OD-004", True),
                             ("OD-002", False), ("OD-005", False)):
                if (od in tail[1]) != want:
                    problems.append(
                        f"{od} is {'missing from' if want else 'listed under'} what binds no "
                        "single product, and it is the opposite of what it declares")

        # WHO ELSE READS THIS STATE -- the question the catalog now tells you to ask before
        # adding a check. A question with no entry cannot appear in a view built from
        # registers, so the view names the decision that declares one.
        r = subprocess.run([sys.executable, str(VALIDATE), "--root", str(root),
                            "--emit-index"], capture_output=True, text=True)
        idx = (root / "products" / "alpha" / "product.index.yaml")
        if idx.exists() and "open_unregistered: [DEC-001]" not in idx.read_text():
            problems.append("the derived view does not name the decision that declares an "
                            "unregistered open point, so 'what is open for this product' "
                            "still means 'what is open and already written down'")

        alpha = (root / "products" / "alpha" / "product.index.yaml").read_text()
        for od, want in (("OD-001", True), ("OD-003", True), ("OD-004", True),
                         ("OD-005", True), ("OD-002", False)):
            if (od in alpha) != want:
                problems.append(
                    f"alpha's derived view {'is missing' if want else 'holds'} {od}, which "
                    "is the opposite of what that entry binds")
    return problems


@check("the vocabularies have a value for the day the question is first asked")
def _vocabularies_cover_not_yet():
    # Three enums were extended on the same day and the omissions were all on one axis: the
    # answer on the day nobody has done the work yet. A vocabulary without it makes omission
    # the honest move, and an omitted field is invisible -- a risk with no likelihood drops
    # out of every ordering, a promise with no feasibility reads as unexamined rather than
    # unexaminable. This asserts the values exist, so removing one is a decision somebody
    # takes rather than a line somebody deletes.
    want = {
        ("risk-register", "risks", "category"): {"security", "organisational"},
        ("risk-register", "risks", "likelihood"): {"C"},
        ("commitments", "commitments", "feasibility"): {"not-yet-assessable",
                                                        "not-applicable"},
        ("open-register", "entries", "status"): {"parked"},
        ("operational-stack", "stack", "status"): {"dropped"},
    }
    problems = []
    for (type_name, field, key), values in sorted(want.items()):
        spec = REGISTRY["types"].get(type_name, {})
        enums = ((spec.get("maps") or {}).get(field, {}).get("fields", {}).get("enums") or {})
        have = set(enums.get(key) or [])
        missing = values - have
        if missing:
            problems.append(
                f"{type_name}.{field}.{key} has no {', '.join(sorted(missing))}. Every one of "
                "these is the answer on the day the question is first asked -- not decided, "
                "not measurable, not applicable, already true -- and a vocabulary without it "
                "leaves two moves: omit the field, which takes the row out of everything "
                "that reads it, or pick the nearest value, which writes something false that "
                "nothing will contradict")
    return problems


@check("a promise and the risk it creates can be joined, now that both are readable")
def _commitments_and_risks():
    # Two markdown tables until 2.1.0, which is why neither of these could be reported: a
    # product carrying commitments and no risk register at all, and a commercial risk about
    # a claim the commitments register does not contain. A check cannot join two tables it
    # cannot read, and the fix was the same one `OPEN.md` and `STACK.md` already had.
    fm = lambda **kw: "---\n" + "\n".join(f"{k}: {v}" for k, v in kw.items()) + "\n---\n\n"

    def repo(cmts: str, risks: str | None) -> dict:
        files = {
            "framework.yaml": f"framework_version: {REGISTRY['version']}\n",
            "products/alpha/product.yaml": (
                "schema: framework/product-manifest/v1\n"
                "artifact_type: product-manifest\nlifecycle: living\nstatus: active\n"
                "products: [alpha]\nname: alpha\none_liner: A thing.\nowners: [o]\n"
                "created: 2026-01-01 09:00\nlast_review: 2026-01-01 09:00\n"
                "stage:\n  phase: F5\n  block: A\n"),
            "COMMITMENTS.md": fm(schema="framework/commitments/v1",
                                 artifact_type="commitments", lifecycle="living",
                                 status="active", owners="[o]", products="[alpha]",
                                 created="2026-01-01 09:00",
                                 last_review="2026-08-01 09:00",
                                 commitments=cmts) + "# Commitments\n",
        }
        if risks is not None:
            files["products/alpha/RSK.md"] = fm(
                schema="framework/risk-register/v1", artifact_type="risk-register",
                lifecycle="living", status="active", products="[alpha]", owners="[o]",
                created="2026-01-01 09:00", last_review="2026-08-01 09:30",
                risks=risks) + ("# Risks\n\n<!-- section: state -->\n## State\n\n"
                                "<!-- section: acceptances -->\n## Acceptances\n\n"
                                "<!-- section: events -->\n## Events\n")
        return files

    one = ("\n  CMT-001:\n    to: a customer\n    status: open\n"
           "    products: [alpha]\n")
    problems = []
    with tempfile.TemporaryDirectory() as tmp:
        root = Path(tmp)

        def codes(files) -> set[str]:
            for stale in root.rglob("*.md"):
                stale.unlink()
            for rel, text in files.items():
                f = root / rel
                f.parent.mkdir(parents=True, exist_ok=True)
                f.write_text(text)
            r = subprocess.run([sys.executable, str(VALIDATE), "--root", str(root),
                                "--json"], capture_output=True, text=True)
            return {f["code"] for f in json.loads(r.stdout)["findings"]}

        if "XP005" not in codes(repo(one, None)):
            problems.append("a product carrying a commitment and no risk register was not "
                            "reported: the exposure lives only in the sentence that made it")
        if "XP005" in codes(repo(one, "\n  RSK-001:\n    category: technical\n"
                                      "    state: open\n")):
            problems.append("a product with a risk register was reported, so the check is "
                            "asking for something other than the register")

        # A PROMISE NOBODY HAS RECEIVED IS NOT AN EXPOSURE. Counting it put a number in the
        # finding that did not survive being checked, and a reader who verifies one row and
        # finds it was never said to anybody stops believing the rest.
        withheld = ("\n  CMT-002:\n    to: nobody yet, it is in the deck\n"
                    "    status: not-yet-issued\n    products: [alpha]\n")
        if "XP005" in codes(repo(withheld, None)):
            problems.append("a product whose only commitment is `not-yet-issued` was "
                            "reported: nothing has been said to anybody, so there is no "
                            "exposure for a risk register to hold")
        def findings(files, code):
            codes(files)
            r = subprocess.run([sys.executable, str(VALIDATE), "--root", str(root),
                                "--json"], capture_output=True, text=True)
            return [f["message"] for f in json.loads(r.stdout)["findings"]
                    if f["code"] == code]

        msg = findings(repo(one + withheld, None), "XP005")
        if not msg:
            problems.append("a product with one commitment made and one not yet issued was "
                            "not reported at all")
        elif "1 commitment(s)" not in msg[0] or "CMT-002" in msg[0].split("--")[1]:
            problems.append(f"the finding counts the promise nobody has received: {msg[0]}")

        # A commitment that names nobody is paired with nothing, and the row still reads as
        # filled in. And `[all]` used to be read as a product called `all`, which no
        # repository has, so a promise about the whole suite bound nothing at all.
        nameless = "\n  CMT-003:\n    to: a customer\n    status: open\n"
        if "XP006" not in codes(repo(nameless, None)):
            problems.append("a commitment naming no products was not reported: nothing can "
                            "pair it with a risk register, a derived view or the check that "
                            "asks whether its exposure has a home")
        binds_all = ("\n  CMT-004:\n    to: everybody\n    status: open\n"
                     "    products: [all]\n")
        if "XP005" not in codes(repo(binds_all, None)):
            problems.append("a commitment declaring `[all]` bound nothing: read as a product "
                            "named `all`, a promise about the whole suite reached no product "
                            "and no check")
        if "XP005" in codes(repo(binds_all, "\n  RSK-001:\n    category: technical\n"
                                            "    state: open\n")):
            problems.append("a commitment declaring `[all]` was still reported after the one "
                            "product in the repository got a risk register")

        commercial = "\n  RSK-001:\n    category: commercial\n    state: open\n"
        if "REF006" not in codes(repo(one, commercial)):
            problems.append("a commercial risk naming no commitment was not reported: a "
                            "claim tracked as a risk and promised nowhere cannot be "
                            "renegotiated, because there is no promise to renegotiate")
        if "REF006" in codes(repo(one, commercial + "    commitment: CMT-001\n")):
            problems.append("a commercial risk naming a declared commitment was reported")
        if "REF006" not in codes(repo(one, commercial + "    commitment: CMT-404\n")):
            problems.append("a risk naming a commitment no register declares was not "
                            "reported: the promise it is about cannot be found")
        # THE CLASS THE FIRST VERSION WAS WRONG ABOUT. A commercial risk that is not about
        # a promise -- and without a way to say so, the only way to stop the check asking was
        # to file the risk under a category that was false.
        if "REF006" in codes(repo(one, commercial + "    commitment: none\n")):
            problems.append("a commercial risk declaring `commitment: none` was reported. "
                            "Not every commercial exposure is a promise, and if the honest "
                            "answer is unavailable the cheap repair is a wrong category")
        if "REF006" in codes(repo(one, "\n  RSK-001:\n    category: commercial\n"
                                       "    state: closed\n")):
            problems.append("a closed commercial risk was reported: the promise behind it "
                            "stopped mattering, and asking for it is asking about history")

        # THE PROMISE THAT CANNOT BE KEPT, AND WHO CARRIES IT UNTIL SOMEBODY SAYS SO.
        # `XP007` is the half of a carve-out: `ICG` §3 passes over a candidate that
        # contradicts a promise already written off *when a risk owns it*, so that the share
        # of commercial promises that were never buildable stops holding up every candidate
        # that touches them. With nothing reporting the rows nobody owns, that carve-out
        # would be a hole -- an unowned promise reads exactly like a live one.
        beyond = ("\n  CMT-010:\n    to: a customer\n    status: open\n"
                  "    feasibility: out-of-reach\n    products: [alpha]\n")
        owner = ("\n  RSK-001:\n    category: commercial\n    state: open\n"
                 "    commitment: CMT-010\n")
        if "XP007" not in codes(repo(beyond, "\n  RSK-001:\n    category: technical\n"
                                             "    state: open\n")):
            problems.append("a commitment out of reach that no risk names was not reported: "
                            "the renegotiation it is owed has nobody on the hook for it")
        if "XP007" in codes(repo(beyond, owner)):
            problems.append("a commitment out of reach whose exposure a risk owns was "
                            "reported: the owner is what the check asks for, and reporting "
                            "it anyway is what makes a triage stop on the same promise every "
                            "cycle")
        # A risk that is over is not an owner. Same pair `REF006` excludes, and the same
        # reason: the row stopped being about anything.
        if "XP007" not in codes(repo(beyond, owner.replace("state: open",
                                                           "state: closed"))):
            problems.append("a closed risk was read as the owner of an exposure that is "
                            "still in the register")
        # The three states where nothing is owed. `not-yet-issued` was said to nobody --
        # `XP005` already declines to count it -- and the other two are rows where the
        # conversation has happened.
        for status, why in [
            ("not-yet-issued", "the promise is in a deck nobody has received, so there is "
                               "no exposure and the remedy is an internal edit"),
            ("renegotiated", "the conversation this check asks for has already happened"),
            ("met", "the promise was kept"),
        ]:
            quiet = beyond.replace("status: open", f"status: {status}")
            if "XP007" in codes(repo(quiet, "\n  RSK-001:\n    category: technical\n"
                                            "    state: open\n")):
                problems.append(f"a `{status}` commitment out of reach was reported: {why}")
        # `unsatisfiable` is the status saying it will not be delivered, and it is the same
        # question from the other side: somebody still has to carry it until the customer
        # has been told.
        wont = beyond.replace("status: open", "status: unsatisfiable").replace(
            "    feasibility: out-of-reach\n", "")
        if "XP007" not in codes(repo(wont, "\n  RSK-001:\n    category: technical\n"
                                           "    state: open\n")):
            problems.append("an `unsatisfiable` commitment that no risk names was not "
                            "reported: the status says it will not be delivered and nobody "
                            "is answerable for having said it")
        if "XP007" in codes(repo("\n  CMT-011:\n    to: a customer\n    status: open\n"
                                 "    feasibility: feasible\n    products: [alpha]\n",
                                 "\n  RSK-001:\n    category: technical\n"
                                 "    state: open\n")):
            problems.append("a feasible commitment was reported: the check is about promises "
                            "declared beyond reach, not about promises with no risk")
    return problems


@check("a shallow framework checkout is not read as a pin that was never true")
def _shallow_checkout():
    # `actions/checkout` clones one commit by default, and the workflow this framework ships
    # checks the framework out that way. In a shallow clone almost every commit is absent, so
    # "this framework does not contain your pin" is true of nearly everything and means
    # nothing -- and it is the more confident of the two messages, which is the wrong way
    # round.
    #
    # The helpers are exercised directly against a shallow clone rather than by running the
    # validator inside it: a clone carries the committed code, and a fix asserted through one
    # cannot fail until after it is committed.
    x = _load(VALIDATE, "validate")
    problems = []
    with tempfile.TemporaryDirectory() as tmp:
        clone = Path(tmp) / "shallow"
        r = subprocess.run(["git", "clone", "--depth", "1", "-q", f"file://{ROOT}",
                            str(clone)], capture_output=True, text=True)
        if r.returncode != 0:
            return ["a shallow clone could not be made here, so the check is not running"]
        old = subprocess.run(["git", "-C", str(ROOT), "rev-parse", "HEAD~3"],
                             capture_output=True, text=True).stdout.strip()
        real = x.FRAMEWORK
        try:
            x.FRAMEWORK = clone
            if not x.framework_is_shallow():
                problems.append("a one-commit clone was not recognised as shallow, so a pin "
                                "it cannot contain reads as a pin that was never true")
            if x.framework_has(old) is not False:
                problems.append("a commit absent from a shallow clone was reported as "
                                "present, which is the fact the message turns on")
        finally:
            x.FRAMEWORK = real
        if x.framework_is_shallow():
            problems.append("this checkout was called shallow, so the distinction collapses "
                            "the other way and a pin that is genuinely wrong stops being "
                            "reported as wrong")
        if x.framework_has(old) is not True:
            problems.append(f"{old[:12]} is in this repository's history and was reported "
                            "absent")
    return problems


@check("the suite that asks git for history is run in a checkout that has one")
def _ci_fetches_history():
    # THE RED THAT IS ABOUT THE CLONE AND NOT ABOUT THE FRAMEWORK. Two checks above ask git
    # for commits that are not HEAD, and `actions/checkout` clones one commit unless it is
    # told otherwise. The workflow was never told, so every push for a week ended in four
    # problems that could not be fixed in the repository, because they were not about it.
    # That state is worse than having no CI: a red that nobody can act on is a red nobody
    # reads, and it is where a real failure arrives unnoticed.
    #
    # Asserted here rather than left to the next person reading the workflow, because this
    # file is the thing that goes red and the workflow is the thing that has to change: the
    # two are a day apart in a directory nobody opens twice.
    wf = ROOT / ".github" / "workflows" / "framework.yml"
    if not wf.exists():
        return [".github/workflows/framework.yml is gone: this suite runs nowhere, and a "
                "suite that runs nowhere reports nothing"]
    jobs = (yaml.safe_load(wf.read_text()) or {}).get("jobs") or {}
    problems, ran = [], False
    for name, job in jobs.items():
        steps = (job or {}).get("steps") or []
        if not any("selfcheck.py" in str(s.get("run", "")) for s in steps):
            continue
        ran = True
        depths = [(s.get("with") or {}).get("fetch-depth") for s in steps
                  if str(s.get("uses", "")).startswith("actions/checkout")]
        if not depths:
            problems.append(f"{name}: runs this suite and checks nothing out")
        elif 0 not in depths:
            problems.append(f"{name}: checks out with `fetch-depth: {depths[0]!r}` and runs "
                            "a suite that asks git for commits before HEAD. One commit is "
                            "the default, `HEAD~3` is absent in a one-commit clone, and the "
                            "pin checks then report the checkout instead of the framework")
    if not ran:
        problems.append("no job in the workflow runs tests/selfcheck.py: whatever is red on "
                        "a push, it is not this file")
    return problems


@check("the pull request body reaches the check as somebody wrote it")
def _pr_body_is_not_expanded():
    # `printf '%b'` expands backslash escapes in text somebody else wrote, and `\c` truncates
    # everything after it: a body carrying one would lose the line citing the contract, and
    # the gate would report a correct pull request for not citing one.
    wf = ROOT / "ci" / "pull-request.yml"
    if not wf.exists():
        return ["ci/pull-request.yml is gone: the gate has no way to be switched on"]
    text = wf.read_text(encoding="utf-8")
    problems = []
    if "printf '%b'" in text or 'printf "%b"' in text:
        problems.append("the workflow writes the pull request body with `printf %b`, which "
                        "expands escapes in text somebody else wrote -- `\\c` truncates the "
                        "rest, and what is lost is the line naming the change contract")
    if "$PR_TEXT" in text and "env:" not in text:
        problems.append("the body reaches the shell without going through the environment: "
                        "a backtick in a description becomes a command")
    return problems


@check("a pin that was never true is told apart from a checkout that moved")
def _pin_shapes():
    # Both produce a hash that does not match, and the two need opposite responses: one is a
    # migration to read, the other is a line nobody ever verified. And git resolves an
    # uppercase hash as readily as a lowercase one, so rejecting it was a finding about
    # somebody's shift key.
    problems = []
    head = subprocess.run(["git", "-C", str(ROOT), "rev-parse", "--short=12", "HEAD"],
                          capture_output=True, text=True)
    if head.returncode != 0:
        return ["git history is not available here, so the check is not running"]
    now = head.stdout.strip()
    old = subprocess.run(["git", "-C", str(ROOT), "rev-parse", "--short=12", "HEAD~3"],
                         capture_output=True, text=True).stdout.strip()
    with tempfile.TemporaryDirectory() as tmp:
        root = Path(tmp)

        def message(pin: str) -> str:
            (root / "framework.yaml").write_text(
                f"framework_version: {REGISTRY['version']}\nframework_commit: {pin}\n")
            r = subprocess.run([sys.executable, str(VALIDATE), "--root", str(root),
                                "--json"], capture_output=True, text=True)
            found = [f["message"] for f in json.loads(r.stdout)["findings"]
                     if f["code"] == "FW003"]
            return found[0] if found else ""

        if message(now):
            problems.append("a pin at the commit being run was reported")
        if message(now.upper()):
            problems.append("a pin written in uppercase was reported: git resolves it, and a "
                            "field that refuses what the tool accepts is a finding about a "
                            "shift key")
        moved = message(old)
        if "no such commit" in moved or not moved:
            problems.append(f"a pin at an older commit of this framework said {moved[:60]!r}: "
                            "the checkout moved, and that is a migration to read")
        never = message("deadbeefcafe")
        if "no such commit" not in never:
            problems.append("a pin at a commit this framework does not have was reported as a "
                            "checkout that moved, which sends somebody looking for a "
                            "migration that never happened")
        if not message("main"):
            problems.append("a branch name was accepted as a pin: both a branch and a tag "
                            "move, which is the state the field exists to leave")
    return problems


@check("the pull request gate reads what somebody wrote, not what the template says")
def _pr_gate_reads_assertions():
    # The template ships with `no-chg: typo in a comment` inside an HTML comment, as the
    # instruction for writing one. A pull request that keeps the template carried that line
    # into its body, matched the exemption, and turned `PR001` off -- so the gate was silent
    # by default in every repository using the template it comes with.
    #
    # Fenced blocks and quotes go with it. `git log --grep CHG-041` in a snippet is not a
    # change contract being cited, and a quoted line is somebody quoting.
    fm = lambda **kw: "---\n" + "\n".join(f"{k}: {v}" for k, v in kw.items()) + "\n---\n\n"
    template = (ROOT / "ci" / "PULL_REQUEST_TEMPLATE.md")
    problems = []
    with tempfile.TemporaryDirectory() as tmp:
        root = Path(tmp)
        (root / "framework.yaml").write_text(f"framework_version: {REGISTRY['version']}\n")
        (root / "products" / "alpha").mkdir(parents=True)
        (root / "products" / "alpha" / "product.yaml").write_text(
            "schema: framework/product-manifest/v1\nartifact_type: product-manifest\n"
            "lifecycle: living\nstatus: active\nproducts: [alpha]\nname: alpha\n"
            "one_liner: A thing.\nowners: [o]\ncreated: 2026-01-01 09:00\n"
            "last_review: 2026-01-01 09:00\nstage:\n  phase: F5\n  block: A\n")
        empty = root / "changed.txt"
        empty.write_text("")

        def codes(pr_text: str) -> set[str]:
            r = subprocess.run([sys.executable, str(VALIDATE), "--root", str(root),
                                "--pr-text", pr_text, "--changed-files", str(empty),
                                "--json"], capture_output=True, text=True)
            return {f["code"] for f in json.loads(r.stdout)["findings"]}

        cases = [
            (template.read_text(encoding="utf-8") if template.exists() else "",
             True, "the template kept unedited, whose own comment shows how to write "
                   "`no-chg:`"),
            ("Fix a typo\n\nno-chg: prose only", False,
             "an exemption somebody wrote"),
            ("> no-chg: <reason>", True, "an exemption inside a quote"),
            ("Fix\n\n```\nno-chg: example\n```", True,
             "an exemption inside a fenced block"),
            ("Fix\n\n```bash\ngit log --grep CHG-041\n```", True,
             "a change contract named inside a fenced block"),
        ]
        for pr_text, want, what in cases:
            got = "PR001" in codes(pr_text)
            if got != want:
                problems.append(f"{what}: PR001 was {'not ' if want else ''}reported")
    return problems


@check("a malformed map is reported and does not take the run down with it")
def _malformed_maps_survive():
    # `terms: [Freshness, Tenant]` is one bracket away from two rows under `terms:`, and it
    # killed the validator: `.items()` on a list raises, the process died on the malformed
    # document, and nothing was said about everything it had not reached. `entries:` has been
    # guarded since the same thing happened to it; three maps added in one week were not,
    # because each was written by copying the line above it.
    #
    # The schema already reports the shape -- that is what `FM002` is -- so what is asserted
    # here is that the run finishes and says so.
    shapes = {
        "glossary": ("GLOSSARY.md", "terms", "[Freshness, Tenant]", ""),
        "commitments": ("COMMITMENTS.md", "commitments", "[CMT-001]", "products: [alpha]\n"),
        "risk-register": ("products/alpha/RSK.md", "risks", "[RSK-001]", "products: [alpha]\n"),
        "open-register": ("OPEN.md", "entries", "[OD-001]", ""),
        "operational-stack": ("STACK.md", "stack", "[postgres]", ""),
    }
    problems = []
    with tempfile.TemporaryDirectory() as tmp:
        root = Path(tmp)
        (root / "framework.yaml").write_text(f"framework_version: {REGISTRY['version']}\n")
        for atype, (rel, field, value, extra) in sorted(shapes.items()):
            f = root / rel
            f.parent.mkdir(parents=True, exist_ok=True)
            f.write_text(f"---\nschema: framework/{atype}/v1\n"
                         f"artifact_type: {atype}\nlifecycle: living\nstatus: active\n"
                         f"owners: [o]\n{extra}created: 2026-01-01 09:00\n"
                         f"last_review: 2026-08-01 09:00\n{field}: {value}\n---\n\n# X\n")
            r = subprocess.run([sys.executable, str(VALIDATE), "--root", str(root),
                                "--json"], capture_output=True, text=True)
            if not r.stdout.strip():
                problems.append(f"{field} written as a list took the validator down: "
                                f"{(r.stderr or '').strip().splitlines()[-1:] or '(silence)'}. "
                                "Every artifact it had not reached goes unreported with it")
            elif "FM002" not in {x["code"] for x in json.loads(r.stdout)["findings"]}:
                problems.append(f"{field} written as a list was neither reported nor fatal, "
                                "so the field reads as filled in and nothing looks at it")
            f.unlink()
    return problems


@check("a citation to a term, and a decision that says what it left open, both resolve")
def _references_with_a_second_end():
    # Two pairs that had one end each. A data contract sending a reader to the glossary for
    # what a column means, and a decision declaring in its prose that something stays open:
    # both were references, and nothing resolved either of them.
    fm = lambda **kw: "---\n" + "\n".join(f"{k}: {v}" for k, v in kw.items()) + "\n---\n\n"

    def repo(terms: str, cite: str, leaves: str | None, sections: str = "") -> dict:
        dec = {"schema": "framework/decision-record/v1",
               "artifact_type": "decision-record", "id": "DEC-001",
               "lifecycle": "immutable", "status": "accepted", "scope": "architecture",
               "products": "[alpha]", "owners": "[o]", "created": "2026-01-01"}
        if leaves is not None:
            dec["leaves_open"] = leaves
        return {
            "framework.yaml": f"framework_version: {REGISTRY['version']}\n",
            "OPEN.md": fm(schema="framework/open-register/v1",
                          artifact_type="open-register", lifecycle="living",
                          status="active", owners="[o]", created="2026-01-01 09:00",
                          last_review="2026-08-01 09:00",
                          entries="\n  OD-001:\n    status: open\n"
                                  "    cost_to_reverse: low\n    products: [all]\n"
                                  "    default_in_force: nothing\n") + "# Open\n",
            "GLOSSARY.md": fm(schema="framework/glossary/v1", artifact_type="glossary",
                              lifecycle="living", status="active", owners="[o]",
                              created="2026-01-01 09:00", last_review="2026-08-01 09:10",
                              terms=terms) + "# Glossary\n\n" + sections,
            "decisions/DEC-001-slug.md": fm(**dec) + "# DEC-001\n",
            "products/alpha/contracts/DC-001-x.md": fm(
                schema="framework/data-contract/v1", artifact_type="data-contract",
                id="DC-001", lifecycle="living", status="active", products="[alpha]",
                consumers="[alpha]", owners="[o]", created="2026-01-01 09:00",
                last_review="2026-08-01 09:20") + f"# DC-001\n\n{cite}\n",
        }

    problems = []
    with tempfile.TemporaryDirectory() as tmp:
        root = Path(tmp)

        def codes(files) -> set[str]:
            for stale in root.rglob("*.md"):
                stale.unlink()
            for rel, text in files.items():
                f = root / rel
                f.parent.mkdir(parents=True, exist_ok=True)
                f.write_text(text)
            r = subprocess.run([sys.executable, str(VALIDATE), "--root", str(root),
                                "--json"], capture_output=True, text=True)
            return {f["code"] for f in json.loads(r.stdout)["findings"]}

        defined = "\n  Freshness:\n    kind: metric\n"
        blocked = "\n  Freshness:\n    kind: metric\n    blocked_by: OD-001\n"
        stale_block = "\n  Freshness:\n    kind: metric\n    blocked_by: OD-404\n"

        if "REF005" not in codes(repo(defined, "`GLOSSARY §Tenant` is the unit", "[]")):
            problems.append("a citation to a term no glossary declares was not reported: it "
                            "reads as defined, and the reader stops looking")
        # A CITATION IN A TABLE CELL, and one that points at a section of the glossary
        # rather than a term. The first ran to the end of the row and produced a finding
        # naming a string with a pipe in it; the second is the one form of reference that
        # cannot be wrong, and it resolved to no term by construction.
        if "REF005" in codes(repo(defined,
                                  "| field | `GLOSSARY §Freshness` | routed |", "[]")):
            problems.append("a citation inside a table cell was reported: the match ran past "
                            "the cell and the term it names is not what the document wrote")
        # A section of the glossary, which is a reference that resolves by construction --
        # and the section sign is what makes it one. Without requiring it, the exemption
        # covered every heading in the file and handed back the hole `terms:` was added to
        # close: a word defined only as `### Freshness` in the body resolved a citation.
        if "REF005" in codes(repo(defined, "see `GLOSSARY §Domain terms`", "[]",
                                  sections="## §Domain terms\n")):
            problems.append("a citation naming a section of the glossary was reported: "
                            "pointing a reader at one is a reference that resolves by "
                            "construction, and a check that only knows terms reports the one "
                            "form that cannot be wrong")
        if "REF005" not in codes(repo(defined, "see `GLOSSARY §Tenant`", "[]",
                                      sections="### Tenant\n")):
            problems.append("a term written only as a heading in the glossary body resolved "
                            "a citation. That is resolving against prose headings, which is "
                            "what `terms:` exists to stop -- two checks here went quiet for "
                            "it once already")

        if "REF005" in codes(repo(defined, "`GLOSSARY §Freshness` is the unit", "[]")):
            problems.append("a citation to a declared term was reported, so the field cannot "
                            "be answered")
        if "REF005" in codes(repo(blocked, "`GLOSSARY §Freshness` is the unit", "[]")):
            problems.append("a citation to a term declared blocked by an open entry was "
                            "reported: the gap is written down, which is the difference "
                            "between a word waiting on a decision and a word nobody defined")
        # `kind` unless `blocked_by`. A term can be blocked precisely because nobody knows
        # yet whether it is a metric or a domain concept, and requiring `kind` there makes
        # the writer choose the thing `blocked_by` exists to avoid choosing.
        bare_block = "\n  Freshness:\n    blocked_by: OD-001\n"
        if "FM002" in codes(repo(bare_block, "# nothing cited here", "[]")):
            problems.append("a term declaring only `blocked_by` was rejected by the schema: "
                            "the word is used, is not defined, and what kind of thing it is "
                            "may be part of what the decision has to settle")
        if "FM002" not in codes(repo("\n  Freshness: {}\n", "# nothing cited here", "[]")):
            problems.append("a term declaring neither `kind` nor `blocked_by` was accepted, "
                            "so an empty row passes as a definition")

        if "REF005" not in codes(repo(stale_block, "# nothing cited here", "[]")):
            problems.append("a term blocked by an entry no register declares was not "
                            "reported, so the reason the word has no definition points at "
                            "nothing")

        if "REG012" not in codes(repo(defined, "# nothing cited here", None)):
            problems.append("a decision saying nothing about what it leaves open was not "
                            "reported: `[]` is a statement and silence is a gap")
        if "REG012" in codes(repo(defined, "# nothing cited here", "[]")):
            problems.append("a decision declaring `leaves_open: []` was reported, so there "
                            "is no way to say a decision settled everything it touched")
        if "REG012" in codes(repo(defined, "# nothing cited here", "[OD-001]")):
            problems.append("a decision naming an entry a register declares was reported")
        # THE THIRD STATE. A decision that left something open whose entry nobody has
        # written is neither a list nor a silence, and written as a silence it produced the
        # finding for "nobody looked" on a decision where somebody had looked and said so.
        if "REG014" not in codes(repo(defined, "# nothing cited here", "[unregistered]")):
            problems.append("a decision declaring `[unregistered]` was not reported: the "
                            "debt is acknowledged and the question is still one nobody can "
                            "find, rank or count")
        if "REG012" in codes(repo(defined, "# nothing cited here", "[unregistered]")):
            problems.append("`unregistered` was resolved as an entry id, so the reserved "
                            "word reads as a dangling reference and the two states collapse "
                            "back into one")
        both = codes(repo(defined, "# nothing cited here", "[OD-001, unregistered]"))
        if "REG014" not in both or "REG012" in both:
            problems.append("a decision naming a real entry and an unregistered one was not "
                            "reported as the second and only the second: the ids it does "
                            "know are not the part that is missing")

        if "REG012" not in codes(repo(defined, "# nothing cited here", "[OD-404]")):
            problems.append("a decision leaving an entry open that no register declares was "
                            "not reported: the open half of a decision is only open if "
                            "somebody can find it")
    return problems


@check("--emit-index refuses to overwrite an index somebody maintains by hand")
def _emit_index_refuses():
    # `--emit-index` writes only what front matter can express, and the reason a project
    # keeps one of these files by hand is always the part it cannot: a column for why a
    # decision still matters, a row for where a source system enters the chain. The skill
    # used to list this under "apply without asking", which is how an agent deletes that
    # content while believing it is tidying up. The marker is the permission to overwrite.
    fm = lambda **kw: "---\n" + "\n".join(f"{k}: {v}" for k, v in kw.items()) + "\n---\n\n"
    dec = fm(schema="framework/decision-record/v1", artifact_type="decision-record",
             id="DEC-001", lifecycle="immutable", status="accepted", scope="architecture",
             products="[alpha]", owners="[owner]", created="2026-01-01") + "# DEC-001\n"
    by_hand = "# Decision index\n\nHand maintained.\n\n| ID | Why it still matters |\n"
    problems = []

    with tempfile.TemporaryDirectory() as tmp:
        root = Path(tmp)
        (root / "decisions").mkdir(parents=True)
        (root / "decisions" / "DEC-001-slug.md").write_text(dec)
        index = root / "decisions" / "INDEX.md"
        index.write_text(by_hand)

        run = lambda *extra: subprocess.run(
            [sys.executable, str(VALIDATE), "--root", str(root), "--json",
             "--emit-index", "--stale-days", "36500", *extra],
            capture_output=True, text=True)

        r = run()
        if index.read_text() != by_hand:
            problems.append("--emit-index overwrote a hand maintained decisions/INDEX.md")
        if r.returncode == 0:
            problems.append("--emit-index refused the write and still exited 0: a caller "
                            "cannot tell the file was not regenerated")
        if r.returncode in (0, 1):
            rel = "decisions/INDEX.md"
            if rel not in json.loads(r.stdout).get("hand_maintained", []):
                problems.append(f"{rel} was not reported under hand_maintained")

        # --check must not call it out of date: that wording is an instruction to run the
        # overwrite, which is the thing being prevented.
        c = run("--check")
        if c.returncode in (0, 1):
            out = json.loads(c.stdout)
            if "decisions/INDEX.md" in out.get("out_of_date", []):
                problems.append("--emit-index --check reports a hand maintained index as "
                                "out of date, which points the reader at the overwrite")

        # The marker is what grants permission, so a generated file still regenerates.
        index.write_text(GENERATED_MARK + "\nstale\n")
        run()
        if GENERATED_MARK not in index.read_text() or "stale" in index.read_text():
            problems.append("a marked index was not regenerated: the guard is too wide")

    return problems


@check("a register vouches for the identifiers it declares, not the ones it mentions")
def _inline_ids_are_scoped():
    # Both directions, because this check exists to hold a line between them. A register
    # has to vouch for its own entries or every OD reference reads as dangling; it must not
    # vouch for the identifiers it merely cites, or REF001 is cleared by writing a sentence
    # about the missing document, which is precisely what the skill tells you to do.
    fm = lambda **kw: "---\n" + "\n".join(f"{k}: {v}" for k, v in kw.items()) + "\n---\n\n"
    dec = lambda i, **extra: fm(
        schema="framework/decision-record/v1", artifact_type="decision-record", id=i,
        lifecycle="immutable", status="accepted", scope="architecture",
        products="[alpha]", owners="[owner]", created="2026-01-01", **extra) + f"# {i}\n"

    files = {
        # Declares OD-001 and KI-001; cites DEC-999, which it does not own.
        "OPEN.md": fm(schema="framework/open-register/v1", artifact_type="open-register",
                      lifecycle="living", status="active", owners="[owner]",
                      created="2026-01-01", last_review="2026-01-01 09:00")
        + "# Open\n\n- OD-001 - where the landing zone goes\n- KI-001 - nightly job is "
          "flaky\n\nDEC-999 was never written, and this sentence must not make it exist.\n",
        # Resolves against an entry inside the register: this is the case the loose rule
        # was there to serve, and it has to keep working.
        "decisions/DEC-001-slug.md": dec("DEC-001", derives_from="OD-001"),
        # Resolves against nothing at all.
        "decisions/DEC-002-slug.md": dec("DEC-002", derives_from="DEC-999"),
    }

    with tempfile.TemporaryDirectory() as tmp:
        for rel, text in files.items():
            p = Path(tmp) / rel
            p.parent.mkdir(parents=True, exist_ok=True)
            p.write_text(text)
        r = subprocess.run(
            [sys.executable, str(VALIDATE), "--root", tmp, "--json",
             "--stale-days", "36500"], capture_output=True, text=True)
        if r.returncode not in (0, 1):
            return [f"the validator crashed: {r.stderr.strip().splitlines()[-1]}"]
        ref001 = {f["path"] for f in json.loads(r.stdout)["findings"]
                  if f["code"] == "REF001"}

    problems = []
    if "decisions/DEC-001-slug.md" in ref001:
        problems.append("OD-001 is declared as an entry in OPEN.md and was still reported "
                        "dangling: the register no longer vouches for its own entries")
    if "decisions/DEC-002-slug.md" not in ref001:
        problems.append("DEC-999 exists nowhere, yet REF001 did not fire: naming it in the "
                        "prose of a register is enough to clear the finding")
    return problems


@check("a change cites the decision and the evaluation its classification demands")
def _change_contracts():
    # `CHG001` and `CHG002` were catalogued and off from the day they were written, because
    # the routing they needed existed only as prose in the `CHG` body and the recovered
    # version matched words in it. The `ICG` made the join a lookup instead.
    #
    # Both bind on `status`, and that is the half worth testing: the artifacts arrive in an
    # order. A `DEC` precedes the contract, an `EVR` follows the build. A check that
    # demanded either at `draft` would forbid the order the framework prescribes, and would
    # be switched off within a week for crying wolf on correct documents.
    fm = lambda **kw: "---\n" + "\n".join(f"{k}: {v}" for k, v in kw.items()) + "\n---\n\n"
    sections = ("<!-- section: what-changes -->\n## A\n"
                "<!-- section: what-must-not-change -->\n## B\n"
                "<!-- section: how-we-know-it-worked -->\n## C\n")

    base = {
        "OPEN.md": fm(schema="framework/open-register/v1", artifact_type="open-register",
                      lifecycle="living", status="active", owners="[o]",
                      created="2026-01-01", last_review="2026-01-01 09:00") + "# Open\n",
        "products/p/cycles/ICG-001.md": fm(
            schema="framework/impact-classification/v1",
            artifact_type="impact-classification", lifecycle="immutable",
            status="accepted", id="ICG-001", products="[p]", owners="[o]",
            created="2026-01-01",
            routing="\n  SIG-001: architecture",
            impacts="\n  SIG-001: [architecture, ai]")
        + "# ICG-001\n\n<!-- section: intake -->\n## I\n"
          "<!-- section: classification -->\n## C\n"
          "<!-- section: open-questions -->\n## O\n",
        "decisions/DEC-001.md": fm(
            schema="framework/decision-record/v1", artifact_type="decision-record",
            id="DEC-001", lifecycle="immutable", status="accepted", scope="architecture",
            products="[p]", owners="[o]", created="2026-01-01") + "# DEC-001\n",
        "products/p/releases/EVR-001.md": fm(
            schema="framework/evaluation-report/v1", artifact_type="evaluation-report",
            id="EVR-001", lifecycle="immutable", status="final", products="[p]",
            owners="[o]", created="2026-01-01") + "# EVR-001\n",
    }

    def chg(status, derives, verified_by="null", icg="ICG-001"):
        meta = dict(schema="framework/change-contract/v1",
                    artifact_type="change-contract", lifecycle="immutable",
                    status=status, id="CHG-001", products="[p]", owners="[o]",
                    created="2026-01-01", derives_from=derives, verified_by=verified_by)
        if icg:
            meta["icg"] = icg
        return fm(**meta) + "# CHG-001\n\n" + sections

    def codes(chg_text):
        files = dict(base, **{"products/p/changes/CHG-001.md": chg_text})
        with tempfile.TemporaryDirectory() as tmp:
            for rel, text in files.items():
                f = Path(tmp) / rel
                f.parent.mkdir(parents=True, exist_ok=True)
                f.write_text(text)
            r = subprocess.run(
                [sys.executable, str(VALIDATE), "--root", tmp, "--json",
                 "--stale-days", "36500"], capture_output=True, text=True)
            if r.returncode not in (0, 1) or not r.stdout.strip():
                return None
            return {f["code"] for f in json.loads(r.stdout)["findings"]
                    if f["code"].startswith("CHG")}

    cases = [
        ("verified, nothing cited", chg("verified", "[SIG-001]"), {"CHG001", "CHG002"}),
        ("verified, both cited",
         chg("verified", "[SIG-001, DEC-001]", "EVR-001"), set()),
        ("implemented, no EVR yet", chg("implemented", "[SIG-001, DEC-001]"), set()),
        ("approved, no DEC", chg("approved", "[SIG-001]"), {"CHG002"}),
        ("draft, no DEC", chg("draft", "[SIG-001]"), set()),
        ("verified_by names something that is not an EVR",
         chg("verified", "[SIG-001, DEC-001]", "DEC-001"), {"CHG001"}),
        # The escape, and the case that used to expect silence here. A change that names no
        # classification cleared both of the checks written for it, and the expectation
        # recorded in this list was the proof that the hole was deliberate rather than
        # accidental. CHG003 is what closes it, and the direction is the reason it exists:
        # without it these checks question whoever filled `icg` in and say nothing at all
        # to whoever left it out.
        ("verified with no icg: nothing to look up, and that is the finding",
         chg("verified", "[SIG-001]", icg=None), {"CHG003"}),
        ("approved with no icg, which is where authorization starts",
         chg("approved", "[SIG-001]", icg=None), {"CHG003"}),
        # A proposal written before its triage is the order the framework prescribes.
        ("draft with no icg is a change waiting to be classified, not one that dodged it",
         chg("draft", "[SIG-001]", icg=None), set()),
        ("icg naming a classification that does not exist",
         chg("approved", "[SIG-001, DEC-001]", icg="ICG-404"), {"CHG003"}),
    ]
    problems = []
    for what, text, want in cases:
        got = codes(text)
        if got is None:
            problems.append(f"{what}: the validator did not complete")
        elif got != want:
            problems.append(f"{what}: expected {sorted(want) or 'nothing'}, got "
                            f"{sorted(got) or 'nothing'}")
    return problems


@check("a repository says which framework it was written against, or is told it does not")
def _framework_version():
    # The point is telling a migration from a mistake. When the registry moves, findings
    # appear in repositories nobody touched, and without this line there is nothing to
    # distinguish "the rules changed" from "we did this wrong" — which need opposite
    # responses. Guessing wrong at that a few times is how a team stops reading the
    # validator.
    current = REGISTRY["version"]
    fm = lambda **kw: "---\n" + "\n".join(f"{k}: {v}" for k, v in kw.items()) + "\n---\n\n"
    base = {"OPEN.md": fm(schema="framework/open-register/v1",
                          artifact_type="open-register", lifecycle="living",
                          status="active", owners="[owner]", created="2026-01-01",
                          last_review="2026-01-01 09:00") + "# Open\n"}

    def codes(config: str | None):
        files = dict(base)
        if config is not None:
            files["framework.yaml"] = config
        with tempfile.TemporaryDirectory() as tmp:
            for rel, text in files.items():
                (Path(tmp) / rel).write_text(text)
            r = subprocess.run(
                [sys.executable, str(VALIDATE), "--root", tmp, "--json",
                 "--stale-days", "36500"], capture_output=True, text=True)
            # Exit 1 is how the validator reports findings, so it cannot also mean
            # "crashed", and a traceback exits 1 too. Unparseable output is the only
            # thing that tells them apart. Without this the test still catches a
            # regression, by exploding on empty stdout, and a selfcheck that crashes
            # says less than one that names what broke.
            if r.returncode not in (0, 1) or not r.stdout.strip():
                last = r.stderr.strip().splitlines()[-1] if r.stderr.strip() else "?"
                return None, last
            return {(f["code"], f["level"]) for f in json.loads(r.stdout)["findings"]
                    if f["code"].startswith("FW")}, None

    major, minor, _ = current.split(".")
    problems = []
    cases = [
        (None, {("FW002", "info")}, "no framework.yaml at all"),
        ("checks:\n  LC002: warn\n", {("FW002", "info")}, "a config with no version"),
        (f"framework_version: {current}\n", set(), "the current version"),
        (f'framework_version: "{current}"\n', set(), "the current version, quoted"),
        (f"framework_version: {int(major) + 1}.0.0\n", {("FW001", "warn")},
         "a major ahead"),
        (f"framework_version: {major}.{int(minor) + 1}.0\n", {("FW001", "warn")},
         "a minor ahead"),
        # A PATCH IS SILENT, AND ASSERTED IN BOTH DIRECTIONS. By this registry's own
        # definition a patch is wording, a message, a fixture: nothing a repository has to
        # act on. Reported, it asks every project to edit a line whenever a sentence here is
        # rephrased, and a finding whose correct response is "change the number until it
        # stops" is one people learn to clear without reading -- on the check that has to be
        # believed the day the rules really do move.
        (f"framework_version: {major}.{minor}.0\n", set(), "a patch behind"),
        (f"framework_version: {major}.{minor}.99\n", set(), "a patch ahead"),
        # The two YAML traps, and they are the reason the version is a string. `1.1` is a
        # decimal and `2` is a whole number, so neither is comparable with a version, and
        # both are what somebody reaches for when shortening it. Quoting, which used to be
        # the trap, is now fine: the value was always meant to be a string.
        (f"framework_version: {major}.{minor}\n", {("FW001", "warn")},
         "two components, which YAML reads as a decimal"),
        (f"framework_version: {major}\n", {("FW001", "warn")},
         "one component, which YAML reads as a whole number"),
    ]
    for config, want, what in cases:
        got, crash = codes(config)
        if crash:
            problems.append(f"{what}: the validator did not complete: {crash}")
        elif got != want:
            problems.append(f"{what}: expected {sorted(want) or 'nothing'}, got "
                            f"{sorted(got) or 'nothing'}")

    # A malformed version has to be told apart from a real skew in the text, not only in
    # the code: reported as a mismatch it sends somebody looking for a migration that does
    # not exist, which is the confusion this check was added to remove rather than cause.
    with tempfile.TemporaryDirectory() as tmp:
        for rel, text in base.items():
            (Path(tmp) / rel).write_text(text)
        (Path(tmp) / "framework.yaml").write_text(f"framework_version: {major}.{minor}\n")
        r = subprocess.run([sys.executable, str(VALIDATE), "--root", tmp, "--json",
                            "--stale-days", "36500"], capture_output=True, text=True)
        msg = next((f["message"] for f in json.loads(r.stdout)["findings"]
                    if f["code"] == "FW001"), "")
        if "three" not in msg and "dots" not in msg:
            problems.append("a version with two components is reported as a version "
                            f"mismatch rather than as a malformed one: {msg[:80]}")

    # `1.10.0` is after `1.9.0`, and string comparison says the opposite. It is the first
    # place a version made of three parts goes wrong quietly, and it goes wrong in the
    # flattering direction: a repository nine minors behind reads as ahead of the framework,
    # which nobody migrates.
    v = _load(VALIDATE, "validate")
    if not (v.semver("1.9.0") < v.semver("1.10.0")):
        problems.append("1.10.0 does not sort after 1.9.0: the version is being compared "
                        "as text, so a repository behind reads as one ahead")
    for bad in ("2", "1.1", "", "v1.1.0", "1.1.0-rc1"):
        if v.semver(bad) is not None:
            problems.append(f"{bad!r} parses as a version and should not")

    # An unrecognised key stops the validator, by design, so the new one has to be on the
    # allowed list or every project that adopts it cannot run the checks at all.
    got, crash = codes(f"framework_version: {current}\nstale_days: 30\n")
    if crash:
        problems.append(f"`framework_version` beside another key is rejected: {crash}")
    return problems


@check("a signal set aside and a signal nobody read are told apart")
def _triage_state_is_derivable():
    # `LOG` is append-only, so no row can ever be marked handled, and triage state had
    # nowhere to live: every cycle re-read the whole log and guessed. It is derived from
    # the `ICG` now, which is only sound if a signal examined and dismissed counts as
    # examined. That is what `not-a-candidate` is for, and if it stopped counting the
    # check would nag forever about signals somebody already dealt with, which is how a
    # check gets switched off.
    fm = lambda **kw: "---\n" + "\n".join(f"{k}: {v}" for k, v in kw.items()) + "\n---\n\n"
    log = fm(schema="framework/signal-log/v1", artifact_type="signal-log",
             lifecycle="append-only", status="active", products="[alpha]",
             owners="[owner]", created="2026-01-01") \
        + "# Log\n\n| SIG-001 | drift |\n| SIG-002 | request |\n"

    def icg(**routing):
        head = fm(schema="framework/impact-classification/v1",
                  artifact_type="impact-classification", lifecycle="immutable",
                  status="accepted", id="ICG-001", products="[alpha]", owners="[owner]",
                  created="2026-01-01",
                  routing="\n" + "\n".join(f"  {k}: {v}" for k, v in routing.items()))
        return head + "# ICG-001\n\n<!-- section: intake -->\n## I\n" \
                      "<!-- section: classification -->\n## C\n" \
                      "<!-- section: open-questions -->\n## O\n"

    def run(files):
        with tempfile.TemporaryDirectory() as tmp:
            for rel, text in files.items():
                p = Path(tmp) / rel
                p.parent.mkdir(parents=True, exist_ok=True)
                p.write_text(text)
            r = subprocess.run(
                [sys.executable, str(VALIDATE), "--root", tmp, "--json",
                 "--stale-days", "36500"], capture_output=True, text=True)
            if r.returncode not in (0, 1):
                return None
            return [f for f in json.loads(r.stdout)["findings"] if f["code"] == "ICG001"]

    problems = []
    base = "products/alpha/LOG.md"
    cyc = "products/alpha/cycles/ICG-001.md"

    got = run({base: log})
    if got is None:
        return ["the validator crashed on a log with no impact classification"]
    if len(got) != 1:
        problems.append("a log with no ICG at all raised no ICG001: nothing in the "
                        "repository has been triaged and the check says so nowhere")
    elif "SIG-001" not in got[0]["message"] or "SIG-002" not in got[0]["message"]:
        problems.append(f"ICG001 does not name the untriaged signals: {got[0]['message']}")

    got = run({base: log, cyc: icg(**{"SIG-001": "architecture"})})
    if not got or "SIG-002" not in got[0]["message"]:
        problems.append("SIG-002 appears in no ICG and was not reported")
    elif "SIG-001" in got[0]["message"]:
        problems.append("SIG-001 was classified and is still reported as untriaged")

    got = run({base: log, cyc: icg(**{"SIG-001": "architecture",
                                      "SIG-002": "not-a-candidate"})})
    if got:
        problems.append("a signal routed `not-a-candidate` is still counted as never "
                        f"looked at: {got[0]['message']}. Examined and set aside is a "
                        "decision, and re-reporting it forever is how the check dies")
    return problems


@check("a gate verdict has a closed vocabulary, and so do the keys it is filed under")
def _maps_are_constrained():
    # `maps` exists so a document can carry a value per identifier as front matter instead
    # of as a table somebody has to read. It is worth nothing if anything validates: a
    # field with an open vocabulary is prose with a colon in it, and an open key is worse,
    # because the keys are what the rest of the framework joins on.
    #
    # Three shapes, and each is checked both ways. `one_of` is one value per key,
    # `any_of` a set, `fields` a record: the code map needs the third, because a
    # repository is a URL and a sentence saying what is in it, not a single word.
    typed = {n: s for n, s in REGISTRY["types"].items() if s.get("maps")}
    if not typed:
        return ["no type declares `maps`: this check is no longer running"]

    ids = "|".join(REGISTRY["id_prefixes"])
    problems = []
    for name, spec in typed.items():
        schema = json.loads((ROOT / "schemas" / "framework" / name / "v1.json")
                            .read_text(encoding="utf-8"))
        v = Draft202012Validator(schema)
        for field, rule in spec["maps"].items():
            # A key the pattern accepts, and some it must not. The accepted one is found
            # rather than assumed: `keys` is an arbitrary pattern, and hardcoding one shape
            # of nickname meant the check broke the first time a map keyed its rows on
            # identifiers instead.
            # THE ONE MAP WHOSE KEYS ARE NOT IDENTIFIERS. A glossary is keyed on the words
            # the business uses -- "Active customer", "Freshness", "Frontend" -- and there
            # is no vocabulary to close: any name a person writes is a name they write.
            # What protects a citation from a typo is not the pattern but `REF005`, which
            # resolves `GLOSSARY §Term` against these keys with case and whitespace
            # normalised and reports the ones that find nothing.
            if (name, field) == ("glossary", "terms"):
                continue

            pattern = rule.get("keys")
            candidates = ["SIG-001", "OD-001", "CMT-001", "RSK-001", "INC-001", "frontend",
                          "query-engine", "product.backend", "platform.access"]
            if pattern:
                accepted = [k for k in candidates if re.match(pattern, k)]
                if not accepted:
                    problems.append(f"{name}.{field}: the `keys` pattern accepts none of "
                                    f"{candidates}. Add one this check can use, or the map "
                                    "is not being checked at all")
                    continue
                good_key = accepted[0]
                bad_keys = [k for k in candidates if k not in accepted][:2] + ["Frontend"]
            else:
                good_key, bad_keys = "SIG-001", ["banana", "SIG_001"]

            if "one_of" in rule:
                ok_val, bad_vals = rule["one_of"][0], ["not-a-real-outcome"]
            elif "any_of" in rule:
                ok_val, bad_vals = [rule["any_of"][0]], [["not-a-real-outcome"], []]
            elif "scalar" in rule:
                # One constrained string per key. The commit map needs it: a hash is a
                # value, not a record.
                ok_val = "9a734ce" if rule["scalar"] else "anything"
                bad_vals = ["", "not a commit at all"] if rule["scalar"] else [""]
            elif "fields" in rule:
                req = list(rule["fields"]["required"])
                enums = rule["fields"].get("enums") or {}
                # A field of the record may carry its own closed vocabulary, so the legal
                # entry cannot be built out of a filler string for every key.
                def legal(k, _e=enums):
                    return _e[k][0] if k in _e else "x"
                ok_val = {k: legal(k) for k in req}
                bad_vals = [{}, dict(ok_val, surprise="x")]
                if len(req) > 1:
                    bad_vals.append({req[0]: legal(req[0])})   # one required key missing
                for k, values in enums.items():
                    if k in req:
                        bad_vals.append(dict(ok_val, **{k: "not-in-the-vocabulary"}))
                    if not values:
                        problems.append(f"{name}.{field}.enums.{k}: empty vocabulary")
            else:
                problems.append(f"{name}.{field}: declares no value shape")
                continue

            base = {"schema": f"framework/{name}/v1", "artifact_type": name,
                    "lifecycle": spec["lifecycle"], "status": spec["status"][0],
                    "owners": ["someone"], "created": "2026-01-01"}
            if spec.get("products") == "exactly-one":
                base["products"] = ["alpha"]
            # The type's own required fields, or the document under test fails on those
            # instead of on the map, and the check stops being about the map.
            for f2 in spec.get("required", []):
                if f2 not in base and f2 not in spec["maps"]:
                    base[f2] = "something"
            for f2, values in (spec.get("enums") or {}).items():
                base[f2] = values[0]
            for f2, r2 in spec["maps"].items():
                if not r2.get("required"):
                    continue
                k2 = "SIG-001" if not r2.get("keys") else "frontend"
                if "one_of" in r2:
                    base[f2] = {k2: r2["one_of"][0]}
                elif "any_of" in r2:
                    base[f2] = {k2: [r2["any_of"][0]]}
                else:
                    base[f2] = {k2: {k: "x" for k in r2["fields"]["required"]}}

            if not v.is_valid(dict(base, **{field: {good_key: ok_val}})):
                err = next(v.iter_errors(dict(base, **{field: {good_key: ok_val}}))).message
                problems.append(f"{name}.{field}: rejected a legal entry: {err}")
            for bad in bad_vals:
                if v.is_valid(dict(base, **{field: {good_key: bad}})):
                    problems.append(f"{name}.{field}: accepted {bad!r} as a value, so the "
                                    "shape is open and the field is prose with a colon")
            for bad in bad_keys:
                if v.is_valid(dict(base, **{field: {bad: ok_val}})):
                    problems.append(f"{name}.{field}: accepted {bad!r} as a key, so a "
                                    "mistyped name passes and joins on nothing")
    return problems


@check("a key one letter from a field that is read is reported, and a real field is not")
def _key_typos():
    # Declaring a field types its value and not its name, so `supercedes:` is an unknown key
    # and an unknown key is silence: the document validates, every check reading `supersedes`
    # sees nothing, and the line sits there looking right. Both directions matter here more
    # than usual, because the check works on a similarity and the false positive is a
    # legitimate field somebody chose.
    fm = lambda **kw: "---\n" + "\n".join(f"{k}: {v}" for k, v in kw.items()) + "\n---\n\n"

    def dec(n: str, extra: dict) -> str:
        return fm(schema="framework/decision-record/v1", artifact_type="decision-record",
                  id=n, lifecycle="immutable", status="accepted", scope="architecture",
                  products="[alpha]", owners="[o]", created="2026-01-01 09:00",
                  leaves_open="[]", **extra) + f"# {n}\n"

    problems = []
    with tempfile.TemporaryDirectory() as tmp:
        root = Path(tmp)
        (root / "framework.yaml").write_text(f"framework_version: {REGISTRY['version']}\n")
        (root / "decisions").mkdir()

        def codes(docs: dict) -> set[str]:
            for f in (root / "decisions").glob("*.md"):
                f.unlink()
            for n, extra in docs.items():
                (root / "decisions" / f"{n}-x.md").write_text(dec(n, extra))
            r = subprocess.run([sys.executable, str(VALIDATE), "--root", str(root),
                                "--json"], capture_output=True, text=True)
            return {f["code"] for f in json.loads(r.stdout)["findings"]}

        cases = [
            ({"DEC-001": {"supercedes": "DEC-002"}}, True, "`supercedes` for `supersedes`"),
            ({"DEC-001": {"owner": "[maria]"}}, True, "`owner` for `owners`"),
            ({"DEC-001": {"product": "alpha"}}, True,
             "`product` for `products`, which this framework renamed once for this reason"),
            ({"DEC-001": {"supersedes": "DEC-002"}}, False, "the field spelled right"),
            ({"DEC-001": {"approvers": "[g]"}}, False,
             "a field nothing declares and nothing resembles"),
            # A typo somebody copies is likelier than one somebody makes once, so repetition
            # must not turn it into vocabulary.
            ({"DEC-001": {"supercedes": "DEC-002"}, "DEC-002": {"supercedes": "DEC-001"}},
             True, "the same typo in two documents"),
            ({"DEC-001": {"approvers": "[g]"}, "DEC-002": {"approvers": "[m]"}}, False,
             "a real field in two documents"),
        ]
        for docs, want, what in cases:
            got = "FM006" in codes(docs)
            if got != want:
                problems.append(f"{what}: FM006 was {'not ' if want else ''}reported")
    return problems


@check("supersedes is a declared field with a shape, and not a convention")
def _supersedes_is_typed():
    # It was read by the validator and written in a template for weeks without being in any
    # schema, so `supersedes: 12` and `supersedes: {id: DEC-004}` were as legal as the thing
    # anybody meant, and the shape was whatever the last person wrote. Declared on every type
    # whose `status` can reach `superseded`, because those are the documents that can be
    # replaced.
    #
    # The three legal forms are all in use and none is a mistake: absent, one id, several.
    # `null` stays legal because the template writes it -- a field shown as existing and
    # holding nothing is how the templates say "this is yours to fill in".
    problems = []
    for name, spec in REGISTRY["types"].items():
        if "superseded" not in (spec.get("status") or []):
            continue
        if "supersedes" not in (spec.get("id_fields") or []):
            problems.append(f"{name} can reach `status: superseded` and does not declare "
                            "`supersedes`: the field that says what replaced it is a "
                            "convention there, and a convention has no shape")
            continue
        schema = json.loads((ROOT / "schemas" / "framework" / name / "v1.json")
                            .read_text(encoding="utf-8"))
        prop = schema["properties"].get("supersedes")
        if not prop or "oneOf" not in prop:
            problems.append(f"{name}: `supersedes` is declared in the registry and absent "
                            "or untyped in the generated schema")
            continue
        # The property on its own, and not a whole document with a `supersedes` in it: a
        # synthetic `ICG` missing its required maps is invalid for reasons that have nothing
        # to do with this field, and a probe that cannot tell the two apart reports the
        # field as broken on every type whose front matter is hard to fake.
        v = Draft202012Validator(prop)
        legal = [None, "DEC-004", ["DEC-004", "DEC-007"], "DEC-NNN"]
        illegal = [12, {"id": "DEC-004"}, "DEC4", "XYZ-001", ["DEC-004", "DEC-004"], []]
        for value in legal:
            if not v.is_valid(value):
                problems.append(f"{name}: rejected {value!r}, which is one of the forms in "
                                "use -- absent, one id, several, and the `null` a template "
                                "writes to show the field is yours to fill in")
        for value in illegal:
            if v.is_valid(value):
                problems.append(f"{name}: accepted {value!r} as a supersedes, so the field "
                                "is declared and still shapeless")
    return problems


@check("supersedes reads as a list, and a diamond is not a cycle")
def _supersedes_takes_a_list():
    # `supersedes: [DEC-001, DEC-004]` is how anyone replacing two decisions at once would
    # write it, and `derives_from` beside it already takes a list. The string-only read cost
    # twice: REF002 and REF003 skipped the list form without a word, so two documents could
    # both claim to be current, and then the cycle walk died on an unhashable type, which
    # takes the whole gate down with no verdict at all.
    fm = lambda **kw: "---\n" + "\n".join(f"{k}: {v}" for k, v in kw.items()) + "\n---\n\n"
    def dec(i, **extra):
        meta = dict(schema="framework/decision-record/v1",
                    artifact_type="decision-record", id=i, lifecycle="immutable",
                    status="accepted", scope="architecture", products="[alpha]",
                    owners="[owner]", created="2026-01-01")
        meta.update(extra)                      # so a case can set its own status
        return fm(**meta) + f"# {i}\n"

    def codes(files):
        with tempfile.TemporaryDirectory() as tmp:
            for name, text in files.items():
                p = Path(tmp) / "decisions" / name
                p.parent.mkdir(parents=True, exist_ok=True)
                p.write_text(text)
            r = subprocess.run(
                [sys.executable, str(VALIDATE), "--root", tmp, "--json",
                 "--stale-days", "36500"], capture_output=True, text=True)
            if r.returncode not in (0, 1) or not r.stdout.strip():
                last = r.stderr.strip().splitlines()[-1] if r.stderr.strip() else "?"
                return None, last
            return [f["code"] for f in json.loads(r.stdout)["findings"]], None

    problems = []

    # Two at once. Both targets are still `accepted`, so both have to be reported.
    got, crash = codes({"DEC-001.md": dec("DEC-001"), "DEC-002.md": dec("DEC-002"),
                        "DEC-003.md": dec("DEC-003", supersedes="[DEC-001, DEC-002]")})
    if crash:
        problems.append(f"a list-valued supersedes brought the validator down: {crash}")
    elif got.count("REF003") != 2:
        problems.append(f"supersedes: [DEC-001, DEC-002] reported REF003 "
                        f"{got.count('REF003')} times out of 2: the list form is being "
                        "read past in silence, and both documents still read as current")

    # A cycle written with lists is still a cycle.
    got, crash = codes({"DEC-001.md": dec("DEC-001", supersedes="[DEC-002]"),
                        "DEC-002.md": dec("DEC-002", supersedes="[DEC-001]")})
    if crash:
        problems.append(f"a cyclic supersedence chain crashed the validator: {crash}")
    elif "REF004" not in got:
        problems.append("DEC-001 and DEC-002 supersede each other and REF004 did not fire")

    # A diamond is not a cycle. DEC-004 replaces two decisions that both replaced DEC-001,
    # so the walk meets DEC-001 twice by different routes. Nothing here loops, and a plain
    # visited set would call that meeting a loop and report a cycle that is not there.
    got, crash = codes({
        "DEC-001.md": dec("DEC-001", status="superseded"),
        "DEC-002.md": dec("DEC-002", status="superseded", supersedes="[DEC-001]"),
        "DEC-003.md": dec("DEC-003", status="superseded", supersedes="[DEC-001]"),
        "DEC-004.md": dec("DEC-004", supersedes="[DEC-002, DEC-003]")})
    if crash:
        problems.append(f"a branching supersedence chain crashed the validator: {crash}")
    elif "REF004" in got:
        problems.append("two decisions replaced by the same later one were reported as a "
                        "cyclic chain: the walk is confusing a diamond with a loop")

    # The scalar form is what FRAMEWORK.md writes, and it has to keep working.
    got, crash = codes({"DEC-001.md": dec("DEC-001"),
                        "DEC-002.md": dec("DEC-002", supersedes="DEC-001")})
    if crash or "REF003" not in (got or []):
        problems.append("supersedes as a bare string stopped being checked")

    return problems


@check("the extractor says which documents gave it nothing")
def _extract_reports_its_gaps():
    # The other executable in this repository, and until now the one nothing ran. It is
    # allowed to fail on a document: a scanned PDF has no text layer and that is not a bug.
    # What it is not allowed to do is leave the failure only in `inventory.json`, because
    # the corpus then reads as complete to whoever classifies it. Nothing here needs
    # markitdown or python-docx: plain files that yield no text exercise the same path.
    corpus = {
        "offerta.md": "# Offerta\n\nRefresh orario dei dati, non giornaliero.\n",
        "empty.md": "",
        "whitespace.txt": "   \n\t\n",
        "notes.rst": "unsupported extension, must not be counted as ingested\n",
    }
    problems = []

    with tempfile.TemporaryDirectory() as tmp:
        src = Path(tmp) / "corpus"
        src.mkdir()
        for name, text in corpus.items():
            (src / name).write_text(text, encoding="utf-8")
        out = Path(tmp) / "out"

        r = subprocess.run([sys.executable, str(EXTRACT), str(src), "-o", str(out)],
                           capture_output=True, text=True)
        if r.returncode != 0:
            return [f"the extractor exited {r.returncode} on a readable corpus: "
                    f"{r.stderr.strip().splitlines()[-1] if r.stderr.strip() else ''}"]

        md = (out / "extract.md").read_text(encoding="utf-8")
        inv = json.loads((out / "inventory.json").read_text(encoding="utf-8"))

        if "Refresh orario" not in md:
            problems.append("the one document with text is missing from extract.md")
        if "offerta.md" not in md:
            problems.append("a block in extract.md does not name the file it came from: "
                            "provenance is the whole point of the extraction")
        for name in ("empty.md", "whitespace.txt"):
            if name not in md:
                problems.append(f"{name} produced no text and extract.md does not say so: "
                                "the gap is only in inventory.json, where nobody looks")

        listed = {d["source"] for d in inv["documents"]}
        if listed != {"offerta.md", "empty.md", "whitespace.txt"}:
            problems.append(f"inventory.json accounts for {sorted(listed)}: every handled "
                            "document belongs there, and nothing else does")

        # An empty corpus is a failure and has to be one: a caller that reads exit 0 goes on
        # to classify an extraction that never happened.
        empty = Path(tmp) / "empty-dir"
        empty.mkdir()
        e = subprocess.run([sys.executable, str(EXTRACT), str(empty), "-o", str(out)],
                           capture_output=True, text=True)
        if e.returncode == 0:
            problems.append("the extractor exited 0 on a corpus with nothing in it")

        f = Path(tmp) / "not-a-directory.md"
        f.write_text("x", encoding="utf-8")
        o = subprocess.run([sys.executable, str(EXTRACT), str(src), "-o", str(f)],
                           capture_output=True, text=True)
        if o.returncode == 0 or "Traceback" in o.stderr:
            problems.append("-o pointed at a file should say so and stop, not raise: "
                            f"rc={o.returncode}, stderr={o.stderr.strip()[:80]!r}")

    # The text-poor pages branch, reached through the pure builder. Every format that can
    # set `visual_review` goes through poppler, markitdown or LibreOffice, none of which a
    # bare CI runner has, so driving the CLI would leave this untested wherever it matters.
    x = _load(EXTRACT, "extract")
    seen = x.build_extract_md(
        2, [x.Block("deck.pdf", "page 1", "Architettura")],
        [], [x.DocInfo("deck.pdf", "pdf", units=9, visual_review=[2, 3, 4])])
    if "deck.pdf" not in seen or "2, 3, 4" not in seen:
        problems.append("a document with text-poor pages is not named in extract.md with "
                        "its page numbers: the warning lives only on stdout again")

    quiet = x.build_extract_md(1, [x.Block("a.md", "document", "hi")], [],
                               [x.DocInfo("a.md", "md", units=1)])
    if "look at" in quiet or "produced no text" in quiet:
        problems.append("extract.md warns about a corpus that had nothing wrong with it")

    return problems


@check("the corpus is found by what is in a folder, not by what the folder is called")
def _corpus_is_found_by_content():
    # `start` has to know where the client's documents are before it can move them, and the
    # answer is not knowable in advance: corpus, docs, documenti, a folder named after the
    # customer, or loose at the root. Matching names works for the first two and fails
    # silently on the rest, and a silent failure here builds a repository around a corpus
    # nobody read.
    x = _load(EXTRACT, "extract")
    problems = []

    def build(tmp: Path, files: dict[str, str]):
        for rel, text in files.items():
            f = tmp / rel
            f.parent.mkdir(parents=True, exist_ok=True)
            f.write_text(text, encoding="utf-8")

    artifact = ("---\nschema: framework/glossary/v1\nartifact_type: glossary\n"
                "lifecycle: living\nstatus: active\nowners: [o]\n"
                "created: 2026-01-01 09:00\n---\n\n# Glossary\n")

    cases = [
        ("a folder nobody would have guessed the name of",
         {"Materiale Cliente 2026/offerta.pdf": "%PDF-1.4\n",
          "Materiale Cliente 2026/note.txt": "appunti"},
         [("Materiale Cliente 2026", 1, 1)]),
        # The case a location rule cannot survive: the corpus sits beside the artifacts.
        ("a corpus loose at the root of a scaffolded repository",
         {"GLOSSARY.md": artifact, "OPEN.md": artifact,
          "products/alpha/PBR.md": artifact, "deck.pdf": "%PDF-1.4\n"},
         [(".", 1, 0)]),
        # Output of a previous run. Counting it would make every second run find a corpus.
        ("the extractor's own output",
         {"_meta/corpus/alpha/deck.pdf": "%PDF-1.4\n",
          "_meta/extract/extract.md": "# Business corpus extraction\n",
          "_meta/extract/alpha/extract.md": "# Business corpus extraction\n"},
         [("_meta/corpus/alpha", 1, 0)]),
        ("a project with code and no documents", {"src/app.py": "x = 1\n"}, []),
    ]
    for what, files, want in cases:
        with tempfile.TemporaryDirectory() as tmp:
            build(Path(tmp), files)
            got = x.find_corpus(Path(tmp))
        if got != want:
            problems.append(f"{what}: found {got}, expected {want}")

    # Two folders of documents is the one case that has to stop and ask. The runner-up is
    # usually last quarter's version of the same deck, and which is current is not a thing
    # a file count knows.
    with tempfile.TemporaryDirectory() as tmp:
        build(Path(tmp), {"corpus/a.pdf": "%PDF-1.4\n", "corpus/b.pdf": "%PDF-1.4\n",
                          "vecchi/c.pptx": "PK\x03\x04"})
        strong = [g for g in x.group_corpus(x.find_corpus(Path(tmp))) if g["documents"]]
    if len(strong) != 2:
        problems.append(f"two folders of documents came back as {len(strong)} candidate(s): "
                        "the skill decides whether to ask on this count")

    # One corpus filed by kind. Reported as three candidates it forces a question with no
    # right answer, and the answer to it would be wrong however it was given: they are all
    # the corpus. The extractor reads a directory recursively, so the group is usable as-is.
    with tempfile.TemporaryDirectory() as tmp:
        build(Path(tmp), {"docs/contracts/a.pdf": "%PDF", "docs/contracts/b.pdf": "%PDF",
                          "docs/decks/c.pptx": "PK", "docs/spreadsheets/d.xlsx": "PK"})
        got = x.group_corpus(x.find_corpus(Path(tmp)))
    if len(got) != 1 or got[0]["path"] != "docs" or got[0]["documents"] != 4:
        problems.append(f"a corpus split across docs/contracts, docs/decks and "
                        f"docs/spreadsheets came back as {[g['path'] for g in got]}")
    elif len(got[0]["children"]) != 3:
        problems.append("the group hid the subdirectories it gathered: one of them being an "
                        "older version of another is exactly what a person has to be shown")
    return problems


@check("an exclusion names a path when a bare name would be too blunt")
def _skips_are_scoped():
    # `_meta/extract` holds the extractor's output and no source document. Written as
    # `extract` it would also exclude the extraction step of every ETL project that keeps
    # one in a directory of that name, which is the mistake the registry already warns
    # about in the paragraph above the list. Bare names still work, and have to: `corpus`
    # means the same thing wherever it turns up.
    #
    # `_meta/business` is the second one of these and the temptation was stronger, because
    # `business` reads like a word no Data/AI project would use for artifacts -- and a
    # `business/` directory of domain models is exactly what several of them keep. The skill
    # writes under `_meta/`, so the narrow form costs nothing and the broad one would have
    # hidden somebody's documents from every check in the catalog.
    x = _load(VALIDATE, "validate")
    skips = set(x.REGISTRY_SCAN["skip_dirs"]) if hasattr(x, "REGISTRY_SCAN") \
        else set(yaml.safe_load((ROOT / "schemas" / "artifact-types.yaml")
                                .read_text())["scan"]["skip_dirs"])
    problems = []
    for parts, want, what in [
        (("_meta", "extract"), True, "the extractor's own output directory"),
        (("_meta", "extract", "gamma"), True, "a per-product folder inside it"),
        (("extract",), False, "an ETL directory called extract at the root"),
        (("pipelines", "extract"), False, "an extract step nested in a project's pipelines"),
        (("_meta", "business"), True, "where the `business` skill writes its updates"),
        (("_meta", "business", "archive"), True, "a folder somebody made inside it"),
        (("business",), False, "a project's own business directory at the root"),
        (("products", "alpha", "business"), False, "the same name under a product, which no "
                                                   "skill writes and nothing may hide"),
        (("_meta", "presentation"), True, "where the `presentation` skill writes its decks"),
        (("presentation",), False, "a project's own presentation directory at the root"),
        (("_meta", "digest"), True, "the clone where the `digest` skill keeps its state"),
        (("_meta", "digest", "DIG-001-atlas-2026-10-06"), True, "a snapshot inside it"),
        (("digest",), False, "a project's own digest directory at the root"),
        (("_meta", "corpus"), True, "the corpus under _meta"),
        (("products", "alpha", "corpus"), True, "a corpus left where older repos keep it"),
        (("_meta",), False, "_meta itself, which can hold artifacts and must be scanned"),
        (("products", "alpha"), False, "an ordinary product directory"),
    ]:
        if x.skipped_dir(parts, skips) != want:
            problems.append(
                f"{what} ({'/'.join(parts)}) was {'not ' if want else ''}excluded"
                + ("" if want else ": an exclusion that protects the framework's "
                   "convenience is paid for by everyone using it"))
    return problems


@check("nothing anybody copies writes a scalar where the schema wants a list")
def _examples_agree_with_the_schema():
    # A template's front matter is validated against its own schema by the check above. Its
    # *body* is not, and neither is a fenced example in a skill, and both get copied: an
    # agent following `used_by: product-a` in a snippet produces a document that fails
    # `FM002`, and the failure surfaces in somebody's project rather than here.
    #
    # It happened. `depends_on` and `used_by` became lists in the registry and the templates
    # kept writing them as scalars for the minutes between the two edits, and a real run
    # landed in that window -- the plugin is a symlink to this working tree, so there is no
    # copy to be stale. Eighteen findings in a project because of a shape written in a file
    # that no check reads.
    want_list: set[str] = set()
    for spec in REGISTRY["types"].values():
        for field, rule in (spec.get("maps") or {}).items():
            want_list.update(rule.get("fields", {}).get("lists") or [])
            if "any_of" in rule:
                want_list.add(field)
    if not want_list:
        return ["no field is declared as a list: this check is no longer running"]

    # `key: value` where value is neither a flow sequence, a block sequence, nor a comment.
    scalar = re.compile(r"^\s*([a-z_]+):[ \t]*([^\s\[#][^#]*?)\s*$")
    problems = []
    for p in sorted(ROOT.rglob("*")):
        rel = p.relative_to(ROOT)
        # `schemas/` states the constraints rather than demonstrating them:
        # `products: exactly-one` in the registry is a cardinality declaration and not
        # a field somebody copies into a document.
        if p.suffix not in (".md", ".yaml", ".yml") or "build" in rel.parts:
            continue
        if rel.parts[0] in ("schemas", "tests"):
            continue
        if any(part.startswith(".") for part in rel.parts):
            continue
        for i, line in enumerate(p.read_text(encoding="utf-8", errors="replace").splitlines(), 1):
            m = scalar.match(line)
            if m and m.group(1) in want_list:
                problems.append(f"{rel}:{i}: `{line.strip()}` — the schema wants a list here, "
                                "and whoever copies this line gets an FM002 in their project")
    return problems


@check("a document that says it was decided has to name the decision")
def _claims_carry_their_record():
    # Two documents make the same claim in a field and used to make it unbacked. A stack row
    # saying `chosen` with no `decided_in` reads as a decision nobody can find, which is the
    # ambiguity the stack was added to remove, arriving back through the field meant to
    # remove it. An open register entry with no `default_in_force` is the same shape: the file
    # calls that field mandatory in its own instructions and nothing was reading it except
    # REG003, and only on a high cost entry.
    def repo(files):
        with tempfile.TemporaryDirectory() as tmp:
            for rel, text in files.items():
                f = Path(tmp) / rel
                f.parent.mkdir(parents=True, exist_ok=True)
                f.write_text(text)
            r = subprocess.run([sys.executable, str(VALIDATE), "--root", tmp, "--json",
                                "--stale-days", "36500"], capture_output=True, text=True)
            if r.returncode not in (0, 1) or not r.stdout.strip():
                return None
            return {f["code"] for f in json.loads(r.stdout)["findings"]}

    def stack(rows: str) -> str:
        return ("---\nschema: framework/operational-stack/v1\n"
                "artifact_type: operational-stack\nlifecycle: living\nstatus: active\n"
                "owners: [o]\ncreated: 2026-01-01 09:00\nlast_review: 2026-01-01 09:00\n"
                + rows + "---\n\n# Stack\n\n<!-- section: chosen -->\n## Chosen\n"
                "<!-- section: unratified -->\n## Unratified\n"
                "<!-- section: ruled-out -->\n## Ruled out\n")

    dec = ("---\nschema: framework/decision-record/v1\nartifact_type: decision-record\n"
           "id: DEC-001\nlifecycle: immutable\nstatus: accepted\nscope: architecture\n"
           "owners: [o]\ncreated: 2026-01-01 09:00\n---\n\n# DEC-001\n")
    draft = dec.replace("id: DEC-001", "id: DEC-002").replace("accepted", "proposed")
    second = dec.replace("DEC-001", "DEC-003")

    problems = []
    # Every `decided_in` here is bracketed, and that is the point of the case rather than a
    # detail of it. 2.8.8 declared the field a list and this block kept writing scalars, so
    # nothing ran the check against the shape the schema now requires -- and the check was
    # reading the raw value into a set membership test, which a list cannot survive. The
    # first repository to write one lost the report and the generated index with it. A test
    # that exercises a field in the shape no document is allowed to carry is not exercising
    # it.
    for rows, want, what in [
        ("stack:\n  db:\n    tool: PostgreSQL\n    status: chosen\n    decided_in: [DEC-001]\n",
         False, "a chosen tool naming an accepted decision"),
        ("stack:\n  db:\n    tool: PostgreSQL\n    status: chosen\n"
         "    decided_in: [DEC-001, DEC-003]\n",
         False, "a chosen tool ratified by two accepted decisions, which is what the list is for"),
        ("stack:\n  db:\n    tool: PostgreSQL\n    status: chosen\n"
         "    decided_in: [DEC-001, DEC-404]\n",
         True, "a list whose second entry is a decision that does not exist"),
        # The shape a repository carries between reading the migration note and finishing it.
        # It is an `FM002` and it has to stay one: a traceback out of `main()` takes the
        # report, and `--emit-index` shares that entry point.
        ("stack:\n  db:\n    tool: PostgreSQL\n    status: chosen\n    decided_in: DEC-001\n",
         False, "the pre-2.8.8 scalar, which is a schema finding and must not be a crash"),
        ("stack:\n  db:\n    tool: PostgreSQL\n    status: chosen\n",
         True, "a chosen tool naming no decision at all"),
        ("stack:\n  db:\n    tool: PostgreSQL\n    status: unratified\n    decided_in: [DEC-001]\n",
         True, "an unratified tool that names a decision, hiding one that was taken"),
        ("stack:\n  db:\n    tool: PostgreSQL\n    status: chosen\n    decided_in: [DEC-404]\n",
         True, "a decision that does not exist"),
        ("stack:\n  db:\n    tool: PostgreSQL\n    status: chosen\n    decided_in: [DEC-002]\n",
         True, "a decision that was never accepted"),
        ("stack:\n  db:\n    tool: PostgreSQL\n    status: unratified\n",
         False, "an unratified tool, which is the honest state and must stay quiet"),
        # The state the vocabulary did not have: tried, not in use, nobody decided. With
        # three values the only way to file it was `ruled-out` with nothing to name, and a
        # real repository wrote two of those -- so the check reported a document whose only
        # fault was that the words ran out.
        ("stack:\n  db:\n    tool: PostgreSQL\n    status: dropped\n",
         False, "an abandoned experiment nobody decided against, which is what `dropped` is"),
        ("stack:\n  db:\n    tool: PostgreSQL\n    status: dropped\n    decided_in: [DEC-001]\n",
         True, "a dropped tool naming a decision, which is a decision filed as an accident"),
    ]:
        got = repo({"STACK.md": stack(rows), "decisions/DEC-001.md": dec,
                    "decisions/DEC-002.md": draft, "decisions/DEC-003.md": second})
        if got is None:
            problems.append(f"the validator crashed on {what}")
        elif ("STK001" in got) != want:
            problems.append(f"{what} was {'not ' if want else ''}reported")

    def register(entries: str) -> str:
        return ("---\nschema: framework/open-register/v1\nartifact_type: open-register\n"
                "lifecycle: living\nstatus: active\nowners: [o]\n"
                "created: 2026-01-01 09:00\nlast_review: 2026-01-01 09:00\n"
                + entries + "---\n\n# Open\n")

    for entries, want, what in [
        ("entries:\n  OD-001:\n    status: open\n    cost_to_reverse: medium\n"
         "    default_in_force: the nightly job, unasked\n",
         False, "an entry with a cost and a default"),
        ("entries:\n  OD-001:\n    status: open\n    cost_to_reverse: medium\n",
         True, "a medium cost entry with no default, which used to pass"),
        ("entries:\n  OD-001:\n    status: open\n    default_in_force: none\n",
         True, "an entry with no cost to reverse, which is what orders the file"),
        ("entries:\n  OD-001:\n    status: decided\n    cost_to_reverse: low\n"
         "    default_in_force: none\n",
         True, "a decided entry naming no closed_by"),
        ("entries:\n  OD-001:\n    status: open\n    cost_to_reverse: low\n"
         "    default_in_force: none\n    depends_on: [OD-404]\n",
         True, "a dependency on an entry the register does not declare"),
        # A known issue is not a choice: it has no default in force and no cost to reverse.
        ("entries:\n  KI-001:\n    status: open\n", False, "a known issue"),
    ]:
        got = repo({"OPEN.md": register(entries)})
        if got is None:
            problems.append(f"the validator crashed on {what}")
        elif ("REG005" in got) != want:
            problems.append(f"{what} was {'not ' if want else ''}reported")
    return problems


@check("an attestation names repositories that exist, and leaves none of them out")
def _attestation_is_complete():
    # `verified_against` was one hash answering two questions, and the release skill had
    # both twenty lines apart: `git show <verified_against>:.../EVP.md` needs a commit of
    # this repository, and "the commit the evaluation actually ran on" is a commit of the
    # code. One repository made those the same thing. Five do not.
    def manifest(code: str) -> str:
        return ("schema: framework/product-manifest/v1\n"
                "artifact_type: product-manifest\nlifecycle: living\nstatus: active\n"
                "products: [alpha]\nname: Alpha\none_liner: A thing.\n"
                "owners: [o]\ncreated: 2026-01-01 09:00\nlast_review: 2026-01-01 09:00\n"
                "stage:\n  phase: F5\n  block: B\n" + code)

    def arc(attested: str) -> str:
        return ("---\nschema: framework/architecture/v1\nartifact_type: architecture\n"
                "lifecycle: living\nstatus: active\nproducts: [alpha]\nowners: [o]\n"
                "created: 2026-01-01 09:00\nlast_review: 2026-01-01 09:00\n"
                + attested + "---\n\n# Alpha\n\n"
                "<!-- section: current -->\n## Current\n\n"
                "<!-- section: target -->\n## Target\n\n"
                "<!-- section: delta -->\n## Delta\n")

    def codes(files):
        with tempfile.TemporaryDirectory() as tmp:
            for rel, text in files.items():
                f = Path(tmp) / rel
                f.parent.mkdir(parents=True, exist_ok=True)
                f.write_text(text)
            r = subprocess.run([sys.executable, str(VALIDATE), "--root", tmp, "--json",
                                "--stale-days", "36500"], capture_output=True, text=True)
            if r.returncode not in (0, 1) or not r.stdout.strip():
                return None
            return {f["code"] for f in json.loads(r.stdout)["findings"]}

    two = ("code:\n  backend:\n    url: git@e.com:o/be.git\n    contains: api\n"
           "    release_relevant: 'true'\n"
           "  frontend:\n    url: git@e.com:o/fe.git\n    contains: ui\n"
           "    release_relevant: 'true'\n")
    problems = []
    for attested, want, what in [
        ("verified_code:\n  product.backend: 9a734ce\n  product.frontend: 2b8d5a1\n",
         set(), "both release-relevant repositories attested"),
        ("verified_code:\n  product.backend: 9a734ce\n",
         {"VER002"}, "one release-relevant repository left out"),
        ("verified_code:\n  product.backend: 9a734ce\n  product.database: 711dc90\n",
         {"VER001", "VER002"}, "a repository no code map declares"),
    ]:
        got = codes({"products/alpha/product.yaml": manifest(two),
                     "products/alpha/ARC.md": arc(attested)})
        if got is None:
            problems.append(f"the validator crashed on {what}")
        elif {c for c in got if c.startswith("VER")} != want:
            problems.append(
                f"{what}: reported {sorted(c for c in got if c.startswith('VER'))}, "
                f"expected {sorted(want)}"
                + ("" if want else ". A complete attestation must be silent")
                + (". An attestation covering part of the system reads as covering all of "
                   "it, which is the failure a single hash had" if "VER002" in want else ""))

    # Nothing declared release-relevant means nothing is owed: the framework does not get to
    # invent the standard it then enforces.
    one = "code:\n  backend:\n    url: git@e.com:o/be.git\n    contains: api\n"
    got = codes({"products/alpha/product.yaml": manifest(one),
                 "products/alpha/ARC.md": arc("verified_code:\n  product.backend: 9a734ce\n")})
    if got and "VER002" in got:
        problems.append("a project that marked no repository `release_relevant` was told its "
                        "attestation was incomplete against a list it never wrote")
    return problems


@check("the derived view of a product is generated, and its drift is reported")
def _manifest_is_derived():
    # `product.yaml` carried sections marked GENERATED and kept by hand, so "which decisions
    # are open" had two answers: the register, and a list beside it. The stale one is what an
    # agent reads first, because AGENTS.md sends it to the manifest.
    #
    # A separate generated file rather than sections rewritten inside `product.yaml`, which
    # is authoritative and full of comments carrying the reasoning behind each field.
    base = {
        "framework.yaml": f"framework_version: {REGISTRY['version']}\n",
        "products/alpha/product.yaml": (
            "schema: framework/product-manifest/v1\n"
            "artifact_type: product-manifest\nlifecycle: living\nstatus: active\n"
            "products: [alpha]\nname: Alpha\none_liner: A thing.\n"
            "owners: [o]\ncreated: 2026-01-01 09:00\nlast_review: 2026-01-01 09:00\n"
            "stage:\n  phase: F1\n  block: A\n"),
        "OPEN.md": ("---\nschema: framework/open-register/v1\n"
                    "artifact_type: open-register\nlifecycle: living\nstatus: active\n"
                    "owners: [o]\ncreated: 2026-01-01 09:00\n"
                    "last_review: 2026-01-01 09:00\n"
                    "entries:\n  OD-001:\n    status: open\n    cost_to_reverse: low\n"
                    "  OD-002:\n    status: decided\n    cost_to_reverse: low\n"
                    "    closed_by: DEC-001\n---\n\n# Open\n"),
    }
    problems = []
    with tempfile.TemporaryDirectory() as tmp:
        for rel, text in base.items():
            f = Path(tmp) / rel
            f.parent.mkdir(parents=True, exist_ok=True)
            f.write_text(text)
        out = Path(tmp) / "products" / "alpha" / "product.index.yaml"

        r = subprocess.run([sys.executable, str(VALIDATE), "--root", tmp, "--emit-index",
                            "--stale-days", "36500"], capture_output=True, text=True)
        if not out.exists():
            return [f"--emit-index wrote no derived view: {r.stdout.strip()[-160:]}"]
        text = out.read_text()
        if GENERATED_MARK not in text:
            problems.append("the derived view carries no generated marker, so nothing "
                            "protects a hand-written file at that path from being replaced")
        if "OD-001" not in text or "OD-002" in text:
            problems.append("`open_decisions` is not what the register says is open: "
                            f"{[l for l in text.splitlines() if 'open_decisions' in l]}")

        # Drift has to be reported without writing, or CI cannot tell a stale file from a
        # fresh one, and a generated file nobody regenerates is worse than a hand-written
        # one because everybody assumes something keeps it true.
        out.write_text(text.replace("OD-001", "OD-999"))
        r = subprocess.run([sys.executable, str(VALIDATE), "--root", tmp, "--emit-index",
                            "--check", "--stale-days", "36500"], capture_output=True, text=True)
        if r.returncode == 0:
            problems.append("--emit-index --check exited 0 on a derived view that had been "
                            "edited by hand")
        if out.read_text() == text:
            problems.append("--check rewrote the file it was asked only to compare")
    return problems


@check("a repository shared by two products has one home, and only one")
def _shared_code_has_one_home():
    # The case this framework is about to be used on: one product across five repositories,
    # and two more products sharing the layer everybody signs in through. A per-product map
    # has nowhere to put the shared one, so the substrate carries `code:` too — and the way
    # that arrangement fails is a repository written into both, agreeing on the day it is
    # written and diverging afterwards.
    # The key of a `code:` entry is what this file calls the repository; the URL is which
    # repository it is. Keyed on the key, the check fired on two products that each own a
    # `backend` and stayed silent on one repository entered as `identity` here and `auth`
    # there -- and the second is how the duplication actually arises, because two teams
    # naming the same thing each use their own word.
    problems = []
    v = _load(VALIDATE, "validate")
    same = "github.com/org/repo"
    for url, want in [
        ("git@github.com:org/repo.git", same),
        ("ssh://git@github.com/org/repo.git", same),
        ("https://github.com/org/repo.git", same),
        ("https://GitHub.com/org/repo/", same),
        ("git@github.com:org/repo", same),
        # A port belongs to the address. Converted as if it were an scp-style path it became
        # a directory, and a self-hosted GitLab stopped comparing equal to itself.
        ("https://gitlab.self.hosted:8443/grp/repo.git", "gitlab.self.hosted:8443/grp/repo"),
    ]:
        if v.canonical_repo(url) != want:
            problems.append(f"{url} canonicalised to {v.canonical_repo(url)!r}, not {want!r}")

    def manifest(product: str, code: str) -> str:
        return ("schema: framework/product-manifest/v1\n"
                "artifact_type: product-manifest\n"
                "lifecycle: living\nstatus: active\n"
                f"products: [{product}]\nname: {product.title()}\n"
                "one_liner: A thing that does a thing.\n"
                "owners: [Someone]\ncreated: 2026-01-01 09:00\n"
                "last_review: 2026-01-01 09:00\n"
                "stage:\n  phase: F4\n  block: B\n" + code)

    shared = ("code:\n  access:\n    url: git@example.com:org/access.git\n"
              "    contains: sign-in and tenancy\n")
    platform = ("---\nschema: framework/platform-architecture/v1\n"
                "artifact_type: platform-architecture\nlifecycle: living\nstatus: active\n"
                "products: [alpha, beta]\nowners: [Someone]\n"
                "created: 2026-01-01 09:00\nlast_review: 2026-01-01 09:00\n"
                + shared + "---\n\n# Substrate\n")

    def codes(files):
        with tempfile.TemporaryDirectory() as tmp:
            for rel, text in files.items():
                f = Path(tmp) / rel
                f.parent.mkdir(parents=True, exist_ok=True)
                f.write_text(text)
            r = subprocess.run([sys.executable, str(VALIDATE), "--root", tmp, "--json",
                                "--stale-days", "36500"], capture_output=True, text=True)
            if r.returncode not in (0, 1) or not r.stdout.strip():
                return None
            return {f["code"] for f in json.loads(r.stdout)["findings"]}

    own = ("code:\n  frontend:\n    url: git@example.com:org/alpha-fe.git\n"
           "    contains: the app\n")

    # Declared once on the substrate, used by both: the arrangement working.
    got = codes({"PLATFORM.md": platform,
                 "products/alpha/product.yaml": manifest("alpha", own),
                 "products/beta/product.yaml": manifest("beta", "")})
    if got is None:
        problems.append("the validator crashed on a substrate declaring shared code")
    elif "XP004" in got:
        problems.append("a repository declared once on the platform was reported as a "
                        "duplicate: the substrate is where a shared repository belongs")

    # The same repository written into a product as well. One question, two answers.
    got = codes({"PLATFORM.md": platform,
                 "products/alpha/product.yaml": manifest("alpha", shared),
                 "products/beta/product.yaml": manifest("beta", "")})
    if got is None:
        problems.append("the validator crashed on a repository declared twice")
    elif "XP004" not in got:
        problems.append("a repository declared by both the platform and a product was not "
                        "reported: the two entries agree today and one of them gets fixed")

    # And between two products, with no platform in the repository at all.
    got = codes({"products/alpha/product.yaml": manifest("alpha", shared),
                 "products/beta/product.yaml": manifest("beta", shared)})
    if got is None:
        problems.append("the validator crashed on two products claiming one repository")
    elif "XP004" not in got:
        problems.append("two products declaring the same repository were not reported, "
                        "which is the duplication the platform map exists to avoid")

    # The two cases that separate an identity from a nickname, and the reason the check was
    # rewritten. Without them a version keyed on the map key passes this whole block, which
    # is what it did.
    def entry(key: str, url: str) -> str:
        return f"code:\n  {key}:\n    url: {url}\n    contains: stuff\n"

    got = codes({"products/alpha/product.yaml":
                 manifest("alpha", entry("backend", "git@example.com:org/alpha-be.git")),
                 "products/beta/product.yaml":
                 manifest("beta", entry("backend", "git@example.com:org/beta-be.git"))})
    if got is None:
        problems.append("the validator crashed on two products that each own a `backend`")
    elif "XP004" in got:
        problems.append("two products that each call their own repository `backend` were "
                        "reported as sharing one: the key is what a file calls a repository, "
                        "not which repository it is")

    got = codes({"products/alpha/product.yaml":
                 manifest("alpha", entry("identity", "git@example.com:org/access.git")),
                 "products/beta/product.yaml":
                 manifest("beta", entry("auth", "https://example.com/org/access.git"))})
    if got is None:
        problems.append("the validator crashed on one repository under two nicknames")
    elif "XP004" not in got:
        problems.append("one repository entered as `identity` under one product and `auth` "
                        "under another was not reported. This is how the duplication "
                        "actually arises: each side calls it what it calls it, and a check "
                        "keyed on the name catches only the tidy case")
    return problems


@check("a product in discovery is not reported for lacking what discovery has not reached")
def _early_products_are_not_findings():
    # Block A is elastic on purpose. A product at F1 has no brief because it has not got
    # there, and reporting that is the framework asking somebody to backfill a document
    # they had no grounds to write. What separates "early" from "undocumented" is the stage
    # the manifest declares, so a product with no manifest at all still has to be reported:
    # nothing said which of the two it was.
    def manifest(phase: str | None) -> str:
        stage = f"stage:\n  phase: {phase}\n  block: A\n" if phase else ""
        return ("schema: framework/product-manifest/v1\n"
                "artifact_type: product-manifest\n"
                "lifecycle: living\nstatus: active\n"
                "products: [alpha]\nname: Alpha\n"
                "one_liner: A thing that does a thing.\n"
                "owners: [Someone]\ncreated: 2026-01-01 09:00\n"
                "last_review: 2026-01-01 09:00\n" + stage)

    def codes(files):
        with tempfile.TemporaryDirectory() as tmp:
            for rel, text in files.items():
                f = Path(tmp) / rel
                f.parent.mkdir(parents=True, exist_ok=True)
                f.write_text(text)
            r = subprocess.run([sys.executable, str(VALIDATE), "--root", tmp, "--json",
                                "--stale-days", "36500"], capture_output=True, text=True)
            if r.returncode not in (0, 1) or not r.stdout.strip():
                return None
            return {f["code"] for f in json.loads(r.stdout)["findings"]}

    problems = []
    for phase, want, why in [
        ("F1", False, "a product in signal and framing"),
        ("F3", False, "a product in solution discovery"),
        ("F4", True, "a product shaping its MVP, which is where the brief belongs"),
        ("BUILD", True, "a product being built"),
        (None, True, "a product whose manifest declares no stage at all"),
    ]:
        got = codes({"products/alpha/product.yaml": manifest(phase)})
        if got is None:
            problems.append(f"the validator crashed on {why}")
        elif ("XP003" in got) != want:
            problems.append(
                f"{why} {'was not' if want else 'was'} reported for having no PBR"
                + ("" if want else ": discovery is elastic and this is the check telling a "
                   "project it is doing it wrong"))
    return problems


@check("the plugin and the framework ship one version, and it is the same number")
def _one_version():
    # They are one number by decision, reversing a note in the registry that said the
    # framework's version was not the plugin's. The decision only holds if something keeps
    # them equal: two files that are supposed to agree and are edited by hand are two files
    # that disagree, which is the failure this whole framework is a set of arguments about.
    # `claude plugin validate` already refuses a `plugin.json` and a marketplace entry that
    # disagree; this is the third corner of that triangle, which it has no way to see.
    problems = []
    registry = REGISTRY["version"]
    for rel, read in (
        (".claude-plugin/plugin.json", lambda d: d["version"]),
        (".claude-plugin/marketplace.json", lambda d: d["plugins"][0]["version"]),
    ):
        declared = read(json.loads((ROOT / rel).read_text()))
        if declared != registry:
            problems.append(
                f"{rel} declares {declared!r} and schemas/artifact-types.yaml declares "
                f"{registry!r}. One artifact, one version: whichever is right, the other "
                "answers 'what am I running' with a number nobody set")
    if not re.fullmatch(r"\d+\.\d+\.\d+", str(registry)):
        problems.append(f"the registry version is {registry!r}, which is not three numbers "
                        "separated by dots: a plugin manifest cannot carry it")
    return problems


@check("nothing states the framework's version except the registry that defines it")
def _version_is_not_restated():
    # `start` wrote a literal version into every repository it set up, in an example whose
    # own next paragraph says to read the number from the registry. The framework moved and
    # the skill kept seeding a value that was already wrong -- an `FW001` on day one, in the
    # file whose whole job is to tell a migration from a mistake, produced by the skill that
    # creates the file.
    #
    # The literal is not quoted anywhere in this comment, and that is this check working on
    # itself rather than an accident of wording: an exemption for the file that holds the
    # rule is the first hole anybody widens.
    #
    # It is the same failure as the count in prose, with a longer fuse: nothing is wrong on
    # the day it is written, and nothing reports it on the day it becomes wrong. So the
    # rule is the flat one -- no literal anywhere but the registry. An example writes `N`,
    # and an `N` that reaches a real `framework.yaml` is an `FW001` saying it is a string
    # and not a whole number, which is a sentence a stale number never produces.
    literal = re.compile(r"framework_version:\s*['\"]?(\d[\d.]*)")
    skip = {"build", "__pycache__", ".git", "node_modules"}
    problems = []
    for path in sorted(ROOT.rglob("*")):
        if path.suffix not in (".md", ".py", ".yaml") or not path.is_file():
            continue
        if set(path.relative_to(ROOT).parts) & skip:
            continue
        if path == ROOT / "schemas" / "artifact-types.yaml":
            continue          # the one place the number is defined
        for n, line in enumerate(path.read_text(encoding="utf-8").splitlines(), 1):
            m = literal.search(line)
            # Inside a test, a version built from the registry is the point of the test:
            # `f"framework_version: {current + 1}"` has no literal for this to find, and a
            # bare digit in one would be the same defect as anywhere else.
            if m:
                problems.append(
                    f"{path.relative_to(ROOT)}:{n} writes framework_version {m.group(1)} "
                    "as a literal. Read it from `version:` in schemas/artifact-types.yaml, "
                    "or write `N` if it is an example: this is the line that goes wrong "
                    "silently the day the framework moves")
    return problems


@check("every product owns its register, and the root composes them without copying them")
def _registers_are_per_product():
    # The arrangement the framework now asks for: one register per product, one for the
    # substrate, one at the root for what belongs to no single product, and a generated
    # union under the last. Six things have to hold at once and each of them used to fail
    # in the direction that reads as "nothing is open".
    fm = lambda body: ("---\nschema: framework/open-register/v1\n"
                       "artifact_type: open-register\nlifecycle: living\nstatus: active\n"
                       "owners: [o]\ncreated: 2026-01-01 09:00\n"
                       "last_review: 2026-01-01 09:00\n" + body + "---\n\n# Open\n")
    man = lambda p: ("schema: framework/product-manifest/v1\n"
                     "artifact_type: product-manifest\nlifecycle: living\nstatus: active\n"
                     f"products: [{p}]\nname: {p}\none_liner: A thing.\n"
                     "owners: [o]\ncreated: 2026-01-01 09:00\nlast_review: 2026-01-01 09:00\n"
                     "stage:\n  phase: F5\n  block: A\n")
    base = {
        "framework.yaml": f"framework_version: {REGISTRY['version']}\n",
        "products/alpha/product.yaml": man("alpha"),
        "products/beta/product.yaml": man("beta"),
        # No `entries:`, and that is legal here and nowhere else: the root of a repository
        # that files per product is a view plus a parking lot. The marker is what says so.
        "OPEN.md": fm("") + "\n<!-- generated: open-union -->\nx\n<!-- /generated -->\n",
        "platform/OPEN.md": fm("entries:\n  OD-001:\n    status: open\n"
                               "    cost_to_reverse: high\n"
                               "    default_in_force: one database for everyone\n"),
        # `depends_on` reaches into the substrate's register, which is the ordinary case
        # once the entries are filed per product and used to be reported as dangling.
        "products/alpha/OPEN.md": fm("entries:\n  OD-002:\n    status: open\n"
                                     "    cost_to_reverse: low\n"
                                     "    default_in_force: three retries\n"
                                     "    depends_on: [OD-001]\n"),
    }
    problems = []
    with tempfile.TemporaryDirectory() as tmp:
        root = Path(tmp)
        for rel, text in base.items():
            f = root / rel
            f.parent.mkdir(parents=True, exist_ok=True)
            f.write_text(text)

        run = lambda *extra: subprocess.run(
            [sys.executable, str(VALIDATE), "--root", str(root), "--json",
             "--stale-days", "36500", *extra], capture_output=True, text=True)
        r = run("--emit-index")
        if r.returncode not in (0, 1):
            return [f"the validator crashed: {r.stderr.strip().splitlines()[-1]}"]
        codes = [f["code"] for f in json.loads(r.stdout)["findings"]]

        if "REG006" not in codes:
            problems.append("a product directory with no register of its own was not "
                            "reported: an agent sent to work there reads no register and "
                            "concludes nothing is open")
        if "REG004" in codes:
            problems.append("the register at the root was reported for holding no "
                            "`entries:` while carrying the union marker, and the reliable "
                            "way to clear that is to paste the entries back in")
        if "REG005" in codes:
            problems.append("a `depends_on` reaching another register was reported as "
                            f"dangling: {[f['message'] for f in json.loads(r.stdout)['findings'] if f['code'] == 'REG005']}")

        # Attribution comes from the directory. Reading it off `products:` -- absent in a
        # product's own register, because the directory already said it -- put alpha's
        # entries into beta's derived view, and that view is where AGENTS.md sends an agent.
        beta = (root / "products" / "beta" / "product.index.yaml").read_text()
        if "OD-002" in beta:
            problems.append("an entry filed in alpha's register reached beta's derived "
                            "view: scope is being read off `products:` and not off where "
                            "the register sits")
        if "OD-001" not in beta:
            problems.append("a substrate entry is missing from a product's derived view: "
                            "the decisions that bind every product are the ones a product "
                            "cannot see it is waiting on")

        union = (root / "OPEN.md").read_text()
        if "## alpha" not in union or "## beta" not in union:
            problems.append("the union at the root has no heading per product")
        if "OD-001" not in union or "OD-002" not in union:
            problems.append("the union does not hold every open entry")
        if "entries:" in union.split("---")[1]:
            problems.append("the union grew an `entries:` map, so every check that reads "
                            "one now reports each entry twice and `depends_on` has two "
                            "rows with the same id to resolve against")

        # `all` IS A WORD AND NOT A PRODUCT, and the two halves of that sentence used to
        # live in different functions: `binds` knew it, the generator did not, and the union
        # grew a `## all` heading for a product no repository has. Worse, the last section
        # selected the entries that named nothing -- the exact state `REG011` exists to
        # remove -- so the more a repository answered the new check, the emptier the section
        # holding what binds everything became.
        (root / "products" / "beta" / "OPEN.md").write_text(fm("entries: {}\n"))
        (root / "OPEN.md").write_text(
            fm("entries:\n  OD-009:\n    status: open\n    cost_to_reverse: high\n"
               "    products: [all]\n    default_in_force: one database for everyone\n"
               "  OD-010:\n    status: open\n    cost_to_reverse: low\n"
               "    products: [alpha]\n    default_in_force: nothing\n")
            + "\n<!-- generated: open-union -->\nx\n<!-- /generated -->\n")
        run("--emit-index")
        union = (root / "OPEN.md").read_text()
        if "## all" in union:
            problems.append("the union grew a heading for `all`, which is the word an entry "
                            "uses to say it binds every product and not the name of one")
        tail = union.split("## Bound to no single product", 1)
        if len(tail) != 2:
            problems.append("the union lost the section for what binds no single product")
        else:
            if "OD-009" not in tail[1]:
                problems.append("an entry declaring `[all]` is missing from the section for "
                                "what binds no single product. That section used to hold the "
                                "entries naming nothing, which is the state `REG011` removes: "
                                "answer the check and the section empties itself")
            if "OD-010" in tail[1]:
                problems.append("an entry naming one product appeared under what binds no "
                                "single product")
        for prod in ("## alpha", "## beta"):
            block = union.split(prod, 1)[1].split("\n## ", 1)[0]
            if "OD-009" not in block:
                problems.append(f"an entry declaring `[all]` is missing from {prod.strip()}")

        # The prose outside the markers is not the generator's to touch.
        (root / "OPEN.md").write_text(union.replace("# Open", "# Open\n\nMine, by hand."))
        run("--emit-index")
        if "Mine, by hand." not in (root / "OPEN.md").read_text():
            problems.append("--emit-index rewrote prose outside the region markers")

        # A duplicate id is the failure a per-product register invites: each one starts at
        # 001 unless something says otherwise, and the collision resolves itself silently.
        (root / "products" / "beta" / "OPEN.md").write_text(
            fm("entries:\n  OD-002:\n    status: open\n    cost_to_reverse: low\n"
               "    default_in_force: three retries\n"))
        codes = [f["code"] for f in json.loads(run().stdout)["findings"]]
        if "REG007" not in codes:
            problems.append("two registers declaring the same entry id were not reported, "
                            "so a `depends_on` naming it resolves to whichever was read last")

        # A product may genuinely have nothing open, and it still has to have a register.
        # `entries: {}` is how it says so: refusing it would leave inventing an entry as the
        # only way to satisfy both checks, and an entry invented to clear a finding is the
        # failure everything here is arranged against. Absent `entries:` stays a finding,
        # because the two are different claims -- read and nothing found, versus nobody
        # filled this in.
        empty = root / "products" / "beta" / "OPEN.md"
        empty.write_text(fm("entries: {}\n"))
        out = json.loads(run().stdout)
        stated = [f for f in out["findings"] if f["code"] in ("REG004", "FM002")
                  and f["path"].endswith("beta/OPEN.md")]
        if stated:
            problems.append("a register declaring `entries: {}` was reported "
                            f"({[f['code'] for f in stated]}): a product with nothing open "
                            "has no way left to have a register except inventing an entry")
        empty.write_text(fm(""))
        silent = [f["code"] for f in json.loads(run().stdout)["findings"]
                  if f["code"] == "REG004" and f["path"].endswith("beta/OPEN.md")]
        if not silent:
            problems.append("a register with no `entries:` at all was not reported, so "
                            "'nobody filled this in' and 'there is nothing open' now read "
                            "the same")

        (root / "products" / "beta" / "OPEN.md").write_text(
            fm("entries:\n  OD-003:\n    status: open\n    cost_to_reverse: low\n"
               "    default_in_force: three retries\n    products: [alpha]\n"))
        codes = [f["code"] for f in json.loads(run().stdout)["findings"]]
        if "REG008" not in codes:
            problems.append("an entry in beta's register naming alpha was not reported: a "
                            "person reading the directory and the derived view disagree "
                            "about who it is about, and neither is told")

        # `trigger` replaced `deadline` in 2.0.0, and the point of the rename is the value
        # rather than the key: what closes the window is an event, and a date on its own is
        # a number somebody picked. Both spellings of a bare date are asserted, because YAML
        # hands back a `date` object for the unquoted form and a string for the quoted one,
        # and a check that only saw the string would be silent on the spelling everybody
        # actually writes.
        entry = lambda trg: fm("entries:\n  OD-004:\n    status: open\n"
                               "    cost_to_reverse: high\n"
                               "    default_in_force: none\n"
                               f"    trigger: {trg}\n")
        for written in ("2026-09-30", '"2026-09-30"', "2026-09-30."):
            (root / "products" / "beta" / "OPEN.md").write_text(entry(written))
            codes = [f["code"] for f in json.loads(run().stdout)["findings"]]
            if "REG009" not in codes:
                problems.append(f"a trigger written as {written} was not reported. A date "
                                "does not force a decision by arriving, and the entry it "
                                "sits on is one nobody has taken")

        # A date beside the event is the same finding. It reads as the careful version --
        # here is the event, and here is when it falls -- and the date is the half that gets
        # quoted back by somebody who never opened the register, on an entry whose whole
        # content is that nobody has decided.
        for written in ("the external audit, 2026-10-31", "Q4 2026", "the audit on 31/10/2026"):
            (root / "products" / "beta" / "OPEN.md").write_text(entry(written))
            codes = [f["code"] for f in json.loads(run().stdout)["findings"]]
            if "REG009" not in codes:
                problems.append(f"a trigger reading {written!r} was not reported: a date "
                                "beside an event is still a date on an undecided entry")

        # And the shape that must stay silent, because a check that reports the thing the
        # template asks for is a check somebody turns off.
        for written in ("before the second customer", "the external audit"):
            (root / "products" / "beta" / "OPEN.md").write_text(entry(written))
            codes = [f["code"] for f in json.loads(run().stdout)["findings"]]
            if "REG009" in codes:
                problems.append(f"a trigger reading {written!r} was reported: the check is "
                                "matching prose rather than digits")

        # Who an entry at the root binds: `[all]` is a statement, absence is a gap, and
        # before this they were written identically. The root register of this fixture is
        # the one with no `entries:` at all, so the entry is added here and removed after.
        rootreg = root / "OPEN.md"
        keep = rootreg.read_text()
        entry_at_root = lambda decl: (
            fm("entries:\n  OD-009:\n    status: open\n    cost_to_reverse: low\n"
               f"    default_in_force: nothing\n{decl}")
            + "\n<!-- generated: open-union -->\nx\n<!-- /generated -->\n")
        for decl, want in (("", True),
                           ("    products: [all]\n", False),
                           ("    products: [alpha]\n", False),
                           ("    products: [all, alpha]\n", True)):
            rootreg.write_text(entry_at_root(decl))
            codes = [f["code"] for f in json.loads(run().stdout)["findings"]]
            if want and "REG011" not in codes:
                problems.append(f"an entry at the root declaring {decl.strip() or 'nothing'} "
                                "was not reported: it binds every product by rule, and "
                                "nothing distinguishes that from nobody having asked")
            if not want and "REG011" in codes:
                problems.append(f"an entry at the root declaring {decl.strip()} was "
                                "reported, so the field cannot be answered")
        # Answering the question and still being in the wrong file. `REG011` asks who an
        # entry binds; `REG013` is about what the root register is for, and green on the
        # first has never implied the second -- which is how seventeen entries came to sit
        # at the root of a real repository, each naming one product, all of them reported by
        # nothing.
        for decl, want, why in (("    products: [alpha]\n", True,
                                 "an entry at the root binding one product that has a "
                                 "register of its own"),
                                ("    products: [all]\n", False,
                                 "an entry binding every product, which is what the root is "
                                 "for"),
                                ("    products: [alpha, beta]\n", False,
                                 "an entry binding two products, which belongs to neither "
                                 "register"),
                                ("    products: [gamma]\n", False,
                                 "an entry naming a product with no directory anywhere, "
                                 "which has nowhere to be moved to")):
            rootreg.write_text(entry_at_root(decl))
            codes = [f["code"] for f in json.loads(run().stdout)["findings"]]
            if ("REG013" in codes) != want:
                problems.append(f"{why} was {'not ' if want else ''}reported")
        rootreg.write_text(keep)

        # And the state the register is written for: a repository with a register and no
        # products yet. Asking which products an entry binds has no available answer there.
        rootreg.write_text(entry_at_root(""))
        for man in (root / "products").glob("*/product.yaml"):
            man.rename(man.with_suffix(".yaml.parked"))
        codes = [f["code"] for f in json.loads(run().stdout)["findings"]]
        for man in (root / "products").glob("*/product.yaml.parked"):
            man.rename(man.with_suffix("").with_suffix(".yaml"))
        if "REG011" in codes:
            problems.append("an entry at the root was asked which products it binds in a "
                            "repository that declares none: at day one the register is the "
                            "file that exists and the products are what has not been "
                            "created yet")
        rootreg.write_text(keep)

        # The manifest answering a question that moved. Absent from the template now, and a
        # repository that kept the field keeps a second answer that nothing recomputes.
        man_path = root / "products" / "alpha" / "product.yaml"
        original = man_path.read_text()
        for field_name in ("open_decisions: [OD-001]", "open_risks: []", "active_changes: []"):
            man_path.write_text(original + field_name + "\n")
            codes = [f["code"] for f in json.loads(run().stdout)["findings"]]
            if "FM005" not in codes:
                problems.append(f"a manifest carrying `{field_name.split(':')[0]}` was not "
                                "reported: it is derived now, and a hand written copy of a "
                                "derived answer is the one that goes stale unnoticed")
        man_path.write_text(original)

        # A date in a heading, which `REG009` structurally cannot see: it reads the
        # `trigger` of an entry, and a heading belongs to no entry while binding every one
        # filed under it. This framework shipped one in its own template.
        head = lambda h: (fm("entries:\n  OD-004:\n    status: open\n"
                             "    cost_to_reverse: high\n"
                             "    default_in_force: none\n"
                             "    trigger: before the second customer\n")
                          + f"\n{h}\n\n### OD-004 - A choice\n")
        for written in ("## Cost to reverse MEDIUM: decide by 2026-09-30",
                        "## Decide these in Q4 2026",
                        "### OD-004 - move by 30/09/2026"):
            (root / "products" / "beta" / "OPEN.md").write_text(head(written))
            codes = [f["code"] for f in json.loads(run().stdout)["findings"]]
            if "REG010" not in codes:
                problems.append(f"a heading reading {written!r} was not reported. It binds "
                                "every entry under it and no `trigger` records it, so "
                                "nothing will ever report it going stale")

        # A `#` inside a fenced block is a comment in somebody's example. A bash snippet
        # saying `# rigenerato il 2026-09-30` was reported as a heading carrying a date, and
        # a false finding costs the trip to the document.
        for written in ("~~~bash\n# rigenerato il 2026-09-30\n~~~",
                        "```bash\n# rigenerato il 2026-09-30\n```",
                        "```yaml\n# entro il 2026-09-30\nentries: {}\n```"):
            (root / "products" / "beta" / "OPEN.md").write_text(head(written))
            codes = [f["code"] for f in json.loads(run().stdout)["findings"]]
            if "REG010" in codes:
                problems.append("a comment inside a fenced block was reported as a heading "
                                "carrying a date: it is somebody's example, and the finding "
                                "sends a reader to a document that was right")

        for written in ("## Cost to reverse MEDIUM: changing it later costs a migration",
                        "## Cost to reverse HIGH"):
            (root / "products" / "beta" / "OPEN.md").write_text(head(written))
            codes = [f["code"] for f in json.loads(run().stdout)["findings"]]
            if "REG010" in codes:
                problems.append(f"a heading reading {written!r} was reported: the check is "
                                "matching headings rather than dates in them")

        # `decide_with`: the pairing an ordering cannot express. Same resolution as
        # `depends_on`, because a pairing with an entry nobody can find reads as one that
        # was honoured, and the self reference on top -- a field that looks filled in and
        # binds the entry to nothing.
        pair = lambda peer: fm("entries:\n  OD-004:\n    status: open\n"
                               "    cost_to_reverse: high\n"
                               "    default_in_force: none\n"
                               f"    decide_with: [{peer}]\n")
        for peer, why in (("OD-404", "an entry no register declares"),
                          ("OD-004", "itself")):
            (root / "products" / "beta" / "OPEN.md").write_text(pair(peer))
            codes = [f["code"] for f in json.loads(run().stdout)["findings"]]
            if "REG005" not in codes:
                problems.append(f"an entry to be decided with {why} was not reported, so "
                                "the register vouches for a pairing that binds it to "
                                "nothing")
        (root / "products" / "beta" / "OPEN.md").write_text(pair("OD-001"))
        codes = [f["code"] for f in json.loads(run().stdout)["findings"]]
        if "REG005" in codes:
            problems.append("a `decide_with` naming an entry that exists in another "
                            "register was reported: the pairing reaches across registers "
                            "by design, exactly as `depends_on` does")
    return problems


@check("the extractor keeps the provenance the converter throws away")
def _extract_keeps_provenance():
    # anydoc has no slide and no page in its document model: a whole deck comes back as one
    # flat stream. Everything that puts a claim back on slide 4 is written here, so it is
    # this repository's to test. All three pieces are pure, which is the point: a CI runner
    # with no converter installed still checks the part that would silently degrade.
    x = _load(EXTRACT, "extract")
    problems = []

    def deck(n: int) -> bytes:
        ids = "".join(f'<p:sldId id="{255 + i}" r:id="rId{i}"/>' for i in range(1, n + 1))
        pres = ('<?xml version="1.0"?><p:presentation xmlns:p="x">'
                f"<p:sldIdLst>{ids}</p:sldIdLst></p:presentation>")
        buf = io.BytesIO()
        with zipfile.ZipFile(buf, "w") as z:
            z.writestr("ppt/presentation.xml", pres)
            for i in range(1, n + 1):
                z.writestr(f"ppt/slides/slide{i}.xml", f"<p:sld>slide {i}</p:sld>")
        return buf.getvalue()

    # Lazy, and asserted as lazy. Every slice is a whole copy of the package, so returning
    # them together costs deck size times slide count: forty slides of a thirty megabyte
    # deck wanted more than a gigabyte. Checking only the slices would pass again the day
    # somebody collects them into a list, which is what this used to do while its own
    # docstring claimed one existed at a time.
    produced = x.slice_pptx(deck(3))
    if not isinstance(produced, types.GeneratorType):
        problems.append("slice_pptx stopped being lazy: each slice is a full copy of the "
                        "package, and holding all of them is deck size times slide count")
    slices = list(produced)
    if len(slices) != 3:
        problems.append(f"a three slide deck sliced into {len(slices)} packages: without "
                        "one per slide, every claim in a deck is located to the file and "
                        "checking it means reading the deck")
    for i, s in enumerate(slices, 1):
        z = zipfile.ZipFile(io.BytesIO(s))
        if (kept := z.read("ppt/presentation.xml").decode().count("<p:sldId ")) != 1:
            problems.append(f"slice {i} kept {kept} slides in its running order: the split "
                            "did not split, and the numbers it reports are wrong")
        # The trap: writestr stamps the offset of the copy onto the ZipInfo it is handed,
        # so copying by ZipInfo corrupts the source archive and every slice after the first
        # comes back short. Slice 3 having all four parts is the assertion.
        if len(z.namelist()) != 4:
            problems.append(f"slice {i} lost parts of the package: {z.namelist()}")

    if list(x.slice_pptx(b"PPT legacy binary, not a zip")):
        problems.append("something that is not a package was sliced anyway: the caller "
                        "reads an empty list as its cue to convert the file whole")

    # In the report the provenance line is itself a heading. A `# Title` arriving inside a
    # block would outrank it, and the boundary between one document and the next is what
    # tells a classifier which file a claim came from.
    body = x.demote("# Title\n\n```\n# not a heading\n```\n\n## Section")
    if not ("\n" + body).count("\n### Title") or "#### Section" not in body:
        problems.append("a heading inside a block was not demoted below the provenance "
                        "heading: extract.md stops saying where a document ends")
    if "\n# not a heading" not in body:
        problems.append("a comment inside a code fence was rewritten as a heading")

    # Found on a deck built to look like a real one. The slide was the architecture diagram
    # with the word "Architettura" on it, and it went unflagged because the speaker note
    # underneath carried it over the threshold. The question the threshold answers is
    # whether the slide is a picture, and a picture with a talkative presenter is one.
    drawn = "Architettura\n\n> Nota del relatore che spiega a voce tutta l'architettura."
    if x.on_slide_chars(drawn) >= 40:
        problems.append("speaker notes count towards the text on a slide: a diagram with a "
                        "chatty note under it stops being flagged, and that is the slide "
                        "the whole visual review exists to catch")
    if x.on_slide_chars("Prezzi 2026 per fascia di volume, con sconti a scaglioni") < 40:
        problems.append("a slide with a sentence on it was called text-poor")

    # A deck exported to PDF, told apart by the geometry poppler already reports. Both
    # halves are pure so a runner with no poppler still checks them, which is where it
    # matters: the branch only runs on a machine that has it.
    for meta, want, what in [
        ("Pages:  9\nPage size:       960 x 540 pts\n", True, "a 16:9 deck"),
        ("Pages:  9\nPage size:       720 x 540 pts\n", True, "a 4:3 deck"),
        ("Pages:  5\nPage size:       595.25 x 842 pts (A4)\n", False, "an A4 report"),
        # Landscape A4 is 1.41 and 4:3 is 1.33, so no ratio separates them. Included on
        # purpose: the flag feeds a measure taken against the document's own pages, and a
        # dense report has nothing thin against its own median, so the cost is nil.
        ("Pages:  5\nPage size:       842 x 595.25 pts\n", True, "a landscape A4"),
        ("Pages:  5\n", False, "a PDF whose geometry pdfinfo did not report"),
    ]:
        if x.is_deck(meta) != want:
            problems.append(f"{what} was {'not ' if want else ''}read as a presentation: "
                            "on a slide a thin page is a diagram, and on a report it is a "
                            "short page, and the two want opposite responses")

    # The distribution measured off the deck that prompted this. Page 7 was the only one a
    # forty character threshold caught; 1, 5 and 8 are the ones it missed, and 5 is the
    # slide naming the source systems and promising real time, all of it inside screenshots.
    got = x.thin_against_median({1: 120, 2: 1589, 3: 1248, 4: 1111, 5: 474,
                                 6: 1055, 7: 11, 8: 267, 9: 2185})
    if got != [1, 5, 7, 8]:
        problems.append(f"the thin pages of a real deck came back as {got}, not [1, 5, 7, 8]")
    if x.thin_against_median({1: 900, 2: 1000, 3: 1100}):
        problems.append("a deck with even pages had some of them called thin: the measure "
                        "is against the document's own pages, so an even one flags nothing")
    if x.thin_against_median({}):
        problems.append("an empty page table produced flagged pages")

    blocks = x.split_on_headings("req.docx", "intro\n\n# One\n\nalpha\n\n## Two\n\nbeta")
    if (got := [b.locator for b in blocks]) != ["preamble", "§ One", "§ Two"]:
        problems.append(f"sections came back located as {got}: a heading is the only "
                        "locator a document with no pages has")
    joined = "".join(b.text for b in blocks)
    if any(w not in joined for w in ("intro", "alpha", "beta")):
        problems.append("splitting a document on its headings dropped text from it")

    return problems




@check("the stages and increments of a roadmap agree with themselves and with the registers")
def _roadmap_stages():
    # The order of delivery is what the daily digest draws its Gantt from, and a hole in it
    # shows nowhere as an error: it shows as an increment missing from the timeline, a stage
    # with less work than it has, a delivery earlier than the plan allows. Each code is
    # asserted by one mutation of a clean repository, and the clean one reports none of them.
    fm = lambda **kw: "---\n" + "".join(f"{k}: {v}\n" for k, v in kw.items()) + "---\n\n"
    rmp = lambda stages, incs, body="": (
        fm(schema="framework/roadmap/v1", artifact_type="roadmap", lifecycle="living",
           status="active", version="1.0.0", products="[p]", owners="[o]",
           created="2026-01-01 09:00", last_review="2026-01-01 09:00")
        .replace("---\n\n", f"delivery_stages:\n{stages}increments:\n{incs}---\n\n", 1)
        + "# Roadmap\n\n## Increments\n\n### INC-001 · First\n\n### INC-002 · Second\n\n"
          "### INC-003 · Third\n" + body)
    stages = ("  t2:\n    name: Second\n    after: [t1]\n    increments: [INC-002, INC-003]\n"
              "  t1:\n    name: First\n    increments: [INC-001]\n"
              "    milestone: Two customers apart\n")
    incs = ("  INC-001:\n    state: committed\n    requires: [OD-001]\n"
            "  INC-003:\n    state: conditional\n    depends_on: [INC-002]\n"
            "  INC-002:\n    state: shaped\n    depends_on: [INC-001]\n"
            "    requires: [KI-002]\n    changes: [CHG-002]\n")
    chg = lambda cid, derives: fm(
        schema="framework/change-contract/v1", artifact_type="change-contract",
        lifecycle="immutable", status="draft", id=cid, products="[p]", owners="[o]",
        approvers="[o]", created="2026-01-01 09:00", derives_from=derives,
        verified_by="null") + f"# {cid}\n"
    base = {
        "framework.yaml": f"framework_version: {REGISTRY['version']}\n",
        "OPEN.md": fm(schema="framework/open-register/v1", artifact_type="open-register",
                      lifecycle="living", status="active", owners="[o]",
                      created="2026-01-01 09:00", last_review="2026-01-01 09:00")
        + "\n<!-- generated: open-union -->\nx\n<!-- /generated -->\n",
        "products/p/product.yaml": (
            "schema: framework/product-manifest/v1\nartifact_type: product-manifest\n"
            "lifecycle: living\nstatus: active\nproducts: [p]\nname: p\none_liner: A thing.\n"
            "owners: [o]\ncreated: 2026-01-01 09:00\nlast_review: 2026-01-01 09:00\n"),
        "products/p/OPEN.md": fm(
            schema="framework/open-register/v1", artifact_type="open-register",
            lifecycle="living", status="active", owners="[o]", created="2026-01-01 09:00",
            last_review="2026-01-01 09:00",
            entries="\n  OD-001:\n    status: open\n    cost_to_reverse: low\n"
                    "    default_in_force: one store\n  KI-002:\n    status: open")
        + "# Open\n\n### OD-001 · A choice\n\n### KI-002 · An issue\n",
        "products/p/changes/CHG-001.md": chg("CHG-001", "[INC-001]"),
        "products/p/changes/CHG-002.md": chg("CHG-002", "[KI-002]"),
        "products/p/RMP.md": rmp(stages, incs),
    }
    cases = [
        ("RMP001", "a stage naming an increment nobody declared",
         lambda t: t.replace("increments: [INC-001]", "increments: [INC-001, INC-009]")),
        ("RMP001", "a heading with no row", lambda t: t + "\n### INC-004 · Fourth\n"),
        ("RMP001", "an increment in two stages",
         lambda t: t.replace("increments: [INC-001]", "increments: [INC-001, INC-002]")),
        ("RMP002", "a stage after one that does not exist",
         lambda t: t.replace("after: [t1]", "after: [t9]")),
        ("RMP002", "two stages waiting for each other",
         lambda t: t.replace("    name: First\n", "    name: First\n    after: [t2]\n")),
        ("RMP002", "an increment depending on one of a later stage",
         lambda t: t.replace("    requires: [OD-001]\n",
                             "    requires: [OD-001]\n    depends_on: [INC-003]\n")),
        ("RMP003", "a prerequisite no register declares",
         lambda t: t.replace("requires: [OD-001]", "requires: [OD-099]")),
        ("RMP003", "a change that names its increment itself, listed again",
         lambda t: t.replace("changes: [CHG-002]", "changes: [CHG-002, CHG-001]")),
        ("RMP004", "delivered while its decision is open",
         lambda t: t.replace("state: committed", "state: delivered")),
        ("RMP005", "a date in a milestone",
         lambda t: t.replace("Two customers apart", "Two customers apart, by 15/10")),
    ]
    problems = []
    with tempfile.TemporaryDirectory() as tmp:
        root = Path(tmp)
        for rel, text in base.items():
            f = root / rel
            f.parent.mkdir(parents=True, exist_ok=True)
            f.write_text(text)
        target = root / "products" / "p" / "RMP.md"

        def codes():
            r = subprocess.run([sys.executable, str(VALIDATE), "--root", str(root), "--json",
                                "--stale-days", "36500"], capture_output=True, text=True)
            try:
                found = json.loads(r.stdout)["findings"]
            except ValueError:
                return None, r.stderr.strip()[-300:]
            return [f["code"] for f in found if f["code"].startswith("RMP")
                    or (f["code"].startswith("FM") and f["path"].endswith("RMP.md"))], found

        got, found = codes()
        if got is None:
            return [f"the validator did not run: {found}"]
        if got:
            problems.append(f"the clean roadmap was reported: "
                            f"{[(f['code'], f['message'][:80]) for f in found if f['code'] in got]}")
        clean = target.read_text()
        for code, what, mutate in cases:
            target.write_text(mutate(clean))
            got, _ = codes()
            if code not in (got or []):
                problems.append(f"{what}: {code} was not reported ({got})")
        target.write_text(clean)

    v = _load(VALIDATE, "validate_roadmap")
    plan = v.roadmap_plan(yaml.safe_load(base["products/p/RMP.md"].split("---\n")[1]))
    if plan["order"] != ["t1", "t2"] or plan["per_stage"]["t2"] != ["INC-002", "INC-003"]:
        problems.append(f"the sequence is not the one the stages and dependencies give: "
                        f"{plan['order']} {plan['per_stage']}")
    tie = v.roadmap_plan({"delivery_stages": {"b": {"name": "B", "increments": ["INC-002"]},
                                              "a": {"name": "A", "increments": ["INC-001"]}}})
    if tie["order"] != ["b", "a"]:
        problems.append(f"stages with no dependency are not read in the order written: "
                        f"{tie['order']}")
    return problems


@check("the extractor does not take the framework's own output for a corpus")
def _extract_skips_own_output():
    # A presentation is a `.pptx`, which is the format a corpus is made of, and it sits under
    # `_meta/` beside `_meta/corpus/`, which `group_corpus` gathers by first component. Found
    # there, the deck this framework generated would come back in as a document the customer
    # wrote: the one loop every other rule in the framework exists to prevent.
    x = _load(EXTRACT, "extract")
    with tempfile.TemporaryDirectory() as tmp:
        root = Path(tmp)
        for rel in ("_meta/corpus/offerta.pptx",
                    "_meta/presentation/PRS-001-atlas-2026-10-02/PRS-001-atlas-2026-10-02.pptx",
                    "_meta/business/SAL-001-atlas-2026-10-02.md",
                    "_meta/digest/DIG-002-atlas-2026-10-07/DIG-002-atlas-2026-10-07.xlsx",
                    "_meta/extract/offerta/blocks.md",
                    "presentation/deck-del-cliente.pptx"):
            (root / rel).parent.mkdir(parents=True, exist_ok=True)
            (root / rel).write_bytes(b"x")
        found = {d for d, _, _ in x.find_corpus(root)}
    problems = []
    for d in ("_meta/presentation/PRS-001-atlas-2026-10-02", "_meta/business", "_meta/extract/offerta",
              "_meta/digest/DIG-002-atlas-2026-10-07"):
        if d in found:
            problems.append(f"{d} was reported as a corpus: it is this framework's own output")
    for d in ("_meta/corpus", "presentation"):
        if d not in found:
            problems.append(f"{d} was not reported: an exclusion of the framework's output "
                            "must not hide a project's own documents")
    return problems


RENDER = ROOT / "skills" / "presentation" / "scripts" / "render.py"


def _outline(**over) -> dict:
    o = {"lang": "it", "title": "Atlas", "subtitle": "Riconcilia i movimenti bancari",
         "customer": None, "sources": ["PBR §One line"],
         "slides": [
             {"kind": "bullets", "title": "Che cosa fa oggi",
              "bullets": ["Abbina i movimenti alle scritture"],
              "sources": ["PBR §Current capabilities"]},
             {"kind": "roadmap", "title": "A che punto siamo",
              "done": ["Abbinamento automatico"],
              "todo": ["Chiusura guidata", {"text": "Più banche", "planned": True}],
              "sources": ["PBR §Current capabilities", "RMP §Increments"]}]}
    o.update(over)
    return o


@check("a presentation outline refuses what a customer must not be handed")
def _presentation_outline_refuses():
    # The checks are pure and need PyYAML alone, so they run whatever is installed. Each
    # mutation below is one of the rules `skills/presentation/SKILL.md` states, and each must
    # be refused by itself: an outline that passes with one of them in it hands it to a customer.
    x = _load(RENDER, "render")
    problems = []
    with tempfile.TemporaryDirectory() as tmp:
        base = Path(tmp)
        if x.check(_outline(), base):
            problems.append(f"a clean outline was refused: {x.check(_outline(), base)}")
        road = lambda **kw: _outline(slides=[_outline()["slides"][0],
                                             {**_outline()["slides"][1], **kw}])
        bullets = lambda b: _outline(slides=[{**_outline()["slides"][0], "bullets": b},
                                             _outline()["slides"][1]])
        for what, o in [
            ("an identifier on a slide", bullets(["Deciso in DEC-012"])),
            ("a repository path on a slide", bullets(["Vedi products/atlas/PBR.md"])),
            ("a backtick on a slide", bullets(["Usa `dbt`"])),
            ("a week on the roadmap", road(todo=["Chiusura guidata entro due settimane"])),
            ("a quarter on the roadmap", road(todo=["Più banche in Q3"])),
            ("a month on the roadmap", road(done=["Abbinamento, rilasciato a marzo"])),
            ("a percentage on the roadmap", road(done=["Abbinamento al 60%"])),
            ("an ISO date on the roadmap", road(todo=["Chiusura 2026-11-30"])),
            ("a label other than planned", road(todo=[{"text": "Più banche",
                                                        "conditional": True}])),
            ("planned in the done column", road(done=[{"text": "Abbinamento",
                                                        "planned": True}])),
            ("a slide without sources", road(sources=[])),
            ("a title slide without sources", _outline(sources=None)),
            ("no roadmap", _outline(slides=[_outline()["slides"][0]])),
            ("six bullets", bullets([f"Punto {i}" for i in range(6)])),
            ("eleven slides", _outline(slides=[_outline()["slides"][0]] * 9
                                       + [_outline()["slides"][1]])),
            ("an image that is not there", _outline(slides=[
                {"kind": "image", "title": "Come funziona", "image": "diagrams/x.png",
                 "sources": ["ARC#current"]}, _outline()["slides"][1]])),
            ("an unknown key", _outline(audience="interno")),
            ("one card", _outline(slides=[{"kind": "cards", "title": "Per chi",
                                           "cards": [{"title": "CFO", "text": "Firma"}],
                                           "sources": ["PBR §Actors"]},
                                          _outline()["slides"][1]])),
            ("a flow whose focal step does not exist", _outline(slides=[
                {"kind": "flow", "title": "Come funziona", "focal": 3,
                 "steps": [{"title": "Riceve"}, {"title": "Abbina"}],
                 "sources": ["ARC#current"]}, _outline()["slides"][1]])),
            ("an identifier inside a card", _outline(slides=[
                {"kind": "cards", "title": "Per chi", "sources": ["PBR §Actors"],
                 "cards": [{"title": "CFO", "text": "Come da DEC-012"},
                           {"title": "AFC", "text": "Chiude"}]}, _outline()["slides"][1]])),
            ("lanes beside done and todo", road(lanes=[{"name": "Atlas", "done": [],
                                                        "todo": []}])),
            ("a colour that is not one", _outline(theme={"accent": "blue"})),
        ]:
            if not x.check(o, base):
                problems.append(f"{what} was accepted")
        lanes = road(lanes=[{"name": "Atlas", "done": ["Abbinamento"], "todo": []},
                            {"name": "Borea", "done": [],
                             "todo": [{"text": "Prima consegna", "planned": True}]}])
        for k in ("done", "todo"):
            lanes["slides"][1].pop(k)
        if x.check(lanes, base):
            problems.append(f"a roadmap with a row per product was refused: {x.check(lanes, base)}")
        lanes["slides"][1]["lanes"][1]["todo"] = ["Prima consegna a giugno"]
        if not any("a month" in p for p in x.check(lanes, base)):
            problems.append("a month inside a product's row of the roadmap was accepted")
        # "Il report di fine mese" is a capability, not a deadline: time words are refused on
        # the roadmap only, and refusing them elsewhere teaches the outline to blur the product.
        if x.check(bullets(["Prepara il report di fine mese"]), base):
            problems.append("a month mentioned outside the roadmap was refused: it is a "
                            "capability there, not a date")
    return problems


@check("a checked outline renders to a .pptx with no source on it")
def _presentation_renders():
    # python-pptx is in requirements.txt so that this runs in CI rather than printing `ok`
    # over a renderer nobody ran. What it asserts is the part the outline check cannot see:
    # the sources and the customer filter stay in the outline, and the file -- notes included,
    # since notes travel with it -- carries the slide text and nothing else.
    x = _load(RENDER, "render")
    from pptx import Presentation
    problems = []
    with tempfile.TemporaryDirectory() as tmp:
        base = Path(tmp) / "PRS-001-atlas-2026-10-02"
        base.mkdir()
        (base / "outline.yaml").write_text(yaml.safe_dump(_outline(), allow_unicode=True),
                                           encoding="utf-8")
        with contextlib.redirect_stdout(io.StringIO()):
            rc = x.main([str(base / "outline.yaml")])
        out = base / "PRS-001-atlas-2026-10-02.pptx"
        if rc != 0 or not out.is_file():
            return [f"render.py exited {rc} and wrote {'a' if out.is_file() else 'no'} file"]
        prs = Presentation(str(out))
        if len(prs.slides) != 3:
            problems.append(f"{len(prs.slides)} slides for a title and two: expected 3")
        text = " ".join(sh.text_frame.text for sl in prs.slides for sh in sl.shapes
                        if sh.has_text_frame)
        notes = " ".join(sl.notes_slide.notes_text_frame.text for sl in prs.slides
                         if sl.has_notes_slide)
        for want in ("Atlas", "Fatto", "Da fare", "Più banche", "previsto"):
            if want not in text:
                problems.append(f"{want!r} is not on any slide")
        for leak in ("PBR", "RMP", "§"):
            if leak in text or leak in notes:
                problems.append(f"{leak!r} reached the .pptx: sources stay in the outline")
    return problems

DIGEST = ROOT / "skills" / "digest" / "scripts" / "digest.py"
WORKBOOK = ROOT / "skills" / "digest" / "scripts" / "workbook.py"
REFERENCE = ROOT / "tests" / "fixtures" / "digest" / "DIG-003-atlas-2026-10-09.xlsx"
# THE REFERENCE WAS EDITED ONCE, IN 4.4.0, WHERE ITS READERS ASKED FOR A DIFFERENT FILE, AND
# NOWHERE ELSE. Everything in it is still what a spreadsheet wrote and computed, but for:
# - Attività, rows 3 to 7: the first table is «1. La prossima settimana», the items whose days
#   on the Gantt touch 12/10 to 16/10, CHG-018, OD-120 and OD-121, each row copied from the
#   item's own row of «Da fare» in the reference, with the Stato of the new rule;
# - Riepilogo, C31:D37: no column «Ottimale» and no «Differenza», nor the conditional format
#   on the second, nor the pie that charted the first.
# Each part of the archive these do not touch is the spreadsheet's, byte for byte.
_GIT_ID = {"GIT_AUTHOR_NAME": "lead", "GIT_AUTHOR_EMAIL": "lead@example.com",
           "GIT_COMMITTER_NAME": "lead", "GIT_COMMITTER_EMAIL": "lead@example.com"}
# WHERE THE WORKBOOK OF THE FIXTURE MAY DIFFER FROM THE REFERENCE, AND WHY. Everything else has
# to be equal, value by value and style by style, in every cell the reference fills.
#
# Since 4.2.0 there is one estimate instead of a best and a worst case, by the person's choice;
# since 4.5.0 every item counts for the middle of its size, and the pace is the average of the
# last five working days instead of three. So where the reference prints a range the workbook
# prints its middle, and the one row of the estimate takes for its hours the mean of the
# reference's two cases, which is the sum of the middles. The pace of five days, and from it
# the days needed, the expected delivery and the delay, are not in the reference: they are
# computed here, from the reference's own day by day, by the rules of the estimate written
# again below, so that the script's arithmetic is checked against a second one.
_DIGEST_DIFFERS = {
    # D9: in the fixture OD-116 is closed by a decision instead of left without an answer, so
    # nothing leaves the workbook in silence; the reason printed for CHG-021 says so.
    ("Attività", "H36"): "Fuori perimetro dal 09/10: decisione presa su OD-116, passa al "
                         "rilascio successivo",
    # 4.5.0: the summary no longer says what the expectation depends on or where the time
    # goes, the estimate is in rows nobody sees, and the sections after it move up by one.
    ("Riepilogo", "A7"): None, ("Riepilogo", "B7"): None,
    ("Riepilogo", "A8"): None, ("Riepilogo", "B8"): None,
    ("Riepilogo", "A10"): "Calcolo stima consegna",
    ("Riepilogo", "D11"): "Ore lavorate al giorno (media ultimi 5 giorni)",
    ("Riepilogo", "A12"): "Stima",
    ("Riepilogo", "A15"): "2. Cosa serve da altri",
    ("Riepilogo", "A22"): "3. Avanzamento per tema",
    ("Riepilogo", "A30"): "4. Ripartizione delle ore di progetto",
    # 4.5.0: nowhere does the file say how many hours a size counts.
    ("Come leggerlo", "B18"): "da 1 a 2 ore.",
    ("Come leggerlo", "B19"): "da 3 a 5 ore.",
    ("Come leggerlo", "B20"): "da 6 a 12 ore.",
    ("Come leggerlo", "B21"): "da 13 a 40 ore. Troppo grande per essere stimata bene: va "
                              "spezzata in voci più piccole.",
    ("Come leggerlo", "B38"): "Ogni voce ha la sua stima in ore, e la consegna prevista è una "
                              "data sola.",
    ("Come leggerlo", "B39"): "Una voce in corso è stimata per quello che le resta da fare. Le "
                              "ore già spese sono indicate a parte.",
    ("Come leggerlo", "B40"): "Nel Riepilogo, al punto 1.",
}
# 4.5.0: the two rows the summary no longer prints were merged across.
_DIGEST_UNMERGED = {("Riepilogo", "B7:J7"), ("Riepilogo", "B8:J8")}
# 4.5.0: the day by day folds by month and by week, and Excel opens no group on a protected
# sheet; the Gantt with phases was the first sheet left open for the same reason.
_DIGEST_UNPROTECTED = {"Ore lavorate"}
# D11 c: the delay is red when there is one, through a conditional format; in the reference it
# was red by hand, and stayed red at zero.
_DIGEST_RED_WHEN_LATE = {("Riepilogo", "B6"), ("Riepilogo", "J12")}


def _workday(d, n: int):
    """Excel's WORKDAY without holidays: n working days after d, or before it when n < 0."""
    import datetime as _dt
    step = 1 if n > 0 else -1
    while n:
        d += _dt.timedelta(days=step)
        if d.weekday() < 5:
            n -= step
    return d


def _networkdays(a, b) -> int:
    """Excel's NETWORKDAYS without holidays: the working days from a to b, both included."""
    import datetime as _dt
    return sum(1 for k in range((b - a).days + 1) if (a + _dt.timedelta(days=k)).weekday() < 5)


def _excel_round(x: float, digits: int) -> float:
    from decimal import ROUND_HALF_UP, Decimal
    return float(Decimal(repr(x)).quantize(Decimal(1).scaleb(-digits), rounding=ROUND_HALF_UP))


def _digest_expected(ref) -> dict:
    """The value every cell of the reference is expected to have in the generated workbook."""
    import datetime as _dt
    import math as _math
    import re as _re

    def middle(m) -> str:
        x = (int(m.group(1)) + int(m.group(2))) / 2
        return f"{x:g}".replace(".", ",") + " ore"

    out = {}
    for name in ref.sheetnames:
        ws = ref[name]
        for row in ws.iter_rows():
            for c in row:
                v = c.value
                if isinstance(v, str):
                    v = _re.sub(r"^tra (\d+) e (\d+) ore$", middle, v)
                out[(name, c.coordinate)] = v
    # The estimate, again: its hours from the reference's two cases, its pace from the
    # reference's day by day, as the summary's formulas compute them.
    legend, summary, hours = ref["Come leggerlo"], ref["Riepilogo"], ref["Ore lavorate"]
    today, delivery, start = (legend[f"B{r}"].value.date() for r in (4, 6, 7))
    days = {}                                  # date -> the row of the day by day
    for r in range(30, hours.max_row + 1):
        if isinstance(hours[f"B{r}"].value, _dt.datetime):
            days[hours[f"B{r}"].value.date()] = r
    lo, hi = _workday(today, -5), today - _dt.timedelta(days=1)
    window = [r for d, r in days.items() if lo <= d <= hi]
    sums = {c: sum(hours[f"{c}{r}"].value or 0 for r in window) for c in "CDEFGH"}
    nd = _networkdays(lo, hi)
    C = (summary["C12"].value + summary["C13"].value) / 2
    D, E, F = (_excel_round(sums[c] / nd, 1) for c in "CGH")
    G = _excel_round(D - E - F, 1)
    H = _math.ceil(_excel_round(C / G, 9))
    I = _workday(today - _dt.timedelta(days=1), max(1, H))
    J = _networkdays(delivery + _dt.timedelta(days=1), I) if I > delivery else 0
    for col, v in zip("CDEFGHIJ", (C, D, E, F, G, H, _dt.datetime.combine(I, _dt.time()), J)):
        out[("Riepilogo", f"{col}12")] = v
    for col in "ABCDEFGHIJ":
        out[("Riepilogo", f"{col}13")] = None
    # An item of the next week that nobody planned says its days in the queue: the hours of the
    # items before it in «Da fare», and its own, at the pace above.
    acts = ref["Attività"]
    mids = {"S": 1.5, "M": 4, "L": 9, "XL": 26.5}
    head = next(r for r in range(1, acts.max_row + 1)
                if str(acts[f"A{r}"].value or "").startswith("2. Da fare"))
    queue = []
    for r in range(head + 2, acts.max_row + 1):
        if str(acts[f"A{r}"].value or "").endswith(" voci"):
            break
        queue.append((acts[f"A{r}"].value, mids[acts[f"E{r}"].value]))
    eve = today - _dt.timedelta(days=1)
    for r in range(1, head):
        said = acts[f"H{r}"].value
        if not _re.fullmatch(r"(dal \S+ al \S+|il \S+), \d+ ore", str(said or "")):
            continue
        k = [i for i, _ in queue].index(acts[f"A{r}"].value)
        cum, own = sum(h for _, h in queue[:k]), queue[k][1]
        a = _workday(eve, _math.floor(_excel_round(cum / G, 9)) + 1)
        b = _workday(eve, max(1, _math.ceil(_excel_round((cum + own) / G, 9))))
        span = f"il {a:%d/%m}" if a == b else f"dal {a:%d/%m} al {b:%d/%m}"
        out[("Attività", f"H{r}")] = f"{span}, {own:g} ore".replace(".", ",")
    out[("Riepilogo", "B5")] = f"il {I:%d/%m}"
    out[("Riepilogo", "B6")] = ("nessuno" if J == 0 else f"{J} giorno lavorativo" if J == 1
                                else f"{J} giorni lavorativi")
    # The first row of each table of Ore lavorate is the window of the pace.
    out[("Ore lavorate", "A5")] = out[("Ore lavorate", "A12")] = "Ultimi 5 giorni lavorativi"
    out[("Ore lavorate", "B5")] = out[("Ore lavorate", "B12")] = \
        _dt.datetime.combine(lo, _dt.time())
    for col, src in zip("DEFGHI", "CDEFGH"):
        out[("Ore lavorate", f"{col}5")] = sums[src]
        if col == "D":
            out[("Ore lavorate", "D12")] = sums[src]
        else:
            out[("Ore lavorate", f"{col}12")] = sums[src] / sums["C"] if sums["C"] else 0
    on_release = _networkdays(max(lo, start), hi)
    out[("Ore lavorate", "J5")] = (sums["D"] + sums["E"] + sums["F"]) / on_release
    note = next(c.coordinate for row in hours.iter_rows(min_row=30) for c in row
                if c.value == "I giorni in rosso sono sabato e domenica.")
    out[("Ore lavorate", note)] = ("Mesi e settimane si aprono con il «+» a sinistra. I giorni "
                                   "in rosso sono sabato e domenica.")
    out.update(_DIGEST_DIFFERS)
    return out


def _digest_where(ref, got) -> dict:
    """Where a cell of the reference sits in the generated workbook, when it moved: the days
    of Ore lavorate, under the rows of their month and week, and the note under them."""
    import datetime as _dt
    hours, mine = ref["Ore lavorate"], got["Ore lavorate"]
    at = {}
    for r in range(30, mine.max_row + 1):
        v = mine[f"B{r}"].value
        if isinstance(v, _dt.datetime):
            at[v.date()] = r
        elif str(mine[f"A{r}"].value or "").startswith("Mesi e settimane"):
            at["note"] = r
    where = {}
    for row in hours.iter_rows(min_row=30):
        for c in row:
            v = hours[f"B{c.row}"].value
            key = (v.date() if isinstance(v, _dt.datetime) else
                   "note" if hours[f"A{c.row}"].value else None)
            if key in at:
                where[("Ore lavorate", c.coordinate)] = f"{c.column_letter}{at[key]}"
    return where


def _digest_copy(tmp: str, which: str = "atlas") -> Path | None:
    """A private copy of the digest fixture: rendering commits and pushes, so never the build."""
    src = built("digest")
    if src is None:
        return None
    dst = Path(tmp) / "fixture"
    shutil.copytree(src, dst, symlinks=True)
    return dst / which


def _digest(root: Path, *args: str) -> tuple[int, str]:
    r = subprocess.run([sys.executable, str(DIGEST), "--root", str(root), "--product", "atlas",
                        *args], capture_output=True, text=True)
    return r.returncode, r.stdout + r.stderr


def _edit(path: Path, old: str, new: str) -> None:
    text = path.read_text(encoding="utf-8")
    if old not in text:
        raise AssertionError(f"{path.name} no longer contains {old[:50]!r}: the fixture moved")
    path.write_text(text.replace(old, new, 1), encoding="utf-8")


def _state(root: Path, change) -> None:
    """Edit the state file of the fixture as data: `change` receives it and edits it in place."""
    path = root / "_meta" / "digest" / "state-atlas.yaml"
    data = yaml.safe_load(path.read_text(encoding="utf-8"))
    change(data)
    path.write_text(yaml.safe_dump(data, allow_unicode=True, sort_keys=False), encoding="utf-8")


def _cell_style(c) -> tuple:
    """What a reader sees of a cell, with the spellings of a default made one."""
    colour = c.font.color.rgb if c.font.color is not None and c.font.color.type == "rgb" else None
    fill = c.fill.fgColor.rgb if c.fill is not None and c.fill.fill_type == "solid" else None
    return (c.font.name, c.font.sz, bool(c.font.b), bool(c.font.i), colour, fill,
            c.number_format, bool(c.alignment.wrap_text), c.alignment.horizontal or "general",
            c.alignment.vertical or "bottom")


def _same_value(a, b) -> bool:
    import datetime as _dt
    a, b = [x.date() if isinstance(x, _dt.datetime) else x for x in (a, b)]
    if isinstance(a, (int, float)) and isinstance(b, (int, float)) \
            and not isinstance(a, bool) and not isinstance(b, bool):
        return abs(a - b) < 1e-9
    return a == b


@check("the workbook of the fixture is the reference, cell by cell, value and style")
def _digest_reproduces_the_reference():
    # THE FORMAT IS FIXED, AND THIS IS WHERE IT IS FIXED. The person who reads the digest every
    # morning made the reference workbook before a line of this code existed, and a spreadsheet
    # computed every formula in it. Here the script renders the same day from a register and a
    # state file: every value it writes beside a formula has to be the value that spreadsheet
    # computed, and every cell has to look the same. A change to the format is a change to the
    # reference first, and `_DIGEST_DIFFERS` lists the cells allowed to differ, with the reason.
    import openpyxl
    with tempfile.TemporaryDirectory() as tmp:
        root = _digest_copy(tmp)
        if root is None:
            return ["the digest fixture could not be built: evals/fixtures/generators/digest.py"]
        rc, out = _digest(root, "--date", "2026-10-09", "--now", "2026-10-09T08:30:00+02:00",
                          "--copy-to", str(Path(tmp) / "copy"))
        snap = root / "_meta" / "digest" / "DIG-003-atlas-2026-10-09"
        xlsx = snap / "DIG-003-atlas-2026-10-09.xlsx"
        if rc != 0 or not xlsx.is_file():
            return [f"digest.py exited {rc} and wrote {'a' if xlsx.is_file() else 'no'} "
                    f"workbook: {out.strip()[-400:]}"]
        problems = []
        ref = openpyxl.load_workbook(REFERENCE, data_only=True)
        got = openpyxl.load_workbook(xlsx, data_only=True)
        want_sheets = ref.sheetnames[:1] + ["Gantt"] + ref.sheetnames[1:]
        if got.sheetnames != want_sheets:
            return [f"the sheets are {got.sheetnames}, not {want_sheets}"]
        expected = _digest_expected(ref)
        where = _digest_where(ref, got)
        for name in ref.sheetnames:
            r, g = ref[name], got[name]
            for row in r.iter_rows():
                for c in row:
                    other = g[where.get((name, c.coordinate), c.coordinate)]
                    want = expected[(name, c.coordinate)]
                    if not _same_value(want, other.value):
                        problems.append(f"{name}!{c.coordinate} is {other.value!r}, expected "
                                        f"{want!r}; the reference says {c.value!r}")
                    if c.value is None or want is None:
                        continue
                    a, b = _cell_style(c), _cell_style(other)
                    if (name, c.coordinate) in _DIGEST_RED_WHEN_LATE:
                        a, b = a[:4] + a[6:], b[:4] + b[6:]
                    if a != b:
                        problems.append(f"{name}!{c.coordinate} looks {b}, the reference {a}")
            for what, mine, theirs in (
                    ("merged cells", {str(m) for m in g.merged_cells.ranges if m.min_row < 40},
                     {str(m) for m in r.merged_cells.ranges
                      if (name, str(m)) not in _DIGEST_UNMERGED}),
                    ("frozen panes", g.freeze_panes, r.freeze_panes),
                    ("protection", g.protection.sheet,
                     r.protection.sheet and name not in _DIGEST_UNPROTECTED),
                    ("filter and sort", (g.protection.autoFilter, g.protection.sort),
                     (r.protection.autoFilter, r.protection.sort))):
                if mine != theirs:
                    problems.append(f"{name}: {what} {mine}, the reference {theirs}")
            for col, dim in r.column_dimensions.items():
                if dim.width and abs((g.column_dimensions[col].width or 0) - dim.width) > 0.75:
                    problems.append(f"{name}: column {col} is {g.column_dimensions[col].width} "
                                    f"wide, the reference {dim.width}")
        summary = openpyxl.load_workbook(xlsx)["Riepilogo"]
        if not all(summary.row_dimensions[r].hidden for r in range(10, 14)):
            problems.append("the estimate of the delivery is in view: its rows are to be hidden")
        conditional = {str(rng.sqref): [rule.formula for rule in rng.rules]
                       for rng in summary.conditional_formatting}
        if not any("$J$12>0" in str(f) for f in conditional.get("B6:J6", [])) or \
                not any("J12>0" in str(f) for f in conditional.get("J12", [])):
            problems.append(f"the delay is not red through a conditional format: {conditional}")
        # The reference has a pie under section 5, and had two until the split «optimal» it
        # charted went in 4.4.0; the person who sends the file asked for the pies to go in
        # 4.3.0, and the summary has no chart at all.
        if got["Riepilogo"]._charts:
            problems.append("the summary carries a chart: the pies were taken out on request")
        if not (Path(tmp) / "copy" / xlsx.name).is_file():
            problems.append("--copy-to did not leave a copy where it was asked")
        if f"Consegna prevista oggi: {expected[('Riepilogo', 'B5')]}" not in out \
                or "spezzata: OD-098 in OD-120, OD-121" not in out:
            problems.append(f"the conversation does not get the summary and the changes: {out}")
        remote = root / ".remote" / "digest.git"
        log = subprocess.run(["git", "-C", str(remote), "log", "--format=%s"],
                             capture_output=True, text=True).stdout.split()
        if "DIG-003-atlas-2026-10-09" not in log:
            problems.append(f"the snapshot was not pushed to the private store: {log}")
    return problems[:12]


def _gantt_rows(ws) -> dict:
    """The rows of the Gantt by what names them: the ID, or the name of a milestone."""
    rows: dict = {}
    for r in range(7, ws.max_row + 1):
        a, b = ws.cell(r, 1).value, ws.cell(r, 2).value
        key = a if a not in (None, "", "—") else b
        if key:
            rows.setdefault(str(key), r)
    return rows


@check("the Gantt sheet places every item, the milestones and today, by day and by week")
def _digest_draws_the_gantt():
    # The Gantt is the one sheet the reference does not have, so it is checked against the plan
    # of the fixture, computed by hand at 3,1 hours a day from 09/10 with every item at the
    # middle of its size: CHG-018, first in the queue, 9 hours, 2,9 days, from 09/10 to the
    # third working day, 13/10, and «in corso dal 08/10», the day it was started, which its bar
    # does not draw; OD-120 from 9/3,1 to 18/3,1 days, so from the third to the sixth; and so
    # on to CHG-024, whose 43rd hour falls on the 14th working day, 28/10, the date of point 1
    # of the summary. CHG-022 and CHG-024 wait for what the infrastructure team was asked.
    import datetime as _dt
    import openpyxl
    with tempfile.TemporaryDirectory() as tmp:
        root = _digest_copy(tmp)
        if root is None:
            return ["the digest fixture could not be built"]
        rc, out = _digest(root, "--date", "2026-10-09", "--now", "2026-10-09T08:30:00+02:00")
        xlsx = root / "_meta/digest/DIG-003-atlas-2026-10-09/DIG-003-atlas-2026-10-09.xlsx"
        if rc != 0:
            return [f"the digest of 09/10 was refused: {out.strip()[:300]}"]
        book = openpyxl.load_workbook(xlsx, data_only=True)
        formulas = openpyxl.load_workbook(xlsx)["Gantt"]
    problems = []
    if book.sheetnames[1] != "Gantt":
        return [f"the Gantt is not the second sheet: {book.sheetnames}"]
    ws = book["Gantt"]
    if book["Riepilogo"].max_row > 40:
        problems.append("the summary still carries a calendar below the pies")
    # The days shown one by one run from four weeks before the digest's week, 07/09, to four
    # weeks after it, Friday 06/11; the project started on 22/09, so they begin on the Monday
    # of that week, and no column of a week is left before them; the last date, 28/10, is
    # inside them, so none after.
    calendar, d = [], _dt.date(2026, 9, 21)
    while d <= _dt.date(2026, 11, 6):
        calendar += [d] if d.weekday() < 5 else []
        d += _dt.timedelta(days=1)
    header = [ws.cell(6, 8 + k).value for k in range(len(calendar) + 1)]
    weekly = [ws.cell(5, 8 + k).value == "sett." for k in range(len(calendar))]
    if [h.date() if h else None for h in header] != calendar + [None] or any(weekly):
        problems.append(f"the columns are not the working days from 21/09 to 06/11: {header}")
    rows = _gantt_rows(ws)

    def day(text: str) -> _dt.date:
        dd, mm = map(int, text.split("/"))
        return _dt.date(2026, mm, dd)

    for label, state, first, last in (
            ("CHG-012", "Chiusa", "24/09", "24/09"), ("DEC-026", "Chiusa", "06/10", "06/10"),
            ("KI-013", "Chiusa", "07/10", "07/10"),
            ("CHG-018", "In corso dal 08/10", "09/10", "13/10"),
            ("OD-120", "Da iniziare", "13/10", "16/10"),
            ("OD-121", "Da iniziare", "16/10", "21/10"),
            ("OD-109", "Rallentata", "21/10", "26/10"),
            ("CHG-022", "Bloccata da risposta sul registro delle immagini (team "
                        "infrastruttura)", "27/10", "27/10"),
            ("CHG-024", "Bloccata da ambiente preprod (team infrastruttura)", "27/10", "28/10")):
        r = rows.get(label)
        got = tuple(ws.cell(r, c).value for c in (4, 5, 6)) if r else None
        if got is None or got[0] != state or got[1].date() != day(first) \
                or got[2].date() != day(last):
            problems.append(f"{label} is {got}, not {state} from {first} to {last}")
    r = rows.get("OD-120", 1)
    if "Attività!$E$" not in str(formulas.cell(r, 5).value) or "Riepilogo!$G$12" \
            not in str(formulas.cell(r, 6).value):
        problems.append("an item to do is not placed by formulas on its size and on the pace: "
                        f"{formulas.cell(r, 5).value!r}")
    for label, on in (("Demo al team funzionale", "15/10"), ("Consegna concordata", "13/10"),
                      ("Consegna prevista", "28/10")):
        r = rows.get(label)
        marks = [calendar[k] for k in range(len(calendar)) if r and ws.cell(r, 8 + k).value == "◆"]
        if marks != [day(on)]:
            problems.append(f"the milestone {label!r} is marked on {marks}, not on {on}")
    rules = {str(c.sqref): [x.formula[0] for x in c.rules] for c in formulas.conditional_formatting}
    bars = next((v for k, v in rules.items() if k.startswith("H8:")), [])
    if len(bars) != 6 or "'Come leggerlo'!$B$4" not in bars[-1]:
        problems.append(f"the bars and the line of today are not conditional formats: {rules}")
    if "OD-098" in rows or "CHG-021" in rows:
        problems.append("an item split or out of the perimeter is on the Gantt")
    if formulas.freeze_panes != "H7":
        problems.append(f"the names and the days do not stay in view: {formulas.freeze_panes}")
    return problems


@check("a new entry in the register stops the digest until it is classified")
def _digest_refuses_the_unclassified():
    # An item the state file does not know has no theme, no scope and no size, so the hours
    # left would be computed without it and the workbook would print a number that is wrong
    # in a way nobody sees. The morning before anybody answered is that state, three times.
    with tempfile.TemporaryDirectory() as tmp:
        root = _digest_copy(tmp, "atlas-asks")
        if root is None:
            return ["the digest fixture could not be built"]
        problems = []
        rc, out = _digest(root, "--date", "2026-10-09", "--check")
        for want in ("KI-014", "OD-120", "OD-121", "no hours declared for 2026-10-07",
                     "no plan for 2026-10-12"):
            if rc != 1 or want not in out:
                problems.append(f"the morning before the answers did not refuse on {want!r} "
                                f"(exit {rc}): {out.strip()[:300]}")
        rc, out = _digest(root, "--date", "2026-10-09")
        if rc == 0 or (root / "_meta/digest/DIG-003-atlas-2026-10-09").exists():
            problems.append("the workbook was written with entries nobody classified")
    return problems


@check("the digest refuses every way a number could enter or leave it unseen")
def _digest_refusals():
    # One mutation each, on a fresh copy, and each must be refused by itself with a sentence
    # that names it: a check that passes with one of these in it prints a count somebody sends.
    def day(d, value):
        return lambda s: s["days"].__setitem__(d, value)

    cases = [
        ("an open item gone from the register without closing",
         lambda r: (_edit(r / "products/atlas/OPEN.md",
                          "  OD-109:\n    status: open\n    cost_to_reverse: low\n"
                          "    default_in_force: a question with no defined metric gets no "
                          "answer\n    trigger: the metric glossary arriving from the "
                          "functional team", ""),
                    _edit(r / "products/atlas/OPEN.md",
                          "### OD-109 · How to answer a question no metric covers", "### Gone")),
         "OD-109"),
        ("a register title changed after the confirmation",
         lambda r: _edit(r / "products/atlas/OPEN.md", "### OD-121 · How often",
                         "### OD-121 · When and how often"), "confirm its classification"),
        ("a change in draft planned for today",
         lambda r: _edit(next((r / "products/atlas/changes").glob("CHG-018-*.md")),
                         "status: approved", "status: draft"), "draft"),
        ("a register status the digest does not map",
         lambda r: _edit(next((r / "products/atlas/changes").glob("CHG-022-*.md")),
                         "status: approved", "status: paused"), "does not map"),
        ("an item closed in the perimeter with no closing date",
         lambda r: _state(r, lambda s: s["items"]["KI-013"].pop("closed_on")), "closed_on"),
        ("an item outside the perimeter with no reason",
         lambda r: _state(r, lambda s: s["items"]["INC-041"].pop("out_reason")), "out_reason"),
        ("an identifier inside a description",
         lambda r: _state(r, lambda s: s["items"]["OD-115"].__setitem__(
             "what", "cosa vede l'utente quando una richiesta fallisce, come in DEC-026.")),
         "an identifier"),
        ("a long dash in a title",
         lambda r: _state(r, lambda s: s["items"]["CHG-012"].__setitem__(
             "title", "repository unica — monorepo")), "long dash"),
        ("a working day with no hours declared",
         lambda r: _state(r, lambda s: s["days"].pop("2026-10-08")),
         "no hours declared for 2026-10-08"),
        ("hours on a day not over yet",
         lambda r: _state(r, day("2026-10-09", {"hours": {"CHG-018": 1}})), "not over yet"),
        ("a category outside the product that does not exist",
         lambda r: _state(r, day("2026-10-08", {"hours": {"CHG-018": 4, "OD-076": 1},
                                                "outside": {"caffè": 2}})),
         "is not a category"),
        ("days rebuilt in bulk that the items do not add up to",
         lambda r: _state(r, lambda s: s["items"]["CHG-012"].__setitem__("hours_before", 10)),
         "have to agree"),
        ("an open item missing from its list",
         lambda r: _state(r, lambda s: s["order"]["todo"].remove("OD-121")),
         "does not list OD-121"),
        ("a planned item with no hours",
         lambda r: _state(r, lambda s: s["plan"]["2026-10-12"][1].pop("hours")), "no hours"),
        ("no milestones declared",
         lambda r: _state(r, lambda s: s.pop("milestones")), "no `milestones`"),
        ("an owner that is not a name",
         lambda r: _state(r, lambda s: s["items"]["CHG-018"].__setitem__("owner", 3)),
         "`owner` is who follows it"),
        ("a Gantt that starts from a month",
         lambda r: _state(r, lambda s: s.__setitem__("gantt", {"from": "mese"})),
         "`gantt.from` is 'mese'"),
        ("a preference of the Gantt that nothing reads",
         lambda r: _state(r, lambda s: s.__setitem__("gantt", {"done": False, "colori": "blu"})),
         "unknown key 'colori' under `gantt`"),
        # The keys of 4.5.0, each refused when its value is not one the Gantt can print.
        ("a Gantt that shows the closed items sometimes",
         lambda r: _state(r, lambda s: s.__setitem__("gantt", {"done": "a volte"})),
         "`gantt.done` is 'a volte'"),
        ("a default owner that is not a name",
         lambda r: _state(r, lambda s: s.__setitem__("gantt", {"owner": 3})),
         "`gantt.owner` is the name"),
        ("an owner of the release that is not a name",
         lambda r: _state(r, lambda s: s["release"].__setitem__("owner", "  ")),
         "`release.owner` is who answers"),
        ("an owner of a milestone that is not a name",
         lambda r: _state(r, lambda s: s["milestones"][0].__setitem__("owner", 5)),
         "`owner` is who answers for it"),
        ("the requests in view named by a sentence, not a list",
         lambda r: _state(r, lambda s: s.__setitem__("gantt", {"requests": "team funzionale"})),
         "`gantt.requests` is a list of names"),
        ("the requests in view named by a number",
         lambda r: _state(r, lambda s: s.__setitem__("gantt", {"requests": [3]})),
         "`gantt.requests` is a list of names"),
        ("a phase to outline in a roadmap without phases",
         lambda r: _state(r, lambda s: s.__setitem__(
             "gantt", {"outline": {"t3": {"from": "2026-11-02", "to": "2026-11-27"}}})),
         "`gantt.outline` names 't3', and the roadmap has no such stage"),
        ("a state file of 4.1.0",
         lambda r: _state(r, lambda s: s.pop("format")), "--migrate"),
        ("the hours of one item declared in two products",
         lambda r: (r / "_meta/digest/state-borea.yaml").write_text(
             "format: 2\nproduct: borea\ndays:\n  2026-10-08:\n    hours:\n      CHG-018: 2\n",
             encoding="utf-8"), "declared once"),
    ]
    problems = []
    for what, mutate, says in cases:
        with tempfile.TemporaryDirectory() as tmp:
            root = _digest_copy(tmp)
            if root is None:
                return ["the digest fixture could not be built"]
            mutate(root)
            rc, out = _digest(root, "--date", "2026-10-09", "--check")
            if rc != 1 or says not in out:
                problems.append(f"{what}: exit {rc}, and the refusal does not say {says!r}: "
                                f"{out.strip()[:200]}")
    with tempfile.TemporaryDirectory() as tmp:
        root = _digest_copy(tmp)
        rc, out = _digest(root, "--date", "2026-10-10", "--check")
        if rc != 1 or "Saturday" not in out:
            problems.append(f"a digest dated on a Saturday was not refused: {out.strip()[:200]}")
    return problems


@check("a day already sent is corrected aloud, and an earlier snapshot cannot be edited")
def _digest_history_is_frozen():
    # Since 4.5.0 a day already sent can be declared anew: the readers know the file changes
    # from one sending to the next. It is counted again, said in the conversation before the
    # file is written, and recorded in the snapshot. A day that disappears from the state file
    # is not a correction, and an earlier snapshot is never edited.
    import datetime as _dt
    import openpyxl
    problems = []
    for what, mutate, says in [
        ("the hours of a day already sent, gone",
         lambda r: _state(r, lambda s: s["days"].pop("2026-10-06")), "no longer declares it"),
        ("an earlier snapshot",
         lambda r: _edit(next((r / "_meta/digest").glob("DIG-001-*/frozen.yaml")),
                         "baseline: true", "baseline: false"), "no longer matches"),
    ]:
        with tempfile.TemporaryDirectory() as tmp:
            root = _digest_copy(tmp)
            if root is None:
                return ["the digest fixture could not be built"]
            mutate(root)
            rc, out = _digest(root, "--date", "2026-10-09", "--check")
            if rc != 1 or says not in out:
                problems.append(f"{what} was edited and nothing said so: {out.strip()[:300]}")
    with tempfile.TemporaryDirectory() as tmp:
        root = _digest_copy(tmp)
        if root is None:
            return ["the digest fixture could not be built"]
        before = {}

        def correct(s):
            before.update(s["days"]["2026-10-06"]["hours"])
            s["days"]["2026-10-06"]["hours"]["OD-114"] = before["OD-114"] + 2

        _state(root, correct)
        note = "Giorni già inviati con DIG-002-atlas-2026-10-07 e corretti, contati di nuovo: 06/10"
        rc, out = _digest(root, "--date", "2026-10-09", "--check")
        if rc != 0 or note not in out:
            problems.append(f"a day already sent and corrected is refused, or not said before "
                            f"the file is written: {out.strip()[:300]}")
        rc, out = _digest(root, "--date", "2026-10-09", "--now", "2026-10-09T08:30:00+02:00")
        snap = root / "_meta/digest/DIG-003-atlas-2026-10-09"
        if rc != 0 or note not in out:
            return problems + [f"the digest with a day corrected: {out.strip()[:300]}"]
        hours = openpyxl.load_workbook(snap / "DIG-003-atlas-2026-10-09.xlsx",
                                       data_only=True)["Ore lavorate"]
        day = next(r for r in range(30, hours.max_row + 1)
                   if hours[f"B{r}"].value == _dt.datetime(2026, 10, 6))
        ref = openpyxl.load_workbook(REFERENCE, data_only=True)["Ore lavorate"]
        was = next(ref[f"C{r}"].value for r in range(30, ref.max_row + 1)
                   if ref[f"B{r}"].value == _dt.datetime(2026, 10, 6))
        if hours[f"C{day}"].value != was + 2:
            problems.append(f"the day corrected counts {hours[f'C{day}'].value} hours, not the "
                            f"{was} it was sent with and the 2 added")
        frozen = yaml.safe_load((snap / "frozen.yaml").read_text(encoding="utf-8"))
        if frozen.get("corrected") != ["2026-10-06"]:
            problems.append(f"the snapshot does not record the day corrected: "
                            f"{frozen.get('corrected')}")
    return problems


@check("the workbook says «non stimabile» and «Nessuna.» where a number or a row would lie")
def _digest_edge_cases():
    # D11 d. No hours on the release in the last five working days, and every list empty: a
    # spreadsheet left to itself prints #DIV/0!, a delivery dated yesterday, or counts the
    # line that says the table is empty. The model is built by hand, so that a case the
    # fixture never reaches is still rendered and read back.
    import openpyxl
    x = _load(DIGEST, "digest_edges")
    from datetime import date as _d
    from fractions import Fraction as Fr
    when, start = _d(2026, 10, 9), _d(2026, 10, 1)
    table = []
    d = start
    while d < when:
        table.append([d] + [Fr(0)] * 4 + [Fr(4), Fr(0), Fr(1), Fr(0), Fr(0)])
        d += x.timedelta(days=1)
    m = x.Model("Borea", when, _d(2026, 10, 8), _d(2026, 10, 20), start, _d(2026, 10, 12), "1.0",
                "linea di base", "08/10", [], [], [], [], {}, [], {t: 0 for t in x.THEMES},
                {t: 0 for t in x.THEMES}, {t: 0 for t in x.THEMES}, table, [], [])
    F = x.figures(m)
    problems = []
    if F["B5"] != "non stimabile" or F["case"]["H"] != "non stimabile":
        problems.append(f"no hours on the release gave a date: {F['B5']!r}")
    if F["B7"] != "non dipende da richieste ad altri.":
        problems.append(f"the conversation says the forecast depends on requests: {F['B7']!r}")
    with tempfile.TemporaryDirectory() as tmp:
        out = Path(tmp) / "edge.xlsx"
        _load(WORKBOOK, "digest_edges_workbook").write(m, F, out, x.datetime(2026, 10, 9, 8, 30))
        wb = openpyxl.load_workbook(out, data_only=True)
    s, a = wb["Riepilogo"], wb["Attività"]
    if [wb["Gantt"].cell(r, 4).value for r in (8, 9, 10, 11)] != \
            ["Concordata", "Non stimabile", None, None]:
        problems.append("a Gantt with nothing to do does not hold the two deliveries alone, the "
                        "expected one «Non stimabile»")
    for cell, want in (("B5", "non stimabile"), ("B6", "non stimabile"), ("I12", "non stimabile"),
                       ("B7", None), ("B12", 0),
                       ("A17", "Nessuna.")):
        if s[cell].value != want:
            problems.append(f"Riepilogo!{cell} is {s[cell].value!r}, not {want!r}")
    nessuna = [c.coordinate for row in a.iter_rows() for c in row if c.value == "Nessuna."]
    totals = [c.value for row in a.iter_rows() for c in row
              if isinstance(c.value, str) and c.value.endswith(" voci")]
    if len(nessuna) != 4 or totals != ["0 voci"] * 3:
        problems.append(f"empty lists show {nessuna} and count {totals}")
    return problems


@check("a state file of 4.1.0 migrates, and what a script cannot know is listed, not guessed")
def _digest_migrates_the_state():
    with tempfile.TemporaryDirectory() as tmp:
        root = Path(tmp)
        store = root / "_meta" / "digest"
        store.mkdir(parents=True)
        (store / "state-atlas.yaml").write_text(
            "product: atlas\nstandard_hours: 8\nrelease:\n  name: \"1.0\"\n"
            "  delivery: 2026-10-13\nitems:\n  CHG-019:\n    title: flusso di rilascio\n"
            "    theme: infrastruttura\n    scope: \"1.0\"\n    size: S\n"
            "waits:\n  - what: Ambiente preprod\n    owner: team infrastruttura\n"
            "    asked: 2026-09-25\n    needed_by: 2026-10-09\n    without: niente prove.\n"
            "    note: blocca il rilascio\nrequests:\n  OD-116:\n    to: responsabile\n"
            "    by: 2026-10-09\nplan:\n  2026-10-07: [CHG-019]\nperiods:\n  2026-10-06:\n"
            "    hours:\n      CHG-019: 1.5\n    outside:\n      Riunioni di stato: 1\n"
            "      Formazione di un junior: 0.5\n", encoding="utf-8")
        r = subprocess.run([sys.executable, str(DIGEST), "--root", str(root), "--migrate"],
                           capture_output=True, text=True)
        try:
            data = yaml.safe_load((store / "state-atlas.yaml").read_text(encoding="utf-8"))
        except yaml.YAMLError:
            data = {}
    problems = []
    if r.returncode or data.get("format") != 2:
        return [f"--migrate exited {r.returncode}: {(r.stdout + r.stderr).strip()[:300]}"]
    if data["items"]["CHG-019"]["theme"] != "deploy":
        problems.append("the theme infrastruttura was not renamed deploy")
    day = (data.get("days") or {}).get("2026-10-06") or {}
    if day.get("outside") != {"riunioni": 1, "formazione": 0.5} or \
            day.get("hours") != {"CHG-019": 1.5}:
        problems.append(f"the hours of 06/10 did not become a day: {day}")
    for word in ("release.start", "requests", "standard_hours", "order", "milestones", "note"):
        if word not in r.stdout:
            problems.append(f"the migration does not tell the person about {word!r}")
    if "standard_hours" in data or "requests" in data or "periods" in data:
        problems.append("a key of 4.1.0 survived the migration")
    return problems


@check("every status the registry declares for what the digest reads is mapped")
def _digest_maps_every_status():
    # The registry says which statuses exist and not which of them are done, so the reading
    # lives in the script; this is what stops a status added there from being read as nothing.
    x = _load(DIGEST, "digest_states")
    types = REGISTRY["types"]
    declared = {
        "open-register": types["open-register"]["maps"]["entries"]["fields"]["enums"]["status"],
        "decision-record": types["decision-record"]["status"],
        "change-contract": types["change-contract"]["status"],
        "evaluation-report": types["evaluation-report"]["status"],
    }
    problems = []
    for t, statuses in declared.items():
        if set(statuses) != set(x.STATES.get(t, {})):
            problems.append(f"{t}: the registry declares {sorted(statuses)} and the digest maps "
                            f"{sorted(x.STATES.get(t, {}))}")
    if x.STATES["change-contract"]["verified"] != "closed" or any(
            v == "closed" for k, v in x.STATES["change-contract"].items() if k != "verified"):
        problems.append("a change is closed at `verified` and at nothing else: the Definition "
                        "of Done of templates/AGENTS.md")
    return problems


@check("the items a commit names are proposed day by day, and a regenerated index is not")
def _digest_reads_the_history():
    with tempfile.TemporaryDirectory() as tmp:
        root = _digest_copy(tmp)
        if root is None:
            return ["the digest fixture could not be built"]
        rc, out = _digest(root, "--date", "2026-10-09", "--inventory")
        inv = yaml.safe_load(out) if rc == 0 else {}
    by_day = (inv or {}).get("worked_by_day") or {}
    problems = []
    for d, want in (("2026-10-07", "KI-013"), ("2026-10-08", "CHG-018"),
                    ("2026-10-08", "OD-076"), ("2026-10-08", "OD-116")):
        if want not in (by_day.get(d) or {}):
            problems.append(f"{want} changed on {d} and is not proposed as worked on that day")
    for noise in ("OD-109", "CHG-024", "OD-115"):
        if any(noise in v for v in by_day.values()):
            problems.append(f"{noise} is proposed as worked on because a regenerated index "
                            "names it")
    if (inv or {}).get("aliases") != {"DEC-026": "OD-114", "DEC-036": "OD-116"}:
        problems.append(f"a decision and the entry it closes are one item: {inv.get('aliases')}")
    closed = {c["id"]: c["proposed_closed_on"] for c in (inv or {}).get("closed_since") or []}
    if closed.get("KI-013") != "2026-10-07":
        problems.append(f"the closing day of KI-013 is not proposed from its commit: {closed}")
    return problems


@check("the store is pushed only where it is private, and never into the documents")
def _digest_store_stays_private():
    x = _load(DIGEST, "digest_store")
    problems = []
    if not x.private("../../.remote/digest.git")[0]:
        problems.append("a local path was refused as a remote: nothing leaves the machine")
    with tempfile.TemporaryDirectory() as tmp:
        fake = Path(tmp) / "gh"
        fake.write_text("#!/bin/sh\necho false\n", encoding="utf-8")
        fake.chmod(0o755)
        old = os.environ.get("PATH", "")
        os.environ["PATH"] = f"{tmp}{os.pathsep}{old}"
        try:
            ok, why = x.private("git@github.com:somebody/atlas-digest.git")
        finally:
            os.environ["PATH"] = old
        if ok:
            problems.append(f"a GitHub remote that is not private was accepted: {why}")
    with tempfile.TemporaryDirectory() as tmp:
        root = _digest_copy(tmp)
        if root is None:
            return problems + ["the digest fixture could not be built"]
        (root / ".gitignore").write_text(".remote/\n", encoding="utf-8")
        rc, out = _digest(root, "--date", "2026-10-09")
        if rc != 1 or "not ignored" not in out:
            problems.append("the store was written while the documentation repository would "
                            f"have committed it: {out.strip()[:200]}")
    return problems


def _stages_copy(tmp: str) -> Path | None:
    """A private copy of the fixture with a roadmap in stages."""
    src = built("digest-stages")
    if src is None:
        return None
    dst = Path(tmp) / "stages"
    shutil.copytree(src, dst, symlinks=True)
    return dst / "borea"


def _borea(root: Path, *args: str) -> tuple[int, str]:
    r = subprocess.run([sys.executable, str(DIGEST), "--root", str(root), "--product", "borea",
                        *args], capture_output=True, text=True)
    return r.returncode, r.stdout + r.stderr


@check("with a roadmap in stages, the digest takes its order and its Gantt from the roadmap")
def _digest_reads_the_stages():
    # Computed by hand from the fixture, at 2,4 hours a day for the release from 07/10 (25 hours
    # in the last five working days, 30/09 to 06/10, 13 of them outside the product). Every item
    # counts for the middle of its size and the items to do follow the stages, each from its
    # place in the queue however early it was started: CHG-302 (9, worked 06/10) 07/10 to
    # 12/10; OD-201 (9, blocked, so last in its increment, worked 06/10) 12/10 to 16/10; OD-202
    # (1,5) 16/10 to 19/10; CHG-303 (4, worked 06/10) 19/10 to 20/10; CHG-304 (4, of a
    # conditional increment, counted) 20/10 to 22/10; KI-021 (1,5, in no increment) 22/10 to
    # 23/10, the expected delivery. INC-101 is worth 4 + 6 + 9 + 9 = 28 hours, its done half
    # included; INC-104 has nothing under it. OD-201 waits for the list of the technicians'
    # skills, asked of the operations office. The milestones of the phases head the sheet with
    # the deliveries, by date, and say what the rows of their phase say.
    import datetime as _dt
    import openpyxl
    with tempfile.TemporaryDirectory() as tmp:
        root = _stages_copy(tmp)
        if root is None:
            return ["the stages fixture could not be built: "
                    "evals/fixtures/generators/digest_stages.py"]
        rc, out = _borea(root, "--date", "2026-10-07", "--now", "2026-10-07T08:30:00+02:00")
        xlsx = root / "_meta/digest/DIG-002-borea-2026-10-07/DIG-002-borea-2026-10-07.xlsx"
        if rc != 0:
            return [f"the digest of 07/10 was refused: {out.strip()[:400]}"]
        book = openpyxl.load_workbook(xlsx, data_only=True)
        formulas = openpyxl.load_workbook(xlsx)["Gantt"]
        frozen = yaml.safe_load((xlsx.parent / "frozen.yaml").read_text(encoding="utf-8"))
    problems = []
    want = ["CHG-302", "OD-201", "OD-202", "CHG-303", "CHG-304", "KI-021"]
    if frozen.get("order") != want:
        problems.append(f"the order of «Da fare» is {frozen.get('order')}, not {want}")
    if not (frozen.get("roadmap") or {}).get("sha256"):
        problems.append("the snapshot does not record which roadmap it was computed from")
    ws = book["Gantt"]
    rows = _gantt_rows(ws)

    def at(key: str) -> tuple:
        r = rows.get(key)
        vals = [ws.cell(r, c).value for c in range(4, 8)] if r else [None] * 4
        return tuple(v.date() if isinstance(v, _dt.datetime) else v for v in vals)

    d = lambda dd, mm: _dt.date(2026, mm, dd)
    for key, expected in (
            ("INC-101", ("In corso", d(1, 10), d(16, 10), 28)),
            ("CHG-302", ("In corso dal 06/10", d(7, 10), d(12, 10), 9)),
            ("OD-201", ("Bloccata da elenco delle competenze dei tecnici (ufficio operativo), "
                        "iniziata il 06/10", d(12, 10), d(16, 10), 9)),
            ("INC-103", ("Da iniziare, condizionale", d(20, 10), d(22, 10), 4)),
            ("KI-021", ("Da iniziare", d(22, 10), d(23, 10), 1.5)),
            ("Prova con la squadra pilota", ("A rischio", d(14, 10), d(14, 10), None)),
            ("Primo avviso inviato a un cliente", ("In tempo", d(23, 10), d(23, 10), None)),
            ("App in prova sul campo", ("Incompleta", d(30, 10), d(30, 10), None))):
        if at(key) != expected:
            problems.append(f"{key} reads {at(key)} in the Gantt, not {expected}")
    if not str(ws.cell(rows.get("INC-104", 1), 4).value).startswith("Da scomporre"):
        problems.append("an increment with nothing under it is not «Da scomporre»")
    r101 = rows.get("INC-101")
    if r101 and not (str(formulas.cell(r101, 5).value).startswith("=IF(COUNT(")
                     and "MAX(" in str(formulas.cell(r101, 6).value)
                     and str(formulas.cell(r101, 7).value).startswith("=SUM(")):
        problems.append("the row of an increment is not the MIN, MAX and SUM of what composes it")
    if "t1" not in rows or rows["t1"] > rows.get("INC-101", 0) \
            or ws.cell(rows["t1"], 2).value != "Il calendario delle squadre":
        problems.append("the phases do not head their increments with the name of the roadmap")
    stones = [ws.cell(r, 2).value for r in range(8, 13)]
    if stones != ["Consegna concordata", "Consegna prevista", "Prova con la squadra pilota",
                  "Primo avviso inviato a un cliente", "App in prova sul campo"]:
        problems.append(f"the milestones of the phases do not head the sheet by date: {stones}")
    for line in ("Consegna prevista oggi: il 23/10",
                 "Non conta 1 incremento ancora da scomporre (vedi Gantt)",
                 "Milestone «Prova con la squadra pilota» (tappa t1, 14/10): a rischio, "
                 "la tappa finisce il 16/10",
                 "Da scomporre, fuori dalla stima: INC-104"):
        if line not in out:
            problems.append(f"the conversation does not say {line!r}")
    return problems


@check("with `gantt: {done: false, from: week}` the Gantt shows the open work from this Monday, "
       "folded by stage")
def _digest_folds_the_gantt():
    # The readers of a roadmap in stages found the Gantt too heavy: every item closed since the
    # first day, every component in view. With the two preferences the same day of the fixture
    # loses DEC-210 and CHG-301, which are done, starts on Monday 05/10 instead of 28/09, and
    # opens on the stages and the milestones alone. INC-101 keeps its status, «in corso», which
    # its closed half decides, but is worth only what is left, CHG-302 and OD-201, 9 + 9.
    import datetime as _dt
    import openpyxl
    with tempfile.TemporaryDirectory() as tmp:
        root = _stages_copy(tmp)
        if root is None:
            return ["the stages fixture could not be built"]
        path = root / "_meta" / "digest" / "state-borea.yaml"
        data = yaml.safe_load(path.read_text(encoding="utf-8"))
        data["gantt"] = {"done": False, "from": "week"}
        path.write_text(yaml.safe_dump(data, allow_unicode=True, sort_keys=False),
                        encoding="utf-8")
        rc, out = _borea(root, "--date", "2026-10-07", "--now", "2026-10-07T08:30:00+02:00")
        xlsx = root / "_meta/digest/DIG-002-borea-2026-10-07/DIG-002-borea-2026-10-07.xlsx"
        if rc != 0:
            return [f"the digest of 07/10 was refused: {out.strip()[:400]}"]
        ws = openpyxl.load_workbook(xlsx, data_only=True)["Gantt"]
        formulas = openpyxl.load_workbook(xlsx)["Gantt"]
    problems = []
    rows = _gantt_rows(ws)
    closed = [k for k, r in rows.items() if ws.cell(r, 4).value == "Chiusa"]
    if closed or "DEC-210" in rows or "CHG-301" in rows:
        problems.append(f"items closed are still on the Gantt: {closed}")
    h6 = ws["H6"].value
    if not h6 or h6.date() != _dt.date(2026, 10, 5) or "WEEKDAY(" not in str(formulas["H6"].value):
        problems.append(f"the first column is {h6}, {formulas['H6'].value!r}, not the Monday "
                        "of the digest's week by formula")
    t1 = rows.get("t1")
    got = tuple(v.date() if isinstance(v, _dt.datetime) else v
                for v in (ws.cell(t1, c).value for c in (5, 6, 7))) if t1 else None
    if got != (_dt.date(2026, 10, 7), _dt.date(2026, 10, 16), 18):
        problems.append(f"the row of stage t1 reads {got}, not from 07/10 to 16/10 for 18 hours")
    elif not (str(formulas.cell(t1, 5).value).startswith("=IF(COUNT(")
              and "MAX(" in str(formulas.cell(t1, 6).value)
              and str(formulas.cell(t1, 7).value).startswith("=G")):
        problems.append("the first day, the last day and the hours of a stage are not formulas")
    if not str(ws.cell(rows.get("INC-104", 1), 4).value).startswith("Da scomporre"):
        problems.append("an increment still to break down left with the closed work")

    def level(key: str) -> tuple:
        dim = formulas.row_dimensions[rows[key]] if key in rows else None
        return (dim.outlineLevel, bool(dim.hidden)) if dim else None

    for keys, want in ((("t1", "t2", "t3", "Fuori dalle fasi"), (0, False)),
                       (("INC-101", "INC-102", "INC-103", "INC-104", "KI-021"), (1, True)),
                       (("CHG-302", "OD-201", "OD-202", "CHG-303", "CHG-304"), (2, True))):
        for key in keys:
            if level(key) != want:
                problems.append(f"{key} is at level and hidden {level(key)}, not {want}")
    if t1 and not formulas.row_dimensions[t1].collapsed:
        problems.append("a stage does not open folded")
    if formulas.sheet_properties.outlinePr is None \
            or formulas.sheet_properties.outlinePr.summaryBelow is not False:
        problems.append("the «+» of a group is not on the row of the stage, above its rows")
    # Excel opens and closes no group on a protected sheet.
    if formulas.protection.sheet:
        problems.append("the Gantt that folds is protected, so its groups do not open")
    return problems


def _stages_rendered(tmp: str, change=None, roadmap=()) -> tuple[Path | None, str]:
    """The workbook of 07/10 of the stages fixture, as the state file declares it, or as
    `change` edits it; `roadmap` is a text of the roadmap and what replaces it."""
    root = _stages_copy(tmp)
    if root is None:
        return None, "the stages fixture could not be built"
    if roadmap:
        _edit(root / "products" / "borea" / "RMP.md", *roadmap)
    if change:
        path = root / "_meta" / "digest" / "state-borea.yaml"
        data = yaml.safe_load(path.read_text(encoding="utf-8"))
        change(data)
        path.write_text(yaml.safe_dump(data, allow_unicode=True, sort_keys=False),
                        encoding="utf-8")
    rc, out = _borea(root, "--date", "2026-10-07", "--now", "2026-10-07T08:30:00+02:00")
    xlsx = root / "_meta/digest/DIG-002-borea-2026-10-07/DIG-002-borea-2026-10-07.xlsx"
    return (xlsx if rc == 0 else None), out


@check("«La prossima settimana» lists the items the Gantt places in the five working days "
       "after the digest's")
def _digest_lists_the_next_week():
    # The digest of Wednesday 07/10 looks at Thursday 08/10 to Wednesday 14/10. On the Gantt
    # CHG-302 runs from 07/10 to 12/10 and OD-201 from 12/10 to 16/10, so both touch it;
    # OD-202 starts on 16/10, after it, and is not there. CHG-302 is in the plan the person
    # declared, for 07/10 and for 08/10, and keeps what the plan says; OD-201 is not, and says
    # its days and its hours.
    import datetime as _dt
    import openpyxl
    with tempfile.TemporaryDirectory() as tmp:
        xlsx, out = _stages_rendered(tmp)
        if xlsx is None:
            return [f"the digest of 07/10 was refused: {out.strip()[:400]}"]
        book = openpyxl.load_workbook(xlsx, data_only=True)
    problems = []
    at, gantt = book["Attività"], book["Gantt"]
    if at["A3"].value != "1. La prossima settimana":
        problems.append(f"the first table of Attività is {at['A3'].value!r}")
    listed, r = [], 5
    while at.cell(r, 1).value:
        listed.append((at.cell(r, 1).value, at.cell(r, 8).value))
        r += 1
    lo, hi = _dt.date(2026, 10, 8), _dt.date(2026, 10, 14)
    placed = []
    for r in range(7, gantt.max_row + 1):
        start, end = gantt.cell(r, 5).value, gantt.cell(r, 6).value
        if str(gantt.cell(r, 1).value).startswith("INC-"):
            continue                            # an increment is what composes it, not an item
        if str(gantt.cell(r, 4).value).startswith(("Da iniziare", "In corso", "Rallentata",
                                                    "Bloccata")) \
                and isinstance(start, _dt.datetime) and isinstance(end, _dt.datetime) \
                and start.date() <= hi and end.date() >= lo:
            placed.append(gantt.cell(r, 1).value)
    ids = [i for i, _ in listed]
    if ids != placed or ids != ["CHG-302", "OD-201"]:
        problems.append(f"the next week lists {ids}; the Gantt places {placed} in 08/10-14/10, "
                        "and by hand it is CHG-302 and OD-201")
    if "OD-202" in ids:
        problems.append("OD-202, which starts on 16/10, is in the next week")
    status = dict(listed)
    if status.get("OD-201") != "dal 12/10 al 16/10, 9 ore":
        problems.append(f"an item not in the plan says {status.get('OD-201')!r}")
    if status.get("CHG-302") != ("In programma il 07/10, 4 ore: prosegue, chiusura prevista il "
                                 "08/10; il 08/10, 4 ore: chiusura"):
        problems.append(f"an item in the plan says {status.get('CHG-302')!r}, not what the "
                        "plan says of it")
    return problems


@check("the row of a group of the Gantt carries the hours of the rows it holds")
def _digest_sums_the_group():
    # «Fuori dalle fasi» holds KI-021 alone, 1,5 hours for the middle of S; the phases have carried
    # the sum of their increments since 4.3.0, and a group the sum of its rows since 4.4.0.
    import openpyxl
    with tempfile.TemporaryDirectory() as tmp:
        xlsx, out = _stages_rendered(tmp)
        if xlsx is None:
            return [f"the digest of 07/10 was refused: {out.strip()[:400]}"]
        values = openpyxl.load_workbook(xlsx, data_only=True)["Gantt"]
        formulas = openpyxl.load_workbook(xlsx)["Gantt"]
    head = next((r for r in range(7, values.max_row + 1)
                 if values.cell(r, 2).value == "Fuori dalle fasi"), None)
    if head is None:
        return ["the Gantt of the stages fixture has no «Fuori dalle fasi»"]
    held, r = [], head + 1
    while values.cell(r, 1).value:
        held.append(values.cell(r, 7).value or 0)
        r += 1
    got, formula = values.cell(head, 7).value, str(formulas.cell(head, 7).value)
    if not held or got != sum(held) or got != 1.5 \
            or formula != f"=SUM(G{head + 1}:G{head + len(held)})":
        return [f"«Fuori dalle fasi» carries {got!r}, {formula!r}; its rows add up to "
                f"{sum(held)}"]
    return []


@check("an item started before the digest is placed by the queue, and says since when")
def _digest_places_a_started_item():
    # A bar from the first day worked to the item's place in the queue covered the days of
    # every item before it, and the rows of its increment and its stage began with it. CHG-303
    # was worked one hour on 06/10 and comes after CHG-302, OD-201 and OD-202, 9 + 9 + 1,5 =
    # 19,5 hours at 2,4 a day, 8,1 days: it starts on the ninth working day after 06/10, 19/10,
    # and its 4 hours end on the tenth, 20/10. INC-102 and stage t2 start with OD-202 on 16/10,
    # and the next week, 08/10 to 14/10, does not list it.
    import datetime as _dt
    import openpyxl
    with tempfile.TemporaryDirectory() as tmp:
        xlsx, out = _stages_rendered(tmp)
        if xlsx is None:
            return [f"the digest of 07/10 was refused: {out.strip()[:400]}"]
        book = openpyxl.load_workbook(xlsx, data_only=True)
        formulas = openpyxl.load_workbook(xlsx)["Gantt"]
    ws, problems = book["Gantt"], []
    rows = _gantt_rows(ws)

    def at(key: str) -> tuple:
        r = rows.get(key)
        vals = [ws.cell(r, c).value for c in (4, 5, 6, 7)] if r else [None] * 4
        return tuple(v.date() if isinstance(v, _dt.datetime) else v for v in vals)

    d = lambda dd, mm: _dt.date(2026, mm, dd)
    if at("CHG-303") != ("In corso dal 06/10", d(19, 10), d(20, 10), 4):
        problems.append(f"CHG-303 reads {at('CHG-303')}, not «In corso dal 06/10» from 19/10 "
                        "to 20/10 for 4 hours")
    r = rows.get("CHG-303")
    if r and not str(formulas.cell(r, 5).value).startswith("=IF(Riepilogo!$G$12<=0"):
        problems.append(f"the first day of a started item is {formulas.cell(r, 5).value!r}, not "
                        "the formula on its place in the queue")
    for key in ("INC-102", "t2"):
        if at(key)[1] != d(16, 10):
            problems.append(f"{key} starts on {at(key)[1]}, not with its first item in the "
                            "queue on 16/10")
    week = [book["Attività"].cell(r, 1).value for r in range(5, 9)]
    if "CHG-303" in week:
        problems.append(f"CHG-303, placed from 19/10, is in the next week: {week}")
    rules = [x.formula[0] for c in formulas.conditional_formatting
             if str(c.sqref).startswith("H8:") for x in c.rules]
    if not any('LEFT($D8,8)="In corso"' in f for f in rules):
        problems.append("a status that says since when loses its bar: the rules match it whole")
    return problems


@check("the Gantt is the plan its readers agreed on: milestones on top, an owner, days then weeks")
def _digest_draws_the_plan():
    # The plan the product owner shared puts the milestones first, the two deliveries and then
    # the others by date, says who follows each row, calls the stages phases by the code and
    # the name the roadmap gives them, shows day by day the four weeks before the digest's and
    # the four after it and the rest by week, and says what a blocked item waits for. On the
    # stages fixture, with the project started on 17/08 instead of 28/09:
    # - the days run from 07/09, four weeks before Monday 05/10, to Friday 06/11; the weeks of
    #   17/08, 24/08 and 31/08 come before them, those of 09/11 to 14/12 after, to the
    #   milestone of the release declared on 15/12, marked in the week of 14/12;
    # - CHG-302 and OD-202 have their own owner, the other items `gantt.owner`; INC-101 the one
    #   declared on it, INC-103 that of CHG-304, its only item, and INC-102 none, since OD-202
    #   and CHG-303 have different ones; the agreed delivery has `release.owner`, the expected
    #   one `gantt.owner`, the milestone of t1 its own;
    # - the milestone of t3, moved to 16/10, the agreed delivery, is that delivery: no row;
    #   those of t1 and t2 follow the deliveries, by date, before the release's of 15/12, and
    #   the first says «A rischio» by formula on the rows of its phase;
    # - the roadmap names t2 «F2 · Gli avvisi ai clienti»: code F2, name «Gli avvisi ai clienti»;
    # - OD-201 waits for two things, needed by 09/10 and 14/10: the list of the skills, asked
    #   now up to a comma, and a VPN, whose name keeps its capitals.
    import datetime as _dt
    import openpyxl

    def plan(s: dict) -> None:
        s["release"].update(start="2026-08-17", owner="Direzione")
        s["gantt"] = {"owner": "Analista"}
        s["items"]["CHG-302"]["owner"] = "Team tecnico"
        s["items"]["OD-202"]["owner"] = "Team funzionale"
        s["items"]["INC-101"]["owner"] = "Sviluppo"
        for m in s["milestones"]:
            if m["stage"] == "t1":
                m["owner"] = "PM"
            if m["stage"] == "t3":
                m["date"] = "2026-10-16"
        s["milestones"].append({"name": "Secondo rilascio", "date": "2026-12-15",
                                "owner": "Direzione"})
        s["waits"][0]["what"] = "Elenco delle competenze dei tecnici, per squadra e per turno"
        s["waits"].append({"what": "VPN di collaudo, con le credenziali", "owner": "team sistemi",
                           "asked": "2026-10-05", "needed_by": "2026-10-14",
                           "without": "nessuna prova dal campo.", "missing": "manca la VPN",
                           "blocks": ["OD-201"]})

    with tempfile.TemporaryDirectory() as tmp:
        xlsx, out = _stages_rendered(tmp, plan, ("name: Gli avvisi ai clienti",
                                                 "name: F2 · Gli avvisi ai clienti"))
        if xlsx is None:
            return [f"the digest of 07/10 was refused: {out.strip()[:400]}"]
        book = openpyxl.load_workbook(xlsx, data_only=True)
        formulas = openpyxl.load_workbook(xlsx)["Gantt"]
    ws, problems = book["Gantt"], []
    head = [ws.cell(6, c).value for c in range(1, 8)]
    if head != ["ID", "Attività", "Owner", "Stato", "Inizio", "Fine", "Ore"]:
        problems.append(f"the columns are {head}")
    rows = _gantt_rows(ws)
    top = [ws.cell(r, 2).value for r in range(8, 13)]
    if ws["A7"].value != "PIETRE MILIARI" or top != [
            "Consegna concordata", "Consegna prevista", "Prova con la squadra pilota",
            "Primo avviso inviato a un cliente", "Secondo rilascio"]:
        problems.append(f"the milestones do not head the sheet, deliveries first: {top}")
    if ws["D9"].value != "Ritardo 5 gg lav." or "Riepilogo!$J$12" not in str(formulas["D9"].value):
        problems.append(f"the expected delivery says {ws['D9'].value!r}, not its delay by formula")
    t1 = rows.get("t1", 0)
    verdict = str(formulas["D10"].value)
    if ws["D10"].value != "A rischio" or f"D{t1 + 1}:" not in verdict:
        problems.append(f"the milestone of t1 says {ws['D10'].value!r}, {verdict[:90]!r}, not "
                        "«A rischio» by formula on the rows of its phase")
    for key, who in (("Consegna concordata", "Direzione"), ("Consegna prevista", "Analista"),
                     ("Prova con la squadra pilota", "PM"), ("Secondo rilascio", "Direzione"),
                     ("CHG-302", "Team tecnico"), ("OD-202", "Team funzionale"),
                     ("CHG-303", "Analista"), ("INC-101", "Sviluppo"), ("INC-102", None),
                     ("INC-103", "Analista"), ("t1", None)):
        got = ws.cell(rows[key], 3).value if key in rows else "missing"
        if got != who:
            problems.append(f"{key} has owner {got!r}, not {who!r}")
    if "F2" not in rows or ws.cell(rows["F2"], 2).value != "Gli avvisi ai clienti":
        problems.append("a phase named «F2 · …» does not read F2 and its name")
    blocked = ("Bloccata da elenco delle competenze dei tecnici (ufficio operativo); VPN di "
               "collaudo (team sistemi), iniziata il 06/10")
    if ws.cell(rows.get("OD-201", 1), 4).value != blocked:
        problems.append(f"a blocked item says {ws.cell(rows.get('OD-201', 1), 4).value!r}, not "
                        f"{blocked!r}")
    before, days, after = [], [], []
    for c in range(8, ws.max_column + 1):
        d = ws.cell(6, c).value
        if isinstance(d, _dt.datetime):
            (days if ws.cell(5, c).value != "sett." else before if not days else
             after).append(d.date())
    weeks = lambda first, n: [first + _dt.timedelta(days=7 * k) for k in range(n)]
    if before != weeks(_dt.date(2026, 8, 17), 3) or len(days) != 45 \
            or days[:1] + days[-1:] != [_dt.date(2026, 9, 7), _dt.date(2026, 11, 6)] \
            or after != weeks(_dt.date(2026, 11, 9), 6):
        problems.append(f"the calendar has weeks {before}, days {days[:1]}..{days[-1:]} "
                        f"({len(days)}) and weeks {after}")
    if "MAX(" not in str(formulas.cell(6, 11).value):
        problems.append(f"the first day shown is not the later of four weeks before and the "
                        f"start: {formulas.cell(6, 11).value!r}")
    r = rows.get("Secondo rilascio")
    marks = [ws.cell(6, c).value.date() for c in range(8, ws.max_column + 1)
             if r and ws.cell(r, c).value == "◆"]
    if marks != [_dt.date(2026, 12, 14)]:
        problems.append(f"a milestone on 15/12 is marked on {marks}, not in the week of 14/12")
    if "App in prova sul campo" in rows:
        problems.append("the milestone of a phase on the agreed delivery has a row of its own")
    words = " ".join(str(ws.cell(r, c).value) for r in range(1, ws.max_row + 1)
                     for c in range(1, 8) if ws.cell(r, c).value is not None)
    if re.search(r"\btapp[ae]\b", words, re.I) or "Fuori dalle fasi" not in words:
        problems.append("the sheet still calls the stages «tappe»")
    if re.search(r"massim|\bmedi[oa]\b", words, re.I):
        problems.append("the Gantt says how many hours a size counts")
    rules = [x.formula[0] for c in formulas.conditional_formatting
             if str(c.sqref).startswith("H8:") for x in c.rules]
    if not any("WORKDAY($E8+1,-1)" in f and "WORKDAY($F8+1,-1)" in f for f in rules):
        problems.append("an item closed on a Saturday has no bar on the Friday before it")
    return problems


@check("the Gantt shows what was closed lately, the phases not yet broken down, and what others "
       "owe")
def _digest_shows_the_outline():
    # On the stages fixture of 07/10, with the three preferences of 4.5.0 that change what the
    # Gantt holds:
    # - `done: recent`: the items closed since the first day shown one by one, 28/09, the
    #   week the project started: DEC-210 and CHG-301, here closed on Saturday 03/10, in a
    #   group «Chiuso dal 28/09» at the head of the work, folded; the phases hold only what is
    #   open. With `from: week` the first day is 05/10, and nothing closed is shown;
    # - `outline` for t3 from 02/11 to 27/11: the phase reads «Da dettagliare» on those days,
    #   with no hours and nothing under it, so INC-104 is not shown, and the calendar runs by
    #   week to 23/11; its milestone on 30/10 stays on top, «Incompleta»;
    # - `requests` for the operations office and the systems team: their two requests under
    #   the milestones, the access asked on 01/10 and needed by 05/10, «Scaduta da 2 giorni»,
    #   then the list of skills needed by 09/10, «Entro 2 giorni».
    import datetime as _dt
    import openpyxl

    def outline(s: dict) -> None:
        s["gantt"] = {"done": "recent", "requests": ["team sistemi", "ufficio operativo"],
                      "outline": {"t3": {"from": "2026-11-02", "to": "2026-11-27"}}}
        s["items"]["CHG-301"]["closed_on"] = "2026-10-03"
        s["waits"][1]["needed_by"] = "2026-10-05"

    with tempfile.TemporaryDirectory() as tmp:
        xlsx, out = _stages_rendered(tmp, outline)
        if xlsx is None:
            return [f"the digest of 07/10 was refused: {out.strip()[:400]}"]
        ws = openpyxl.load_workbook(xlsx, data_only=True)["Gantt"]
        formulas = openpyxl.load_workbook(xlsx)["Gantt"]
    with tempfile.TemporaryDirectory() as tmp:
        later, out2 = _stages_rendered(tmp, lambda s: (
            outline(s), s["gantt"].update({"from": "week"})))
        weekly = openpyxl.load_workbook(later, data_only=True)["Gantt"] if later else None
    problems = []
    rows = _gantt_rows(ws)
    col = lambda r: [ws.cell(r, c).value for c in range(1, 8)]
    d = lambda dd, mm: _dt.datetime(2026, mm, dd)
    head = next((r for r in range(7, ws.max_row + 1)
                 if ws.cell(r, 1).value == "COSA SERVE DA ALTRI"), None)
    asked = [col(r)[1:6] for r in (head + 1, head + 2)] if head else []
    if asked != [["Accesso al server di collaudo", "team sistemi", "Scaduta da 2 giorni",
                  d(1, 10), d(5, 10)],
                 ["Elenco delle competenze dei tecnici", "ufficio operativo", "Entro 2 giorni",
                  d(29, 9), d(9, 10)]]:
        problems.append(f"the requests to others read {asked}")
    elif head < rows.get("Consegna prevista", 99) \
            or "Scaduta da" not in str(formulas.cell(head + 1, 4).value):
        problems.append("the requests are not under the milestones, with their status by formula")
    if head:
        rules = [x.formula[0] for c in formulas.conditional_formatting
                 if str(c.sqref).startswith(f"H{head + 1}:") for x in c.rules]
        if len(rules) != 2 or "$F" not in rules[0] or "WORKDAY($E" not in rules[1]:
            problems.append(f"a request has no bar from its asking to its deadline: {rules}")
    group = rows.get("Chiuso dal 28/09")
    held = [ws.cell(r, 1).value for r in range(group + 1, group + 3)] if group else []
    if held != ["DEC-210", "CHG-301"] or (rows.get("t1") or 0) < group:
        problems.append(f"the items closed lately are not in a group at the head of the work: "
                        f"{held}")
    else:
        dims = [formulas.row_dimensions[r] for r in (group, group + 1)]
        if (dims[0].outlineLevel, bool(dims[0].collapsed), dims[1].outlineLevel,
                bool(dims[1].hidden)) != (0, True, 1, True):
            problems.append("the group of the items closed lately does not open folded")
    if sum(1 for r in rows.values() if ws.cell(r, 4).value == "Chiusa") != 2:
        problems.append("an item closed shows twice, or in its phase")
    t3 = rows.get("t3")
    if not t3 or col(t3)[3:7] != ["Da dettagliare", d(2, 11), d(27, 11), None] \
            or "INC-104" in rows:
        problems.append(f"the phase not broken down reads {col(t3) if t3 else None}")
    if ws.cell(rows.get("App in prova sul campo", 1), 4).value != "Incompleta":
        problems.append("the milestone of a phase not broken down is not «Incompleta»")
    last = [ws.cell(6, c).value for c in range(8, ws.max_column + 1)][-1]
    if last != d(23, 11):
        problems.append(f"the calendar ends on {last}, not on the week of 23/11 the outline "
                        "reaches")
    if weekly is None:
        problems.append(f"the digest with `from: week` was refused: {out2.strip()[:300]}")
    elif any(weekly.cell(r, 4).value == "Chiusa" for r in range(8, weekly.max_row + 1)) \
            or any(str(weekly.cell(r, 2).value).startswith("Chiuso dal")
                   for r in range(8, weekly.max_row + 1)):
        problems.append("items closed before the first day shown are on the Gantt")
    return problems


@check("the day by day of Ore lavorate folds by month and by week, and its sums stay the days'")
def _digest_folds_the_hours():
    # The stages fixture of 07/10 has the days from Monday 28/09 to Tuesday 06/10: September
    # holds the week of 28/09 up to 30/09, October the rest of it, 01/10 to 04/10, and the week
    # of 05/10. A month and a week are SUBTOTALs of the rows under them and have no date, so the
    # sums by date at the head of the sheet, and the total since the start, count each day once.
    import datetime as _dt
    import openpyxl
    with tempfile.TemporaryDirectory() as tmp:
        xlsx, out = _stages_rendered(tmp)
        if xlsx is None:
            return [f"the digest of 07/10 was refused: {out.strip()[:400]}"]
        ws = openpyxl.load_workbook(xlsx, data_only=True)["Ore lavorate"]
        formulas = openpyxl.load_workbook(xlsx)["Ore lavorate"]
    problems = []
    shape, days = [], {}
    for r in range(30, ws.max_row + 1):
        a, b = ws.cell(r, 1).value, ws.cell(r, 2).value
        if isinstance(b, _dt.datetime):
            days[r] = b.date()
        elif a and not str(a).startswith("Mesi e settimane"):
            shape.append((r, a, formulas.row_dimensions[r].outlineLevel))
    if [(a, lvl) for _, a, lvl in shape] != [
            ("Settembre 2026", 0), ("settimana dal 28/09 al 30/09", 1), ("Ottobre 2026", 0),
            ("settimana dal 01/10 al 04/10", 1), ("settimana dal 05/10 al 06/10", 1)]:
        problems.append(f"the months and weeks are {[(a, lvl) for _, a, lvl in shape]}")
    for r, a, lvl in shape:
        under = [x for x in days if x > r and all(not (r < y < x) for y, _, l in shape
                                                  if l <= lvl)]
        want = sum(ws.cell(x, 3).value for x in under)
        if ws.cell(r, 3).value != want or ws.cell(r, 2).value is not None \
                or not str(formulas.cell(r, 3).value).startswith("=SUBTOTAL(9,"):
            problems.append(f"«{a}» carries {ws.cell(r, 3).value}, "
                            f"{formulas.cell(r, 3).value!r}; its days add up to {want}")
    if any(not formulas.row_dimensions[r].hidden or formulas.row_dimensions[r].outlineLevel != 2
           for r in days):
        problems.append("a day is not one level under its week, closed when the file opens")
    total = sum(ws.cell(r, 3).value for r in days)
    outside = sum(ws.cell(r, c).value for r in days for c in range(9, 14))
    if ws["D8"].value != total or ws["B25"].value != outside:
        problems.append(f"the sums since the start count {ws['D8'].value} and "
                        f"{ws['B25'].value} outside the product, the days {total} and {outside}")
    if not str(formulas["B20"].value).startswith("=SUBTOTAL(9,"):
        problems.append(f"the hours outside the product since the start are "
                        f"{formulas['B20'].value!r}, which counts the months and the weeks too")
    if formulas.protection.sheet:
        problems.append("the sheet that folds is protected, so its groups do not open")
    return problems


@check("a stage milestone needs its date, and the stages place what `order.todo` no longer must")
def _digest_stage_refusals():
    problems = []
    cases = [
        ("a stage milestone with no date",
         lambda s: s["milestones"].pop(2), 1, "has no date"),
        ("a date for a stage the roadmap does not have",
         lambda s: s["milestones"].append({"stage": "t9", "date": "2026-11-02"}), 1,
         "no such stage"),
        ("a pinned item that is not open",
         lambda s: s["order"].__setitem__("first", ["CHG-301"]), 1, "`order.first`"),
        ("an item the stages place still listed by hand",
         lambda s: s["order"]["todo"].insert(0, "CHG-302"), 0, "now place"),
        ("a phase to outline that the roadmap does not have",
         lambda s: s.__setitem__("gantt", {"outline": {"t9": {"from": "2026-11-02",
                                                               "to": "2026-11-27"}}}),
         1, "`gantt.outline` names 't9', and the roadmap has no such stage"),
        ("a phase to outline that ends before it starts",
         lambda s: s.__setitem__("gantt", {"outline": {"t3": {"from": "2026-11-27",
                                                               "to": "2026-11-02"}}}),
         1, "`gantt.outline.t3` ends on 2026-11-02, before it starts on 2026-11-27"),
        ("a phase to outline without its last day",
         lambda s: s.__setitem__("gantt", {"outline": {"t3": {"from": "2026-11-02"}}}),
         1, "`gantt.outline.t3` is {from, to}"),
        ("a phase to outline with work to do in it",
         lambda s: s.__setitem__("gantt", {"outline": {"t1": {"from": "2026-11-02",
                                                               "to": "2026-11-27"}}}),
         1, "`gantt.outline.t1` draws a phase not yet broken down, and CHG-302, OD-201 are"),
    ]
    for what, change, rc_want, says in cases:
        with tempfile.TemporaryDirectory() as tmp:
            root = _stages_copy(tmp)
            if root is None:
                return ["the stages fixture could not be built"]
            path = root / "_meta" / "digest" / "state-borea.yaml"
            data = yaml.safe_load(path.read_text(encoding="utf-8"))
            change(data)
            path.write_text(yaml.safe_dump(data, allow_unicode=True, sort_keys=False),
                            encoding="utf-8")
            args = ("--check",) if rc_want else ("--inventory",)
            rc, out = _borea(root, "--date", "2026-10-07", *args)
            if rc != rc_want or says not in out:
                problems.append(f"{what}: exit {rc}, and it does not say {says!r}: "
                                f"{out.strip()[:300]}")
    with tempfile.TemporaryDirectory() as tmp:
        root = _stages_copy(tmp)
        rmp = root / "products" / "borea" / "RMP.md"
        rmp.write_text(rmp.read_text(encoding="utf-8").replace(
            "increments: [INC-102, INC-103]", "increments: [INC-103, INC-102]"),
            encoding="utf-8")
        env = {**os.environ, **_GIT_ID, "GIT_AUTHOR_DATE": "2026-10-06T16:00:00+02:00",
               "GIT_COMMITTER_DATE": "2026-10-06T16:00:00+02:00"}
        subprocess.run(["git", "commit", "-qam", "stages reordered"], cwd=root, env=env)
        rc, out = _borea(root, "--date", "2026-10-07", "--inventory")
        inv = yaml.safe_load(out) if rc == 0 else {}
        worked = (inv.get("worked_by_day") or {}).get("2026-10-06") or {}
        if "INC-102" in worked or "INC-103" in worked:
            problems.append("reordering the stages was proposed as work on the increments: "
                            f"{sorted(worked)}")
        if inv.get("increments_to_break_down") != ["INC-104"]:
            problems.append(f"the inventory does not name the increment still to break down: "
                            f"{inv.get('increments_to_break_down')}")
    return problems


@check("the digest fixtures validate with nothing to report")
def _digest_fixtures_validate():
    src = built("digest")
    if src is None:
        return ["the digest fixture could not be built"]
    stages = built("digest-stages")
    if stages is None:
        return ["the stages fixture could not be built"]
    problems = []
    for name, where in (("atlas", src / "atlas"), ("atlas-asks", src / "atlas-asks"),
                        ("borea", stages / "borea")):
        r = subprocess.run([sys.executable, str(VALIDATE), "--root", str(where), "--json",
                            "--stale-days", "36500"], capture_output=True, text=True)
        try:
            found = [(f["code"], f["path"]) for f in json.loads(r.stdout)["findings"]
                     if f["level"] in ("error", "warn")]
        except (ValueError, KeyError):
            problems.append(f"{name}: the validator did not run: {r.stderr.strip()[-200:]}")
            continue
        if found:
            problems.append(f"{name}: {found}")
    return problems


@check("a count of findings written in prose is the count the fixture produces")
def _fixture_counts_are_measured():
    # Three times the number of findings in `audit/dirty-repo` was written into prose and
    # three times it went stale: eleven, then thirteen against a fixture producing twelve,
    # then thirteen again on the day `REG006` was added, when thirteen was already the count
    # before it. Every one of those was written in a file whose subject is documents going
    # out of date, and nothing reported any of them.
    #
    # The count stays in prose rather than being deleted, because "four errors and fourteen
    # warnings, by construction" is what tells a reader the fixture is planted and not
    # merely dirty, and a pointer to a command they have to run says much less. What was
    # missing is the thing that makes a number in prose safe: something that reads it.
    #
    # The fixture is the authority, not this file. It is built into a temporary directory
    # and measured, so the check cannot be cleared by editing a number here.
    words = {"zero": 0, "one": 1, "two": 2, "three": 3, "four": 4, "five": 5, "six": 6,
             "seven": 7, "eight": 8, "nine": 9, "ten": 10, "eleven": 11, "twelve": 12,
             "thirteen": 13, "fourteen": 14, "fifteen": 15, "sixteen": 16,
             "seventeen": 17, "eighteen": 18, "nineteen": 19, "twenty": 20}
    # Past twenty the fixture kept growing and the vocabulary did not, so the sentence
    # stopped matching and this check reported that nobody states the count any more --
    # correct, and for the wrong reason. The compound forms first in the alternation, or
    # `twenty` matches the first half of `twenty-one` and the number reads as 20.
    words.update({f"twenty-{w}": 20 + n for w, n in
                  (("one", 1), ("two", 2), ("three", 3), ("four", 4), ("five", 5),
                   ("six", 6), ("seven", 7), ("eight", 8), ("nine", 9))})
    words["thirty"] = 30
    ordered = sorted(words, key=len, reverse=True)
    stated = re.compile(r"\b(" + "|".join(ordered) + r")\s+errors?\s+and\s+"
                        r"(" + "|".join(ordered) + r")\s+warnings?\b", re.I)

    problems = []
    with tempfile.TemporaryDirectory() as tmp:
        gen = ROOT / "evals" / "fixtures" / "generators" / "audit.py"
        r = subprocess.run([sys.executable, str(gen), tmp],
                           capture_output=True, text=True)
        if r.returncode != 0:
            return [f"the audit fixture would not build, so its documented count cannot be "
                    f"checked: {r.stderr.strip() or r.stdout.strip()}"]

        measured = {}
        for name in ("dirty-repo", "clean-repo"):
            v = subprocess.run([sys.executable, str(VALIDATE), "--root",
                                str(Path(tmp) / name), "--json"],
                               capture_output=True, text=True)
            try:
                d = json.loads(v.stdout)
            except json.JSONDecodeError:
                problems.append(f"the validator returned nothing readable on {name}")
                continue
            measured[name] = (d["errors"], d["warnings"])

        # The fixture that must report nothing. Its README says a checker finding something
        # here is inventing, which is a claim about output and belongs with the other one.
        if measured.get("clean-repo", (0, 0)) != (0, 0):
            e, w = measured["clean-repo"]
            problems.append(f"audit/clean-repo reports {e} error(s) and {w} warning(s). "
                            "Its whole description is that anything reported there is "
                            "invented, and a fixture that contradicts its own description "
                            "teaches you to read past its output")

        if "dirty-repo" not in measured:
            return problems
        want = measured["dirty-repo"]

        found = 0
        for rel in ("evals/fixtures/README.md", "evals/behaviour/audit/cases.yaml"):
            text = (ROOT / rel).read_text(encoding="utf-8")
            for m in stated.finditer(text):
                found += 1
                got = (words[m.group(1).lower()], words[m.group(2).lower()])
                line = text[:m.start()].count("\n") + 1
                if got != want:
                    problems.append(
                        f"{rel}:{line} says {got[0]} errors and {got[1]} warnings; "
                        f"audit/dirty-repo produces {want[0]} and {want[1]}. Either the "
                        "fixture changed and the sentence did not, or the sentence was "
                        "wrong when it was written")
        if not found:
            problems.append(
                "no file states what audit/dirty-repo produces any more. The count is "
                "what tells a reader the fixture is planted rather than merely dirty; if "
                "it is gone on purpose, this check is what has to go with it")
    return problems


@check("a change set is checked against the contract that authorizes it")
def _pull_request_binding():
    # The four `PR` checks are the only ones that read something outside the repository,
    # and that is what this pins: they must stay silent without that context. A check that
    # fired on a plain `--root` run would report every desk in the project as unauthorized
    # work, and would be switched off within the day.
    def fm(**kw):
        return "---\n" + "\n".join(f"{k}: {v}" for k, v in kw.items()) + "\n---\n\n"

    chg = lambda status: fm(
        schema="framework/change-contract/v1", artifact_type="change-contract",
        lifecycle="immutable", status=status, id="CHG-001", owners="[o]",
        created="2026-01-01 09:00", icg="ICG-001", derives_from="[SIG-001]") + (
        "<!-- section: what-changes -->\n# 1\n"
        "<!-- section: what-must-not-change -->\n# 2\n"
        "<!-- section: how-we-know-it-worked -->\n# 3\n")

    icg = fm(schema="framework/impact-classification/v1",
             artifact_type="impact-classification", lifecycle="immutable",
             status="accepted", id="ICG-001", owners="[o]", created="2026-01-01 09:00",
             routing="\n  SIG-001: architecture",
             impacts="\n  SIG-001: [data]") + (
        "<!-- section: intake -->\n# 1\n<!-- section: classification -->\n# 2\n"
        "<!-- section: open-questions -->\n# 3\n")

    dc = fm(schema="framework/data-contract/v1", artifact_type="data-contract",
            lifecycle="living", status="active", id="DC-001", owners="[o]",
            created="2026-01-01 09:00", last_review="2026-01-01 09:00", version="1.0.0",
            products="[alpha]", consumers="[alpha]") + "# DC\n"

    def run(status, text=None, changed=None):
        with tempfile.TemporaryDirectory() as tmp:
            for rel, body in {"CHG-001.md": chg(status), "ICG-001.md": icg,
                              "products/alpha/DC-001.md": dc}.items():
                f = Path(tmp) / rel
                f.parent.mkdir(parents=True, exist_ok=True)
                f.write_text(body)
            extra = [] if text is None else ["--pr-text", text]
            if changed is not None:
                cf = Path(tmp) / "changed.txt"
                cf.write_text("\n".join(changed))
                extra += ["--changed-files", str(cf)]
            r = subprocess.run(
                [sys.executable, str(VALIDATE), "--root", tmp, "--json",
                 "--stale-days", "36500", *extra], capture_output=True, text=True)
            if r.returncode not in (0, 1) or not r.stdout.strip():
                return None
            return {f["code"] for f in json.loads(r.stdout)["findings"]
                    if f["code"].startswith("PR")}

    problems = []
    for want, args, what in [
        (set(), dict(status="approved"),
         "no pull request context at all, which must report nothing"),
        ({"PR001"}, dict(status="approved", text="tidy up the logging"),
         "a change set citing no contract"),
        (set(), dict(status="approved", text="tidy up\n\nno-chg: a typo in a comment"),
         "the declared exception, with its reason"),
        ({"PR001"}, dict(status="approved", text="tidy up\n\nno-chg:"),
         "the exception with no reason, which is the check deleted with extra steps"),
        ({"PR002"}, dict(status="approved", text="implements CHG-404"),
         "a contract that is not in the repository"),
        ({"PR003"}, dict(status="draft", text="implements CHG-001"),
         "implementing a draft"),
        ({"PR003"}, dict(status="rolled-back", text="implements CHG-001"),
         "citing a contract that was rolled back"),
        (set(), dict(status="approved", text="implements CHG-001"),
         "an approved contract with no diff supplied: PR004 cannot run and must not guess"),
        ({"PR004"}, dict(status="approved", text="implements CHG-001",
                         changed=["src/app.py"]),
         "a data impact whose change set touches no data contract"),
        (set(), dict(status="approved", text="implements CHG-001",
                     changed=["src/app.py", "products/alpha/DC-001.md"]),
         "the same change set once the contract moves with it"),
        (set(), dict(status="approved", text="implements CHG-001",
                     changed=["./products/alpha/DC-001.md"]),
         "a diff that writes its paths with a leading ./"),
    ]:
        got = run(**args)
        if got is None:
            problems.append(f"the validator crashed on {what}")
        elif got != want:
            problems.append(f"{what}: expected {sorted(want) or 'nothing'}, got "
                            f"{sorted(got) or 'nothing'}")
    return problems


@check("the migration tool reconstructs the version a project pins")
def _migration_is_executable():
    # `P-12` is a procedure nobody can execute unless the old validator can be got hold of.
    # It is not kept anywhere: it is rebuilt from this repository's history, at the commit
    # where the registry last declared the version the project pinned. So what this asserts
    # is the one thing that would silently stop being true -- that every version this
    # framework has ever declared is still reachable, and that the note explaining it is
    # still beside the number.
    migrate = _load(ROOT / "skills" / "audit" / "scripts" / "migrate.py", "migrate")
    registry = (ROOT / "schemas" / "artifact-types.yaml").read_text(encoding="utf-8")
    notes = {to for _, to, _ in migrate.migration_notes(registry)}
    problems = []

    # A check reports the same code on the same file many times -- one per entry of a
    # register, one per unresolved citation. Keyed on the pair alone, three findings and
    # five findings are the same key, and a version that made a check noisier on a file it
    # already reported would come out of the tool as "nothing new".
    seen = collections.Counter()
    keys = [migrate.key({"code": "REG005", "path": "OPEN.md"}, seen) for _ in range(3)]
    if len(set(keys)) != 3:
        problems.append("three findings with the same code and path collapse to "
                        f"{len(set(keys))} key(s): a check getting noisier would read as "
                        "nothing having changed")

    # The two the tool clears by writing a number rather than by anybody editing a
    # document. They are separated in the report and skipped by the adopt gate, and those
    # two have to be the same list: when they were not, the report told the reader to go
    # and fix something that `--adopt` was already ignoring.
    # A PROJECT THAT PINS HAS TO BE ABLE TO ADOPT. `FW003` fires the moment the framework
    # moves past the pin, it landed in NEW, `--adopt` refuses while NEW is non-empty, and the
    # only thing that clears `FW003` is `--adopt` writing the pin. Pinning made adopting
    # impossible, and the way out was deleting the pin -- the field's purpose undone by the
    # tool that exists to move it. Asserted end to end, because the deadlock was invisible in
    # every list: each name was a real check and each check was in the catalog.
    with tempfile.TemporaryDirectory() as tmp:
        proj = Path(tmp)
        # A commit whose registry declares a version BEHIND this one, found rather than
        # counted back to: one of the numbers in this history was declared and withdrawn, so
        # `HEAD~4` picked a project that was ahead of the framework and the tool refused it
        # for the right reason -- which tested the refusal instead of the deadlock.
        # NOT WHILE THE BUMP IS UNCOMMITTED. Between editing `version:` and committing it,
        # the number in the working tree belongs to no commit, and everything downstream of
        # that -- what `--adopt` writes, which commit declares what -- describes a framework
        # state no consumer can be in. This check went red exactly there, on every bump, and
        # a check that is red while you work is one you commit through.
        head_registry = subprocess.run(
            ["git", "-C", str(ROOT), "show", f"HEAD:{migrate.REGISTRY_REL}"],
            capture_output=True, text=True)
        committed = (migrate.registry_version(head_registry.stdout)
                     if head_registry.returncode == 0 else None)
        if committed != REGISTRY["version"]:
            return problems + [f"NOT RUN: the working tree declares "
                               f"{REGISTRY['version']} and the last commit declares "
                               f"{committed}. Commit the bump and run this again -- what "
                               "`--adopt` writes cannot be asserted against a version no "
                               "commit has."]

        now = migrate.semver(REGISTRY["version"])
        old_commit = old_version = None
        for sha in subprocess.run(["git", "-C", str(ROOT), "rev-list", "-40", "HEAD"],
                                  capture_output=True, text=True).stdout.split():
            v = migrate.registry_version(subprocess.run(
                ["git", "-C", str(ROOT), "show", f"{sha}:{migrate.REGISTRY_REL}"],
                capture_output=True, text=True).stdout)
            sv = migrate.semver(v) if v else None
            if sv and now and sv < now:
                old_commit, old_version = sha, v
                break
        if old_version:
            # --adopt must bind committed bytes, not this possibly dirty working tree.
            # Exercise it against an isolated clean checkout of the committed runtime;
            # test_adoption.py separately covers the complete proposed source export.
            trusted_target = proj / "trusted-framework"
            clone = subprocess.run(["git", "clone", "--local", "--no-hardlinks", str(ROOT),
                                    str(trusted_target)], capture_output=True, text=True)
            if clone.returncode:
                return problems + ["could not create the isolated migration target: " + clone.stderr]
            proj = proj / "project"
            proj.mkdir()
            cfg = proj / "framework.yaml"
            cfg.write_text(f'framework_version: "{old_version}"\n'
                           f'framework_commit: "{old_commit}"\n')
            subprocess.run([sys.executable,
                            str(ROOT / "skills" / "audit" / "scripts" / "migrate.py"),
                            "--root", str(proj), "--framework", str(trusted_target), "--adopt"],
                           capture_output=True, text=True)
            after = cfg.read_text()
            if REGISTRY["version"] not in after:
                problems.append("a project that pins could not adopt: `FW003` is what the "
                                "pin produces when the framework moves, and if it counts as "
                                "migration work the only way to adopt is to stop pinning")
            elif old_commit[:12] in after:
                problems.append("adopting moved the version and left the pin behind, which "
                                "is `FW003` for the rest of the week")

    for code in migrate.ADOPT_CLEARS:
        if code not in CHECKS["checks"]:
            problems.append(f"migrate.py treats {code} as cleared by --adopt, and it is "
                            "not in the catalog")

    log = subprocess.run(["git", "log", "--format=%H", "--", "schemas/artifact-types.yaml"],
                         cwd=ROOT, capture_output=True, text=True)
    if log.returncode != 0:
        return ["git history is not available here, so the check is not running"]

    seen = []
    for sha in log.stdout.split():
        show = subprocess.run(["git", "show", f"{sha}:schemas/artifact-types.yaml"],
                              cwd=ROOT, capture_output=True, text=True)
        v = migrate.registry_version(show.stdout) if show.returncode == 0 else None
        if v and v not in seen:
            seen.append(v)

    for v in seen:
        if migrate.semver(v) is None:
            continue          # the integers this file used before it carried three numbers
        sha, why = migrate.commit_declaring(v, ROOT)
        if sha is None:
            problems.append(f"{v}: {why}")
        # THE CURRENT VERSION USED TO BE EXEMPT, AND THAT IS WHERE A MALFORMED NOTE HIDES.
        # The parser wants a full stop after the number; a note written with a comma is not
        # a note, and nothing said so until the next bump made it somebody else's problem --
        # by which time the note that explains the version a project is pinned to is the one
        # that silently does not exist. The note is written with the bump, so it can be
        # required with the bump.
        if v not in notes:
            problems.append(f"{v}: no `X -> {v}` note in the registry. The number moved and "
                            "nothing says what it asks a project to do, which is the half "
                            "of `FW001` a version comparison cannot supply")
    return problems


@check("a run that could not have read the repository says so instead of reporting it clean")
def _wrong_arguments_are_not_a_clean_report():
    # `audit/SKILL.md` has said from its first version that "running it against the wrong
    # directory produces a clean report, and a clean report on the wrong repository is worse
    # than an error". Nothing enforced it. Pointed at a path that is not there, the validator
    # scanned nothing, said the repository declares no framework version, and exited 0 --
    # three true sentences whose only available reading was that the project is fine.
    #
    # The other half is the same failure with a stack trace on top: `framework.yaml` is the
    # first file read, so a stray quote in it killed both tools before either had looked at a
    # document, with a message naming a unicode string and no path at all.
    migrate = ROOT / "skills" / "audit" / "scripts" / "migrate.py"
    problems = []

    def run(script, *args):
        r = subprocess.run([sys.executable, str(script), *args],
                           capture_output=True, text=True)
        return r.returncode, (r.stdout + r.stderr)

    with tempfile.TemporaryDirectory() as tmp:
        missing = str(Path(tmp) / "not-here")
        broken = Path(tmp) / "broken"
        broken.mkdir()
        # The version is `N` because a literal one here is the defect `_version_is_not_restated`
        # exists to catch, and it caught this line. What is being tested is a quote that never
        # closes; which number it fails to close is not part of it.
        (broken / "framework.yaml").write_text('framework_version: "N\nchecks: [\n')

        for script, args, what in [
            (VALIDATE, ["--root", missing], "the validator pointed at no directory"),
            (migrate, ["--root", missing], "the migration tool pointed at no directory"),
            (migrate, ["--root", str(broken), "--framework", tmp],
             "a --framework that holds no registry"),
            (VALIDATE, ["--root", str(broken)],
             "a framework.yaml that does not parse, read by the validator"),
            (migrate, ["--root", str(broken)],
             "a framework.yaml that does not parse, read by the migration tool"),
        ]:
            code, out = run(script, *args)
            if code == 0:
                problems.append(f"{what}: exited 0. Nothing distinguishes it from a run "
                                "that had something to check and found it correct")
            if "Traceback" in out:
                problems.append(f"{what}: died with a traceback rather than a sentence")
            # The path is what turns the message into a repair. Either of the two directories
            # involved will do: what must not happen is a message that names neither.
            if not any(p in out for p in (missing, str(broken), tmp)):
                problems.append(f"{what}: the message names no path, so it says which of "
                                f"the two arguments was wrong to nobody: {out.strip()[:120]}")
    return problems


@check("a pinned framework commit is checked against the one doing the checking")
def _the_pin_is_read():
    # An unchecked pin is a comment, and this is the whole value of the field: a project
    # that writes the commit down is claiming its report was produced by that code, and
    # nothing but this compares the claim with the process actually running.
    #
    # The two silences are asserted as hard as the two findings. No pin at all is the normal
    # state and must stay quiet, or the field becomes a checklist item; and a framework
    # installed rather than cloned has no history to ask, where unverifiable is a different
    # claim from violated -- that one cannot be exercised here, because this repository is
    # always a checkout when the self check runs, so it is stated rather than tested.
    head = subprocess.run(["git", "-C", str(ROOT), "rev-parse", "HEAD"],
                          capture_output=True, text=True)
    if head.returncode != 0:
        return ["git history is not available here, so the check is not running"]
    sha = head.stdout.strip()
    version = REGISTRY["version"]
    problems = []

    def codes(config: str):
        with tempfile.TemporaryDirectory() as tmp:
            (Path(tmp) / "framework.yaml").write_text(config)
            r = subprocess.run([sys.executable, str(VALIDATE), "--root", tmp, "--json",
                                "--stale-days", "36500"], capture_output=True, text=True)
            if r.returncode not in (0, 1) or not r.stdout.strip():
                return None
            return {f["code"] for f in json.loads(r.stdout)["findings"]}

    declared = f'framework_version: "{version}"\n'
    for config, want, what in [
        (declared, False, "a project that pins nothing, which is the ordinary state"),
        (declared + f'framework_commit: "{sha}"\n', False, "a pin that is being honoured"),
        (declared + f'framework_commit: "{sha[:10]}"\n', False, "the same pin, abbreviated"),
        (declared + 'framework_commit: "0123456789abcdef"\n', True,
         "a pin naming a commit that is not the one running"),
        (declared + 'framework_commit: "main"\n', True,
         "a branch name, which is a pin that moves and therefore not a pin"),
    ]:
        got = codes(config)
        if got is None:
            problems.append(f"the validator crashed on {what}")
        elif ("FW003" in got) != want:
            problems.append(f"{what} was {'not ' if want else ''}reported")
    return problems

@check("a register's two halves carry the same ids, and a generated view is not one of them")
def _register_halves():
    # THE FAILURE THAT PUT THIS HERE WAS A SILENCE, AND SILENCES ARE WHAT THIS SUITE IS FOR.
    # Until 3.0.0 the three registers kept a map in front matter and repeated its fields in
    # the body, and `risks: {}` under a full §state table validated clean: every check that
    # joins risks reads the map, so all of them went quiet at once while the document still
    # showed the rows to a person. Four ways of looking, all at the blank half.
    #
    # The repair was to stop keeping the fields twice, so what is asserted is what is left:
    # the same ids in both halves, and the three exemptions that make the check survivable --
    # a generated view is not the body, a decided entry keeps its row and loses its prose,
    # and either body shape counts.
    def repo(files: dict[str, str]) -> set[str] | None:
        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp)
            (root / "framework.yaml").write_text(
                f"framework_version: {REGISTRY['version']}\n")
            for rel, text in files.items():
                f = root / rel
                f.parent.mkdir(parents=True, exist_ok=True)
                f.write_text(text)
            r = subprocess.run([sys.executable, str(VALIDATE), "--root", str(root),
                                "--json", "--stale-days", "36500"],
                               capture_output=True, text=True)
            if not r.stdout.strip():
                return None
            return {f["code"] for f in json.loads(r.stdout)["findings"]}

    def rsk(risks: str, body: str) -> dict[str, str]:
        return {"products/alpha/product.yaml":
                    ("---\nschema: framework/product-manifest/v1\n"
                     "artifact_type: product-manifest\nlifecycle: living\nstatus: active\n"
                     "products: [alpha]\nowners: [o]\ncreated: 2026-01-01 09:00\n"
                     "last_review: 2026-08-01 09:00\n---\n\n# alpha\n"),
                "products/alpha/RSK.md":
                    ("---\nschema: framework/risk-register/v1\n"
                     "artifact_type: risk-register\nlifecycle: living\nstatus: active\n"
                     "products: [alpha]\nowners: [o]\ncreated: 2026-01-01 09:00\n"
                     "last_review: 2026-08-01 09:00\n" + risks + "---\n\n# Risks\n\n"
                     "<!-- section: state -->\n## State\n\n"
                     "| ID | Risk | Mitigation | Owner | Reviewed |\n|---|---|---|---|---|\n"
                     + body)}

    one = "risks:\n  RSK-001:\n    category: technical\n    state: open\n"
    row = "| RSK-001 | the thing | the mitigation | o | 2026-08-01 |\n"
    problems = []

    for what, files, want in [
        ("a risk in the map and in §state", rsk(one, row), False),
        ("a map with a row §state does not carry", rsk(one, ""), True),
        ("a §state row the map does not declare", rsk("risks: {}\n", row), True),
        ("an empty map under an empty §state", rsk("risks: {}\n", ""), False),
    ]:
        got = repo(files)
        if got is None:
            problems.append(f"the validator produced nothing on {what}")
        elif ("REG015" in got) != want:
            problems.append(f"{what} was {'not ' if want else ''}reported")

    # `risks:` absent entirely is `REG004` and not `REG015`: there is no map to compare, and
    # the finding has to say the register is unreadable rather than that a row is missing.
    got = repo(rsk("", row))
    if got is not None:
        if "REG004" not in got:
            problems.append("a risk register with no `risks:` at all was not reported by REG004")
        if "REG015" in got:
            problems.append("a register with no map was reported as a half-written row, "
                            "which sends somebody to add a row rather than the map")

    # THE GENERATED UNION IS NOT THE BODY OF THE REGISTER IT SITS IN. `--emit-index` composes
    # every register in the repository into §5 at the root, ids in the first column, and
    # those entries belong to the registers under each product. Read as body, the root
    # register would be reported for entries it correctly does not declare -- a finding on
    # exactly the arrangement the framework asks for.
    union = ("---\nschema: framework/open-register/v1\nartifact_type: open-register\n"
             "lifecycle: living\nstatus: active\nowners: [o]\ncreated: 2026-01-01 09:00\n"
             "last_review: 2026-08-01 09:00\nentries:\n  OD-001:\n    status: open\n"
             "    cost_to_reverse: low\n    products: [all]\n"
             "    default_in_force: nothing\n---\n\n# Open\n\n"
             "### OD-001 - the one at the root\n\n"
             "<!-- generated: open-union -->\n"
             "| ID | Register |\n|---|---|\n| OD-014 | products/alpha/OPEN.md |\n"
             "<!-- /generated -->\n")
    got = repo({"OPEN.md": union})
    if got is not None and "REG015" in got:
        problems.append("an id inside the generated union was read as the root register's "
                        "own body, so the composed view reports the registers it composes")

    # A decided entry keeps its row in the map and loses its prose from §1, which is the one
    # document in the framework that is supposed to get shorter.
    decided = ("---\nschema: framework/open-register/v1\nartifact_type: open-register\n"
               "lifecycle: living\nstatus: active\nowners: [o]\ncreated: 2026-01-01 09:00\n"
               "last_review: 2026-08-01 09:00\nentries:\n  OD-001:\n    status: decided\n"
               "    cost_to_reverse: low\n    products: [all]\n"
               "    default_in_force: nothing\n    closed_by: DEC-001\n---\n\n# Open\n\n"
               "# §4 - closed\n\n- 2026-02-01 OD-001 -> DEC-001\n")
    got = repo({"OPEN.md": decided})
    if got is not None and "REG015" in got:
        problems.append("a decided entry was reported for having left §1, which is what "
                        "taking a decision is supposed to do")

    # EITHER BODY SHAPE COUNTS, AND THIS IS THE ONE THE HEADING BRANCH TOLERATES. An open
    # register writing its entries as a table instead of `### OD-001 - title` is the shape
    # `rows_of` exists for, and the prefix filter must not take it away: the ids in that
    # first column are the ones this type declares.
    as_table = ("---\nschema: framework/open-register/v1\nartifact_type: open-register\n"
                "lifecycle: living\nstatus: active\nowners: [o]\ncreated: 2026-01-01 09:00\n"
                "last_review: 2026-08-01 09:00\nentries:\n  OD-001:\n    status: open\n"
                "    cost_to_reverse: low\n    products: [all]\n"
                "    default_in_force: nothing\n---\n\n# Open\n\n"
                "| ID | Question |\n|---|---|\n| OD-001 | the one written as a row |\n")
    got = repo({"OPEN.md": as_table})
    if got is not None and "REG015" in got:
        problems.append("an open register writing its entries as a table was reported, so "
                        "the shape `rows_of` exists for is the shape that fails")

    # And a row the map does not declare is still reported, or the tolerated shape would be
    # a way to stop being checked rather than a way to write the register.
    got = repo({"OPEN.md": as_table + "| OD-002 | the row nothing declares |\n"})
    if got is not None and "REG015" not in got:
        problems.append("a table row the map does not declare was not reported, so the "
                        "tolerated shape is a way out of the check")

    # A FOREIGN PREFIX IN A FIRST CELL IS A CITATION AND NOT A HALF-WRITTEN ENTRY, AND THE
    # COST IS ASSERTED HERE RATHER THAN LEFT IMPLICIT. `entries:` is schema-forbidden to
    # hold a `DEC`, so a `DEC` in the body can only be cited -- and the same silence covers
    # a mistyped prefix, which is what this filter is paid for.
    got = repo({"OPEN.md": as_table + "| DEC-001 | the decision this row argues about |\n"})
    if got is not None and "REG015" in got:
        problems.append("a `DEC` cited in the first cell of a table was read as an entry "
                        "the map is missing, which `entries:` cannot hold")
    return problems


@check("a reference to a per-product identifier names the product")
def _qualified_references():
    # `DEC` lives at the root and its template prescribes `derives_from: [..., SIG-NNN]`,
    # while signal logs are one per product: two products numbering their signals from 001
    # made every such reference ambiguous, and it resolved itself -- whichever register was
    # read last won. From 3.0.0 the reference names the product and the *declaration* does
    # not, which is the asymmetry worth asserting: a log under `products/alpha/` has already
    # said whose its rows are by sitting there.
    def repo(dec_ref: str, dec_products: str = "[alpha]") -> set[str] | None:
        files = {
            "products/alpha/product.yaml":
                ("---\nschema: framework/product-manifest/v1\n"
                 "artifact_type: product-manifest\nlifecycle: living\nstatus: active\n"
                 "products: [alpha]\nowners: [o]\ncreated: 2026-01-01 09:00\n"
                 "last_review: 2026-08-01 09:00\n---\n\n# alpha\n"),
            # The declaration, bare, in the register whose directory says whose it is.
            "products/alpha/LOG.md":
                ("---\nschema: framework/signal-log/v1\nartifact_type: signal-log\n"
                 "lifecycle: append-only\nstatus: active\nproducts: [alpha]\n"
                 "owners: [o]\ncreated: 2026-01-01 09:00\n---\n\n# Log\n\n"
                 "| ID | What |\n|---|---|\n| SIG-001 | the observation |\n"),
            "decisions/DEC-001-slug.md":
                ("---\nschema: framework/decision-record/v1\n"
                 "artifact_type: decision-record\nid: DEC-001\nlifecycle: immutable\n"
                 "status: accepted\nscope: architecture\n"
                 f"products: {dec_products}\nowners: [o]\ncreated: 2026-01-01 09:00\n"
                 f"derives_from: [{dec_ref}]\nleaves_open: []\n---\n\n# DEC-001\n"),
        }
        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp)
            (root / "framework.yaml").write_text(
                f"framework_version: {REGISTRY['version']}\n")
            for rel, text in files.items():
                f = root / rel
                f.parent.mkdir(parents=True, exist_ok=True)
                f.write_text(text)
            r = subprocess.run([sys.executable, str(VALIDATE), "--root", str(root),
                                "--json", "--stale-days", "36500"],
                               capture_output=True, text=True)
            if not r.stdout.strip():
                return None
            return {f["code"] for f in json.loads(r.stdout)["findings"]}

    problems = []
    cases = [
        # the reference, the code that must appear, the code that must not, what it is
        ("alpha:SIG-001", None, {"FM002", "REF001", "REF007", "REF008"},
         "a qualified reference to a signal its product declares"),
        ("SIG-001", "FM002", set(),
         "a bare reference to a signal, which the schema's reference pattern rejects"),
        ("beta:SIG-001", "REF007", set(),
         "a qualifier naming a product this repository does not have"),
        ("alpha:SIG-404", "REF001", {"REF007"},
         "a qualified reference to a signal that product's register does not declare"),
    ]
    for ref, want, forbidden, what in cases:
        got = repo(ref)
        if got is None:
            problems.append(f"the validator produced nothing on {what}")
            continue
        if want and want not in got:
            problems.append(f"{what} was not reported by {want}: got {sorted(got)}")
        for code in forbidden & got:
            problems.append(f"{what} was reported by {code}, which is a different claim")

    # THE QUALIFIER AGAINST `products:`, WHICH IS THE HALF NO PATTERN CAN SEE. A decision
    # resolving a signal in a product it does not claim to bind is either missing a product
    # from the field every downstream view is built from, or citing a signal copied out of
    # another document.
    got = repo("alpha:SIG-001", dec_products="[gamma]")
    if got is not None and "REF008" not in got:
        problems.append("a qualifier outside the document's own `products:` was not reported")
    got = repo("alpha:SIG-001", dec_products="[all]")
    if got is not None and "REF008" in got:
        problems.append("`products: [all]` was read as naming no product, so a decision "
                        "that binds every product cannot cite any product's signal")
    return problems


@check("the eval case files load, and the counts the README prints come from the same parse")
def _eval_cases_load():
    # THE MEASURING EQUIPMENT TURNS ON. This is not a check about the semantics of a
    # document, which is what the rest of this suite is for and what the freeze covers: it is
    # a check that the instruments switch on, and it is here because they did not. `business`
    # arrived in 2.8.11 carrying `per il management: dove siamo` as an unquoted scalar, so
    # `yaml.safe_load` at `evals/trigger/run.py:177` raised, the runner died before reaching
    # a model, and it did so on every skill rather than only on the new one. It stayed broken
    # for four days and nothing could have said so: the eval suite needs a model, which is
    # why it is not in CI, and the cost of that is a data file whose breakage is invisible
    # until somebody runs it by hand. Parsing needs no model.
    #
    # AND THE NUMBERS FALL OUT OF THE SAME PARSE. `evals/README.md` states how many prompts
    # there are and, in the results table, how many cases each skill has as the denominator
    # of its score. Both were written by hand and both were wrong the moment the set grew --
    # it said 112 while holding 118. A number in prose that nothing derives is the thing this
    # framework spends its own rules forbidding, and this is the cheap half of deriving it:
    # not generating the line, but failing when it stops being true.
    problems = []
    trigger = ROOT / "evals" / "trigger" / "cases.yaml"
    files = [trigger, ROOT / "evals" / "trigger" / "fixtures.yaml"]
    files += sorted((ROOT / "evals" / "behaviour").glob("*/cases.yaml"))

    loaded = {}
    for f in files:
        rel = f.relative_to(ROOT)
        try:
            loaded[rel] = yaml.safe_load(f.read_text(encoding="utf-8"))
        except yaml.YAMLError as e:
            mark = getattr(e, "problem_mark", None)
            where = f" at line {mark.line + 1}" if mark else ""
            problems.append(f"{rel} does not parse{where}: {getattr(e, 'problem', e)}. "
                            "The runner loads it with `yaml.safe_load` and dies here, on "
                            "every case rather than the one that broke it. An unquoted "
                            "value containing `: ` is the way this has happened.")
    if problems:
        return problems                      # counts cannot be read from a file that broke

    # The shape the runners actually index into, asserted per file rather than assumed.
    for rel, doc in loaded.items():
        if rel.name == "fixtures.yaml":
            if not isinstance(doc, dict) or not isinstance(doc.get("fixtures"), dict):
                problems.append(f"{rel}: no `fixtures:` map, which `run.py --fixtures` reads")
            continue
        cases = doc.get("cases") if isinstance(doc, dict) else None
        if not isinstance(cases, list) or not cases:
            problems.append(f"{rel}: no `cases:` list")
            continue
        # THE TWO RUNNERS INDEX DIFFERENTLY, AND THE CHECK ASKS WHAT EACH ONE READS. A
        # behaviour case names a `fixture` and may inherit the prompt from the file --
        # `behaviour/run.py:197` does `c.setdefault("prompt", spec.get("prompt"))` -- while a
        # trigger case carries its own prompt and the label to score it against. Asserting
        # one shape over both reported six correct files as broken, which is the failure this
        # check is supposed to be the opposite of.
        behaviour = rel.parts[1] == "behaviour"
        for i, c in enumerate(cases):
            if not isinstance(c, dict):
                problems.append(f"{rel}: case {i + 1} is not a mapping")
                continue
            where = c.get("prompt") or c.get("fixture") or f"case {i + 1}"
            if behaviour:
                if "fixture" not in c:
                    problems.append(f"{rel}: case {i + 1} names no `fixture`, which is the "
                                    "repository the run happens in and its name")
                if not (c.get("prompt") or doc.get("prompt")):
                    problems.append(f"{rel}: {where!r} has no `prompt`, on the case or on "
                                    "the file it inherits from")
            else:
                if "prompt" not in c:
                    problems.append(f"{rel}: case {i + 1} has no `prompt`")
                elif "expect" not in c:
                    problems.append(f"{rel}: {str(where)[:40]!r} has no `expect`, so a run "
                                    "of it cannot be scored either way")
    if problems:
        return problems

    cases = loaded[trigger.relative_to(ROOT)]["cases"]
    readme = (ROOT / "evals" / "README.md").read_text(encoding="utf-8")

    stated = re.search(r"holds (\d+) prompts", readme)
    if not stated:
        problems.append("evals/README.md no longer says how many prompts the set holds, in "
                        "the form `holds N prompts`, so nothing here can be compared")
    elif int(stated.group(1)) != len(cases):
        problems.append(f"evals/README.md says the set holds {stated.group(1)} prompts and "
                        f"`cases.yaml` holds {len(cases)}")

    # The denominator of each row of the results table is that skill's case count. A row
    # whose denominator has drifted is a score being read against a set that no longer
    # exists, which is how 112 survived a set of 118.
    # A case marked `unmeasured` has not been run, so no table can have scored it: it counts
    # in the set and in no denominator, until the run that scores it removes the mark.
    per_skill = collections.Counter(c.get("expect") for c in cases if not c.get("unmeasured"))
    for label, expect in [(r"`(\w+)`", None), (r"negatives", "none")]:
        for m in re.finditer(r"^\| " + label + r" \| (\d+)/(\d+) \|", readme, re.M):
            skill = expect or m.group(1)
            denom = int(m.group(len(m.groups())))
            if skill not in per_skill and expect is None:
                continue                     # a row about something that is not a skill
            if denom != per_skill[skill]:
                problems.append(
                    f"evals/README.md scores {skill!r} out of {denom} and `cases.yaml` "
                    f"labels {per_skill[skill]} case(s) with that `expect`")
    return problems


# The repository every annotation check below is run against: two living documents that go
# stale on demand, one file whose front matter does not parse, and nothing else. Two
# warnings on two different paths and one error, which is the smallest shape that can
# exercise a join, a stale entry and a refusal at once.
def _annotation_fixture(tmp: Path, annotations: str | None = None,
                        at: str = ".framework/expected-findings.yaml",
                        also: dict[str, str] | None = None) -> None:
    fm = lambda **kw: "---\n" + "\n".join(f"{k}: {v}" for k, v in kw.items()) + "\n---\n\n"
    files = {
        "framework.yaml": f"framework_version: {REGISTRY['version']}\n",
        "OPEN.md": fm(schema="framework/open-register/v1", artifact_type="open-register",
                      lifecycle="living", status="active", owners="[owner]",
                      created="2026-01-01", last_review="2026-01-01 09:00",
                      entries="\n  OD-001:\n    status: open\n"
                              "    cost_to_reverse: low\n    products: [all]\n"
                              "    default_in_force: nothing is scheduled\n")
                   + "# Open\n\n## Cost to reverse LOW\n\n### OD-001 - the one question\n",
        "products/alpha/PBR.md": fm(
            schema="framework/product-brief/v1", artifact_type="product-brief",
            lifecycle="living", status="active", products="[alpha]", owners="[owner]",
            created="2026-01-01", last_review="2026-01-01 09:00") + "# Brief\n",
        # Front matter that does not parse, which is `FM001` and an error. The one thing an
        # annotation is not allowed to touch needs to be present to be refused.
        "broken.md": "---\nartifact_type: [oops\n---\n\n# broken\n",
    }
    if annotations is not None:
        files[at] = annotations
    files.update(also or {})
    for rel, text in files.items():
        f = tmp / rel
        f.parent.mkdir(parents=True, exist_ok=True)
        f.write_text(text, encoding="utf-8")


def _annotated_run(annotations: str | None, **kw) -> tuple[dict | None, str]:
    """The validator over that repository, as parsed JSON and whatever went to stderr."""
    with tempfile.TemporaryDirectory() as tmp:
        _annotation_fixture(Path(tmp), annotations, **kw)
        r = subprocess.run([sys.executable, str(VALIDATE), "--root", tmp, "--json",
                            "--stale-days", "0"], capture_output=True, text=True)
        if r.returncode not in (0, 1) or not r.stdout.strip():
            return None, r.stderr
        return json.loads(r.stdout), r.stderr


ANNOTATES_A_WARNING = """expected:
  - code: LC002
    path: OPEN.md
    reason: >
      the register is read at every cycle and the date is what nobody moves.
    clears_when: >
      the next reading of the register, stamped with the instant it finished.
"""


@check("a finding a project has examined is reported, annotated, and never reclassified")
def _annotations_are_read():
    # THE POINT OF THE MECHANISM IS THE OPPOSITE OF SILENCING, AND THAT IS WHAT IS ASSERTED
    # HERE. A project already had two ways to make a finding go away -- lower the severity,
    # widen the scan -- and `audit/SKILL.md` forbids both because the finding stops being
    # reported and nothing records that anybody decided. So the annotated finding has to
    # still be in `findings`, still at `warn`, and carrying the two sentences; if any of
    # those three slipped, this would have become the third way rather than the alternative
    # to them.
    problems = []
    plain, _ = _annotated_run(None)
    if plain is None:
        return ["the validator crashed on a repository with no annotation file"]
    if plain["warnings"] != 2 or plain["errors"] != 1:
        return [f"the fixture no longer produces 2 warnings and 1 error: "
                f"{plain['errors']} error(s), {plain['warnings']} warning(s)"]
    if plain["annotated"] != 0 or plain["unannotated"] != 2:
        problems.append(f"with no file at all the counts read {plain['annotated']} "
                        f"annotated and {plain['unannotated']} unannotated")
    if any("accepted" in f for f in plain["findings"]):
        problems.append("a finding carries an `accepted` key in a repository that has "
                        "annotated nothing: the key has to be absent, not null, or every "
                        "consumer of this output grows a branch it never needed")

    got, _ = _annotated_run(ANNOTATES_A_WARNING)
    if got is None:
        return problems + ["the validator crashed on a repository with an annotation"]
    marked = [f for f in got["findings"] if f.get("accepted")]
    if len(marked) != 1 or marked[0]["code"] != "LC002" or marked[0]["path"] != "OPEN.md":
        problems.append(f"the annotation matched {len(marked)} finding(s) instead of the "
                        "one it names")
    elif marked[0]["level"] != "warn":
        problems.append(f"the annotated finding is now at {marked[0]['level']!r}. It stays "
                        "at `warn`: a project checking its own warnings keys on `level`, "
                        "and a finding that quietly became something else drops out of that "
                        "check without a word")
    elif not marked[0]["accepted"].get("clears_when"):
        problems.append("the annotation reached the finding without what clears it")
    if got["warnings"] != 2:
        problems.append(f"annotating one warning changed the warning count to "
                        f"{got['warnings']}: the count is of what was reported, not of "
                        "what is still unexplained")
    if (got["annotated"], got["unannotated"]) != (1, 1):
        problems.append(f"the counts read {got['annotated']} annotated and "
                        f"{got['unannotated']} unannotated, and the fixture has one of each")
    if any(f["code"].startswith("AN") for f in got["findings"]):
        problems.append("a correct annotation produced an `AN` finding")
    return problems


@check("an annotation that explains nothing, or explains away an error, is refused")
def _annotations_are_policed():
    # The two ways this file rots, and both are errors because the file is load-bearing on
    # every run now rather than only during a migration. A stale entry is a reason left
    # standing for a finding that is gone, which the next reader takes as current; an
    # annotation on an `error` is the level that blocks being talked out of blocking, one
    # project at a time.
    problems = []
    stale, _ = _annotated_run("""expected:
  - code: REG005
    path: nowhere.md
    reason: >
      a finding that is not reported here.
    clears_when: >
      nothing, which is the whole problem.
""")
    if stale is None:
        return ["the validator crashed on a stale annotation"]
    an001 = [f for f in stale["findings"] if f["code"] == "AN001"]
    if len(an001) != 1:
        problems.append(f"a stale annotation produced {len(an001)} AN001")
    elif an001[0]["level"] != "error":
        problems.append(f"AN001 is at {an001[0]['level']!r}. It is the one line that keeps "
                        "this file from becoming a list of permanent exemptions with a "
                        "better name, and a warning is not that line")

    on_error, _ = _annotated_run("""expected:
  - code: FM001
    path: broken.md
    reason: >
      an error, which this must not be able to explain away.
    clears_when: >
      the front matter parsing.
""")
    if on_error is None:
        return problems + ["the validator crashed on an annotation naming an error"]
    if not [f for f in on_error["findings"] if f["code"] == "AN002"]:
        problems.append("an annotation on an `error` finding was accepted or ignored. "
                        "Ignoring is the worse of the two: the project believes it has "
                        "annotated and finds out when the gate stays red")
    if [f for f in on_error["findings"] if f["code"] == "FM001" and f.get("accepted")]:
        problems.append("the error was marked as accepted anyway")

    strict, _ = _annotated_run("require_all: true\n" + ANNOTATES_A_WARNING)
    if strict is None:
        return problems + ["the validator crashed with require_all"]
    an003 = [f for f in strict["findings"] if f["code"] == "AN003"]
    if len(an003) != 1 or an003[0]["path"] != "products/alpha/PBR.md":
        problems.append(f"`require_all` reported {len(an003)} unannotated warning(s), and "
                        "the fixture leaves exactly one")
    loose, _ = _annotated_run(ANNOTATES_A_WARNING)
    if loose and any(f["code"] == "AN003" for f in loose["findings"]):
        problems.append("AN003 fired without `require_all`: permissive is the default, and "
                        "a check that reddened every repository on arrival would be off in "
                        "a week")
    return problems


@check("the annotation file has one home, is never an artifact, and says so on stderr")
def _the_annotation_file_has_one_home():
    # Where this file sits was forced rather than chosen: it has to be readable by the
    # tooling arriving and invisible to every validator already released, and a hidden path
    # is the only slot with that property. Which makes two things worth asserting: that a
    # project cannot turn it back into an artifact, and that the move from the place the
    # first project had to invent is announced rather than silent.
    problems = []
    # THE GRACE PERIOD ENDED WHERE THE NOTE SAID IT WOULD, AND SAYING SO IS THE HALF THAT
    # STAYS. 3.1.0 read the old path and wrote down that 3.2.0 would stop; a note that says
    # when something ends is a promise or it is decoration. What must not happen is the
    # silent version: a file sitting in a repository doing nothing, discovered six months
    # later by somebody wondering why an annotation stopped applying.
    legacy, err = _annotated_run(ANNOTATES_A_WARNING, at=".claude/expected-findings.yaml")
    if legacy is None:
        return ["the validator crashed on a repository carrying only the old path"]
    if [f for f in legacy["findings"] if f.get("accepted")]:
        problems.append("the old path is still being read, and the version note says it is "
                        "not: one of the two is wrong and the note is the promise")
    if ".claude" not in err:
        problems.append("a file at the old path was ignored in silence. It has to be "
                        "reported precisely because it is not read any more: nothing in it "
                        "is applying, and that is invisible from the report otherwise")

    both, err = _annotated_run(ANNOTATES_A_WARNING, also={
        ".claude/expected-findings.yaml": """expected:
  - code: LC002
    path: products/alpha/PBR.md
    reason: >
      an annotation in the file that is not read.
    clears_when: >
      nothing.
"""})
    if both is None:
        problems.append("the validator crashed with both files present")
    else:
        marked = {f["path"] for f in both["findings"] if f.get("accepted")}
        if marked != {"OPEN.md"}:
            problems.append(f"with both files present the annotations applied were "
                            f"{sorted(marked)}: only the new path is read")
        if ".claude" not in err:
            problems.append("the leftover file was ignored in silence, which is how a "
                            "project ends up editing the one that is not read")

    # `skip_hidden: false` is a legal thing for a project to write, and it is what would
    # otherwise turn this file into an `FM001`, at `error`, on the run it exists to inform.
    exposed, _ = _annotated_run(
        ANNOTATES_A_WARNING,
        also={"framework.yaml": f"framework_version: {REGISTRY['version']}\n"
                                "scan:\n  skip_hidden: false\n"})
    if exposed is None:
        problems.append("the validator crashed with skip_hidden: false")
    elif [f for f in exposed["findings"] if "expected-findings" in f["path"]]:
        problems.append("the annotation file was reported as a document once the project "
                        "stopped skipping hidden paths. The validator reads it by an "
                        "explicit path, and that alone does not stop `discover` finding it: "
                        "`scan.skip_files` is the half that holds, because `load_scan` "
                        "unions and never subtracts")
    return problems


@check("the pin is said at the moment the framework is invoked, not only in the report")
def _the_pin_is_said_before_the_scan():
    # `FW003` reports this and cannot be relied on to: a project whose own gate runs the
    # validator of the version it declares stops running that gate the moment the checkout
    # moves, correctly, and the thing that makes it refuse is the divergence itself. So the
    # report carrying the finding is the one nobody is running, and it is not being run
    # because the fault is present. The line on stderr is the same fact said where the
    # circularity cannot swallow it.
    head = subprocess.run(["git", "-C", str(ROOT), "rev-parse", "HEAD"],
                          capture_output=True, text=True)
    if head.returncode != 0:
        return ["git history is not available here, so the check is not running"]
    problems = []
    declared = f"framework_version: {REGISTRY['version']}\n"
    for config, want, what in [
        (declared, False, "a project that pins nothing"),
        (declared + f'framework_commit: "{head.stdout.strip()}"\n', False,
         "a pin that is being honoured"),
        (declared + 'framework_commit: "0123456789abcdef"\n', True,
         "a pin naming a commit that is not the one running"),
    ]:
        with tempfile.TemporaryDirectory() as tmp:
            _annotation_fixture(Path(tmp), also={"framework.yaml": config})
            r = subprocess.run([sys.executable, str(VALIDATE), "--root", tmp, "--json",
                                "--stale-days", "0"], capture_output=True, text=True)
        said = "this checkout is at" in r.stderr
        if said != want:
            problems.append(f"{what} was {'not ' if want else ''}announced on stderr")
        # The line has to be on stderr and nowhere else. stdout carries the JSON every
        # caller parses, `migrate.py` included, so a line printed there would break the
        # tool the message exists to send somebody to.
        try:
            json.loads(r.stdout)
        except json.JSONDecodeError:
            problems.append(f"stdout stopped being JSON for {what}: the line went to the "
                            "stream the callers parse")
    return problems


def _vocabulary_repo(tmp: Path, entries: str, extra: dict[str, str] | None = None) -> None:
    """A root register with a product beside it, which is what `REG011` needs to ask at all."""
    fm = lambda **kw: "---\n" + "\n".join(f"{k}: {v}" for k, v in kw.items()) + "\n---\n\n"
    files = {
        "framework.yaml": f"framework_version: {REGISTRY['version']}\n",
        "OPEN.md": fm(schema="framework/open-register/v1", artifact_type="open-register",
                      lifecycle="living", status="active", owners="[owner]",
                      created="2026-01-01", last_review="2026-01-01 09:00",
                      entries=entries)
                   + "# Open\n\n<!-- generated: open-union -->\n\n<!-- /generated -->\n",
        "products/alpha/product.yaml": fm(
            schema="framework/product-manifest/v1", artifact_type="product-manifest",
            lifecycle="living", status="active", products="[alpha]", owners="[owner]",
            created="2026-01-01", last_review="2026-01-01 09:00",
            stage="\n  phase: F4\n  since: 2026-01-01"),
    }
    files.update(extra or {})
    for rel, text in files.items():
        f = tmp / rel
        f.parent.mkdir(parents=True, exist_ok=True)
        f.write_text(text, encoding="utf-8")


def _vocabulary_run(entries: str, extra: dict[str, str] | None = None, emit: bool = False):
    with tempfile.TemporaryDirectory() as tmp:
        root = Path(tmp)
        _vocabulary_repo(root, entries, extra)
        args = ["--emit-index"] if emit else []
        r = subprocess.run([sys.executable, str(VALIDATE), "--root", tmp, "--json",
                            "--stale-days", "36500", *args], capture_output=True, text=True)
        if r.returncode not in (0, 1) or not r.stdout.strip():
            return None, ""
        return json.loads(r.stdout), (root / "OPEN.md").read_text(encoding="utf-8")


BINDS_NOTHING = """
  OD-001:
    status: open
    cost_to_reverse: low
    products: [all]
    default_in_force: nothing is scheduled
  KI-002:
    status: open
    cost_to_reverse: low
    products: [none]
    default_in_force: the tooling defect stands
"""


@check("an entry that binds no product says so, and reaches the view without reaching a product")
def _binds_no_product():
    # `[none]` IS THE FOURTH ANSWER AND THE ONE THIS FILE'S OWN NOTE NAMED WITHOUT GIVING IT
    # A VALUE: "not decided, not measurable, not applicable, not ours". The last one is the
    # entry whose subject is the repository or the tooling it is checked with, and the two
    # ways of writing it before were both wrong -- `[all]` says something untrue in the field
    # everything joins on, absence says the thing `REG011` exists to remove.
    #
    # The half that is easy to get wrong is not the check, it is the view. An entry that
    # binds nothing and appears nowhere would validate, satisfy every check, and be missing
    # from the only composed list of what is open, which is the same disappearance from the
    # other side.
    problems = []
    got, open_md = _vocabulary_run(BINDS_NOTHING, emit=True)
    if got is None:
        return ["the validator crashed on a register carrying `[none]`"]
    if [f for f in got["findings"] if f["code"] in ("REG011", "REG013")]:
        problems.append("`[none]` was reported as a missing or misfiled answer: it is an "
                        "answer, and the whole point is that it is a different one from "
                        "`[all]` and from silence")
    region = open_md.split("<!-- generated: open-union -->", 1)[-1]
    if "## Bound to no product at all" not in region:
        problems.append("the generated union has no section for entries that bind nothing, "
                        "so an entry saying its subject is not a product is absent from the "
                        "one view that composes what is open")
    else:
        tail = region.split("## Bound to no product at all", 1)[1]
        if "KI-002" not in tail:
            problems.append("the entry that binds nothing did not reach its own section")
        if "KI-002" in region.split("## Bound to no product at all", 1)[0]:
            problems.append("the entry that binds nothing also appears under a product or "
                            "among the ones that bind every product, which is the finding "
                            "this value exists to remove, moved by a metre")
    if "## none" in region:
        problems.append("`none` was emitted as a product heading. It is a word, not a "
                        "product, and this is the shape `all` had before it was one")

    # A reserved word answers how many products are bound, so it cannot share the field.
    beside = _vocabulary_run(BINDS_NOTHING.replace("products: [none]",
                                                   "products: [none, alpha]"))[0]
    if beside is None or not [f for f in beside["findings"] if f["code"] == "REG011"]:
        problems.append("a reserved word sitting beside a named product was not reported, "
                        "and the reader has to guess which half was the afterthought")

    # And a product that carries one of the two names makes every use of the field ambiguous.
    fm = lambda **kw: "---\n" + "\n".join(f"{k}: {v}" for k, v in kw.items()) + "\n---\n\n"
    named = _vocabulary_run(BINDS_NOTHING, extra={"products/none/product.yaml": fm(
        schema="framework/product-manifest/v1", artifact_type="product-manifest",
        lifecycle="living", status="active", products="[none]", owners="[owner]",
        created="2026-01-01", last_review="2026-01-01 09:00",
        stage="\n  phase: F4\n  since: 2026-01-01")})[0]
    said = [f for f in (named or {}).get("findings", []) if f["code"] == "XP008"]
    if named is None or not said:
        problems.append("a product actually called `none` was not reported, so every "
                        "`products:` naming it says two things and nothing says which")
    elif "open_decisions" not in said[0]["message"]:
        # The check is `warn`, so the message carries the whole weight. What it costs is a
        # silent disappearance -- no heading in the generated union, an empty derived view --
        # and a warning that says only "reserved word" gets read as a naming preference by
        # somebody who then never connects the missing section to it.
        problems.append("XP008 says the name is reserved and not what it costs. A warning "
                        "whose damage is a product vanishing from the composed view has to "
                        "name the damage: whoever reads it is deciding what ignoring it "
                        "costs them")
    return problems


COMMITMENTS_WITH = """---
schema: framework/commitments/v1
artifact_type: commitments
lifecycle: living
status: active
owners: [owner]
products: [alpha]
created: 2026-01-01
last_review: 2026-01-01 09:00
commitments:
  CMT-001:
    to: the committee
    status: open
    products: [alpha]
    unanswerable:
      feasibility:
        reason: >
          each capability is buildable alone and the set is not.
        settled_by: >
          the decision on the scope.
%s
---

# Commitments

### CMT-001 - the one that cannot be answered yet
%s
"""


def _commitments(rows: str = "", headings: str = ""):
    return {"COMMITMENTS.md": COMMITMENTS_WITH % (rows, headings)}


@check("a field left empty on purpose is declared, counted, and cannot be declared and filled")
def _unanswerable_is_declared_and_guarded():
    # The registry's note on closed vocabularies describes this failure and left it without a
    # repair for as long as it has existed: the careful writer omits the field, and the
    # omission is the honest thing to write and the most invisible one. What is asserted here
    # is that declaring it is visible -- the count -- and that the declaration cannot outlive
    # the state it describes, which is the only part of the rule a script can hold. Whether
    # the event in `settled_by` has happened is not knowable from here, and `REG009` says the
    # same about a trigger one field over.
    problems = []
    clean = _vocabulary_run(BINDS_NOTHING, extra=_commitments())[0]
    if clean is None:
        return ["the validator crashed on a record carrying `unanswerable`"]
    if clean["unanswerable"] != 1:
        problems.append(f"the report counts {clean['unanswerable']} declared field(s) and "
                        "the fixture declares one: the count is what makes a deliberate "
                        "emptiness visible, and without it this is an omission again")
    if [f for f in clean["findings"] if f["code"].startswith("UNA")]:
        problems.append("a correct declaration produced a finding")

    both = _vocabulary_run(BINDS_NOTHING, extra=_commitments(
        rows="""    feasibility: feasible"""))[0]
    if both is None or not [f for f in both["findings"] if f["code"] == "UNA001"]:
        problems.append("a field declared unanswerable and filled in anyway was accepted. "
                        "That is the declaration outliving its own state, and it is the one "
                        "half of the rule a check can see")
    elif [f for f in both["findings"] if f["code"] == "UNA001"][0]["level"] != "error":
        problems.append("UNA001 is not an error. It is to this what AN001 is to the "
                        "annotation file: without it the key is a permit not to answer")

    for swap, what in (("      fesibility:", "a field the map does not have"),
                       ("      status:", "a field the map requires")):
        got = _vocabulary_run(BINDS_NOTHING, extra={
            "COMMITMENTS.md": COMMITMENTS_WITH.replace("      feasibility:", swap)
                              % ("", "")})[0]
        if got is None or not [f for f in got["findings"] if f["code"] == "UNA002"]:
            problems.append(f"{what} was accepted as a declaration: it reads as an answer "
                            "given, and the field it was meant for is still silent")

    dated = _vocabulary_run(BINDS_NOTHING, extra={
        "COMMITMENTS.md": COMMITMENTS_WITH.replace(
            "          the decision on the scope.",
            "          2026-12-31, the end of the year.") % ("", "")})[0]
    if dated is None or not [f for f in dated["findings"] if f["code"] == "UNA003"]:
        problems.append("a date in `settled_by` was accepted. It takes the event after "
                        "which the field would have a true value, for the reason REG009 "
                        "gives one field over")
    return problems


@check("every artifact sits where the registry says its type lives, and the patterns parse")
def _placement_is_declared_and_true():
    # THE INVARIANT THAT CATCHES THE CLASS INSTEAD OF THE INSTANCE, which is the rule this
    # suite is written to. Two of the thirty placements were false about this repository's own
    # files -- `CHG-NNN.md` and `DC-NNN.md`, while every real one carries a slug -- and they
    # were false for months, because the field was read by one line of the generator and by
    # nothing else. Writing the patterns out found twelve more of the same mistake.
    #
    # So what is asserted is not "those two are fixed", it is that no fixture artifact sits
    # anywhere its type does not describe. That fails on the day somebody writes a placement
    # that does not match the files, and it fails on the day somebody files a document
    # somewhere new without saying so in the registry, which are the two directions the
    # mistake arrives from.
    v = _load(VALIDATE, "validate")
    problems = []
    for name, spec in sorted(REGISTRY["types"].items()):
        where = spec.get("path")
        if not isinstance(where, list) or not where:
            problems.append(f"{name}: `path` is {where!r} and has to be a list of placements. "
                            "As a string it was a sentence, and a sentence is what nothing "
                            "could read")
            continue
        for one in where:
            try:
                v.placement_pattern(one)
            except re.error as e:
                problems.append(f"{name}: {one!r} does not compile as a placement: {e}")
        if "path_not_enforced" in spec and not str(spec["path_not_enforced"]).strip():
            problems.append(f"{name}: declares `path_not_enforced` with no reason. The point "
                            "of the key is the reason: without it, it is the check switched "
                            "off for one type and nothing saying why")

    roots = corpus()
    if not roots:
        return problems + ["no fixture roots could be built: the check is not running"]
    seen = 0
    for root in roots:
        scan = v.load_scan(REGISTRY, v.load_project(root))
        for a in v.discover(root, scan, REGISTRY, v.Report({})):
            spec = REGISTRY["types"].get(a.type or "")
            if not spec or not spec.get("path") or spec.get("path_not_enforced"):
                continue
            seen += 1
            rel = a.rel.replace("\\", "/")
            if not any(v.placement_pattern(one).match(rel) for one in spec["path"]):
                problems.append(
                    f"{root.name}/{rel}: a {a.type} sits where the registry does not say one "
                    f"lives ({' or '.join(spec['path'])}). Either the file is misfiled or the "
                    "placement is wrong, and the second is how this field came to be false "
                    "about two types for months")
    if seen < 100:
        problems.append(f"only {seen} typed artifacts were examined: the fixtures are not "
                        "built, so this ran against almost nothing")
    return problems


@check("a document that declares a type in the wrong place is reported, and an exempt type is not")
def _misplaced_artifact_is_reported():
    # The case, rebuilt: a derived planning document in a working directory, declaring
    # `artifact_type: roadmap`, `living`, a product, and a `last_review` fresher than the real
    # roadmap's. Everything about it validated. It was counted as an artifact and it came out
    # first in the derived view of that product, because that list is ordered by path and a
    # working directory sorts before `products/`.
    v = _load(VALIDATE, "validate")
    fm = lambda **kw: "---\n" + "\n".join(f"{k}: {v}" for k, v in kw.items()) + "\n---\n\n"
    stray = fm(schema="framework/roadmap/v1", artifact_type="roadmap", lifecycle="living",
               status="active", products="[alpha]", owners="[owner]", created="2026-01-01",
               last_review="2026-01-01 09:00") + "# A draft derived from the real roadmap\n"
    problems = []
    with tempfile.TemporaryDirectory() as tmp:
        root = Path(tmp)
        for rel, text in {
            "framework.yaml": f"framework_version: {REGISTRY['version']}\n",
            "products/alpha/RMP.md": stray,
            "_work/planning/draft.md": stray,
        }.items():
            f = root / rel
            f.parent.mkdir(parents=True, exist_ok=True)
            f.write_text(text, encoding="utf-8")
        r = subprocess.run([sys.executable, str(VALIDATE), "--root", tmp, "--json",
                            "--stale-days", "36500"], capture_output=True, text=True)
        if r.returncode not in (0, 1) or not r.stdout.strip():
            return ["the validator crashed on a misplaced artifact"]
        got = json.loads(r.stdout)
    said = [f for f in got["findings"] if f["code"] == "LOC001"]
    if len(said) != 1 or said[0]["path"] != "_work/planning/draft.md":
        problems.append(f"the misplaced roadmap produced {len(said)} LOC001, and the one in "
                        "`products/alpha/` must not be among them: a check that reports the "
                        "correct file too is a check that gets switched off")
    elif "roadmap" not in said[0]["message"]:
        problems.append("the finding does not name the type, so a reader cannot tell which "
                        "of the two claims about the file is the wrong one")

    # A type the registry declares unenforceable is skipped, and counted rather than silent.
    fake = {"types": {"roadmap": {**REGISTRY["types"]["roadmap"],
                                  "path_not_enforced": "a reason"}}}
    rep = v.Report({})
    art = v.Artifact(Path("x"), "anywhere/else.md",
                     {"artifact_type": "roadmap"}, "")
    exempt = v.check_placement([art], fake, rep)
    if rep.findings:
        problems.append("a type declared `path_not_enforced` was still reported: the key is "
                        "the framework saying it cannot write the pattern, and reporting "
                        "anyway makes the declaration a comment")
    if exempt != 1:
        problems.append(f"the count of unenforced types came back {exempt}: the report says "
                        "how many types the check does not run on, or the hole is invisible")
    return problems


def _reg016(root: Path) -> list[dict]:
    r = subprocess.run([sys.executable, str(VALIDATE), "--root", str(root), "--json",
                        "--stale-days", "36500"], capture_output=True, text=True)
    if r.returncode not in (0, 1) or not r.stdout.strip():
        return []
    return [f for f in json.loads(r.stdout)["findings"] if f["code"] == "REG016"]


def _kind(message: str) -> str:
    if "names it and carries" in message:
        return "both"
    return "label" if "not what the map holds" in message else "value"


@check("the body repeating a field is found by the label and by the value, and both are needed")
def _reg016_looks_twice():
    # NEITHER MECHANISM IS REDUNDANT, AND THAT IS ASSERTED ON THE CORPUS RATHER THAN ARGUED.
    # The label finds a column headed with a field name whose cells say the same thing in
    # other words -- `medium` and `high` where the map holds `M` and `H` -- and the value
    # finds a column headed `To whom` whose cells are the `to:` of the map verbatim, which no
    # vocabulary of field names can recognise because the label is a paraphrase. Drop either
    # half and one of those two goes silent, in English, with nobody the wiser.
    roots = corpus()
    kinds = collections.Counter()
    for root in roots:
        for f in _reg016(root):
            kinds[_kind(f["message"])] += 1
    problems = []
    for kind, what in (("label", "a column named for a field whose cells are not the map's value"),
                       ("value", "a cell that is the map's value under a header that is not the "
                                 "field name"),
                       ("both", "a column named for a field and carrying its value")):
        if not kinds[kind]:
            problems.append(f"no finding of the `{kind}` kind anywhere in the fixtures, so "
                            f"nothing measures {what}. Either the corpus stopped containing it "
                            "or half the check stopped working, and the two are "
                            "indistinguishable from here")
    if sum(kinds.values()) < 10:
        problems.append(f"only {sum(kinds.values())} findings across the corpus: the fixtures "
                        "are not built")
    return problems


@check("a register written in another language is found by the value, and its prose is not")
def _reg016_does_not_need_english():
    # THE FIXTURE EXISTS BECAUSE WITHOUT IT THIS DEFECT COULD NOT BE MEASURED AT ALL. Every
    # fixture wrote English labels, so a check that only understood English reported clean on
    # the whole corpus and the bench reproduced the flaw of the thing it was testing.
    #
    # What is asserted is both halves of the answer: the value carries the whole finding when
    # the labels are in another language, and the residue stays a residue. A sentence restating
    # a likelihood in Italian words is found by neither mechanism, which is written in the
    # catalog as a limit and has to stay true or the limit is a guess.
    root = built("requirement")
    if root is None:
        return ["the requirement fixture would not build"]
    found = _reg016(root)
    problems = []
    if not found:
        return ["the Italian register reports nothing: the value half is not running"]
    if any(_kind(f["message"]) != "value" for f in found):
        problems.append("a finding in the Italian register came through the label. Its column "
                        "headers are Italian, so the label can only be matching something that "
                        "is not a header, and the fixture has stopped measuring what it is for")
    rsk = (root / "products" / "riconciliazione" / "RSK.md").read_text(encoding="utf-8")
    if "La probabilita' che accada" not in rsk:
        problems.append("the fixture no longer carries the sentence that restates a field in "
                        "prose, so the residue this check documents is no longer measured")
    else:
        line = next(n for n, l in enumerate(rsk.splitlines(), 1)
                    if "La probabilita' che accada" in l)
        if any(str(line) in f["message"] for f in found):
            problems.append(f"line {line} is reported. It restates a likelihood in words and "
                            "repeats no value: reporting it would make this check about "
                            "meaning, which it cannot be")
    return problems


@check("the duplication the framework deliberately kept is not reported, and prose is not searched")
def _reg016_leaves_what_was_left():
    # TWO EXEMPTIONS, AND BOTH WERE PAID FOR. §1 of an open register groups entries under
    # `## Cost to reverse HIGH`, which repeats a map field on purpose: 3.0.0 left that one
    # because the heading is the reading order of the document, and a check reporting it would
    # be reporting the arrangement the framework asks for. And a value must be compared whole:
    # searching for it inside a line reported `default_in_force: Italian only` against a body
    # line asking whether the interface should be `Italian only, or Italian and English?`,
    # where the question is about that choice and repeats nothing.
    fm = lambda **kw: "---\n" + "\n".join(f"{k}: {v}" for k, v in kw.items()) + "\n---\n\n"
    body = ("# Open\n\n"
            "## Cost to reverse LOW: changing it later costs an afternoon\n\n"
            "### OD-001 - which language the interface is in\n\n"
            "- **Question:** Italian only, or Italian and English?\n")
    with tempfile.TemporaryDirectory() as tmp:
        root = Path(tmp)
        (root / "framework.yaml").write_text(f"framework_version: {REGISTRY['version']}\n")
        (root / "OPEN.md").write_text(
            fm(schema="framework/open-register/v1", artifact_type="open-register",
               lifecycle="living", status="active", owners="[owner]", created="2026-01-01",
               last_review="2026-01-01 09:00",
               entries="\n  OD-001:\n    status: open\n    cost_to_reverse: low\n"
                       "    products: [all]\n    default_in_force: Italian only\n") + body,
            encoding="utf-8")
        found = _reg016(root)
    problems = []
    for f in found:
        if "`cost_to_reverse`" in f["message"]:
            problems.append("the `## Cost to reverse` heading of §1 was reported. That "
                            "duplication was left standing on purpose and the templates say "
                            "which of the two a check reads: reporting it asks a repository to "
                            "undo the arrangement the framework prescribes")
        if "`default_in_force`" in f["message"]:
            problems.append("a body line was reported for containing the value inside a longer "
                            "sentence. A value is compared whole, as a cell is, or this check "
                            "reports every question that mentions its own subject")
    return problems


@check("a document changed after it was read is reported, and one reread after being changed is not")
def _review_gap():
    # THE BENCH FOR THIS HAD TO BE BUILT, AND THE MEASUREMENT THAT SAYS SO IS THE POINT. The
    # six `release/` fixtures are git repositories with dated commits and looked like a bench:
    # thirty of their living documents carried a gap between the instant they attest and the
    # commit that last touched them. Twenty-nine were the first commit importing everything
    # after those instants, in repositories where nothing had been edited. A check validated
    # against that number would have been validated against its own worst failure mode.
    #
    # `review/gap` plants three commits that can each be said in a sentence, and the three
    # properties are asserted here rather than in the shape of the fixture: the import is not a
    # change, the reading is not a change, and a change that was never read is.
    root = built("review")
    if root is None:
        return ["the review fixture would not build"]
    r = subprocess.run([sys.executable, str(VALIDATE), "--root", str(root), "--json",
                        "--stale-days", "36500"], capture_output=True, text=True)
    if r.returncode not in (0, 1) or not r.stdout.strip():
        return ["the validator crashed on the review fixture"]
    out = json.loads(r.stdout)
    said = [f for f in out["findings"] if f["code"] == "LC006"]
    problems = []
    paths = {f["path"] for f in said}
    want = {"products/atlas/PBR.md", "products/atlas/contracts/DC-001-orders.md"}
    if paths != want:
        problems.append(f"the fixture reports {sorted(paths)} and plants {sorted(want)}. "
                        "`ARC.md` is changed by the same commit and reread by the next one: "
                        "reporting it means a reading is being counted as a change, and then "
                        "every honest review carries a gap of the minutes it took to commit")
    if len(out["findings"]) != len(said):
        others = [f["code"] for f in out["findings"] if f["code"] != "LC006"]
        problems.append(f"the fixture also reports {sorted(set(others))}. It has to be clean in "
                        "every direction but the one it is planted in, or the finding it exists "
                        "to produce arrives inside a list of unrelated ones")
    for f in said:
        if "11 day" not in f["message"]:
            problems.append(f"{f['path']}: the gap is not stated as eleven days. The size is the "
                            "finding: there is no floor, so what separates a document nobody "
                            "reread from one reread a minute before committing is the number")
    # Largest first, because an ordering that carries no meaning gets read as though it did.
    order = [f["path"] for f in out["findings"] if f["code"] == "LC006"]
    if len(order) > 1:
        gaps = [int(f["message"].split(" day")[0].split()[-1]) for f in said]
        if gaps != sorted(gaps, reverse=True):
            problems.append(f"the findings come out in the order {gaps}, and they are ordered "
                            "by gap descending or thirty of them are unreadable")

    # And silence where there is no history. Not a fixture: those live inside this repository,
    # so git answers with the framework's own history. A directory outside any repository is
    # the case, and it is the one a project distributed as an archive is in.
    with tempfile.TemporaryDirectory() as tmp:
        outside = Path(tmp)
        (outside / "framework.yaml").write_text(f"framework_version: {REGISTRY['version']}\n")
        (outside / "OPEN.md").write_text(
            "---\nschema: framework/open-register/v1\nartifact_type: open-register\n"
            "lifecycle: living\nstatus: active\nowners: [owner]\ncreated: 2026-01-01\n"
            "last_review: 2026-01-01 09:00\nentries:\n  OD-001:\n    status: open\n"
            "    products: [all]\n---\n\n# Open\n\n### OD-001 - one question\n")
        r = subprocess.run([sys.executable, str(VALIDATE), "--root", str(outside), "--json",
                            "--stale-days", "36500"], capture_output=True, text=True)
        if r.returncode not in (0, 1) or not r.stdout.strip():
            problems.append("the validator crashed outside a repository")
        elif [f for f in json.loads(r.stdout)["findings"] if f["code"] == "LC006"]:
            problems.append("a directory outside any git repository was reported. There is no "
                            "history to ask there, and unverifiable is not the same claim as "
                            "violated: that is the position FW003 already takes")
    return problems


@check("a table in a register that argues in headings is not read as its entries")
def _tables_are_the_other_shape():
    # "A register adopted from before the framework, or one written by somebody who preferred a
    # table, puts the same ids in a first column" was written about a register that has no
    # headings, and was applied as "read the tables too". What that cost: a register keeping a
    # section to declare an identifier space retired from another repository had those ids read
    # as its own, and the finding landed on the one section written to prevent exactly the
    # confusion the finding described. The prefixes are identical, because the collision is
    # between two repositories, so no filter could have separated them.
    fm = lambda **kw: "---\n" + "\n".join(f"{k}: {v}" for k, v in kw.items()) + "\n---\n\n"
    head = fm(schema="framework/open-register/v1", artifact_type="open-register",
              lifecycle="living", status="active", owners="[owner]", created="2026-01-01 09:00",
              last_review="2026-01-01 09:00",
              entries="\n  OD-011:\n    status: open\n    cost_to_reverse: low\n"
                      "    products: [all]\n    default_in_force: nothing is scheduled\n")

    def codes(body: str, entries: str | None = None) -> set[str] | None:
        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp)
            (root / "framework.yaml").write_text(f"framework_version: {REGISTRY['version']}\n")
            (root / "OPEN.md").write_text((entries or head) + body, encoding="utf-8")
            r = subprocess.run([sys.executable, str(VALIDATE), "--root", tmp, "--json",
                                "--stale-days", "36500"], capture_output=True, text=True)
            if r.returncode not in (0, 1) or not r.stdout.strip():
                return None
            return {f["code"] for f in json.loads(r.stdout)["findings"]}

    retired = ("# Open\n\n### OD-011 - the one open question here\n\n"
               "## The retired register\n\n"
               "The identifiers below are not this file's.\n\n"
               "| There | Here |\n|---|---|\n| OD-001 | OD-011 |\n| OD-002 | none |\n")
    problems = []
    got = codes(retired)
    if got is None:
        return ["the validator crashed on a register carrying a retired identifier space"]
    if "REG015" in got:
        problems.append("the retired table was read as this register's entries. The headings "
                        "already said what the entries are, so the table is the other shape "
                        "and not a second source")

    # And the shape the table harvest exists for still works: a register with no headings at
    # all keeps its rows read, or a register that argues in tables reports every row as missing.
    rows_only = ("# Open\n\n| ID | Question |\n|---|---|\n"
                 "| OD-011 | the one open question here |\n")
    got = codes(rows_only)
    if got is None or "REG015" in got:
        problems.append("a register that argues in tables and has no headings stopped being "
                        "read. The fallback is the whole reason the table shape is recognised: "
                        "without it such a register is reported for its formatting")

    # Every id in a heading, not the first: an entry moved down to a product register leaves a
    # line naming two, and taking the first made a finding depend on the order of the prose.
    two = fm(schema="framework/open-register/v1", artifact_type="open-register",
             lifecycle="living", status="active", owners="[owner]", created="2026-01-01 09:00",
             last_review="2026-01-01 09:00",
             entries="\n  OD-038:\n    status: open\n    products: [all]\n"
                     "  OD-039:\n    status: open\n    products: [all]\n")
    got = codes("# Open\n\n### OD-038 and OD-039 moved down to the product register\n", two)
    if got is None or "REG015" in got:
        problems.append("a heading naming two entries declared only one of them, so which of "
                        "the two is reported depends on the order somebody wrote them in")
    return problems


@check("a document filed under one product and declaring another is reported, once")
def _product_of_the_directory():
    # `REG008` written for a document instead of for an entry. It is a code of its own and not
    # that one widened, because an annotation covers every finding sharing `(code, path)`
    # including ones that arrive later: growing `REG008` with a class nobody had in mind when
    # they annotated it would make those born explained by a reason about something else.
    root = built("cycle")
    if root is None:
        return ["the cycle fixture would not build"]
    stray = (root / "products" / "atlas" / "RMP.md").read_text(encoding="utf-8")
    problems = []
    with tempfile.TemporaryDirectory() as tmp:
        work = Path(tmp) / "repo"
        shutil.copytree(root, work)
        (work / "products" / "beta").mkdir(parents=True, exist_ok=True)
        (work / "products" / "beta" / "RMP.md").write_text(stray, encoding="utf-8")
        r = subprocess.run([sys.executable, str(VALIDATE), "--root", str(work), "--json",
                            "--stale-days", "36500"], capture_output=True, text=True)
        if r.returncode not in (0, 1) or not r.stdout.strip():
            return ["the validator crashed on a misfiled artifact"]
        found = [f for f in json.loads(r.stdout)["findings"] if f["code"] == "LOC002"]
    if len(found) != 1 or found[0]["path"] != "products/beta/RMP.md":
        problems.append(f"a roadmap under one product declaring another produced "
                        f"{[f['path'] for f in found]}: one finding, on the misfiled file")
    elif "beta" not in found[0]["message"] or "atlas" not in found[0]["message"]:
        problems.append("the finding does not name both the directory and the field, so a "
                        "reader cannot see which two statements disagree")
    if built("review") and [f for f in found if f["code"] != "LOC002"]:
        problems.append("something else was reported as well")
    return problems


@check("a change set that edits a living document without rereading it is one finding, not thirty")
def _change_set_review():
    # `LC006` measures the state; this stops it being reached, at the moment the edit is
    # proposed and addressed to whoever made it. One finding for the change set, because a
    # migration touches thirty documents and thirty findings on the pull request with the best
    # reason to exist is a wall in the worst place. And `LC006` goes quiet about what this
    # names, or the same document is reported twice on one run.
    root = built("review")
    if root is None:
        return ["the review fixture would not build"]
    problems = []
    with tempfile.TemporaryDirectory() as tmp:
        listing = Path(tmp) / "changed.txt"
        listing.write_text("products/atlas/PBR.md\n"
                           "products/atlas/contracts/DC-001-orders.md\n", encoding="utf-8")
        r = subprocess.run([sys.executable, str(VALIDATE), "--root", str(root), "--json",
                            "--stale-days", "36500", "--changed-files", str(listing)],
                           capture_output=True, text=True)
        if r.returncode not in (0, 1) or not r.stdout.strip():
            return ["the validator crashed with a change set"]
        out = json.loads(r.stdout)
    pr = [f for f in out["findings"] if f["code"] == "PR005"]
    lc = [f for f in out["findings"] if f["code"] == "LC006"]
    if len(pr) != 1:
        problems.append(f"a change set touching two unreread documents produced {len(pr)} "
                        "PR005: it is one finding for the change set, whatever the count")
    elif not all(n in pr[0]["message"] for n in ("PBR.md", "DC-001-orders.md")):
        problems.append("the finding does not name the documents, so it says a number and "
                        "sends nobody anywhere")
    if lc:
        problems.append(f"LC006 also reported {[f['path'] for f in lc]}. A document the change "
                        "set touches belongs to PR005 alone, or the pull request carries the "
                        "same document twice")
    # And without the flag it stays silent, like every other check in its family.
    r = subprocess.run([sys.executable, str(VALIDATE), "--root", str(root), "--json",
                        "--stale-days", "36500"], capture_output=True, text=True)
    if r.returncode in (0, 1) and r.stdout.strip():
        out = json.loads(r.stdout)
        if [f for f in out["findings"] if f["code"] == "PR005"]:
            problems.append("PR005 fired without a change set. It cannot know what a change "
                            "set touched without being told, and a check that guesses at that "
                            "reports every document in the repository")
        if len([f for f in out["findings"] if f["code"] == "LC006"]) != 2:
            problems.append("without the flag LC006 no longer reports both documents, so the "
                            "suppression is leaking into runs that have no change set")
    return problems


# ─────────────────────────────────────────────────────────────────────────────
# The attestations of rereading, and the history behind each line.
#
# THE BENCH IS A REPOSITORY BUILT IN THE TEST, ONE COMMIT AT A TIME, WITH DATES AND OFFSETS
# SAID OUT LOUD. The fixtures under `evals/` have histories too, but each was planted for its
# own story and none of them says "this value was written after its commit" or "this commit
# was rebased": those are the cases here, and a case that can be said in one sentence is
# written as one sentence of git. The framework's own checkout is never the subject, for the
# reason `_shallow_checkout` gives: a clone carries the committed code, so a fix asserted
# through one cannot fail until after it is committed.

ATTEST = ROOT / "skills" / "audit" / "scripts" / "attest.py"


def _fm(**kw) -> str:
    return "---\n" + "\n".join(f"{k}: {v}" for k, v in kw.items()) + "\n---\n\n"


def _brief(product: str, review: str, scope: str | None = None, body: str = "It does a thing.",
           created: str = "2026-06-01 09:00", comment: str | None = None) -> str:
    fields = dict(schema="framework/product-brief/v1", artifact_type="product-brief",
                  lifecycle="living", status="active", products=f"[{product}]", owners="[o]",
                  created=created, last_review=review)
    if scope is not None:
        fields["review_scope"] = scope
    text = _fm(**fields)
    if comment is not None:
        # A YAML comment beside the stamp, at column zero: text, by the rule under test.
        text = text.replace("\n---\n\n", f"\n# {comment}\n---\n\n", 1)
    return text + f"# Brief\n\n<!-- section: what -->\n## What it does\n\n{body}\n"


class _Repo:
    """A git repository with dated commits, for the checks that read the history behind a line."""

    def __init__(self, path: Path):
        self.path = path
        path.mkdir(parents=True, exist_ok=True)
        self.git("init", "-q", "-b", "main")
        self.git("config", "commit.gpgsign", "false")
        self.write("framework.yaml", f"framework_version: {REGISTRY['version']}\n")

    def git(self, *args: str, env: dict | None = None) -> str:
        e = dict(os.environ, GIT_AUTHOR_NAME="t", GIT_AUTHOR_EMAIL="t@example.com",
                 GIT_COMMITTER_NAME="t", GIT_COMMITTER_EMAIL="t@example.com",
                 # The person's own git configuration stays out: `merge.ff=only` or a merge
                 # driver there would change what these histories are.
                 GIT_CONFIG_GLOBAL=os.devnull, GIT_CONFIG_NOSYSTEM="1")
        e.update(env or {})
        r = subprocess.run(["git", "-C", str(self.path), *args], capture_output=True,
                           text=True, env=e)
        if r.returncode:
            raise RuntimeError(f"git {' '.join(args)} failed:\n{r.stderr}")
        return r.stdout.strip()

    def write(self, rel: str, text: str) -> None:
        p = self.path / rel
        p.parent.mkdir(parents=True, exist_ok=True)
        p.write_text(text, encoding="utf-8")

    def commit(self, message: str, authored: str, committed: str | None = None) -> str:
        self.git("add", "-A")
        self.git("commit", "-q", "-m", message,
                 env={"GIT_AUTHOR_DATE": authored, "GIT_COMMITTER_DATE": committed or authored})
        return self.git("rev-parse", "HEAD")

    def validate(self, *extra: str, env: dict | None = None):
        """(exit code, parsed JSON or None, stderr)."""
        r = subprocess.run([sys.executable, str(VALIDATE), "--root", str(self.path), "--json",
                            "--stale-days", "36500", *extra], capture_output=True, text=True,
                           env={**os.environ, **(env or {})})
        try:
            out = json.loads(r.stdout) if r.stdout.strip() else None
        except json.JSONDecodeError:
            out = None
        return r.returncode, out, r.stderr

    def findings(self, code: str, env: dict | None = None) -> list[dict]:
        rc, out, err = self.validate(env=env)
        if out is None:
            raise RuntimeError(f"the validator returned nothing readable (exit {rc}): {err}")
        return [f for f in out["findings"] if f["code"] == code]

    def attestation_codes(self, env: dict | None = None) -> set[str]:
        rc, out, err = self.validate(env=env)
        if out is None:
            raise RuntimeError(f"the validator returned nothing readable (exit {rc}): {err}")
        return {f["code"] for f in out["findings"]} & {"LC006", "LC007", "LC008"}


def _previous_version(before: str | None = None) -> str | None:
    """The version declared before the current one, read out of the registry's history.

    Never typed: a literal here is the line that goes wrong silently the day the framework
    moves, and the check on version literals below this file's own rule would report it.
    `before` asks for the version declared before another one instead: a case about the
    version that retired a code needs the one before *that*, which stops being the previous
    one the first time a patch ships on top of it -- as 4.0.1 did.
    """
    migrate = _load(ROOT / "skills" / "audit" / "scripts" / "migrate.py", "migrate_for_versions")
    now = migrate.semver(before or REGISTRY["version"])
    for sha in subprocess.run(["git", "-C", str(ROOT), "rev-list", "-60", "HEAD"],
                              capture_output=True, text=True).stdout.split():
        v = migrate.registry_version(subprocess.run(
            ["git", "-C", str(ROOT), "show", f"{sha}:{migrate.REGISTRY_REL}"],
            capture_output=True, text=True).stdout)
        if v and migrate.semver(v) and now and migrate.semver(v) < now:
            return v
    return None


@check("a last_review later than the commit that carries it is an error, compared at the minute in the commit's offset")
def _review_not_ahead():
    # THE CASE THIS IS FOR. Four living documents attested with instants after their own
    # commits, chosen apart from each other so that the check on shared minutes stayed quiet,
    # and nothing looked in that direction. `LC006` asks whether the text changed after the
    # reading; this asks whether the reading is claimed before its own record.
    problems = []
    with tempfile.TemporaryDirectory() as tmp:
        repo = _Repo(Path(tmp) / "r")
        doc = "products/p/PBR.md"
        repo.write(doc, _brief("p", "2026-06-01 09:00"))
        repo.commit("first", "2026-06-01T09:00:00+02:00")

        # 1. Written at 10:30, committed at 10:00:39: the line did not exist at 10:30.
        repo.write(doc, _brief("p", "2026-06-02 10:30", scope="the whole file"))
        repo.commit("attest ahead of the commit", "2026-06-02T10:00:39+02:00")
        rc, out, _ = repo.validate()
        found = [f for f in (out or {}).get("findings", []) if f["code"] == "LC007"]
        if rc != 1 or len(found) != 1:
            problems.append(f"a value thirty minutes after its own commit produced exit {rc} and "
                            f"{len(found)} LC007: a date that cannot be true is an error, once")
        elif "+02:00" not in found[0]["message"] or "commit" not in found[0]["message"]:
            problems.append("the finding does not say in which offset it read the bare value, "
                            "or which commit it compared with: a reader in another timezone "
                            "cannot check it")
        elif found[0]["level"] != "error":
            problems.append(f"LC007 came out at {found[0]['level']!r}: it is an error, because "
                            "there is no reading to redo")

        # 2. Equality at the minute is not "after": 10:00 against a commit at 10:00:39 is the
        # honest procedure, and reading the value as the end of its minute would report it.
        repo.write(doc, _brief("p", "2026-06-02 10:00", scope="the whole file"))
        repo.commit("attest in the same minute", "2026-06-02T10:00:39+02:00")
        if repo.findings("LC007"):
            problems.append("a value in the same minute as its commit was reported: the "
                            "comparison is at the minute and equality is not after")

        # 6. A bare value is read in the offset of the commit that carries it, wherever the
        # validator runs from: 18:00 written at +09:00, checked from UTC, is not 09:00.
        repo.write(doc, _brief("p", "2026-06-03 18:00", scope="the whole file, in Tokyo"))
        repo.commit("attest at +09:00", "2026-06-03T18:00:20+09:00")
        if repo.findings("LC007", env={"TZ": "UTC"}):
            problems.append("a bare value equal to its commit's local time was reported when "
                            "the validator ran at UTC: the value was read on the machine's "
                            "clock instead of the commit's")
        repo.write(doc, _brief("p", "2026-06-03 18:01", scope="the whole file, in Tokyo"))
        repo.commit("attest one minute ahead at +09:00", "2026-06-03T18:00:20+09:00")
        if not repo.findings("LC007", env={"TZ": "UTC"}):
            problems.append("a bare value one minute after its commit at +09:00 was not "
                            "reported from UTC: read on the machine's clock it looks nine "
                            "hours early, which is the verdict moving with the machine")

        # 5. A value with its own offset is an instant, and no clock reads it.
        repo.write(doc, _brief("p", "2026-06-04 18:00 +09:00", scope="the whole file"))
        repo.commit("attest with an offset", "2026-06-04T18:00:20+09:00")
        if repo.findings("LC007", env={"TZ": "UTC"}):
            problems.append("a value carrying its own offset, equal to its commit, was "
                            "reported from UTC")
        repo.write(doc, _brief("p", "2026-06-04 18:05 +09:00", scope="the whole file"))
        repo.commit("attest ahead with an offset", "2026-06-04T18:00:20+09:00")
        if not repo.findings("LC007", env={"TZ": "Asia/Tokyo"}):
            problems.append("a value with an offset five minutes after its commit was not "
                            "reported: the offset made the value an instant and the "
                            "comparison did not use it")

        # 12. A rebase moves the committer date forward and leaves the value where it was:
        # the bound loosens, and a true value stays true. `--amend` is the smallest rebase.
        repo.write(doc, _brief("p", "2026-06-05 10:00 +02:00", scope="the whole file"))
        repo.commit("attest, then rebase", "2026-06-05T10:00:10+02:00")
        repo.git("commit", "-q", "--amend", "--no-edit",
                 env={"GIT_COMMITTER_DATE": "2026-06-20T08:00:00+02:00"})
        if repo.findings("LC007"):
            problems.append("a rebase that moved the committer date forward produced LC007: "
                            "the value can only become further behind its commit")

        # 3 and 4. A line not yet committed is compared with the clock, and the finding says
        # so and names the machine's offset. 2099 rather than "now plus a minute": the test
        # has to hold whatever the clock says.
        repo.write(doc, _brief("p", "2099-01-01 09:00", scope="the whole file"))
        found = repo.findings("LC007", env={"TZ": "UTC"})
        if len(found) != 1:
            problems.append(f"an uncommitted value in 2099 produced {len(found)} LC007: the "
                            "reference for a line not yet committed is the current time")
        elif "not committed" not in found[0]["message"] or "+00:00" not in found[0]["message"]:
            problems.append("the finding on an uncommitted line does not say that it compared "
                            "with the clock, in this machine's offset: the container case is "
                            "then a mystery instead of a diagnosis")
        repo.write(doc, _brief("p", "2026-06-05 10:00 +02:00", scope="the whole file"))
        if repo.findings("LC007"):
            problems.append("an uncommitted value in the past was reported")
        repo.write(doc, _brief("p", "2099-01-01 09:00 +09:00", scope="the whole file"))
        found = repo.findings("LC007")
        if len(found) != 1 or "own offset +09:00" not in found[0]["message"]:
            problems.append(f"an uncommitted value carrying its own offset produced {found}: "
                            "one LC007 saying the offset is the value's and not the machine's")

        # A bare date is midnight, and midnight the day after a commit at 23:00 is after it.
        repo.write(doc, _brief("p", "2026-06-06", scope="the whole file"))
        repo.commit("attest with a date at 23:00 the evening before", "2026-06-05T23:00:00+02:00")
        if not repo.findings("LC007"):
            problems.append("a bare date one day after the commit that carries it was not "
                            "reported: the date alone counts as midnight, which is after 23:00")
        repo.write(doc, _brief("p", "2026-06-05", scope="the whole file, again"))
        repo.commit("attest with the date of the commit", "2026-06-05T23:30:00+02:00")
        if repo.findings("LC007"):
            problems.append("a bare date equal to the commit's day was reported")

        # A bare `.yaml` artifact is front matter throughout, and its stamp is read the same.
        repo.write("products/p/product.yaml",
                   "schema: framework/product-manifest/v1\nartifact_type: product-manifest\n"
                   "lifecycle: living\nstatus: active\nproducts: [p]\nname: P\n"
                   "one_liner: one\nowners: [o]\ncreated: 2026-06-01\n"
                   "last_review: 2026-06-07 12:00\nreview_scope: the whole manifest\n"
                   "stage:\n  phase: F4\n  since: 2026-06-01\n")
        repo.commit("attest the manifest ahead of its commit", "2026-06-07T11:00:00+02:00")
        found = [f for f in repo.findings("LC007") if f["path"].endswith("product.yaml")]
        if len(found) != 1:
            problems.append(f"a bare .yaml artifact attested an hour ahead of its commit "
                            f"produced {len(found)} LC007: its stamp is read like any other")
    return problems


@check("a file HEAD does not hold is compared with the clock, and a reformatting commit does not write the stamp")
def _history_edges():
    # TWO EVERYDAY STATES THAT ARE NOT IN ANY COMMIT: a file added to the index and not yet
    # committed, and a repository with files staged and no commit at all. Both are in
    # `git ls-files` and in no commit, and `git blame` on either fails; the first version
    # read the index and stopped there with "nothing was checked", where the docstring
    # promised "simply not committed". And two everyday commits that do not write the stamp:
    # one that re-indents the line, one that moves it among the other keys.
    problems = []
    with tempfile.TemporaryDirectory() as tmp:
        repo = _Repo(Path(tmp) / "r")
        repo.write("products/p/PBR.md", _brief("p", "2026-06-01 09:00"))
        repo.git("add", "-A")            # staged, and no commit yet
        rc, out, err = repo.validate()
        if rc not in (0, 1) or out is None:
            problems.append(f"a repository with staged files and no commit stopped the "
                            f"validator (exit {rc}): {err.strip()[-200:]}")
        elif [f for f in out["findings"] if f["code"] in ("LC006", "LC007", "LC008")]:
            problems.append("a day-one set staged and not committed was reported")
        repo.write("products/p/PBR.md", _brief("p", "2099-01-01 09:00"))
        if not repo.findings("LC007"):
            problems.append("a value in 2099 in a repository with no commit was not compared "
                            "with the clock")
        repo.write("products/p/PBR.md", _brief("p", "2026-06-01 09:00"))
        repo.commit("first", "2026-06-01T09:00:00+02:00")
        repo.write("products/q/PBR.md", _brief("q", "2026-06-02 09:00"))
        repo.git("add", "-A")            # a new file staged above a history
        rc, out, err = repo.validate()
        if rc not in (0, 1) or out is None:
            problems.append(f"a new file staged and not committed stopped the validator "
                            f"(exit {rc}): {err.strip()[-200:]}")
        elif [f for f in out["findings"] if f["code"] == "LC007"]:
            problems.append("a staged file with a past stamp produced LC007")
        repo.commit("second product", "2026-06-02T09:00:00+02:00")

        # A reformatting commit: the stamp line re-indented and moved among the keys. Neither
        # wrote the value, so neither is read as the attestation, and the day-one document
        # owes no note afterwards.
        text = (repo.path / "products/q/PBR.md").read_text(encoding="utf-8")
        head, body = text.split("\n---\n", 1)
        fields = head.split("\n")[1:]
        moved = [l for l in fields if not l.startswith("last_review:")]
        moved.insert(2, next(l for l in fields if l.startswith("last_review:")))
        repo.write("products/q/PBR.md", "---\n" + "\n".join(moved) + "\n---\n" + body)
        repo.commit("reorder the keys", "2026-06-20T09:00:00+02:00")
        found = repo.findings("LC008")
        if [f for f in found if f["path"] == "products/q/PBR.md"]:
            problems.append("a commit that moved the `last_review` line among the keys was "
                            "read as the commit that wrote it, and a day-one document was "
                            "asked for a note")
        # And blame's rename following: the file moves, and is not born again at the move.
        repo.git("mv", "products/q/PBR.md", "products/q/PBR-renamed.md")
        repo.commit("rename", "2026-06-21T09:00:00+02:00")
        if [f for f in repo.findings("LC008") if "renamed" in f["path"]]:
            problems.append("a renamed day-one document was asked for a note: the rename did "
                            "not write the stamp")
    return problems


@check("a change to the attested text after the stamp is reported, an attestation alone is not, and a YAML comment is text")
def _review_gap_on_the_block():
    # THE BLOCK IS STRUCTURAL. `last_review` and `review_scope`, with their continuation
    # lines, are what a reading may write; everything else is text, and a comment beside the
    # stamp is text. In the repository that produced this version the comment was where the
    # coverage of each reading was written, so every honest attestation reported itself, and
    # the way out that was taken -- committing the comment first and the stamp after -- cost
    # two commits per reading. The field costs none.
    problems = []
    with tempfile.TemporaryDirectory() as tmp:
        repo = _Repo(Path(tmp) / "r")
        doc = "products/p/PBR.md"
        repo.write(doc, _brief("p", "2026-06-01 09:00"))
        repo.commit("first", "2026-06-01T09:00:00+02:00")

        # 8. The text changes and nobody rereads it.
        repo.write(doc, _brief("p", "2026-06-01 09:00", body="It does two things."))
        repo.commit("change the text", "2026-06-10T09:00:00+02:00")
        found = repo.findings("LC006")
        if len(found) != 1 or "9 day" not in found[0]["message"]:
            problems.append(f"a text changed nine days after its stamp produced {found}: one "
                            "LC006 stating the gap")

        # 7. The honest procedure: content committed, then the stamp and the note alone.
        repo.write(doc, _brief("p", "2026-06-10 10:00 +02:00", scope="the whole file",
                               body="It does two things."))
        repo.commit("attest", "2026-06-10T10:05:00+02:00")
        if repo.attestation_codes():
            problems.append(f"the honest procedure -- content committed, then the stamp and "
                            f"the note in a commit of their own -- produced "
                            f"{sorted(repo.attestation_codes())}: it has to produce nothing, or "
                            "there is no honest way to attest")

        # 9. A YAML comment edited after the reading is a change to text nobody reread.
        repo.write(doc, _brief("p", "2026-06-10 10:00 +02:00", scope="the whole file",
                               body="It does two things.", comment="covered §1 only"))
        repo.commit("edit the comment", "2026-06-11T09:00:00+02:00")
        found = repo.findings("LC006")
        if len(found) != 1:
            problems.append(f"a YAML comment added after the stamp produced {len(found)} LC006: "
                            "a comment is text, because the check cannot tell a comment about "
                            "the reading from a comment about the document")

        # Reread after the comment, then 10: change the note alone. The walk steps over it.
        repo.write(doc, _brief("p", "2026-06-11 12:00 +02:00", scope="the whole file",
                               body="It does two things.", comment="covered §1 only"))
        repo.commit("attest again", "2026-06-11T12:05:00+02:00")
        repo.write(doc, _brief("p", "2026-06-11 12:00 +02:00",
                               scope="the whole file, and §1 in particular",
                               body="It does two things.", comment="covered §1 only"))
        repo.commit("reword the note", "2026-06-12T09:00:00+02:00")
        if repo.findings("LC006"):
            problems.append("a commit that changed `review_scope` alone was counted as a change "
                            "to the attested text: the field is the attestation block")
        # 11. Two attestations in a row above old content: the walk has to take two steps.
        # With a new sentence, or `LC008` correctly reports the note copied forward.
        repo.write(doc, _brief("p", "2026-06-12 10:00 +02:00",
                               scope="the whole file again, a day later",
                               body="It does two things.", comment="covered §1 only"))
        repo.commit("attest a third time", "2026-06-12T10:05:00+02:00")
        if repo.attestation_codes():
            problems.append(f"three attestation-only commits above one content commit produced "
                            f"{sorted(repo.attestation_codes())}: the comparison steps back "
                            "over every attestation, not only the newest")

        # A commit that moves the stamp AND the text counts, and the finding says why.
        repo.write(doc, _brief("p", "2026-06-13 08:00 +02:00",
                               scope="the whole file, and §1 in particular",
                               body="It does three things.", comment="covered §1 only"))
        repo.commit("change and stamp together", "2026-06-13T09:00:00+02:00")
        found = repo.findings("LC006")
        if len(found) != 1 or "changed something else as well" not in found[0]["message"]:
            problems.append("a commit that moved the stamp and changed the text was not "
                            "reported with the sentence saying why it was not excluded")

        # 12. A rebase: committer dates move past the stamp, author dates stay. Author date
        # is what says when the text changed, and the test proves the other choice would fire.
        repo.write(doc, _brief("p", "2026-06-13 08:00 +02:00", scope="the whole file",
                               body="It does four things.", comment="covered §1 only"))
        repo.commit("change the text", "2026-06-14T09:00:00+02:00")
        repo.write(doc, _brief("p", "2026-06-14 10:00 +02:00", scope="the whole file, reread",
                               body="It does four things.", comment="covered §1 only"))
        repo.commit("attest", "2026-06-14T10:05:00+02:00")
        repo.git("rebase", "-q", "-f", "HEAD~2", env={"GIT_COMMITTER_DATE": "2026-06-20T08:00:00+02:00"})
        content_committed = repo.git("log", "-1", "--format=%cI", "HEAD~1")
        if not content_committed.startswith("2026-06-20"):
            problems.append(f"the rebase did not move the committer date ({content_committed}), "
                            "so the case below proves nothing")
        if repo.attestation_codes():
            problems.append(f"a rebase after the attestation produced "
                            f"{sorted(repo.attestation_codes())}: the committer dates moved and "
                            "the author dates did not, and LC006 reads the author's while LC007 "
                            "reads the committer's")

        # An indented comment under `last_review:` is text, however it is indented.
        repo.write(doc, _brief("p", "2026-06-14 10:00 +02:00", scope="the whole file, reread",
                               body="It does four things.", comment="covered §1 only")
                   .replace("\nreview_scope:", "\n  # an indented comment under the stamp\nreview_scope:", 1))
        repo.commit("add an indented comment under the stamp", "2026-06-15T09:00:00+02:00")
        if not repo.findings("LC006"):
            problems.append("an indented comment added under `last_review` was read as part "
                            "of the attestation block: a comment is text wherever it sits")
        repo.write(doc, _brief("p", "2026-06-15 10:00 +02:00", scope="the whole file, after the comment",
                               body="It does four things.", comment="covered §1 only")
                   .replace("\nreview_scope:", "\n  # an indented comment under the stamp\nreview_scope:", 1))
        repo.commit("attest after the comment", "2026-06-15T10:05:00+02:00")
        # A block scalar written by hand, with a blank line after it, changed alone.
        hand = (repo.path / doc).read_text(encoding="utf-8").replace(
            "review_scope: the whole file, after the comment\n",
            "review_scope: >-\n  the whole file, after the comment,\n  written by hand\n\n", 1)
        repo.write(doc, hand)
        repo.commit("reword the note as a block", "2026-06-16T09:00:00+02:00")
        if repo.attestation_codes():
            problems.append(f"a `review_scope` rewritten by hand as a block scalar produced "
                            f"{sorted(repo.attestation_codes())}: the block is the note")

        # 19. Adoption: the note arrives later, in a commit that moves nothing else -- the
        # indented comment planted above stays where it is, or its removal is a change.
        repo.write(doc, _brief("p", "2026-06-15 10:00 +02:00", scope="the whole file, reread, "
                               "written after the fact", body="It does four things.",
                               comment="covered §1 only")
                   .replace("\nreview_scope:", "\n  # an indented comment under the stamp\nreview_scope:", 1))
        repo.commit("add the note", "2026-06-21T09:00:00+02:00")
        if repo.attestation_codes():
            problems.append(f"adding `review_scope` without moving `last_review` produced "
                            f"{sorted(repo.attestation_codes())}: the migration step has to be "
                            "free of findings or nobody takes it")
    return problems


WORK = ROOT / "skills" / "audit" / "scripts" / "work.py"


def _open_register(entries: dict, review: str = "2026-06-01 09:00", parked: str = "",
                   union: str = "old union") -> str:
    """A root register: the map, a parking lot with room at its tail, and the union."""
    rows = "".join(f"  {k}:\n    status: {s}\n    cost_to_reverse: low\n    default_in_force: none\n"
                   f"    products: [all]\n" for k, s in entries.items())
    return ("---\nschema: framework/open-register/v1\nartifact_type: open-register\n"
            "lifecycle: living\nstatus: active\nproducts: [all]\nowners: [o]\n"
            f"created: 2026-06-01 09:00\nlast_review: {review}\nclassification: internal\n"
            "entries:\n" + rows + "---\n\n# Open\n\n## Parking lot\n\n- first line\n" + parked +
            f"\n## Union\n\n<!-- generated: open-union -->\n{union}\n<!-- /generated -->\n")


def _decision(i: str, derives: str, **extra) -> str:
    meta = dict(schema="framework/decision-record/v1", artifact_type="decision-record", id=i,
                lifecycle="immutable", status="accepted", scope="architecture",
                products="[all]", owners="[o]", created="2026-06-02",
                derives_from=f"[{derives}]")
    meta.update(extra)
    return _fm(**meta) + f"# {i}\n"


@check("a merge records no reading: what it produced is a change, a side taken whole is not, and the union the generator owns is not attested text")
def _review_gap_on_merges():
    # Two branches that both touch a living document meet in a merge, and before `--cc` the
    # merge listed no file: the walk went past it to an older commit on either side, so the
    # combination of two texts -- or one side's prose under the other side's stamp -- was
    # attested by whichever stamp happened to survive. And read in log order, the two sides
    # interleaved: a side whose version a resolution threw away was named as the change, and a
    # side that deleted the file silenced the check on the version that was kept.
    problems = []
    v = _load(VALIDATE, "validate_for_merges")
    attest = _load(ATTEST, "attest_for_merges")
    body = "Alpha.\n\n\n\nBeta.\n\n\n\nGamma."
    doc = "products/p/PBR.md"
    D = lambda when: {"GIT_AUTHOR_DATE": when, "GIT_COMMITTER_DATE": when}

    def text(review: str, scope: str, words: str = body) -> str:
        return _brief("p", review, scope=scope, body=words)

    def edit(repo: _Repo, rel: str, *pairs: tuple[str, str]) -> None:
        t = (repo.path / rel).read_text(encoding="utf-8")
        for old, new in pairs:
            if old not in t:
                raise RuntimeError(f"the case is not what it says: {old!r} is not in {rel}")
            t = t.replace(old, new, 1)
        repo.write(rel, t)

    def fresh(tmp: str, name: str) -> _Repo:
        repo = _Repo(Path(tmp) / name)
        repo.write(doc, text("2026-06-01 09:00", "the whole file"))
        repo.write("products/p/notes.txt", "one\n")
        repo.commit("first", "2026-06-01T09:00:00+02:00")
        return repo

    def attest_main(repo: _Repo, rel: str, scope: str):
        return subprocess.run([sys.executable, str(ATTEST), "--root", str(repo.path), rel,
                               "--scope", scope], capture_output=True, text=True)

    with tempfile.TemporaryDirectory() as tmp:
        # 1. The branch changes and attests; the main line moves elsewhere. The merge takes the
        # document whole from the branch, and the branch attested it.
        repo = fresh(tmp, "taken-whole")
        repo.git("checkout", "-q", "-b", "work")
        edit(repo, doc, ("Alpha.", "Alpha, on the branch."))
        repo.commit("branch: text", "2026-06-02T09:00:00+02:00")
        edit(repo, doc, ("last_review: 2026-06-01 09:00", "last_review: 2026-06-02 10:00 +02:00"),
             ("review_scope: the whole file", "review_scope: the whole file, after Alpha"))
        repo.commit("branch: attest", "2026-06-02T10:05:00+02:00")
        repo.git("checkout", "-q", "main")
        repo.write("products/p/notes.txt", "two\n")
        repo.commit("main: another file", "2026-06-03T09:00:00+02:00")
        repo.git("checkout", "-q", "work")
        repo.git("merge", "-q", "--no-edit", "main", env=D("2026-06-04T09:00:00+02:00"))
        if repo.attestation_codes():
            problems.append(f"a merge that took the document whole from the branch produced "
                            f"{sorted(repo.attestation_codes())}")

        # 2. The main line changes the same document elsewhere; the merge is clean and what it
        # produced is neither side's text. LC006 names the merge, attest.py names the same
        # commit, and PR005 says that it is a merge.
        repo.git("checkout", "-q", "main")
        edit(repo, doc, ("Gamma.", "Gamma, on main."))
        repo.commit("main: text", "2026-06-05T09:00:00+02:00")
        repo.git("checkout", "-q", "work")
        repo.git("merge", "-q", "--no-edit", "main", env=D("2026-06-06T09:00:00+02:00"))
        merged = repo.git("rev-parse", "HEAD")
        found = repo.findings("LC006")
        if len(found) != 1 or "is a merge" not in found[0]["message"] \
                or merged[:12] not in found[0]["message"]:
            problems.append(f"a clean merge of two changes to one attested document produced "
                            f"{found}: one LC006 naming the merge")
        r = v.Repository.open(repo.path)
        named = attest.content_commit(r, doc)
        r.close()
        if not named or named[0] != merged or "a merge" not in named[2]:
            problems.append(f"attest.py would say it attests {named}: the merge LC006 names")
        changed = Path(tmp) / "changed.txt"
        changed.write_text(doc + "\n")
        rc, out, _ = repo.validate("--changed-files", str(changed))
        pr005 = [f for f in (out or {}).get("findings", []) if f["code"] == "PR005"]
        if len(pr005) != 1 or "changed by a merge" not in pr005[0]["message"]:
            problems.append(f"PR005 on a change set whose document a merge changed said "
                            f"{pr005}: that it was a merge, which records no reading")

    with tempfile.TemporaryDirectory() as tmp:
        # 3. One side's prose under the other side's stamp, in a clean merge: the branch changes
        # the text and does not reread, the main line attests its own version and changes no
        # text. Neither side moved what the other did, so git merges without a word.
        repo = fresh(tmp, "two-halves")
        repo.git("checkout", "-q", "-b", "work")
        edit(repo, doc, ("Beta.", "Beta, changed on the branch and not reread."))
        repo.commit("branch: text", "2026-06-02T09:00:00+02:00")
        repo.git("checkout", "-q", "main")
        edit(repo, doc, ("last_review: 2026-06-01 09:00", "last_review: 2026-06-03 10:00 +02:00"),
             ("review_scope: the whole file", "review_scope: the whole file, read on main"))
        repo.commit("main: attest", "2026-06-03T10:05:00+02:00")
        repo.git("checkout", "-q", "work")
        repo.git("merge", "-q", "--no-edit", "main", env=D("2026-06-04T09:00:00+02:00"))
        t = (repo.path / doc).read_text(encoding="utf-8")
        if "2026-06-03 10:00" not in t or "not reread" not in t:
            problems.append("the clean merge did not produce the branch's text under main's stamp")
        elif not repo.findings("LC006"):
            problems.append("a merge that put main's stamp on the branch's unread text was read "
                            "as a reading")
        r = attest_main(repo, doc, "the whole file, after the merge")
        if r.returncode:
            problems.append(f"attest.py refused a document after a committed merge: {r.stderr}")
        repo.git("checkout", "-q", "--", doc)

        # attest.py refuses while a merge is in progress.
        repo.git("checkout", "-q", "main")
        edit(repo, doc, ("Gamma.", "Gamma, on main again."))
        repo.commit("main: text", "2026-06-05T09:00:00+02:00")
        repo.git("checkout", "-q", "work")
        repo.git("merge", "-q", "--no-commit", "--no-ff", "main", env=D("2026-06-06T09:00:00+02:00"))
        r = attest_main(repo, doc, "the whole file, in the middle of a merge")
        if r.returncode != 2 or "merge is in progress" not in r.stderr:
            problems.append(f"attest.py during a merge gave exit {r.returncode}: it has to refuse")
        repo.git("merge", "--abort")

    with tempfile.TemporaryDirectory() as tmp:
        # 4. A resolution that keeps the branch's version whole, while the main line's newer
        # change is thrown away: the file came down the branch, and main's commit is not in it.
        repo = fresh(tmp, "thrown-away")
        repo.git("checkout", "-q", "-b", "work")
        edit(repo, doc, ("Beta.", "Beta, on the branch."))
        repo.commit("branch: text", "2026-06-02T09:00:00+02:00")
        edit(repo, doc, ("last_review: 2026-06-01 09:00", "last_review: 2026-06-02 10:00 +02:00"),
             ("review_scope: the whole file", "review_scope: the whole file, after Beta"))
        repo.commit("branch: attest", "2026-06-02T10:05:00+02:00")
        repo.git("checkout", "-q", "main")
        edit(repo, doc, ("Beta.", "Beta, on main, later."))
        repo.commit("main: text", "2026-06-05T09:00:00+02:00")
        repo.git("checkout", "-q", "work")
        repo.git("merge", "-q", "-s", "ours", "--no-edit", "main", env=D("2026-06-06T09:00:00+02:00"))
        if repo.attestation_codes():
            problems.append(f"a merge that kept the branch's version whole named the main line's "
                            f"discarded commit: {sorted(repo.attestation_codes())}")

    with tempfile.TemporaryDirectory() as tmp:
        # 5. The main line deletes the document; the branch changed it and did not reread; the
        # merge keeps the branch's version. The deletion is not on the line the text came down.
        repo = fresh(tmp, "kept-after-delete")
        repo.git("checkout", "-q", "-b", "work")
        edit(repo, doc, ("Beta.", "Beta, on the branch, not reread."))
        repo.commit("branch: text", "2026-06-02T09:00:00+02:00")
        repo.git("checkout", "-q", "main")
        repo.git("rm", "-q", doc)
        repo.commit("main: delete", "2026-06-05T09:00:00+02:00")
        repo.git("checkout", "-q", "work")
        try:
            repo.git("merge", "-q", "--no-edit", "main", env=D("2026-06-06T09:00:00+02:00"))
        except RuntimeError:
            repo.git("add", doc)
            repo.git("commit", "-q", "--no-edit", env=D("2026-06-06T09:00:00+02:00"))
        if not (repo.path / doc).exists():
            problems.append("the modify/delete case did not keep the document")
        elif not repo.findings("LC006"):
            problems.append("a document kept after the other side deleted it, changed and not "
                            "reread, was silent: the deletion is not on its line")

    with tempfile.TemporaryDirectory() as tmp:
        # 6. The real union: an open register at the root, regenerated by `--emit-index` when
        # an entry is added to another register. A branch attested the root; the main line
        # added an entry elsewhere; the merge brings main's union and nothing else.
        repo = _Repo(Path(tmp) / "union")
        repo.write("OPEN.md", _open_register({"OD-001": "open"}))
        repo.write("platform/OPEN.md", _open_register({"OD-002": "open"}).replace(
            "products: [all]\nowners", "products: [none]\nowners", 1))
        repo.validate("--emit-index")
        repo.commit("first", "2026-06-01T09:00:00+02:00")
        repo.git("checkout", "-q", "-b", "work")
        edit(repo, "OPEN.md", ("last_review: 2026-06-01 09:00", "last_review: 2026-06-02 10:00 +02:00"))
        repo.commit("branch: attest", "2026-06-02T10:05:00+02:00")
        repo.git("checkout", "-q", "main")
        edit(repo, "platform/OPEN.md", ("  OD-002:\n", "  OD-003:\n    status: open\n    cost_to_reverse: low\n"
                                         "    default_in_force: none\n    products: [all]\n  OD-002:\n"))
        repo.validate("--emit-index")
        repo.commit("main: an entry, and the union regenerated", "2026-06-03T09:00:00+02:00")
        repo.git("checkout", "-q", "work")
        repo.git("merge", "-q", "--no-edit", "main", env=D("2026-06-04T09:00:00+02:00"))
        root = (repo.path / "OPEN.md").read_text(encoding="utf-8")
        if "OD-003" not in root:
            problems.append("the merge did not bring main's union, so the case proves nothing")
        elif [f for f in repo.findings("LC006") if f["path"] == "OPEN.md"]:
            problems.append("a merge equal to the branch in everything but main's union made the "
                            "root register stale")
        if repo.findings("GEN001"):
            problems.append("GEN001 on a union that is what the registers derive to")
        repo.write("OPEN.md", root.replace("OD-003", "OD-999", 1))
        repo.commit("edit the union by hand", "2026-06-05T09:00:00+02:00")
        if not repo.findings("GEN001"):
            problems.append("a union edited by hand was checked by nothing on the ordinary run")

    with tempfile.TemporaryDirectory() as tmp:
        # 7. Only the union, in the file the generator writes it into, when that file is the
        # register the generator writes into.
        repo = _Repo(Path(tmp) / "names")
        region = lambda name, body: f"<!-- generated: {name} -->\n{body}\n<!-- /generated -->"
        repo.write("OPEN.md", _open_register({"OD-001": "open"}).replace(
            "<!-- /generated -->\n", "<!-- /generated -->\n\n" + region("invented", "x") + "\n", 1))
        repo.write("products/p/OPEN.md", _open_register({"OD-002": "open"}))
        repo.commit("first", "2026-06-01T09:00:00+02:00")
        edit(repo, "OPEN.md", ("x\n<!-- /generated -->", "written by hand\n<!-- /generated -->"))
        edit(repo, "products/p/OPEN.md", ("old union", "written by hand"))
        repo.commit("edit an invented region and a product's copy of the union", "2026-06-02T09:00:00+02:00")
        found = {f["path"] for f in repo.findings("LC006")}
        if found != {"OPEN.md", "products/p/OPEN.md"}:
            problems.append(f"LC006 fired on {sorted(found)}: text between markers the generator "
                            "does not own, or in a file it does not write into, is attested text")
        brief = _brief("p", "2026-06-01 09:00", body=region("open-union", "one"))
        other = _Repo(Path(tmp) / "not-a-register")
        other.write("OPEN.md", brief)
        other.commit("first", "2026-06-01T09:00:00+02:00")
        other.write("OPEN.md", brief.replace("\none\n", "\ntwo\n", 1))
        other.commit("edit", "2026-06-02T09:00:00+02:00")
        if not other.findings("LC006"):
            problems.append("the union's markers in a root OPEN.md that is not an open register "
                            "were left out of the attested text: the generator writes nothing there")
        r = subprocess.run([sys.executable, str(ATTEST), "--root", str(repo.path), "OPEN.md",
                            "--scope", "the whole file"], capture_output=True, text=True)
        repo.git("checkout", "-q", "--", "OPEN.md")
        edit(repo, "OPEN.md", ("old union", "a union not yet committed"))
        r = subprocess.run([sys.executable, str(ATTEST), "--root", str(repo.path), "OPEN.md",
                            "--scope", "the whole file, again"], capture_output=True, text=True)
        if r.returncode != 2:
            problems.append("attest.py accepted a working tree whose union differs from HEAD: "
                            "commit the regenerated union first, then attest")

    # The list the exclusion reads and the list the generator writes, both ways.
    with tempfile.TemporaryDirectory() as tmp:
        root = Path(tmp)
        register = v.Artifact(root / "OPEN.md", "OPEN.md",
                              {"artifact_type": "open-register", "entries": {}}, "")
        written = {str(p.relative_to(root)): set(names)
                   for p, names in v.build_regions(root, [register]).items()}
        declared = {rel: set(names) for rel, (_, names) in v.GENERATED_REGIONS.items()}
        if written != declared:
            problems.append(f"build_regions writes {written} and GENERATED_REGIONS names "
                            f"{declared}: an excluded region the generator does not write is "
                            "text no check reads")
    return problems


@check("several documents attested in one commit with their own notes produce nothing, and the same note produces one finding")
def _honest_batch_and_the_pasted_note():
    # PROPERTY 2, WITH SEVERAL DOCUMENTS. The old check on shared minutes made this impossible:
    # three documents attested in one commit share an instant by construction. What separates
    # a reading from a stamp now is the sentence, and this is the shape it takes.
    problems = []
    with tempfile.TemporaryDirectory() as tmp:
        repo = _Repo(Path(tmp) / "r")
        docs = [f"products/p{i}/PBR.md" for i in range(3)]
        for i, d in enumerate(docs):
            repo.write(d, _brief(f"p{i}", "2026-06-01 09:00"))
        repo.commit("first", "2026-06-01T09:00:00+02:00")
        for i, d in enumerate(docs):
            repo.write(d, _brief(f"p{i}", "2026-06-01 09:00", body=f"It does thing {i}."))
        repo.commit("change all three", "2026-06-10T09:00:00+02:00")
        # 7. One commit, three stamps at the same instant, three sentences.
        notes = ["§1 and the new line on thing 0", "the whole file", "§2 only; §1 not reread"]
        for i, d in enumerate(docs):
            repo.write(d, _brief(f"p{i}", "2026-06-10 10:00 +02:00", scope=notes[i],
                                 body=f"It does thing {i}."))
        repo.commit("attest all three", "2026-06-10T10:05:00+02:00")
        if repo.attestation_codes():
            problems.append(f"three documents attested honestly in one commit produced "
                            f"{sorted(repo.attestation_codes())}: a shared minute is not a "
                            "finding any more, the sentence is what is read")
        # 16. One sentence pasted onto three documents: one finding, on the first path, naming
        # the others, so that an annotation has one place to go.
        for i, d in enumerate(docs):
            repo.write(d, _brief(f"p{i}", "2026-06-11 10:00 +02:00", scope="reread",
                                 body=f"It does thing {i}."))
        repo.commit("attest all three with one sentence", "2026-06-11T10:05:00+02:00")
        found = repo.findings("LC008")
        if len(found) != 1 or found[0]["path"] != docs[0]:
            problems.append(f"one sentence on three documents in one commit produced "
                            f"{[(f['path']) for f in found]}: one finding, on the first path")
        elif not all(d in found[0]["message"] for d in docs[1:]):
            problems.append("the finding does not name the other documents sharing the sentence")
        elif "proxy" not in found[0]["message"]:
            problems.append("the finding does not say it is a proxy: what it cannot catch has "
                            "to be said where it is read")
    return problems


@check("a review_scope that is empty after a reading, or copied forward, is reported, and day one is not")
def _review_scope():
    # `LC008` reads the sentence beside the stamp. Not owed on day one: `start` writes the
    # whole first set in one session and every document is born attesting itself. In the
    # history that is a stamp written by the commit that created the file; outside any
    # history it is a stamp on the same day as `created`.
    problems = []
    with tempfile.TemporaryDirectory() as tmp:
        repo = _Repo(Path(tmp) / "r")
        doc = "products/p/PBR.md"
        # 15. Day one: created and attested by the same commit, no note owed.
        repo.write(doc, _brief("p", "2026-06-01 10:00", created="2026-06-01"))
        repo.commit("first", "2026-06-01T10:00:30+02:00")
        if repo.findings("LC008"):
            problems.append("a document born attesting itself was asked for a note: a creation "
                            "is not a reading, and this fires on every new repository")
        # 14. Reread without a sentence.
        repo.write(doc, _brief("p", "2026-06-05 10:00 +02:00", created="2026-06-01"))
        repo.commit("attest without a note", "2026-06-05T10:05:00+02:00")
        found = repo.findings("LC008")
        if len(found) != 1 or "empty" not in found[0]["message"]:
            problems.append(f"a stamp with no `review_scope` after a reading produced {found}: "
                            "one finding saying the note is empty")
        # 18. The note arrives for the first time: absent at the parent is not "unchanged".
        repo.write(doc, _brief("p", "2026-06-06 10:00 +02:00", created="2026-06-01",
                               scope="the whole file"))
        repo.commit("attest with a note", "2026-06-06T10:05:00+02:00")
        if repo.findings("LC008"):
            problems.append("a note written for the first time was reported as copied forward")
        # 17. The stamp moves and the sentence does not. A bare value, so that the case below
        # can add the offset to it.
        repo.write(doc, _brief("p", "2026-06-07 10:00", created="2026-06-01",
                               scope="the whole file"))
        repo.commit("attest with the same note", "2026-06-07T10:05:00+02:00")
        found = repo.findings("LC008")
        if len(found) != 1 or "left `review_scope` as it was" not in found[0]["message"]:
            problems.append(f"a stamp moved with its sentence unchanged produced {found}: one "
                            "finding saying the note was copied forward")
        # The same minute with an offset added, same sentence: the stamp did not move.
        repo.write(doc, _brief("p", "2026-06-07 10:00 +02:00", created="2026-06-01",
                               scope="the whole file"))
        repo.commit("add the offset to the stamp", "2026-06-07T10:10:00+02:00")
        if repo.findings("LC008"):
            problems.append("a commit that only added the offset to the same minute was "
                            "reported as the note copied forward: naive and aware values "
                            "naming one minute are the same stamp")
        # 18 again. A different sentence, and nothing to report.
        repo.write(doc, _brief("p", "2026-06-08 10:00 +02:00", created="2026-06-01",
                               scope="§1 only, after the change of the 7th"))
        repo.commit("attest with a new note", "2026-06-08T10:05:00+02:00")
        if repo.findings("LC008"):
            problems.append("a stamp moved with a new sentence was reported")

    # Outside any repository the history half cannot run and the empty note still can, with
    # the day rule standing in for the creation commit.
    with tempfile.TemporaryDirectory() as tmp:
        outside = Path(tmp)
        (outside / "framework.yaml").write_text(f"framework_version: {REGISTRY['version']}\n")

        def codes(text: str) -> set[str]:
            (outside / "products" / "p").mkdir(parents=True, exist_ok=True)
            (outside / "products" / "p" / "PBR.md").write_text(text, encoding="utf-8")
            r = subprocess.run([sys.executable, str(VALIDATE), "--root", str(outside), "--json",
                                "--stale-days", "36500"], capture_output=True, text=True)
            return {f["code"] for f in json.loads(r.stdout)["findings"]}

        if "LC008" not in codes(_brief("p", "2026-08-01 09:00", created="2026-01-01")):
            problems.append("outside a repository a stamp months after `created` with no note "
                            "was not reported: the empty note needs no history")
        if "LC008" in codes(_brief("p", "2026-01-01 09:00", created="2026-01-01")):
            problems.append("outside a repository a stamp on the day of `created` was reported: "
                            "`created` is written as a day, and the finer rule would fire on "
                            "every new repository")
        if "LC008" in codes(_brief("p", "2026-08-01 09:00", created="2026-01-01",
                                   scope="the whole file")):
            problems.append("a stamp with a note was reported outside a repository")
    return problems


@check("a shallow clone stops the validator with exit 2 and the command that fetches, and reports nothing")
def _shallow_clone_stops():
    # `actions/checkout` clones one commit unless told otherwise, and in that clone the commit
    # behind every attestation is absent. Reading that as "nothing to compare" would make the
    # checks above pass exactly where they run for everybody. 2 and not 1: 1 is a repository
    # with errors, and a caller reading the code has to tell the two apart.
    problems = []
    with tempfile.TemporaryDirectory() as tmp:
        repo = _Repo(Path(tmp) / "origin")
        repo.write("products/p/PBR.md", _brief("p", "2026-06-01 09:00"))
        repo.commit("first", "2026-06-01T09:00:00+02:00")
        repo.write("products/p/PBR.md", _brief("p", "2026-06-02 10:00 +02:00",
                                              scope="the whole file"))
        repo.commit("attest", "2026-06-02T10:05:00+02:00")
        clone = Path(tmp) / "shallow"
        r = subprocess.run(["git", "clone", "-q", "--depth", "1", f"file://{repo.path}",
                            str(clone)], capture_output=True, text=True)
        if r.returncode != 0:
            return ["a shallow clone could not be made here, so the check is not running"]
        r = subprocess.run([sys.executable, str(VALIDATE), "--root", str(clone), "--json"],
                           capture_output=True, text=True)
        if r.returncode != 2:
            problems.append(f"the validator exited {r.returncode} in a shallow clone: 2 is the "
                            "code for a run that did not happen, and 1 would read as errors")
        if r.stdout.strip():
            problems.append("the validator printed a report in a shallow clone: a partial JSON "
                            "gets parsed, and this one would say the attestations are fine")
        if "fetch" not in r.stderr or "fetch-depth" not in r.stderr:
            problems.append("the message does not say how to fetch the history, in the words "
                            "the CI configuration uses")
        r = subprocess.run([sys.executable, str(ATTEST), "--root", str(clone),
                            "products/p/PBR.md", "--scope", "the whole file, again"],
                           capture_output=True, text=True)
        if r.returncode != 2 or "fetch" not in r.stderr:
            problems.append(f"attest.py ran in a shallow clone (exit {r.returncode}): the "
                            "stamp binds to a commit in a history that is not there")
        # And the same clone made whole is a repository like any other.
        subprocess.run(["git", "-C", str(clone), "fetch", "-q", "--unshallow"],
                       capture_output=True, text=True)
        r = subprocess.run([sys.executable, str(VALIDATE), "--root", str(clone), "--json",
                            "--stale-days", "36500"], capture_output=True, text=True)
        if r.returncode not in (0, 1) or not r.stdout.strip():
            problems.append("the validator did not run in the clone once its history was fetched")
    return problems


@check("attest.py stamps the instant it runs at with its offset, writes the note, commits nothing, and refuses what it must")
def _attest_command():
    # THE COMMAND IS THE HONEST PROCEDURE MADE ONE STEP. A person commits the content, runs
    # this once per document reread, and commits the attestations. It writes the clock and
    # not the content commit's date: the link to the text read is the commit that will carry
    # the stamp and its parent, which `git blame` finds, and the content commit is printed so
    # that the person knows what they are attesting.
    problems = []
    with tempfile.TemporaryDirectory() as tmp:
        repo = _Repo(Path(tmp) / "r")
        doc = "products/p/PBR.md"
        inline = lambda text: text.replace("last_review: 2026-06-01 09:00\n",
                                           "last_review: 2026-06-01 09:00   # on the line\n", 1)
        repo.write(doc, inline(_brief("p", "2026-06-01 09:00", comment="a comment that stays")))
        repo.commit("first", "2026-06-01T09:00:00+02:00")
        repo.write(doc, inline(_brief("p", "2026-06-01 09:00", body="It does two things.",
                                      comment="a comment that stays")))
        content = repo.commit("change the text", "2026-06-10T09:00:00+02:00")

        def attest(rel: str, scope: str, env: dict | None = None):
            return subprocess.run([sys.executable, str(ATTEST), "--root", str(repo.path), rel,
                                   "--scope", scope], capture_output=True, text=True,
                                  env={**os.environ, **(env or {})})

        r = attest(doc, "the whole file: both things it does now", env={"TZ": "Asia/Tokyo"})
        text = (repo.path / doc).read_text(encoding="utf-8")
        meta, _, _ = _load(VALIDATE, "validate_for_attest").parse_front_matter(text)
        stamp = str(meta.get("last_review"))
        if r.returncode != 0:
            problems.append(f"attest.py exited {r.returncode}: {r.stderr.strip()}")
        elif not re.fullmatch(r"\d{4}-\d\d-\d\d \d\d:\d\d \+09:00", stamp):
            problems.append(f"the stamp written is {stamp!r}: the instant the command ran at, "
                            "to the minute, with the offset of the clock it read")
        else:
            written = datetime.strptime(stamp, "%Y-%m-%d %H:%M %z")
            if abs((datetime.now().astimezone() - written).total_seconds()) > 180:
                problems.append(f"the stamp {stamp} is not the moment the command ran")
        if meta.get("review_scope") != "the whole file: both things it does now":
            problems.append(f"`review_scope` reads {meta.get('review_scope')!r}: the sentence "
                            "given, folded back whole")
        if "# a comment that stays" not in text:
            problems.append("a YAML comment in the front matter was lost: the command edits "
                            "lines and does not reserialise")
        if not re.search(r"^last_review: .*   # on the line$", text, re.M):
            problems.append("a comment on the `last_review` line itself was lost: it is text, "
                            "and the command does not delete text")
        order = [k for k in ("created", "last_review", "review_scope", "# a comment") if k in text]
        if [text.index(k) for k in order] != sorted(text.index(k) for k in order):
            problems.append("the attestation block did not land where `last_review` was")
        if content[:12] not in r.stdout or "attests the text as of" not in r.stdout:
            problems.append("the command does not print the commit whose text it attests")
        if content[:12] in text:
            problems.append("the content commit was written into the document: it is printed, "
                            "and the attestation commit's parent carries it")
        status = repo.git("status", "--porcelain")
        if not status or "M " not in status and " M" not in status:
            problems.append(f"after attesting, `git status` reads {status!r}: the command commits "
                            "nothing and leaves the file modified")
        if repo.attestation_codes():
            problems.append(f"the working tree after attest.py reports "
                            f"{sorted(repo.attestation_codes())}: the stamp is the clock and the "
                            "line is not committed, so nothing is ahead and the note is there")
        repo.commit("attest", datetime.now().astimezone().isoformat(timespec="seconds"))
        if repo.attestation_codes():
            problems.append(f"the attestation committed reports "
                            f"{sorted(repo.attestation_codes())}: this is the honest procedure "
                            "and it has to be free of findings")

        # Refusals, each with exit 2 and the file untouched.
        before = (repo.path / doc).read_text(encoding="utf-8")
        r = attest(doc, "   ")
        if r.returncode != 2 or "empty" not in r.stderr:
            problems.append(f"an empty --scope was accepted (exit {r.returncode})")
        r = attest(doc, "the whole file: both things it does now")
        if r.returncode != 2 or "already there" not in r.stderr:
            problems.append(f"the sentence already in `review_scope` was accepted again (exit "
                            f"{r.returncode}): LC008 would report it copied forward")
        repo.write(doc, before.replace("It does two things.", "It does three things."))
        repo.git("add", "-A")            # staged, not committed: the index differs from HEAD
        repo.write(doc, before)          # and the working tree does not, so only the index does
        r = attest(doc, "the whole file, staged")
        if r.returncode != 2 or "the index" not in r.stderr:
            problems.append(f"content staged and not committed was attested (exit "
                            f"{r.returncode}): the commit that follows would carry both")
        repo.git("reset", "-q")
        repo.write(doc, before)
        # A sentence YAML would read as a boolean is written so that it stays a sentence.
        r = attest(doc, "no")
        after = (repo.path / doc).read_text(encoding="utf-8")
        note, _, _ = _load(VALIDATE, "validate_for_attest_words").parse_front_matter(after)
        if r.returncode != 0 or note.get("review_scope") != "no":
            problems.append(f"`--scope no` was written so that YAML reads "
                            f"{note.get('review_scope')!r}: a word YAML reads as a boolean has "
                            "to be written as a string")
        repo.write(doc, before)
        repo.write(doc, before.replace("It does two things.", "It does three things."))
        r = attest(doc, "the whole file")
        if r.returncode != 2 or "differs from HEAD" not in r.stderr:
            problems.append(f"uncommitted content was attested (exit {r.returncode}): the commit "
                            "that follows would carry both the change and the stamp")
        repo.write(doc, before)
        repo.write("decisions/DEC-001-x.md", _fm(
            schema="framework/decision-record/v1", id="DEC-001", artifact_type="decision-record",
            lifecycle="immutable", status="accepted", scope="product", products="[p]",
            owners="[o]", approvers="[o]", created="2026-06-01") + "# DEC-001\n")
        repo.commit("a decision", "2026-06-12T09:00:00+02:00")
        r = attest("decisions/DEC-001-x.md", "the whole file")
        if r.returncode != 2 or "living" not in r.stderr:
            problems.append(f"an immutable was attested (exit {r.returncode})")
        repo.write("products/q/PBR.md", _brief("q", "2026-06-01 09:00"))
        r = attest("products/q/PBR.md", "the whole file")
        if r.returncode != 2 or "not tracked" not in r.stderr:
            problems.append(f"an untracked file was attested (exit {r.returncode})")
        if (repo.path / doc).read_text(encoding="utf-8") != before:
            problems.append("a refused invocation changed the document")
    return problems


@check("a retired code is never emitted, stays documented, and an annotation or a pin on it is told so")
def _retired_codes():
    # A code that changed meaning under an annotation would make a reason written about one
    # thing cover another, because the annotation joins on the code. So a code retires and its
    # number is never reused; the catalog keeps the entry with the version and the reason.
    problems = []
    retired = {c: s for c, s in CHECKS["checks"].items() if s.get("level") == "retired"}
    if "LC005" not in retired:
        problems.append("LC005 is not retired in the catalog: 4.0.0 retired it")
    for code, spec in retired.items():
        if not spec.get("retired_in") or not spec.get("retired_because"):
            problems.append(f"{code}: retired without `retired_in` and `retired_because`, so "
                            "nobody can tell when or why")
    with tempfile.TemporaryDirectory() as tmp:
        root = Path(tmp)
        (root / "framework.yaml").write_text(f"framework_version: {REGISTRY['version']}\n")
        (root / "products" / "p").mkdir(parents=True)
        (root / "products" / "p" / "PBR.md").write_text(
            _brief("p", "2026-01-01 09:00", created="2026-01-01"), encoding="utf-8")
        (root / ".framework").mkdir()
        (root / ".framework" / "expected-findings.yaml").write_text(
            "expected:\n  - code: LC005\n    path: products/p/PBR.md\n"
            "    reason: five documents shared a minute\n    clears_when: the framework repair\n")
        r = subprocess.run([sys.executable, str(VALIDATE), "--root", str(root), "--json",
                            "--stale-days", "36500"], capture_output=True, text=True)
        out = json.loads(r.stdout) if r.stdout.strip() else {"findings": []}
        an = [f for f in out["findings"] if f["code"] == "AN001"]
        if len(an) != 1 or "retired" not in an[0]["message"]:
            problems.append(f"an annotation on LC005 produced {an}: one AN001 saying the code "
                            "is retired, so that the adopter removes the line instead of "
                            "looking for the finding")
        (root / "framework.yaml").write_text(f"framework_version: {REGISTRY['version']}\n"
                                             "checks:\n  LC005: warn\n")
        r = subprocess.run([sys.executable, str(VALIDATE), "--root", str(root), "--json"],
                           capture_output=True, text=True)
        if r.returncode != 2 or "retired" not in r.stderr or r.stdout.strip():
            problems.append(f"a framework.yaml declaring this version and pinning LC005 ran "
                            f"(exit {r.returncode}): a pin on a retired code switches nothing "
                            "on, and the stop is a run that did not happen, which is exit 2")
        # Declaring the version before the one that retired it, the same line is not wrong
        # yet: the run goes on, says so on stderr, and `migrate.py` can compare.
        previous = _previous_version(before=retired["LC005"]["retired_in"])
        if previous is None:
            return problems + ["no earlier version in the last sixty commits: the case cannot "
                               "be built"]
        (root / "framework.yaml").write_text(f'framework_version: "{previous}"\n'
                                             "checks:\n  LC005: warn\n")
        r = subprocess.run([sys.executable, str(VALIDATE), "--root", str(root), "--json",
                            "--stale-days", "36500"], capture_output=True, text=True)
        if r.returncode not in (0, 1) or not r.stdout.strip():
            problems.append(f"a framework.yaml declaring {previous} and pinning LC005 was "
                            f"stopped (exit {r.returncode}): the project has not adopted the "
                            "version that retired the code, and stopping it hides the migration")
        elif "retired" not in r.stderr or "ignored" not in r.stderr:
            problems.append("the ignored pin was not said on stderr: a line that switches "
                            "nothing on has to be announced somewhere")
        r = subprocess.run([sys.executable, str(VALIDATE), "--root", str(root),
                            "--list-checks"], capture_output=True, text=True)
        (root / "framework.yaml").write_text(f"framework_version: {REGISTRY['version']}\n")
        r = subprocess.run([sys.executable, str(VALIDATE), "--root", str(root),
                            "--list-checks"], capture_output=True, text=True)
        if not re.search(r"^retired\s+LC005", r.stdout, re.M):
            problems.append("--list-checks does not print LC005 as retired")
    return problems


@check("migrate.py files a pin on a retired code under GONE and --adopt removes the line, instead of stopping")
def _retired_pin_migrates():
    # THE PROJECT THAT PINNED LC005 IS THE ONE THAT TOOK IT MOST SERIOUSLY, and the new
    # validator stops on the pin once the project declares 4.0.0. A migration tool that
    # stopped on the same line would leave that project no way to see what the move asks: its
    # job is to say what changes, and the pin is one of the things that go.
    migrate = _load(ROOT / "skills" / "audit" / "scripts" / "migrate.py", "migrate_for_pins")
    problems = []
    kept = migrate.without_pins('framework_version: N\n# kept\nchecks:\n'
                                '  LC005: error   # bitten once\n  XP003: off\n', ["LC005"])
    if "LC005" in kept or "# kept" not in kept or "XP003: off" not in kept:
        problems.append(f"without_pins produced {kept!r}: the pinned line goes, the comment "
                        "and the other pins stay")
    kept = migrate.without_pins('checks:\n  LC005:\n    level: error\n  XP003: off\n',
                                ["LC005"])
    if "level" in kept or "XP003: off" not in kept:
        problems.append(f"the long form left an orphan or took a neighbour: {kept!r}")
    kept = migrate.without_pins('checks:\n  LC005:\n  # why\n    level: error\n'
                                '  # the next one\n  XP003: off\n', ["LC005"])
    if "level" in kept or "# the next one" not in kept or "XP003: off" not in kept:
        problems.append(f"a comment between the key and its `level:` was left as an orphan, "
                        f"or the next key's comment was taken: {kept!r}")
    with tempfile.TemporaryDirectory() as tmp:
        cfg = Path(tmp) / "framework.yaml"
        flow = 'framework_version: N\nchecks: {LC005: warn, XP003: off}\n'
        cfg.write_text(flow, encoding="utf-8")
        why = migrate.adopt(cfg, REGISTRY["version"], None, drop=["LC005"])
        if not why or "LC005" not in why:
            problems.append("a flow style pin was not refused by --adopt: the number would be "
                            "written with the pin in place, and the validator would stop")
        if cfg.read_text(encoding="utf-8") != flow:
            problems.append("--adopt wrote framework.yaml although it could not remove the pin")

    previous = _previous_version(before=CHECKS["checks"]["LC005"]["retired_in"])
    if previous is None:
        return problems + ["no earlier version in the last sixty commits: the case cannot be built"]

    def project(where: Path) -> Path:
        where.mkdir(parents=True)
        (where / "framework.yaml").write_text(
            f'framework_version: "{previous}"\n# why the pin is here\nchecks:\n'
            "  LC005: error   # bitten once\n", encoding="utf-8")
        (where / "products" / "p").mkdir(parents=True)
        (where / "products" / "p" / "PBR.md").write_text(
            _brief("p", "2026-01-01 09:00", created="2026-01-01"), encoding="utf-8")
        return where

    with tempfile.TemporaryDirectory() as tmp:
        proj = project(Path(tmp) / "project")
        r = subprocess.run([sys.executable, str(ROOT / "skills/audit/scripts/migrate.py"),
                            "--root", str(proj), "--framework", str(ROOT), "--json"],
                           capture_output=True, text=True)
        try:
            out = json.loads(r.stdout)
        except json.JSONDecodeError:
            return problems + [f"migrate.py returned nothing readable: {r.stderr.strip()[-300:]}"]
        if out["problems"]:
            problems.append(f"migrate.py could not compare a project pinning LC005: "
                            f"{out['problems']}")
        pins = [g for g in out["gone"] if g["code"] == "LC005" and g["path"] == "framework.yaml"]
        if len(pins) != 1 or "--adopt" not in pins[0]["message"]:
            problems.append(f"the pin is not under GONE with what removes it: {pins}")
        if any(g["code"] == "LC005" for g in out["new"]):
            problems.append("LC005 appeared under NEW: a retired code is not migration work")

        # `--adopt` needs a clean committed framework to bind the pin to, like the check on
        # the migration being executable: an isolated clone of what is committed here.
        dirty = subprocess.run(["git", "-C", str(ROOT), "status", "--porcelain"],
                               capture_output=True, text=True).stdout.strip()
        if dirty:
            problems.append("NOT RUN: the working tree is dirty, so `--adopt` against a clone "
                            "of it would test committed bytes that are not these. Commit and "
                            "run again")
            return problems
        clone = Path(tmp) / "trusted"
        c = subprocess.run(["git", "clone", "-q", "--local", "--no-hardlinks", str(ROOT),
                            str(clone)], capture_output=True, text=True)
        if c.returncode:
            return problems + ["could not clone the framework for the adoption: " + c.stderr]
        proj2 = project(Path(tmp) / "project-adopting")
        r = subprocess.run([sys.executable, str(ROOT / "skills/audit/scripts/migrate.py"),
                            "--root", str(proj2), "--framework", str(clone), "--adopt",
                            "--json"], capture_output=True, text=True)
        after = (proj2 / "framework.yaml").read_text(encoding="utf-8")
        if REGISTRY["version"] not in after:
            problems.append(f"the project pinning LC005 could not adopt: {r.stdout[-400:]}")
        if "LC005" in after:
            problems.append("`--adopt` wrote the number and left the pin, which stops the "
                            "validator of the version just adopted")
        if "# why the pin is here" not in after:
            problems.append("`--adopt` dropped a comment from framework.yaml")
        v = subprocess.run([sys.executable, str(VALIDATE), "--root", str(proj2), "--json",
                            "--stale-days", "36500"], capture_output=True, text=True)
        if v.returncode not in (0, 1) or not v.stdout.strip():
            problems.append(f"after adopting, the validator does not run on the project "
                            f"(exit {v.returncode}): {v.stderr.strip()[-200:]}")
    # And the adoption note names the pin where it names the annotations.
    note = (ROOT / "references" / "adoption.md").read_text(encoding="utf-8")
    step0 = note[note.find("0. **Remove"):note.find("1. **Correct")]
    if "pin" not in step0 or "stops the validator" not in step0:
        problems.append("adoption.md step 0 does not say that a pin on LC005 stops the "
                        "validator once the project declares 4.0.0")
    return problems


@check("the static fixtures carry no attestation ahead of its commit and no reading without a note")
def _static_fixtures_attest_cleanly():
    # They sit inside this repository, so the history behind their lines is this
    # repository's own: a stamp moved by a later commit of the framework owes the sentence,
    # and three of them carry it for that reason. Nothing else validated them in place.
    problems = []
    roots = sorted(p.parent for p in (ROOT / "evals" / "fixtures" / "static").rglob("AGENTS.md"))
    if not roots:
        return ["no static fixture found: the check is not running"]
    for root in roots:
        r = subprocess.run([sys.executable, str(VALIDATE), "--root", str(root), "--json",
                            "--stale-days", "36500"], capture_output=True, text=True)
        if r.returncode not in (0, 1) or not r.stdout.strip():
            problems.append(f"{root.relative_to(ROOT)}: the validator did not run "
                            f"(exit {r.returncode}): {r.stderr.strip()[-200:]}")
            continue
        found = [(f["code"], f["path"]) for f in json.loads(r.stdout)["findings"]
                 if f["code"] in ("LC007", "LC008")]
        if found:
            problems.append(f"{root.relative_to(ROOT)}: reports {found}. A fixture whose stamp "
                            "this repository's history contradicts teaches you to read past it")
    return problems


@check("the attestation fields are declared once, in the registry, and the validator and the living schemas agree")
def _attestation_fields_agree():
    x = _load(VALIDATE, "validate_for_fields")
    declared = tuple(REGISTRY.get("attestation_fields") or [])
    problems = []
    if declared != tuple(x.ATTESTATION_FIELDS):
        problems.append(f"the registry declares {declared} and the validator reads "
                        f"{x.ATTESTATION_FIELDS}: one authority, and the other reads it")
    for name, spec in REGISTRY["types"].items():
        schema = json.loads((ROOT / "schemas" / "framework" / name / "v1.json").read_text())
        props = set(schema.get("properties", {}))
        if spec["lifecycle"] == "living" and not set(declared) <= props:
            problems.append(f"{name}: living, and its schema does not declare {declared}")
        if spec["lifecycle"] != "living" and set(declared) & props:
            problems.append(f"{name}: not living, and its schema declares an attestation field")
    return problems


@check("the CI assets check the framework out beside the project, never inside the root they validate")
def _ci_assets_keep_the_framework_beside():
    # Inside the project the validator read the framework's own files as the project's -- its
    # fixtures numbered DEC-001, its prose without front matter -- and every pull request job
    # went red; inside `.framework/` the checkout emptied the project's folder of that name.
    problems = []
    for name in ("pull-request.yml", "branches.yml"):
        wf = yaml.safe_load((ROOT / "ci" / name).read_text(encoding="utf-8"))
        for job, spec in (wf.get("jobs") or {}).items():
            project = framework = None
            for step in spec.get("steps") or []:
                if str(step.get("uses", "")).startswith("actions/checkout"):
                    w = step.get("with") or {}
                    if "framework-data-ai" in str(w.get("repository", "")):
                        framework = w.get("path")
                    else:
                        project = w.get("path")
            if not project or not framework or framework.startswith(project + "/") \
                    or framework.startswith("."):
                problems.append(f"ci/{name}, job {job}: project in {project!r}, framework in "
                                f"{framework!r}: they go side by side, neither inside the other "
                                "nor in a hidden folder of the project's")
            # Every call names its root: without one the validator reads `.`, the workspace that
            # holds both checkouts, which is the failure this check is for.
            lines = [line for step in spec.get("steps") or []
                     for line in str(step.get("run", "")).replace("\\\n", " ").splitlines()]
            names = {m.group(1) for line in lines
                     for m in [re.match(r"\s*([A-Za-z_]\w*)=\S*validate\.py\s*$", line)] if m}
            calls = [line for line in lines
                     if not re.match(r"\s*[A-Za-z_]\w*=\S*validate\.py\s*$", line)
                     and ("validate.py" in line
                          or any(re.search(rf"\$\{{?{n}\b", line) for n in names))]
            if not calls:
                problems.append(f"ci/{name}, job {job}: no validator call found, so the check "
                                "below proves nothing")
            for call in calls:
                root = re.search(r"--root(?:\s+|=)(\S+)", call)
                if not root or root.group(1) not in {project, '"/tmp/at-$sha"', '"/tmp/before-$sha"'}:
                    problems.append(f"ci/{name}, job {job}: a validator call with "
                                    f"{root.group(0) if root else 'no --root'}, not the project's "
                                    f"checkout {project!r}")
    return problems


@check("every file a workflow installs from is in the framework")
def _ci_install_files_exist():
    # The CI assets asked for `constraints.txt` for weeks after it was left behind, and `pip`
    # failed in every project that copied them; nothing here read the line.
    problems = []
    seen = [0]
    flows = [*(ROOT / "ci").glob("*.yml"), *(ROOT / ".github" / "workflows").glob("*.yml")]
    for wf in flows:
        jobs = (yaml.safe_load(wf.read_text(encoding="utf-8")) or {}).get("jobs") or {}
        for job, spec in jobs.items():
            for step in spec.get("steps") or []:
                installs = [l for l in str(step.get("run", "")).splitlines() if "pip install" in l]
                seen[0] += len(installs)
                for m in re.finditer(r"\s(?:-r|-c|--requirement|--constraint)[=\s]+(\S+)",
                                     " ".join(installs)):
                    path = m.group(1)
                    for prefix in ("framework-data-ai/", ".framework/"):
                        path = path.removeprefix(prefix)
                    if not (ROOT / path).is_file():
                        problems.append(f"{wf.relative_to(ROOT)}, job {job}: installs from "
                                        f"{m.group(1)}, which the framework does not have")
    if not flows or not seen[0]:
        problems.append("no workflow, or no `pip install` in one: the check is not running")
    return problems


@check("every checkout in the CI assets a project copies fetches the full history")
def _ci_assets_fetch_history():
    # `_ci_fetches_history` above asserts it for this framework's own suite. These are the
    # workflows a project copies, and from 4.0.0 the validator stops in a clone without
    # history: a project that copied the file before this line was written has a gate that
    # exits 2 on every pull request, with a message about the clone. The assertion is here so
    # that the file cannot lose the setting without this going red.
    problems = []
    for wf in sorted([*(ROOT / "ci").glob("*.yml"), *(ROOT / "ci").glob("*.yaml")]):
        jobs = (yaml.safe_load(wf.read_text(encoding="utf-8")) or {}).get("jobs") or {}
        for name, job in jobs.items():
            for step in (job or {}).get("steps") or []:
                if not str(step.get("uses", "")).startswith("actions/checkout"):
                    continue
                depth = (step.get("with") or {}).get("fetch-depth")
                if depth != 0:
                    what = step.get("with", {}).get("repository", "the project")
                    problems.append(f"ci/{wf.name}, job {name}: checks out {what} with "
                                    f"`fetch-depth: {depth!r}`. The validator reads the commit "
                                    "behind every attestation and stops in a clone without it")
    if not problems and not [*(ROOT / "ci").glob("*.yml"), *(ROOT / "ci").glob("*.yaml")]:
        problems.append("no workflow under ci/: the check is not running")
    return problems


EXCHANGE = ROOT / "skills" / "audit" / "scripts" / "exchange.py"

_EXCHANGE_SECTIONS = ("<!-- section: what-changes -->\n## A\n"
                      "<!-- section: what-must-not-change -->\n## B\n"
                      "<!-- section: how-we-know-it-worked -->\n## C\n")


def _exchange_docs(tmp: Path, exchange: str | None = "../_scambio") -> tuple[_Repo, Path, str]:
    """A documentation repository beside its exchange: what every check on the place reads.

    `CHG-001` implemented and `OD-002` decided are closed; `CHG-002` and `CHG-003` approved and
    `OD-001` open are not. The exchange is not created: each check makes what it needs.
    """
    repo = _Repo(tmp / "docs")
    if exchange is not None:
        repo.write(".framework/sessions.yaml", f"exchange: {exchange}\n")
    for ident, status in (("CHG-001", "implemented"), ("CHG-002", "approved"),
                          ("CHG-003", "approved")):
        repo.write(f"products/p/changes/{ident}.md", _fm(
            schema="framework/change-contract/v1", artifact_type="change-contract",
            lifecycle="immutable", status=status, id=ident, products="[p]", owners="[o]",
            created="2026-10-01") + f"# {ident}\n\n" + _EXCHANGE_SECTIONS)
    repo.write("OPEN.md", _fm(
        schema="framework/open-register/v1", artifact_type="open-register",
        lifecycle="living", status="active", owners="[o]", created="2026-10-01",
        last_review="2026-10-01 09:00",
        entries="\n  OD-001:\n    status: open\n    cost_to_reverse: low"
                "\n  OD-002:\n    status: decided\n    cost_to_reverse: low"
                "\n    closed_by: DEC-001") + "# Open\n")
    sha = repo.commit("the documents", "2026-10-01T09:00:00+02:00")
    return repo, tmp / "_scambio", sha


def _put(base: Path, files: dict[str, str]) -> None:
    for rel, text in files.items():
        p = base / rel
        if rel.endswith("/"):
            p.mkdir(parents=True, exist_ok=True)
            continue
        p.parent.mkdir(parents=True, exist_ok=True)
        p.write_text(text, encoding="utf-8")


def _exc(repo: _Repo, code: str) -> list[str]:
    """The paths, under the exchange, at which `code` is reported."""
    return sorted(f["path"].split("_scambio/", 1)[-1] for f in repo.findings(code))


_RESOURCE_HEAD = "| Tipo | Percorso o nome | Comando che la toglie |\n|---|---|---|\n"
_DECLARATION = ("# Chiusura\n\n## Committato\n\nniente\n\n## Non committato\n\n"
                "- `OPEN.md`: OD-001, le righe del default\n\n## Fuori dai repository\n\n"
                "vedi `risorse.md`\n")


@check("`exchange:` in .framework/sessions.yaml is a place relative to the documentation and outside it, and optional")
def _exchange_is_declared():
    # The value is committed and read on every machine, so a path that names one machine is
    # refused, and one inside the repository is refused because what waits there is not
    # committed. Without the file nothing changes, and a declared place that is not on this
    # machine -- every CI runner -- runs no check and says so rather than warning on each push.
    # Not a key of `framework.yaml`: one there stops every older validator.
    problems = []
    with tempfile.TemporaryDirectory() as tmp:
        for n, (value, says) in enumerate((
                ("/srv/_scambio", "a place on one machine"),
                ("~/_scambio", "a place on one machine"),
                ("'C:\\\\scambio'", "a place on one machine"),
                ("scambio", "inside the documentation repository"),
                ("'.'", "inside the documentation repository"),
                ("[a, b]", "has to be a path"),
                ("../_scambio\nwhere: here", "unknown key"))):
            repo, _, _ = _exchange_docs(Path(tmp) / f"refused-{n}", value)
            rc, out, err = repo.validate()
            if rc == 0 or out is not None or says not in err:
                problems.append(f"`exchange: {value}` was not refused with {says!r} "
                                f"(exit {rc}): {err.strip()[-160:]}")
        if "exchange" in _load(VALIDATE, "validate_exchange_keys").PROJECT_KEYS:
            problems.append("`framework.yaml` accepts `exchange`: every validator before this "
                            "one would stop on it")
        repo, ex, _ = _exchange_docs(Path(tmp) / "none", None)
        repo.write("framework.yaml", f"framework_version: {REGISTRY['version']}\n")
        _put(ex, {"in-arrivo/CHG-999/x.txt": "x"})
        rc, out, err = repo.validate()
        if out is None or out.get("exchange") != "not-declared" or any(
                f["code"].startswith("EXC") for f in out["findings"]):
            problems.append("a project without `exchange:` is not read as before: "
                            f"{(out or {}).get('exchange')!r}, {err.strip()[-120:]}")
        repo, ex, _ = _exchange_docs(Path(tmp) / "absent")
        rc, out, err = repo.validate()
        if out is None or out.get("exchange") != "absent":
            problems.append(f"a declared exchange not on this machine reads as "
                            f"{(out or {}).get('exchange')!r}, not `absent`")
        r = subprocess.run([sys.executable, str(VALIDATE), "--root", str(repo.path),
                            "--stale-days", "36500"], capture_output=True, text=True)
        if "is not on this machine" not in r.stdout:
            problems.append("the report does not say that the exchange checks did not run "
                            "where the exchange is absent")
        ex.mkdir()
        rc, out, err = repo.validate()
        if out is None or out.get("exchange") != "read":
            problems.append(f"an exchange on this machine reads as {(out or {}).get('exchange')!r}")
    return problems


@check("EXC001: a folder in the exchange names an element the documentation does not declare")
def _exc001():
    with tempfile.TemporaryDirectory() as tmp:
        repo, ex, _ = _exchange_docs(Path(tmp))
        _put(ex, {"in-arrivo/CHG-002/brief.md": "b", "in-arrivo/OD-001/verifica.md": "v",
                  "in-arrivo/CHG-009/brief.md": "b", "integrati/KI-404/brief.md": "b"})
        got = _exc(repo, "EXC001")
    want = ["in-arrivo/CHG-009", "integrati/KI-404"]
    return [] if got == want else [f"EXC001 reported at {got}, expected {want}: a change "
                                   "and an entry of a register both name an element"]


@check("EXC002: an element done in the documents still has its folder in in-arrivo/")
def _exc002():
    with tempfile.TemporaryDirectory() as tmp:
        repo, ex, _ = _exchange_docs(Path(tmp))
        _put(ex, {f"in-arrivo/{i}/verifica.md": "v"
                  for i in ("CHG-001", "CHG-002", "OD-001", "OD-002")})
        _put(ex, {"integrati/CHG-001/brief.md": "b"})
        got = _exc(repo, "EXC002")
    want = ["in-arrivo/CHG-001", "in-arrivo/OD-002"]
    return [] if got == want else [f"EXC002 reported at {got}, expected {want}: implemented "
                                   "and decided close an element, approved and open do not, "
                                   "and integrati/ is where a closed one belongs"]


@check("EXC003: a folder in integrati/ without an INTEGRATO that says where each part went")
def _exc003():
    problems = []
    with tempfile.TemporaryDirectory() as tmp:
        repo, ex, sha = _exchange_docs(Path(tmp))
        good = (f"commit: {sha[:12]}\ndate: 2026-10-14\nparts:\n"
                "  verifica.md: CHG-001 §Verification\n  evidenze/: CHG-001 §Verification\n")
        cases = {
            "CHG-001": (good, []),
            "CHG-002": (None, ["integrati/CHG-002"]),
            "CHG-003": (good.replace(sha[:12], "0123456789ab"), ["integrati/CHG-003/INTEGRATO"]),
            "OD-001": (good.replace("  evidenze/: CHG-001 §Verification\n", ""),
                       ["integrati/OD-001/INTEGRATO"]),
            "OD-002": (good + "  mr.md: CHG-001\n", ["integrati/OD-002/INTEGRATO"]),
        }
        for ident, (mark, _) in cases.items():
            _put(ex, {f"integrati/{ident}/verifica.md": "v",
                      f"integrati/{ident}/evidenze/log.txt": "l"})
            if mark is not None:
                _put(ex, {f"integrati/{ident}/INTEGRATO": mark})
        got = _exc(repo, "EXC003")
        want = sorted(p for _, paths in cases.values() for p in paths)
        if got != want:
            problems.append(f"EXC003 reported at {got}, expected {want}: a missing file, a "
                            "commit the history does not hold, a part with no destination and "
                            "a destination with no part")
        # A second delivery after the first was integrated: one block per integration, and
        # the parts of all of them cover the folder.
        _put(ex, {"integrati/CHG-001/INTEGRATO":
                  f"- commit: {sha[:12]}\n  date: 2026-10-14\n  parts:\n"
                  "    verifica.md: CHG-001 §Verification\n"
                  f"- commit: {sha}\n  date: 2026-10-20\n  parts:\n"
                  "    evidenze: CHG-001 §Verification\n",
                  "integrati/CHG-002/INTEGRATO": "commit: [unreadable\n"})
        got = _exc(repo, "EXC003")
        if "integrati/CHG-001/INTEGRATO" in got or "integrati/CHG-002/INTEGRATO" not in got:
            problems.append(f"two integration blocks, or an unreadable INTEGRATO, read wrong: {got}")
    return problems


@check("EXC004: a resource listed by an element is still there after the element is done")
def _exc004():
    with tempfile.TemporaryDirectory() as tmp, tempfile.TemporaryDirectory() as away:
        repo, ex, _ = _exchange_docs(Path(tmp))
        there = Path(away) / "app-chg"
        there.mkdir()
        rows = (_RESOURCE_HEAD
                + f"| worktree | `{there}` | `git -C /src/app worktree remove {there}` |\n"
                + f"| worktree | {Path(away) / 'gone'} | git worktree remove gone |\n"
                + "| ramo locale | chg-001 in /src/app | git -C /src/app branch -D chg-001 |\n")
        _put(ex, {"in-arrivo/CHG-002/risorse.md": rows, "integrati/CHG-001/risorse.md": rows,
                  "integrati/CHG-001/INTEGRATO": "x: y\n"})
        got = [(f["path"].split("_scambio/", 1)[-1], f["message"].split(":", 1)[0])
               for f in repo.findings("EXC004")]
    want = [("integrati/CHG-001/risorse.md", "line 3"),
            ("integrati/CHG-001/risorse.md", "line 5")]
    return [] if sorted(got) == want else [
        f"EXC004 reported {sorted(got)}, expected {want}: a worktree whose path exists and a "
        "branch still listed are there; a path that is gone is a resource removed; an "
        "element still approved keeps its resources"]


@check("EXC005: a copy of credentials is under the projects folder, or the list carries a value")
def _exc005():
    with tempfile.TemporaryDirectory() as tmp, tempfile.TemporaryDirectory() as scratch:
        repo, ex, _ = _exchange_docs(Path(tmp))
        near, far = Path(tmp) / ".env.backup", Path(scratch) / ".env.misure"
        near.write_text("x")
        far.write_text("x")
        _put(ex, {"in-arrivo/CHG-002/risorse.md": _RESOURCE_HEAD
                  + f"| copia di credenziali | {far} (DB_URL, DB_PASSWORD) | rm {far} |\n"
                  + f"| copia di credenziali | {near} (DB_URL) | rm {near} |\n"
                  + f"| copia di credenziali | {far} (DB_PASSWORD=hunter2) | rm {far} |\n"})
        got = [f["message"].split(":", 1)[0] for f in repo.findings("EXC005")]
    return [] if sorted(got) == ["line 4", "line 5"] else [
        f"EXC005 reported on {sorted(got)}, expected lines 4 and 5: a copy beside the "
        "repositories, and a value written in the list; a copy in a scratchpad is where it "
        "belongs"]


@check("EXC006: a name in the exchange that the framework does not read")
def _exc006():
    problems = []
    with tempfile.TemporaryDirectory() as tmp:
        repo, ex, _ = _exchange_docs(Path(tmp))
        _put(ex, {"in-arrivo/CHG-002-verifica.md": "the old shape",
                  "in-arrivo/CHG-002/verifica.md": "v", "in-arrivo/CHG-002/mr.md": "m",
                  "in-arrivo/CHG-002/mr-link.txt": "l", "in-arrivo/CHG-002/evidenze/": "",
                  "in-arrivo/CHG-002/contribution-policy.json": "{}",
                  "in-arrivo/CHG-002/brief.md": "b", "in-arrivo/CHG-002/risorse.md": _RESOURCE_HEAD,
                  "in-arrivo/CHG-002/sessioni/2026-10-14-ab12cd34.md": _DECLARATION,
                  "in-arrivo/CHG-003/note.txt": "n", "in-arrivo/CHG-003/evidenze": "a file",
                  "in-arrivo/CHG-003/INTEGRATO": "x", "copie/": "",
                  "sessioni/2026-10-14-ab12cd34.md": _DECLARATION})
        got = _exc(repo, "EXC006")
        want = ["copie", "in-arrivo/CHG-002-verifica.md", "in-arrivo/CHG-003/INTEGRATO",
                "in-arrivo/CHG-003/evidenze", "in-arrivo/CHG-003/note.txt"]
        if got != want:
            problems.append(f"EXC006 reported at {got}, expected {want}")
    with tempfile.TemporaryDirectory() as tmp:
        outer = _Repo(Path(tmp) / "outer")
        repo, ex, _ = _exchange_docs(Path(tmp) / "work", "../../outer/_scambio")
        (outer.path / "_scambio").mkdir()
        got = [f["message"] for f in repo.findings("EXC006")]
        if len(got) != 1 or "inside the git repository" not in got[0]:
            problems.append(f"an exchange inside a git repository reads as {got}")
    return problems


@check("EXC007: a resource row or a closing declaration that cannot be read")
def _exc007():
    with tempfile.TemporaryDirectory() as tmp:
        repo, ex, _ = _exchange_docs(Path(tmp))
        _put(ex, {"in-arrivo/CHG-002/risorse.md": _RESOURCE_HEAD
                  + "| worktree | /srv/app-chg002 | git worktree remove /srv/app-chg002 |\n"
                  + "| container | app-db | docker rm -f app-db |\n"
                  + "| volume | app-db-data | |\n"
                  + "| worktree | ../app-chg002 | git worktree remove ../app-chg002 |\n"
                  + "| copia di credenziali | /tmp/s/.env | rm /tmp/s/.env |\n",
                  "in-arrivo/CHG-002/sessioni/2026-10-14-ab12cd34.md": _DECLARATION,
                  "in-arrivo/CHG-002/sessioni/chiusura.md": _DECLARATION,
                  "sessioni/2026-10-13-ef56.md": _DECLARATION.replace("Non committato", "Altro")})
        got = sorted((f["path"].split("_scambio/", 1)[-1], f["message"].split(":", 1)[0]
                      if f["message"].startswith("line") else "")
                     for f in repo.findings("EXC007"))
    want = sorted([("in-arrivo/CHG-002/risorse.md", "line 4"),
                   ("in-arrivo/CHG-002/risorse.md", "line 5"),
                   ("in-arrivo/CHG-002/risorse.md", "line 6"),
                   ("in-arrivo/CHG-002/risorse.md", "line 7"),
                   ("in-arrivo/CHG-002/sessioni/chiusura.md", ""),
                   ("sessioni/2026-10-13-ef56.md", "")])
    return [] if got == want else [
        f"EXC007 reported {got}, expected {want}: a type not in the list, a missing command, a "
        "relative path, a copy of credentials without its keys, a declaration misnamed and "
        "one without a part"]


@check("the critic ships with the plugin, read only, and every rule that runs it names the same agent")
def _the_critic_ships_with_the_plugin():
    # It lived in this repository's `.claude/agents/`, which Claude Code loads only for a
    # session working here: every project that adopted the framework went without it, and
    # nothing said so. A plugin's agents are the ones in `agents/` at its root, named
    # `<plugin>:<agent>`. One copy, so that the sessions working here and the projects that
    # adopt the framework are reviewed by the same text.
    problems = []
    agent = ROOT / "agents" / "critic.md"
    if not agent.exists():
        return ["agents/critic.md is missing: the plugin ships no critic"]
    try:
        meta = front_matter(agent.read_text(encoding="utf-8")) or {}
    except yaml.YAMLError as e:
        return [f"agents/critic.md: the front matter is not YAML, so Claude Code reads no agent "
                f"from it: {str(e).splitlines()[0]}"]
    if meta.get("name") != "critic":
        problems.append(f"agents/critic.md is named {meta.get('name')!r}")
    listed = meta.get("tools", "")
    tools = {str(t).strip() for t in (listed if isinstance(listed, list)
                                      else str(listed).split(",")) if str(t).strip()}
    if tools != {"Read", "Grep", "Glob"}:
        problems.append(f"the critic's tools are {sorted(tools)}: it reads, and nothing else, so "
                        "that it cannot repair what it criticises")
    ignored = {"permissionMode", "hooks", "mcpServers", "initialPrompt"} & set(meta)
    if ignored:
        problems.append(f"{sorted(ignored)} in the critic's front matter, which a plugin agent "
                        "ignores: a line that reads as applied and is not")
    if (ROOT / ".claude" / "agents" / "critic.md").exists():
        problems.append(".claude/agents/critic.md exists again: here it would win over the "
                        "plugin's, and the two copies would drift")
    plugin = json.loads((ROOT / ".claude-plugin" / "plugin.json").read_text())["name"]
    name = f"{plugin}:critic"
    for rel in ("references/preamble.md", "templates/AGENTS.md", "skills/cycle/SKILL.md",
                "skills/audit/SKILL.md"):
        if name not in (ROOT / rel).read_text(encoding="utf-8"):
            problems.append(f"{rel} does not name `{name}`: the rule that runs the critic and "
                            "the agent it runs have to be the same one")
    # The rule is written twice, for sessions running a skill and for those running none, and
    # the two have to exempt the same writes or one of them asks for what the other forbids.
    for rel in ("references/preamble.md", "templates/AGENTS.md"):
        text = (ROOT / rel).read_text(encoding="utf-8")
        section = text[text.index(name):][:2500] if name in text else ""
        for exempt in ("parking", "SIG", "script"):
            if exempt not in section:
                problems.append(f"{rel}: the review rule does not exempt the {exempt} write")
    return problems


@check("the reviser ships with the plugin, runs the critic, and writes without running anything")
def _the_reviser_ships_with_the_plugin():
    # It stands between a session and the critic, so that the critic's reports do not fill the
    # session: it runs the critic, fixes the objections on the proposal and hands back what
    # needs a person. Edit and Write to fix, Agent to launch the critic, and no Bash: a plugin
    # agent ignores `permissionMode`, so a shell would give it every command. Agent could
    # still launch an agent that has one: that it launches the critic and nothing else is its
    # text's, not a tool's.
    problems = []
    agent = ROOT / "agents" / "reviser.md"
    if not agent.exists():
        return ["agents/reviser.md is missing: the rules that run the reviser run nothing"]
    text = agent.read_text(encoding="utf-8")
    try:
        meta = front_matter(text) or {}
    except yaml.YAMLError as e:
        return [f"agents/reviser.md: the front matter is not YAML, so Claude Code reads no agent "
                f"from it: {str(e).splitlines()[0]}"]
    if meta.get("name") != "reviser":
        problems.append(f"agents/reviser.md is named {meta.get('name')!r}")
    listed = meta.get("tools", "")
    tools = {str(t).strip() for t in (listed if isinstance(listed, list)
                                      else str(listed).split(",")) if str(t).strip()}
    if tools != {"Read", "Grep", "Glob", "Edit", "Write", "Agent"}:
        problems.append(f"the reviser's tools are {sorted(tools)}: it reads, fixes the text it "
                        "was given and launches the critic, and runs no command")
    ignored = {"permissionMode", "hooks", "mcpServers", "initialPrompt"} & set(meta)
    if ignored:
        problems.append(f"{sorted(ignored)} in the reviser's front matter, which a plugin agent "
                        "ignores: a line that reads as applied and is not")
    if (ROOT / ".claude" / "agents" / "reviser.md").exists():
        problems.append(".claude/agents/reviser.md exists: here it would win over the "
                        "plugin's, and the two copies would drift")
    plugin = json.loads((ROOT / ".claude-plugin" / "plugin.json").read_text())["name"]
    if f"{plugin}:critic" not in text.split("\n---", 1)[-1]:
        problems.append(f"agents/reviser.md does not name `{plugin}:critic` in its body: the "
                        "agent it launches has to be the one the rules name")
    name = f"{plugin}:reviser"
    for rel in ("references/preamble.md", "templates/AGENTS.md", "skills/cycle/SKILL.md"):
        if name not in (ROOT / rel).read_text(encoding="utf-8"):
            problems.append(f"{rel} does not name `{name}`: the rule that runs the reviser and "
                            "the agent it runs have to be the same one")
    # The pull request somebody else wrote is read by the critic alone: the reviser would fix it.
    if name in (ROOT / "skills" / "audit" / "SKILL.md").read_text(encoding="utf-8"):
        problems.append("skills/audit/SKILL.md names the reviser: a pull request somebody else "
                        "wrote is read by the critic, not fixed")
    # One set of classes: an objection of a class the reviser does not sort is one nobody acts on.
    critic = (ROOT / "agents" / "critic.md").read_text(encoding="utf-8")
    for cls in ("proposal", "substance", "out of scope"):
        if f"[{cls}]" not in critic or f"*{cls}*" not in text:
            problems.append(f"the class {cls!r} is not both in the critic's output and among "
                            "those the reviser sorts")
    return problems


@check("the write guard stops a write outside the working area, and only in the working folder")
def _the_write_guard_holds():
    # A plugin's hooks run in every session of everyone who has the plugin enabled, and inside
    # every subagent. In the working folder -- a repository with `.framework/` and the ones
    # beside it -- the guard lets a session write its repository and that repository's
    # worktrees, the exchange, what Claude Code keeps for the session and the folders the
    # person lists; anywhere else it says nothing. Only exit 2 stops a call: an import the
    # machine lacks, or one a `PYTHONPATH` breaks, would exit 1 and let every write through, so
    # the guard has the standard library alone and runs isolated.
    problems = []
    guard = ROOT / "hooks" / "guard.py"
    hooks = json.loads((ROOT / "hooks" / "hooks.json").read_text(encoding="utf-8"))
    run_on = [(e.get("matcher", ""), h.get("command", ""))
              for e in hooks.get("hooks", {}).get("PreToolUse", []) for h in e.get("hooks", [])]
    if not any(set(m.split("|")) >= {"Write", "Edit", "NotebookEdit"}
               and "python3 -I \"${CLAUDE_PLUGIN_ROOT}/hooks/guard.py\"" in c for m, c in run_on):
        problems.append(f"hooks/hooks.json runs no isolated guard on Write, Edit and "
                        f"NotebookEdit: {run_on}")
    imported = set()
    for node in ast.walk(ast.parse(guard.read_text(encoding="utf-8"))):
        if isinstance(node, ast.Import):
            imported |= {a.name.split(".")[0] for a in node.names}
        elif isinstance(node, ast.ImportFrom) and node.module and not node.level:
            imported.add(node.module.split(".")[0])
    if imported - set(sys.stdlib_module_names):
        problems.append(f"hooks/guard.py imports {sorted(imported - set(sys.stdlib_module_names))}"
                        ": where it is missing the guard exits 1, and the write goes through")
    option = json.loads((ROOT / ".claude-plugin" / "plugin.json").read_text()).get("userConfig")
    if "allowed_dirs" not in (option or {}):
        problems.append("plugin.json declares no `allowed_dirs` option: the person has no way "
                        "to allow a folder but to switch the guard off")

    with tempfile.TemporaryDirectory() as tmp:
        tmp = Path(os.path.realpath(tmp))
        work, scratch, config, temp = (tmp / "work", tmp / "scratch", tmp / "config",
                                       tmp / "temp")
        docs, code = _Repo(work / "docs"), _Repo(work / "code")
        elsewhere = _Repo(tmp / "elsewhere" / "project")
        for r in (docs, code, elsewhere):
            r.commit("start", "2026-01-01T10:00:00")
        (docs.path / ".framework").mkdir()
        docs.git("worktree", "add", "-q", str(work / "docs--x"))
        # `a-locked` sorts before the repositories: a session in `code` looks into it first.
        for d in (work / "_scambio", work / "other", work / "a-locked", scratch):
            d.mkdir()
        (docs.path / "link").symlink_to(work / "other")
        slug = lambda p: re.sub(r"[^A-Za-z0-9]", "-", str(p))
        uid = getattr(os, "getuid", lambda: None)()
        session_scratch = temp / f"claude-{uid}" / slug(docs.path) / "s1" / "scratchpad"

        def run(session: Path, payload: dict, allowed: str = "",
                env: dict | None = None) -> tuple[int, str]:
            r = subprocess.run([sys.executable, "-I", str(guard)], input=json.dumps(payload),
                               capture_output=True, text=True,
                               env={**os.environ, "CLAUDE_PROJECT_DIR": str(session),
                                    "CLAUDE_CONFIG_DIR": str(config),
                                    "CLAUDE_CODE_TMPDIR": str(temp),
                                    "CLAUDE_PLUGIN_OPTION_ALLOWED_DIRS": allowed, **(env or {})})
            return r.returncode, r.stderr

        def write(path, **extra) -> dict:
            return {"tool_name": "Write", "tool_input": {"file_path": str(path)}, **extra}

        memory = config / "projects"
        cases = (
            ("its own repository", docs.path, write(docs.path / "a.md"), "", 0),
            ("a worktree of its repository", docs.path, write(work / "docs--x" / "a.md"), "", 0),
            ("the main checkout, from a worktree", work / "docs--x", write(docs.path / "a.md"),
             "", 0),
            ("the exchange", code.path, write(work / "_scambio" / "a.md"), "", 0),
            ("the scratchpad it is given", docs.path,
             write(scratch / "a.md", scratchpad_dir=str(scratch)), "", 0),
            *((("the scratchpad of its session, not given", docs.path,
                 write(session_scratch / "a.md", session_id="s1"), "", 0),
               ("the scratchpad of another session", docs.path,
                 write(session_scratch / "a.md", session_id="s2"), "", 2)) if uid is not None
              else ()),
            ("its project's memory", docs.path,
             write(memory / slug(docs.path) / "memory" / "m.md"), "", 0),
            ("the memory of its main checkout, from a worktree", work / "docs--x",
             write(memory / slug(docs.path) / "memory" / "m.md"), "", 0),
            ("another project's memory", docs.path, write(memory / "p" / "memory" / "m.md"),
             "", 2),
            ("a plan", docs.path, write(config / "plans" / "p.md"), "", 0),
            ("a folder of the option, as a JSON list", docs.path,
             write(work / "other" / "a.md"), json.dumps(["other"]), 0),
            ("a folder of the option, one per line", docs.path, write(work / "other" / "a.md"),
             f"{scratch}\nother", 0),
            ("another repository of the working folder", code.path,
             write(docs.path / "a.md"), "", 2),
            ("a folder beside the repositories", docs.path, write(work / "other" / "a.md"),
             "", 2),
            ("a link inside the repository to a folder outside", docs.path,
             write(docs.path / "link" / "a.md"), "", 2),
            ("a path that climbs out with `..`", docs.path,
             write(f"{docs.path}/../other/a.md"), "", 2),
            ("Claude Code's settings", docs.path, write(config / "settings.json"), "", 2),
            ("a notebook outside", docs.path, {"tool_name": "NotebookEdit", "tool_input": {
                "notebook_path": str(work / "other" / "n.ipynb")}}, "", 2),
            ("a call with no path it knows", docs.path,
             {"tool_name": "NotebookEdit", "tool_input": {"path": str(docs.path / "n.ipynb")}},
             "", 2),
            ("anything, from a repository outside the working folder", elsewhere.path,
             write(work / "other" / "a.md"), "", 0),
            ("anything, from the parent folder, which is no repository", work,
             write(work / "other" / "a.md"), "", 0))
        locked = work / "a-locked"
        try:
            # A folder beside the repositories nobody can look into is not an error that
            # turns the guard off. Run as root, the case is not exercised: root looks anyway.
            locked.chmod(0)
            for what, session, payload, allowed, want in cases:
                got, said = run(session, payload, allowed)
                if got != want:
                    problems.append(f"the guard exits {got} on {what}, expected {want}: "
                                    f"{said.strip()[-200:]}")
                elif got == 2 and "ask the person" not in said:
                    problems.append(f"the guard stops {what} without saying to ask the person")
        finally:
            locked.chmod(0o755)

        # The place `.framework/sessions.yaml` declares is the exchange of the repository that
        # declares it, of no other, once committed -- the file on disk is one the session may
        # write -- and only as a folder beside it that is not a repository.
        shared = write(work / "shared" / "a.md")
        docs.write(".framework/sessions.yaml", "exchange: ../shared  # beside the others\n")
        if run(docs.path, shared)[0] != 2:
            problems.append("the guard takes the exchange from a sessions.yaml nobody committed")
        docs.commit("the exchange", "2026-01-02T10:00:00")
        if run(docs.path, shared)[0] != 0:
            problems.append("the guard stops a write in the exchange the repository declares")
        if run(code.path, shared)[0] != 2:
            problems.append("the guard lets a code repository write in the place another "
                            "repository declares as its exchange")
        for value, target in (("..", work / "x.md"), ("../code", code.path / "a.md"),
                              (str(work / "other"), work / "other" / "a.md"),
                              ("../other/..", work / "x.md")):
            docs.write(".framework/sessions.yaml", f"exchange: {value}\n")
            docs.commit(f"exchange {value}", "2026-01-03T10:00:00")
            if run(docs.path, write(target))[0] != 2:
                problems.append(f"the guard takes `exchange: {value}` as a place to write: "
                                "only a folder beside the repository that is not one counts")
        # The scratchpad where Claude Code puts it with no `CLAUDE_CODE_TMPDIR`, in the
        # temporary folder, and under a `CLAUDE_CODE_TMPDIR` without `claude-<uid>`.
        if uid is not None:
            plain = tmp / "plain"           # `tempfile` passes over a TMPDIR that is not there
            plain.mkdir()
            for env, place in (({"CLAUDE_CODE_TMPDIR": "", "TMPDIR": str(plain)},
                                plain / f"claude-{uid}"), ({}, temp)):
                target = place / slug(docs.path) / "s1" / "scratchpad" / "a.md"
                if run(docs.path, write(target, session_id="s1"), env=env)[0] != 0:
                    problems.append(f"the guard stops a write in the session's scratchpad at "
                                    f"{target.relative_to(tmp)}")
        # Nor is `_scambio` a way out when it is a link to the parent folder.
        (work / "_scambio").rmdir()
        (work / "_scambio").symlink_to(work)
        if run(code.path, write(work / "x.md"))[0] != 2:
            problems.append("the guard takes `_scambio`, a link to the parent folder, as the "
                            "exchange: every folder beside the repositories is open")
    return problems


@check("every status a change or an entry can have is either done or not, for the exchange")
def _exchange_maps_every_status():
    # Two meanings of "closed" already exist, the digest's and this one, and the second is
    # earlier on purpose. A status the schemas gain and neither set names would read as not
    # done for ever, and its resources would never be proposed for removal.
    x = _load(VALIDATE, "validate_exchange_states")._exchange
    problems = []
    for kind, path, done, not_done in (
            ("change-contract", ("properties", "status", "enum"), x.CLOSED_CHANGE, x.OPEN_CHANGE),
            ("open-register", ("properties", "entries", "additionalProperties", "properties",
                               "status", "enum"), x.CLOSED_ENTRY, x.OPEN_ENTRY)):
        node = json.loads((ROOT / "schemas" / "framework" / kind / "v1.json").read_text())
        for key in path:
            node = node[key]
        if set(node) != done | not_done or done & not_done:
            problems.append(f"{kind}: the schema allows {sorted(node)}, the exchange calls "
                            f"{sorted(done)} done and {sorted(not_done)} not")
    return problems


@check("a value written in risorse.md is never repeated, by the validator or by the listing")
def _exchange_keeps_values_out():
    # The list names a copy of credentials and its keys. A row that carries a value anyway is
    # `EXC005`, and every other place that prints a row -- `EXC004`, `EXC007`, the listing and
    # its JSON -- must not become the second copy of the secret.
    problems = []
    with tempfile.TemporaryDirectory() as tmp, tempfile.TemporaryDirectory() as scratch:
        repo, ex, _ = _exchange_docs(Path(tmp))
        copy = Path(scratch) / ".env.misure"
        copy.write_text("x")
        _put(ex, {f"in-arrivo/{i}/risorse.md": _RESOURCE_HEAD
                  + f"| copia di credenziali | {copy} (DB_PASSWORD=hunter2) | rm {copy} |\n"
                  + f"| copia di credenziali | {copy} (DB_URL) | rm {copy} -e TOKEN=hunter2 |\n"
                  + "| copia di credenziali | .env (DB_PASSWORD=hunter2) | rm .env |\n"
                  for i in ("CHG-001", "CHG-002")})
        rc, out, err = repo.validate()
        outputs = {"validate --json": json.dumps(out or {}) + err}
        for extra in ((), ("--json",)):
            r = subprocess.run([sys.executable, str(EXCHANGE), "--root", str(repo.path), *extra],
                               capture_output=True, text=True)
            outputs[f"exchange.py {' '.join(extra)}"] = r.stdout + r.stderr
        for name, text in outputs.items():
            if "hunter2" in text:
                problems.append(f"{name} repeats a value written in risorse.md")
        if not any(f["code"] == "EXC004" for f in (out or {}).get("findings", [])):
            problems.append("the row with a value is not reported as present at all: the "
                            "test proves nothing")
        if len([f for f in (out or {}).get("findings", []) if f["code"] == "EXC005"]) != 4:
            problems.append("each row carrying a value, in either cell, is an EXC005")
    return problems


@check("an annotation on an EXC finding is not obsolete where the exchange is not on the machine")
def _exchange_annotations_travel():
    # Annotations join on (code, path), and the exchange is on the machines where sessions
    # work and nowhere else. Written there, an annotation on an `EXC` would read as `AN001`,
    # obsolete, on every CI run, where nothing about the exchange ran.
    problems = []
    with tempfile.TemporaryDirectory() as tmp:
        repo, ex, _ = _exchange_docs(Path(tmp))
        repo.write(".framework/expected-findings.yaml",
                   "expected:\n  - code: EXC001\n    path: ../_scambio/in-arrivo/CHG-009\n"
                   "    reason: a delivery for a change still being written\n"
                   "    clears_when: CHG-009 is written\n")
        _put(ex, {"in-arrivo/CHG-009/brief.md": "b", "Thumbs.db": "", "in-arrivo/desktop.ini": "",
                  "integrati/CHG-001/brief.md": "b",
                  "integrati/CHG-001/INTEGRATO": "commit: HEAD\ndate: 2026-10-14\nparts: {}\n"})
        sha = repo.git("rev-parse", "HEAD")
        (ex / "integrati/CHG-001/INTEGRATO").write_text(
            f"commit: {sha[:12]}\ndate: 2026-10-14\nparts: {{}}\n")
        rc, out, err = repo.validate()
        codes = sorted((f["code"], bool(f.get("accepted"))) for f in (out or {}).get("findings", [])
                       if f["code"].startswith(("EXC", "AN")))
        if codes != [("EXC001", True)]:
            problems.append(f"with the exchange here: {codes}, expected the annotated EXC001 "
                            "alone -- litter is not a name, and a folder that delivered nothing "
                            "moves with `parts: {}`")
        shutil.rmtree(ex)
        rc, out, err = repo.validate()
        if any(f["code"].startswith(("EXC", "AN")) for f in (out or {}).get("findings", [])):
            problems.append(f"without the exchange: {[f['code'] for f in out['findings']]}, "
                            "expected no EXC and no AN001")
    return problems


@check("BR001-BR003 compare what a branch adds with the main line, and nothing they share")
def _branch_checks():
    # Two people resolving one open decision, each with the next free number: the case the
    # checks exist for, and the history both sides share, which they must never report.
    problems = []
    with tempfile.TemporaryDirectory() as tmp:
        repo = _Repo(Path(tmp) / "r")
        repo.write("OPEN.md", _open_register({"OD-001": "open", "OD-003": "open"}))
        # History nobody can undo: one entry closed by two accepted decisions.
        repo.write("decisions/DEC-010.md", _decision("DEC-010", "OD-003"))
        repo.write("decisions/DEC-011.md", _decision("DEC-011", "OD-003"))
        repo.commit("base", "2026-06-01T09:00:00+02:00")
        repo.git("checkout", "-q", "-b", "work/a")
        repo.write("decisions/DEC-012-branch.md", _decision("DEC-012", "OD-001"))
        repo.write("OPEN.md", _open_register({"OD-001": "open", "OD-002": "open", "OD-003": "open"}))
        repo.commit("branch", "2026-06-02T09:00:00+02:00")
        repo.git("checkout", "-q", "main")
        repo.write("decisions/DEC-012-main.md", _decision("DEC-012", "OD-001"))
        repo.write("OPEN.md", _open_register({"OD-001": "decided", "OD-002": "open", "OD-003": "open"}))
        repo.commit("main decides", "2026-06-02T10:00:00+02:00")
        repo.git("checkout", "-q", "work/a")

        def codes(*extra):
            rc, out, err = repo.validate(*extra)
            if out is None:
                raise RuntimeError(f"no report (exit {rc}): {err}")
            return sorted((f["code"], f["path"]) for f in out["findings"]
                          if f["code"][:2] in ("BR", "WK"))

        if codes():
            problems.append(f"without --base the branch checks ran: {codes()}")
        got = codes("--base", "main")
        want = [("BR001", "OPEN.md"), ("BR001", "decisions/DEC-012-branch.md"),
                ("BR002", "decisions/DEC-012-branch.md")]
        if got != want:
            problems.append(f"two sides creating DEC-012 and OD-002 and deciding OD-001 gave "
                            f"{got}, expected {want}; and OD-003, closed twice before the "
                            "branch existed, must not appear")

        # After the main line is merged in, the collision is gone from the comparison --
        # `work.py sync` is what refuses to make that merge -- and BR002 still holds.
        repo.git("checkout", "-q", "main")
        repo.git("branch", "-q", "work/b", "HEAD~1")
        repo.git("checkout", "-q", "work/b")
        repo.write("decisions/DEC-013.md", _decision("DEC-013", "OD-001"))
        repo.commit("another decision on OD-001", "2026-06-03T09:00:00+02:00")
        repo.git("merge", "-q", "--no-edit", "main", env={"GIT_AUTHOR_DATE": "2026-06-04T09:00:00+02:00",
                                                         "GIT_COMMITTER_DATE": "2026-06-04T09:00:00+02:00"})
        got = codes("--base", "main")
        if got != [("BR002", "decisions/DEC-013.md")]:
            problems.append(f"a branch that merged the main line and still decides an entry "
                            f"the main line closed gave {got}: BR002 is read by state")
        # Reopening on purpose is a supersession, and is not reported.
        repo.write("decisions/DEC-013.md", _decision("DEC-013", "OD-001", supersedes="[DEC-012]"))
        repo.commit("supersede", "2026-06-05T09:00:00+02:00")
        if codes("--base", "main"):
            problems.append(f"a decision that supersedes the one that closed the entry gave "
                            f"{codes('--base', 'main')}")

        # BR003: only where the project declares work.yaml, and a tag clears it.
        repo.write(".framework/work.yaml", "people:\n  Owner: {}\nlater_key: 1\n")
        repo.commit("declare", "2026-06-05T10:00:00+02:00")
        got = codes("--base", "main")
        if got != [("BR003", "decisions/DEC-013.md"), ("WK001", ".framework/work.yaml")]:
            problems.append(f"a new number with no reservation, and an unknown key, gave {got}")
        repo.git("tag", "id/DEC-013")
        if ("BR003", "decisions/DEC-013.md") not in codes("--base", "main"):
            problems.append("a lightweight tag, which says for whom nothing, cleared BR003")
        repo.git("tag", "-d", "id/DEC-013")
        repo.git("tag", "-a", "id/DEC-013", "-m", "reserve DEC-013\n\nwho: Owner\nbranch: work/b\n")
        if codes("--base", "main") != [("WK001", ".framework/work.yaml")]:
            problems.append("a reservation made for this branch did not clear BR003")

        rc, out, err = repo.validate("--base", "no-such-ref")
        if rc != 2 or out is not None:
            problems.append(f"--base naming a ref that does not exist gave exit {rc}: it has "
                            "to refuse, not report")
    return problems


_WORK_ENV = {"GIT_AUTHOR_NAME": "t", "GIT_AUTHOR_EMAIL": "t@example.com",
             "GIT_COMMITTER_NAME": "t", "GIT_COMMITTER_EMAIL": "t@example.com",
             "GIT_CONFIG_GLOBAL": os.devnull, "GIT_CONFIG_NOSYSTEM": "1"}


def _g(where: Path, *args: str, ok: bool = True) -> subprocess.CompletedProcess:
    r = subprocess.run(["git", "-C", str(where), *args], capture_output=True, text=True,
                       env={**os.environ, **_WORK_ENV})
    if ok and r.returncode:
        raise RuntimeError(f"git {' '.join(args)} in {where.name}: {r.stderr}")
    return r


def _work_origin(tmp: Path) -> tuple[Path, Path, Path]:
    """A remote and two clones of it, side by side under different parents, whose root
    register carries the union as the generator writes it."""
    origin = tmp / "origin.git"
    subprocess.run(["git", "init", "-q", "--bare", "-b", "main", str(origin)], check=True,
                   env={**os.environ, **_WORK_ENV})
    seed = _Repo(tmp / "seed")
    seed.write(".framework/work.yaml", "people:\n  Owner: {}\n  Junior: {}\n")
    seed.write("OPEN.md", _open_register({"OD-001": "open", "OD-002": "open"}))
    seed.write("platform/OPEN.md", _open_register({"OD-003": "open"}).replace(
        "products: [all]\nowners", "products: [none]\nowners", 1))
    seed.validate("--emit-index")
    seed.commit("base", "2026-06-01T09:00:00+02:00")
    seed.git("remote", "add", "origin", str(origin))
    seed.git("push", "-q", "origin", "main")
    clones = []
    for parent in ("projects", "other"):
        c = tmp / parent / "docs"
        c.parent.mkdir(parents=True)
        subprocess.run(["git", "clone", "-q", str(origin), str(c)], check=True,
                       env={**os.environ, **_WORK_ENV})
        clones.append(c)
    return origin, clones[0], clones[1]


def _work(root: Path, *args: str, env: dict | None = None) -> tuple[int, str]:
    r = subprocess.run([sys.executable, str(WORK), "--root", str(root), "--retries", "1",
                        "--pause", "0", *args], capture_output=True, text=True,
                       env={**os.environ, **_WORK_ENV, **(env or {})})
    return r.returncode, r.stdout + r.stderr


def _commit_all(where: Path, msg: str) -> None:
    _g(where, "add", "-A")
    _g(where, "commit", "-qm", msg)


@check("work.py: one claim per element and one holder per number, refused to whoever pushes second")
def _work_claims_and_reserves():
    problems = []
    with tempfile.TemporaryDirectory() as tmp:
        _, a, b = _work_origin(Path(tmp))
        rc, out = _work(a, "open", "--who", "Owner", "--branch", "od-001", "OD-001")
        tree = a.parent / "docs--od-001"
        if rc or not tree.is_dir():
            problems.append(f"the first claim failed or made no worktree beside the checkout "
                            f"(exit {rc}): {out}")
        rc, out = _work(b, "open", "--who", "Junior", "--branch", "mine", "OD-001", "--no-worktree")
        if rc != 1 or "claimed by Owner on work/od-001" not in out:
            problems.append(f"a second claim on OD-001 gave exit {rc}: {out!r}. It has to be "
                            "refused with who holds it and where")
        rc, out = _work(b, "open", "--who", "Junior", "--branch", "od-001", "OD-002", "--no-worktree")
        if rc == 0:
            problems.append("a second `open` naming an existing branch went through: the branch "
                            "is created with a lease that requires it not to exist")
        for args, why in ((("open", "--who", "Junior", "--branch", "x", "OD-009",
                            "--no-worktree"), "an element nobody declared"),
                          (("open", "--who", "Nobody", "--branch", "x", "OD-002",
                            "--no-worktree"), "a person work.yaml does not name"),
                          (("reserve", "DEC", "--who", "Junior"), "a reservation from the main line"),
                          (("reserve", "RSK", "--who", "Owner"), "a per-product family without --product")):
            rc, out = _work(b if why != "a per-product family without --product" else tree, *args)
            if rc != 2:
                problems.append(f"{why} gave exit {rc}: refused, with exit 2: {out[-200:]!r}")
        rc, out = _work(b, "open", "--who", "Junior", "--branch", "od-002", "OD-002", "--no-worktree")
        _g(b, "checkout", "-q", "work/od-002")
        rc1, out1 = _work(tree, "reserve", "DEC", "--who", "Owner", "--count", "2")
        rc2, out2 = _work(b, "reserve", "DEC", "--who", "Junior")
        if rc1 or rc2 or "DEC-001, DEC-002" not in out1 or "DEC-003" not in out2:
            problems.append(f"two people reserving DEC got {out1!r} and {out2!r}: three "
                            "numbers, none shared")
        rc, out = _work(tree, "reserve", "RSK", "--product", "platform", "--who", "Owner")
        if rc or "platform:RSK-001" not in out or _g(a, "rev-parse", "-q", "--verify",
                                                      "refs/tags/id/platform/RSK-001", ok=False).returncode:
            problems.append(f"a per-product reservation was not made under its scope: {out!r}")
        rc, out = _work(tree, "reserve", "RSK", "--product", "atlsa", "--who", "Owner")
        if rc != 2:
            problems.append("a reservation under a product nobody declares was made")
        (tree / "products" / "atlas").mkdir(parents=True)
        (tree / "products" / "atlas" / "product.yaml").write_text(
            "schema: framework/product-manifest/v1\nartifact_type: product-manifest\n"
            "products: [atlas]\n")
        _commit_all(tree, "a product created on the branch")
        rc, out = _work(tree, "reserve", "RSK", "--product", "atlas", "--who", "Owner")
        if rc or "atlas:RSK-002" not in out or _g(a, "rev-parse", "-q", "--verify",
                                                   "refs/tags/id/atlas/RSK-002", ok=False).returncode:
            problems.append(f"a product created on the branch could not reserve its first number, "
                            f"or it was not filed under the product: {out!r}")
        rc, out = _work(b, "list")
        if rc or "work/od-001" not in out or "claims: OD-001" not in out:
            problems.append(f"list does not show the branch and its claim: {out!r}")
        rc, out = _work(b, "release", "OD-001", "--who", "Junior")
        if rc != 2:
            problems.append("somebody else's claim was released without `--theirs`")
        # A third clone sees the claim before it is released and taken again: its tag then
        # points at an object the remote no longer has, and a plain fetch refuses to move it.
        c = Path(tmp) / "third" / "docs"
        c.parent.mkdir(parents=True)
        subprocess.run(["git", "clone", "-q", str(Path(tmp) / "origin.git"), str(c)], check=True,
                       env={**os.environ, **_WORK_ENV})
        _work(c, "list")
        rc, out = _work(tree, "release", "OD-001", "--who", "Owner")
        rc2, out2 = _work(b, "open", "--who", "Junior", "--branch", "after", "OD-001",
                          "--no-worktree")
        if rc or rc2:
            problems.append(f"after its holder released it, OD-001 could not be claimed: {out2!r}")
        rc, out = _work(c, "list")
        rows, current = {}, None
        for line in out.splitlines():
            if line.startswith("work/"):
                current = line.split(":", 1)[0]
                rows[current] = []
            elif current and line.startswith("  claims:"):
                rows[current] += [t.strip() for t in line.split(":", 1)[1].split(",")]
        if rc or "OD-001" not in rows.get("work/after", []) or "OD-001" in rows.get("work/od-001", []):
            problems.append(f"a clone that saw a released claim taken again could not list, or "
                            f"listed the old holder: {rows or out!r}")
        rc, out = _work(tree, "reserve", "DEC", "--who", "Owner")
        if "DEC-004" not in out:
            problems.append(f"a released claim gave back a number, or reservations were "
                            f"deleted: {out!r}")
    return problems


def _work_sync_modify_delete(tree: Path, b: Path) -> list[str]:
    """A file one side deleted and the other changed: left for a person, and `--continue`
    refuses until it is resolved and no marker is staged."""
    problems = []
    _g(b, "pull", "-q", "origin", "main")
    (b / "notes.md").write_text("one\n")
    _commit_all(b, "a note")
    _g(b, "push", "-q", "origin", "main")
    _work(tree, "sync", "--no-push")
    (tree / "notes.md").write_text("one, changed on the branch\n")
    _commit_all(tree, "branch changes the note")
    _g(b, "rm", "-q", "notes.md")
    _commit_all(b, "main deletes the note")
    _g(b, "push", "-q", "origin", "main")
    rc, out = _work(tree, "sync", "--no-push")
    if rc != 1 or "deleted on the side merged in" not in out:
        problems.append(f"a modify/delete conflict was resolved by the script (exit {rc}): {out!r}")
    rc, out = _work(tree, "sync", "--continue", "--no-push")
    if rc != 2:
        problems.append("`--continue` with a file still unmerged went through")
    _g(tree, "add", "notes.md")
    (tree / "OPEN.md").write_text((tree / "OPEN.md").read_text() + "\n<<<<<<< HEAD\nx\n")
    _g(tree, "add", "OPEN.md")
    rc, out = _work(tree, "sync", "--continue", "--no-push")
    if rc != 2 or "conflict markers" not in out:
        problems.append(f"`--continue` with markers staged went through: {out!r}")
    # Resolved for real, while another session moved `origin/main` under this merge: the merge
    # `sync` began is still the one `--continue` finishes.
    text = (tree / "OPEN.md").read_text()
    (tree / "OPEN.md").write_text(text[:text.rindex("\n<<<<<<< HEAD\nx\n")] + "\n")
    _g(tree, "add", "OPEN.md")
    _g(b, "pull", "-q", "origin", "main", ok=False)
    (b / "later.md").write_text("later\n")
    _commit_all(b, "the main line moves on")
    _g(b, "push", "-q", "origin", "main")
    _g(tree, "fetch", "-q", "origin")
    rc, out = _work(tree, "sync", "--continue", "--no-push")
    if rc != 0 or "Merged" not in out:
        problems.append(f"`--continue` refused the merge `sync` began once origin/main had moved: "
                        f"{out!r}")
    return problems


@check("work.py sync merges the main line, resolves only what is mechanical, and leaves every other conflict whole")
def _work_sync():
    problems = []
    x = _load(WORK, "work_for_sync")
    with tempfile.TemporaryDirectory() as tmp:
        _, a, b = _work_origin(Path(tmp))
        _work(a, "open", "--who", "Owner", "--branch", "od-001", "OD-001")
        tree = a.parent / "docs--od-001"
        rc, out = _work(tree, "sync", "--no-push")
        if rc or "already contains" not in out:
            problems.append(f"a branch already up to date was not said to be: {out!r}")

        # The branch decides OD-001, adds a parking line and attests; the main line adds a
        # parking line and an entry elsewhere, regenerates the union, and attests too. Every
        # conflict is in the root register, which carries the generator's mark in its union.
        text = (tree / "OPEN.md").read_text()
        text = text.replace("  OD-001:\n    status: open", "  OD-001:\n    status: decided", 1)
        text = text.replace("- first line\n", "- first line\n- a line from the branch\n"
                            "- **Cost to reverse:** low\n", 1)
        text = text.replace("last_review: 2026-06-01 09:00", "last_review: 2026-06-02 10:00 +02:00", 1)
        (tree / "OPEN.md").write_text(text)
        _commit_all(tree, "branch works")
        _g(b, "pull", "-q", "origin", "main")
        text = (b / "OPEN.md").read_text()
        text = text.replace("- first line\n", "- first line\n- a line from main\n"
                            "- **Cost to reverse:** low\n", 1)
        text = text.replace("last_review: 2026-06-01 09:00", "last_review: 2026-06-03 10:00 +02:00", 1)
        (b / "OPEN.md").write_text(text)
        plat = (b / "platform" / "OPEN.md").read_text().replace(
            "  OD-003:\n", "  OD-004:\n    status: open\n    cost_to_reverse: low\n"
            "    default_in_force: none\n    products: [all]\n  OD-003:\n", 1)
        (b / "platform" / "OPEN.md").write_text(plat)
        subprocess.run([sys.executable, str(VALIDATE), "--root", str(b), "--emit-index",
                        "--stale-days", "36500"], capture_output=True)
        _commit_all(b, "main moves")
        _g(b, "push", "-q", "origin", "main")
        rc, out = _work(tree, "sync", "--no-push")
        text = (tree / "OPEN.md").read_text()
        if rc or "<<<<<<<" in text:
            problems.append(f"conflicts that are an attestation block, two parking lines and the "
                            f"union were not resolved (exit {rc}): {out}")
        elif "status: decided" not in text or "a line from the branch" not in text:
            problems.append("sync lost the branch's version of the root register: the union "
                            "carries the generator's mark, and the file is not the generator's")
        elif not ("- a line from the branch\n- **Cost to reverse:** low\n- a line from main\n"
                  "- **Cost to reverse:** low" in text):
            problems.append("lines both sides added were not all kept, in order: a repeated "
                            "label is not a duplicate of the other side's")
        elif "2026-06-02 10:00" not in text or "OD-004" not in text:
            problems.append("sync kept the wrong stamp, or did not regenerate the union")
        elif "OPEN.md" not in out or "attest" not in out:
            problems.append("sync did not name the root register, whose text came from both "
                            f"sides, among the readings owed: {out!r}")

        # A file one side deleted and the other changed is a person's.
        try:
            problems += _work_sync_modify_delete(tree, b)
        except RuntimeError as e:
            problems.append(f"the modify/delete case did not run to its end: {e}")

        # A collision is refused before anything is merged.
        _g(b, "pull", "-q", "origin", "main")
        (b / "decisions").mkdir(exist_ok=True)
        (b / "decisions" / "DEC-009-main.md").write_text(_decision("DEC-009", "OD-002"))
        _commit_all(b, "main creates DEC-009")
        _g(b, "push", "-q", "origin", "main")
        (tree / "decisions").mkdir(exist_ok=True)
        (tree / "decisions" / "DEC-009-branch.md").write_text(_decision("DEC-009", "OD-001"))
        _commit_all(tree, "branch creates DEC-009")
        head = _g(tree, "rev-parse", "HEAD").stdout
        rc, out = _work(tree, "sync", "--no-push")
        if rc != 1 or "DEC-009" not in out or head != _g(tree, "rev-parse", "HEAD").stdout:
            problems.append(f"both sides creating DEC-009 was merged, or not named (exit {rc}): "
                            f"{out!r}")

    # Content in conflict is left with its three versions; unreadable markers leave the file.
    text = "\n".join(["# Doc", "", "<<<<<<< HEAD", "ours", "||||||| base", "base", "=======",
                      "theirs", ">>>>>>> origin/main", ""])
    resolved, left, _ = x.resolve(text, "products/p/PBR.md")
    if left != 1 or "||||||| base" not in resolved:
        problems.append(f"a conflict in content was resolved by the script: {resolved!r}")
    fm = "\n".join(["---", "<<<<<<< HEAD", "status: a", "||||||| base", "=======",
                    "status: b", ">>>>>>> origin/main", "---", ""])
    if x.resolve(fm, "products/p/PBR.md")[1] != 1:
        problems.append("two keys added at one place of the front matter were both kept")
    spans = "\n".join(["---", "entries:", "<<<<<<< HEAD", "  - x", "last_review: 1", "||||||| base",
                       "  - y", "last_review: 0", "=======", "  - z", "last_review: 2",
                       ">>>>>>> origin/main", "---", ""])
    if x.resolve(spans, "products/p/PBR.md")[1] != 1:
        problems.append("a conflict spanning a list and the stamp was resolved as an attestation")
    yml = "\n".join(["a: 1", "<<<<<<< HEAD", "b: 2", "||||||| base", "=======", "c: 3",
                     ">>>>>>> origin/main", ""])
    if x.resolve(yml, ".framework/expected-findings.yaml")[1] != 1:
        problems.append("lines added at one place of a YAML file were joined by the script")
    for broken in ("<<<<<<< HEAD\nx\n", "<<<<<<< HEAD\nx\n=======\ny\n>>>>>>> m\n"):
        try:
            x.resolve(broken, "products/p/PBR.md")
            problems.append(f"markers it cannot read were resolved: {broken!r}")
        except x.Unreadable:
            pass
    return problems


@check("work.py: a push whose answer was lost is not read as somebody else's")
def _work_lost_answers():
    # The server of one project drops a share of its connections. A push that landed and whose
    # answer was lost, tried again, is refused -- the tag is there, as ours -- and that refusal
    # read as "taken" burned a number and claimed somebody had the element. The remote is asked.
    problems = []
    x = _load(WORK, "work_for_lost_answers")
    real = x.subprocess.run

    def lossy(cmd, *a, **k):
        r = real(cmd, *a, **k)
        if "push" in cmd and r.returncode == 0:
            return subprocess.CompletedProcess(cmd, 128, "", "fatal: the remote end hung up unexpectedly")
        return r

    with tempfile.TemporaryDirectory() as tmp:
        _, a, _ = _work_origin(Path(tmp))
        _work(a, "open", "--who", "Owner", "--branch", "od-001", "OD-001", "--no-worktree")
        _g(a, "checkout", "-q", "work/od-001")
        args = argparse.Namespace(root=a, remote="origin", main="main", retries=3, pause=0,
                                  prefix="DEC", who="Owner", count=1, product=None)
        x.subprocess.run = lossy
        out = io.StringIO()
        try:
            with contextlib.redirect_stdout(out):
                rc = x.cmd_reserve(x.Project(args), args)
        except SystemExit as e:
            rc = e.code
        finally:
            x.subprocess.run = real
        held = _g(a, "ls-remote", "origin", "refs/tags/id/*").stdout
        if rc != 0 or "DEC-001" not in out.getvalue() or "DEC-002" in held:
            problems.append(f"a reservation whose answer was lost gave exit {rc} and "
                            f"{out.getvalue()!r}, with {held.split() if held else 'nothing'} on "
                            "the remote: one number, reported as reserved")
        args = argparse.Namespace(root=a, remote="origin", main="main", retries=3, pause=0,
                                  ids=["OD-002"], who="Owner", branch="od-002", no_worktree=True)
        x.subprocess.run = lossy
        out = io.StringIO()
        try:
            with contextlib.redirect_stdout(out):
                rc = x.cmd_open(x.Project(args), args)
        except SystemExit as e:
            rc = e.code
        finally:
            x.subprocess.run = real
        if rc != 0 or "Somebody pushed first" in out.getvalue():
            problems.append(f"a claim whose answer was lost was reported as another's: "
                            f"{out.getvalue()!r}")

        # The other way round: the first attempt never reaches the remote, the second does. The
        # tag made here has to be there for the second attempt to push.
        attempts = []

        def dropped_once(cmd, *a, **k):
            if "push" in cmd and not attempts:
                attempts.append(cmd)
                return subprocess.CompletedProcess(cmd, 128, "", "fatal: unable to access: Connection reset")
            return real(cmd, *a, **k)

        args = argparse.Namespace(root=a, remote="origin", main="main", retries=3, pause=0,
                                  prefix="DEC", who="Owner", count=1, product=None)
        x.subprocess.run = dropped_once
        out = io.StringIO()
        try:
            with contextlib.redirect_stdout(out):
                rc = x.cmd_reserve(x.Project(args), args)
        except SystemExit as e:
            rc = e.code
        finally:
            x.subprocess.run = real
        if rc != 0 or "Reserved: DEC-002" not in out.getvalue():
            problems.append(f"a push that never landed was not tried again (exit {rc}): "
                            f"{out.getvalue()!r}")

        # A remote that answers neither the push nor the question: the number may be there, and
        # the person is told so rather than that it is not.
        def silent(cmd, *a, **k):
            if "push" in cmd or "ls-remote" in cmd:
                return subprocess.CompletedProcess(cmd, 128, "", "fatal: Could not read from remote")
            return real(cmd, *a, **k)

        x.subprocess.run = silent
        err = io.StringIO()
        try:
            with contextlib.redirect_stderr(err), contextlib.redirect_stdout(io.StringIO()):
                rc = x.cmd_reserve(x.Project(args), args)
        except SystemExit as e:
            rc = e.code
        finally:
            x.subprocess.run = real
        if rc != 2 or "may or may not have landed" not in err.getvalue():
            problems.append(f"a remote that never answered was reported as having refused "
                            f"(exit {rc}): {err.getvalue()!r}")
        for cmd, ns in (("open", argparse.Namespace(root=a, remote="origin", main="main",
                                                     retries=1, pause=0, ids=["OD-003"],
                                                     who="Owner", branch="unanswered",
                                                     no_worktree=True)),
                        ("release", argparse.Namespace(root=a, remote="origin", main="main",
                                                        retries=1, pause=0, ids=["OD-001"],
                                                        who="Owner", theirs=False))):
            x.subprocess.run = silent
            err = io.StringIO()
            try:
                with contextlib.redirect_stderr(err), contextlib.redirect_stdout(io.StringIO()):
                    rc = getattr(x, f"cmd_{cmd}")(x.Project(ns), ns)
            except SystemExit as e:
                rc = e.code
            finally:
                x.subprocess.run = real
            said = err.getvalue()
            if rc != 2 or "may" not in said or (cmd == "open" and "worktree add" in said):
                problems.append(f"`{cmd}` against a remote that never answered said {said!r}")
            if cmd == "open":
                _g(a, "branch", "-D", "work/unanswered", ok=False)
    return problems


def _work_sync_edges_cases() -> list[str]:
    problems = []
    x = _load(WORK, "work_for_sync_edges")
    with tempfile.TemporaryDirectory() as tmp:
        _, a, b = _work_origin(Path(tmp))
        _work(a, "open", "--who", "Owner", "--branch", "od-001", "OD-001")
        tree = a.parent / "docs--od-001"
        _g(b, "pull", "-q", "origin", "main")
        (b / ".gitattributes").write_text("*.csv binary\n")
        (b / "data.csv").write_text("a,b\n1,2\n")
        _commit_all(b, "a table git does not merge by line")
        _g(b, "push", "-q", "origin", "main")
        _work(tree, "sync", "--no-push")
        (tree / "data.csv").write_text("a,b\n1,3\n")
        _commit_all(tree, "branch changes the table")
        (b / "data.csv").write_text("a,b\n1,4\n")
        _commit_all(b, "main changes the table")
        _g(b, "push", "-q", "origin", "main")
        rc, out = _work(tree, "sync", "--no-push")
        if rc != 1 or "without markers" not in out or "1,4" in (tree / "data.csv").read_text() \
                and "1,3" not in (tree / "data.csv").read_text():
            problems.append(f"a file in conflict without markers was resolved by the script "
                            f"(exit {rc}): {out!r}")
        _g(tree, "merge", "--abort", ok=False)
        here = _g(tree, "rev-parse", "--abbrev-ref", "HEAD").stdout.strip()
        _g(tree, "checkout", "-q", "-b", "elsewhere", "HEAD~1")
        (tree / "other.md").write_text("on another branch\n")
        _commit_all(tree, "a commit somewhere else")
        _g(tree, "checkout", "-q", here)
        _g(tree, "merge", "--no-commit", "--no-ff", "-q", "elsewhere", ok=False)
        if _g(tree, "rev-parse", "-q", "--verify", "MERGE_HEAD", ok=False).returncode:
            problems.append("the case did not produce a merge in progress, so it proves nothing")
        else:
            rc, out = _work(tree, "sync", "--continue", "--no-push")
            if rc != 2 or "not the one `sync` began" not in out:
                problems.append(f"`--continue` finished a merge `sync` did not begin: {out!r}")
            _g(tree, "merge", "--abort", ok=False)
    commented = "\n".join(["---", "<<<<<<< HEAD", "last_review: 2026-06-02 10:00", "||||||| base",
                            "last_review: 2026-06-01 09:00", "=======",
                            "last_review: 2026-06-03 10:00", "  # covered only §1", ">>>>>>> origin/main",
                            "---", ""])
    if x.resolve(commented, "products/p/PBR.md")[1] != 1:
        problems.append("a comment under `last_review` on one side was dropped as part of the "
                        "stamp: it is text, by the validator's own reading of the block")
    return problems


@check("work.py sync leaves a file without markers, and a merge it did not begin, to a person")
def _work_sync_edges():
    try:
        return _work_sync_edges_cases()
    except RuntimeError as e:
        return [f"the cases did not run to their end: {e}"]


@check("work.py sync, refused because the branch moved while it merged, says to sync again, and that works")
def _work_sync_raced():
    problems = []
    x = _load(WORK, "work_for_sync_race")
    real = x.subprocess.run
    with tempfile.TemporaryDirectory() as tmp:
        _, a, b = _work_origin(Path(tmp))
        _work(a, "open", "--who", "Owner", "--branch", "od-001", "OD-001")
        tree = a.parent / "docs--od-001"
        (b / "main.md").write_text("the main line moves\n")
        _commit_all(b, "main moves")
        _g(b, "push", "-q", "origin", "main")
        _g(b, "fetch", "-q", "origin")
        _g(b, "checkout", "-q", "-b", "work/od-001", "origin/work/od-001")
        raced = []

        def race(cmd, *a2, **k):
            if "push" in cmd and any("refs/heads/work/od-001" in c for c in cmd) and not raced:
                raced.append(1)
                (b / "theirs.md").write_text("pushed while the other clone merged\n")
                _commit_all(b, "a commit on the same branch, meanwhile")
                _g(b, "push", "-q", "origin", "work/od-001")
            return real(cmd, *a2, **k)

        args = argparse.Namespace(root=tree, remote="origin", main="main", retries=1, pause=0,
                                  cont=False, no_push=False)
        x.subprocess.run = race
        out = io.StringIO()
        try:
            with contextlib.redirect_stdout(out):
                rc = x.cmd_sync(x.Project(args), args)
        except SystemExit as e:
            rc = e.code
        finally:
            x.subprocess.run = real
        if rc != 1 or "pushed from elsewhere" not in out.getvalue():
            problems.append(f"a push refused because the branch moved said {out.getvalue()!r}")
        rc, said = _work(tree, "sync")
        if rc or not (tree / "theirs.md").exists():
            problems.append(f"the sync it says to run did not bring the other clone's commit in: {said!r}")
    return problems


@check("work.py sync on one branch shared by two clones: a number both created is refused, a main line both merged is not")
def _work_sync_shared_branch():
    problems = []
    with tempfile.TemporaryDirectory() as tmp:
        _, a, b = _work_origin(Path(tmp))
        _work(a, "open", "--who", "Owner", "--branch", "od-001", "OD-001")
        tree = a.parent / "docs--od-001"
        # The branch has a commit of its own before the main line moves, so that the two merges
        # below have two best common ancestors: that commit and the main line's tip.
        (tree / "branch.md").write_text("on the branch\n")
        _commit_all(tree, "a commit on the branch")
        _g(tree, "push", "-q", "origin", "work/od-001")
        _g(b, "fetch", "-q", "origin")
        _g(b, "checkout", "-q", "-b", "work/od-001", "origin/work/od-001")
        # The main line creates a number; both clones merge it, one without pushing.
        m = Path(tmp) / "m"
        subprocess.run(["git", "clone", "-q", str(Path(tmp) / "origin.git"), str(m)], check=True,
                       env={**os.environ, **_WORK_ENV})
        (m / "decisions").mkdir()
        (m / "decisions" / "DEC-020.md").write_text(_decision("DEC-020", "OD-002"))
        _commit_all(m, "main creates DEC-020")
        _g(m, "push", "-q", "origin", "main")
        # Dated apart, so that the two merges of the same parents are two commits.
        rc1, out1 = _work(b, "sync", "--no-push", env={"GIT_AUTHOR_DATE": "2026-06-10T09:00:00+02:00",
                                                       "GIT_COMMITTER_DATE": "2026-06-10T09:00:00+02:00"})
        rc2, out2 = _work(tree, "sync", env={"GIT_AUTHOR_DATE": "2026-06-10T10:00:00+02:00",
                                             "GIT_COMMITTER_DATE": "2026-06-10T10:00:00+02:00"})
        rc3, out3 = _work(b, "sync")
        if rc1 or rc2 or rc3 or "both created" in out3 or "come in first" not in out3:
            problems.append(f"two clones that both merged the main line were refused as having "
                            f"created its numbers (exits {rc1} {rc2} {rc3}): {out3!r}")
        # Both write the same new number on the branch itself: refused before anything merges.
        for where in (tree, b):
            if _g(where, "pull", "-q", "--ff-only", "origin", "work/od-001", ok=False).returncode:
                problems.append(f"{where.name}: could not fast-forward to the shared branch")
                return problems
        (tree / "decisions" / "DEC-030-a.md").write_text(_decision("DEC-030", "OD-001"))
        _commit_all(tree, "DEC-030 here")
        _g(tree, "push", "-q", "origin", "work/od-001")
        (b / "decisions" / "DEC-030-b.md").write_text(_decision("DEC-030", "OD-001"))
        _commit_all(b, "DEC-030 there")
        head = _g(b, "rev-parse", "HEAD").stdout
        rc, out = _work(b, "sync")
        if rc != 1 or "both created DEC-030" not in out or head != _g(b, "rev-parse", "HEAD").stdout:
            problems.append(f"two clones creating DEC-030 on one branch were merged (exit {rc}): {out!r}")
    return problems


@check("work.py says which refusal it met, and the CI script tells a new error from a standing one")
def _work_refusals_and_ci_comparison():
    problems = []
    with tempfile.TemporaryDirectory() as tmp:
        origin, a, b = _work_origin(Path(tmp))
        # A remote that refuses claims by rule: not a race, and not reported as one.
        hook = origin / "hooks" / "pre-receive"
        hook.write_text("#!/bin/sh\nwhile read old new ref; do case \"$ref\" in refs/tags/claim/*) "
                        "echo 'claims are not accepted here' >&2; exit 1;; esac; done\nexit 0\n")
        hook.chmod(0o755)
        rc, out = _work(a, "open", "--who", "Owner", "--branch", "od-001", "OD-001", "--no-worktree")
        if rc != 2 or "a hook, a protection or a permission" not in out:
            problems.append(f"a claim refused by a hook was reported as a race (exit {rc}): {out!r}")
        hook.unlink()

        # A work branch someone else pushed to: `sync` brings their commits in first.
        rc, out = _work(a, "open", "--who", "Owner", "--branch", "od-001", "OD-001")
        if rc:
            problems.append(f"the claim after the hook was removed failed: {out!r}")
            return problems
        tree = a.parent / "docs--od-001"
        _g(b, "fetch", "-q", "origin")
        _g(b, "checkout", "-q", "-b", "work/od-001", "origin/work/od-001")
        (b / "theirs.md").write_text("from the other clone\n")
        _commit_all(b, "a commit on the same work branch")
        _g(b, "push", "-q", "origin", "work/od-001")
        _g(b, "checkout", "-q", "main")
        (b / "main.md").write_text("the main line moves\n")
        _commit_all(b, "main moves")
        _g(b, "push", "-q", "origin", "main")
        rc, out = _work(tree, "sync")
        if rc or "moved forward" not in out or not (tree / "theirs.md").exists() \
                or not (tree / "main.md").exists():
            problems.append(f"a sync on a work branch someone else pushed to did not move forward "
                            f"to their commits and then bring the main line in (exit {rc}): {out!r}")
        # Two clones taking turns on one branch converge: the other clone, behind by the merge
        # just pushed, moves forward to it without a merge of its own.
        _g(b, "checkout", "-q", "work/od-001")
        _g(b, "pull", "-q", "--ff-only", "origin", "work/od-001")
        if _g(b, "rev-parse", "HEAD").stdout != _g(tree, "rev-parse", "HEAD").stdout:
            problems.append("the second clone could not fast-forward to the first's merge")
        _g(b, "checkout", "-q", "main")
        merges = _g(tree, "rev-list", "--merges", "--count", f"HEAD~{0}").stdout.strip()
        if int(merges) != 1:
            problems.append(f"{merges} merge commits after one fast-forward and one merge of the "
                            "main line: a fast-forward makes none")
        # The remote refusing the branch by rule is a hook's, and said to be.
        hook.write_text("#!/bin/sh\nwhile read old new ref; do case \"$ref\" in refs/heads/work/*) "
                        "echo 'work branches are frozen' >&2; exit 1;; esac; done\nexit 0\n")
        hook.chmod(0o755)
        (b / "main2.md").write_text("again\n")
        _commit_all(b, "main moves again")
        _g(b, "push", "-q", "origin", "main")
        rc, out = _work(tree, "sync")
        if rc != 1 or "a hook or a protection rule" not in out:
            problems.append(f"a work branch refused by a hook was not said to be: {out!r}")
        hook.unlink()

        # A claim on a product that has since gone is still released by its holder.
        _g(a, "tag", "-a", "claim/gone/RSK-001", "-m", "claim gone/RSK-001\n\nwho: Owner\nbranch: work/x\n",
           "origin/main")
        _g(a, "push", "-q", "origin", "refs/tags/claim/gone/RSK-001")
        rc, out = _work(a, "release", "gone:RSK-001", "--who", "Owner")
        if rc:
            problems.append(f"a claim on a product no longer declared could not be released: {out!r}")

    # The comparison the main-line job makes, run on reports written here.
    wf = yaml.safe_load((ROOT / "ci" / "branches.yml").read_text(encoding="utf-8"))
    run = next(st["run"] for st in wf["jobs"]["main"]["steps"]
               if str(st.get("name", "")).startswith("how each commit"))
    script = run[run.index("<<'PY'\n") + 7:run.index("\nPY\n")]
    with tempfile.TemporaryDirectory() as tmp:
        t = Path(tmp)
        err = {"code": "REF001", "path": "a.md", "message": "standing", "level": "error"}
        new = {"code": "REF001", "path": "b.md", "message": "new", "level": "error"}
        (t / "s.py").write_text(script)
        (t / "after.json").write_text(json.dumps({"findings": [err, new]}))
        (t / "before.json").write_text(json.dumps({"findings": [err]}))
        (t / "empty.json").write_text("")

        def compare(before: str, parent: str) -> str:
            return subprocess.run([sys.executable, str(t / "s.py"), str(t / "after.json"),
                                   str(t / before), "abc123def456", parent],
                                  capture_output=True, text=True).stdout
        out = compare("before.json", "p")
        if "b.md" not in out or "a.md" in out:
            problems.append(f"the main-line job did not keep to the error the commit brought: {out!r}")
        out = compare("empty.json", "p")
        if "could not be validated" not in out or "b.md" in out:
            problems.append(f"a tree before that could not be validated was not said, or its "
                            f"errors were posted as new: {out!r}")
        out = compare("empty.json", "")
        if "could not be validated" in out or "a.md" not in out:
            problems.append(f"a first commit, which has no tree before it, got a false notice or "
                            f"lost its errors: {out!r}")
    return problems


@check("BR003 asks that the reservation be this branch's")
def _branch_reservations_are_owned():
    problems = []
    with tempfile.TemporaryDirectory() as tmp:
        repo = _Repo(Path(tmp) / "r")
        repo.write(".framework/work.yaml", "people:\n  Owner: {}\n")
        repo.write("OPEN.md", _open_register({"OD-001": "open"}))
        repo.commit("base", "2026-06-01T09:00:00+02:00")
        repo.git("checkout", "-q", "-b", "work/mine")
        repo.write("decisions/DEC-001.md", _decision("DEC-001", "OD-001"))
        repo.commit("a decision", "2026-06-02T09:00:00+02:00")
        repo.git("tag", "-a", "id/DEC-001", "-m", "reserve DEC-001\n\nwho: Junior\nbranch: work/theirs\n")
        codes = [f["code"] for f in repo.validate("--base", "main")[1]["findings"]]
        if codes.count("BR003") != 1:
            problems.append(f"a number reserved for another branch cleared BR003: {codes}")
        repo.git("tag", "-d", "id/DEC-001")
        repo.git("tag", "-a", "id/DEC-001", "-m", "reserve DEC-001, for nobody in particular")
        codes = [f["code"] for f in repo.validate("--base", "main")[1]["findings"]]
        if codes.count("BR003") != 1:
            problems.append(f"an annotated tag naming no branch cleared BR003: {codes}")
        repo.git("tag", "-d", "id/DEC-001")
        repo.git("tag", "-a", "id/DEC-001", "-m", "reserve DEC-001\n\nwho: Owner\nbranch: work/mine\n")
        codes = [f["code"] for f in repo.validate("--base", "main")[1]["findings"]]
        if "BR003" in codes:
            problems.append(f"a reservation for this branch did not clear BR003: {codes}")
    return problems


@check("exchange.py lists what waits, the elements done first, the oldest next, with the days")
def _exchange_lists_what_waits():
    problems = []

    def run(root: Path, *extra: str):
        return subprocess.run([sys.executable, str(EXCHANGE), "--root", str(root), *extra],
                              capture_output=True, text=True)

    with tempfile.TemporaryDirectory() as tmp:
        repo, ex, _ = _exchange_docs(Path(tmp))
        r = run(repo.path)
        if r.returncode != 0 or "is not on this machine" not in r.stdout:
            problems.append(f"an exchange not yet created: exit {r.returncode}, {r.stdout!r}")
        now = datetime.now().timestamp()
        _put(ex, {"in-arrivo/CHG-001/verifica.md": "v", "in-arrivo/CHG-002/verifica.md": "v",
                  "in-arrivo/OD-001/verifica.md": "v", "in-arrivo/CHG-003/brief.md": "b",
                  "in-arrivo/CHG-001/risorse.md": _RESOURCE_HEAD
                  + "| ramo locale | chg-001 | git -C /src/app branch -D chg-001 |\n",
                  "sessioni/2026-10-14-ab12cd34.md": _DECLARATION})
        for rel, days in (("in-arrivo/CHG-002/verifica.md", 5), ("in-arrivo/OD-001/verifica.md", 2),
                          ("in-arrivo/CHG-003/brief.md", 1)):
            os.utime(ex / rel, (now - days * 86400, now - days * 86400))
        r = run(repo.path)
        out = r.stdout
        order = [out.find(s) for s in ("DONE IN THE DOCUMENTS", "CHG-001", "WAITING",
                                       "CHG-002", "OD-001", "AT A CODE SESSION", "CHG-003",
                                       "RESOURCES OF ELEMENTS DONE", "branch -D chg-001",
                                       "CLOSING DECLARATIONS", "2026-10-14-ab12cd34.md")]
        if r.returncode != 0 or -1 in order or order != sorted(order):
            problems.append(f"the listing is not in the order it is worth reading (exit "
                            f"{r.returncode}): {order}\n{out}")
        for says in ("waiting 5 days", "waiting 2 days", "brief written 1 day ago"):
            if says not in out:
                problems.append(f"the listing does not say {says!r}")
        r = run(repo.path, "--json")
        try:
            data = json.loads(r.stdout)
        except json.JSONDecodeError:
            data = {}
        if ([e["id"] for e in data.get("waiting", [])] != ["CHG-002", "OD-001"]
                or [e["id"] for e in data.get("done_in_incoming", [])] != ["CHG-001"]
                or [e["id"] for e in data.get("at_a_code_session", [])] != ["CHG-003"]
                or [r["command"] for r in data.get("resources_to_remove", [])]
                != ["git -C /src/app branch -D chg-001"]):
            problems.append(f"--json does not carry the same listing: {r.stdout[:300]}")
        if any(f["code"] != "EXC002" and f["code"] != "EXC004" for f in data.get("findings", [])):
            problems.append(f"the listing reports findings beyond the place: {data.get('findings')}")
        before = sorted(p.relative_to(ex).as_posix() for p in ex.rglob("*"))
        if sorted(p.relative_to(ex).as_posix() for p in ex.rglob("*")) != before:
            problems.append("the listing moved or wrote something")
        bare, _, _ = _exchange_docs(Path(tmp) / "bare", None)
        r = run(bare.path)
        if r.returncode != 2 or "no `exchange` declared" not in r.stderr:
            problems.append(f"without `exchange:` the command exited {r.returncode}: "
                            f"{r.stderr.strip()[-120:]}")
    return problems


@check("a reading changes no derived view: product.index.yaml does not carry last_review")
def _a_reading_regenerates_nothing():
    # Every attestation used to cost two commits, the reading and the regeneration of a view
    # that copied its instant. The view lists the living documents and not when they were
    # last read, which the document says and `LC001`, `LC002` and `LC004` report on.
    problems = []
    with tempfile.TemporaryDirectory() as tmp:
        repo = _Repo(Path(tmp) / "r")
        repo.write("products/p/product.yaml", _fm(
            schema="framework/product-manifest/v1", artifact_type="product-manifest",
            lifecycle="living", status="active", products="[p]", name="P", one_liner="A thing.",
            owners="[o]", created="2026-06-01 09:00", last_review="2026-06-01 09:00",
            stage="\n  phase: F1\n  block: A").removeprefix("---\n").split("\n---\n")[0] + "\n")
        repo.write("products/p/PBR.md", _brief("p", "2026-06-01 09:00", scope="the whole file"))
        rc, out, err = repo.validate("--emit-index")
        view = repo.path / "products/p/product.index.yaml"
        if not view.exists():
            return [f"--emit-index wrote no view (exit {rc}): {err.strip()[-160:]}"]
        text = view.read_text(encoding="utf-8")
        if "2026-06-01" in text or "never reviewed" in text:
            problems.append(f"the view still carries the instant of a reading: {text}")
        if "products/p/PBR.md" not in text:
            problems.append("the view no longer lists the living documents")
        repo.commit("the documents", "2026-06-01T09:00:00+02:00")
        repo.write("products/p/PBR.md", _brief("p", "2026-06-02 10:00",
                                               scope="the whole file, again"))
        repo.commit("a reading", "2026-06-02T10:00:00+02:00")
        rc, out, err = repo.validate("--emit-index", "--check")
        if rc != 0 or (out or {}).get("out_of_date"):
            problems.append(f"after a reading the views are out of date: "
                            f"{(out or {}).get('out_of_date')}")
        if repo.findings("LC006"):
            problems.append("the commit of the reading is read as a change: LC006 no longer "
                            "steps over an attestation")
    return problems


@check("migrate.py names the checks the new catalog has and the declared one did not")
def _migrate_names_new_checks():
    migrate = _load(ROOT / "skills" / "audit" / "scripts" / "migrate.py", "migrate_new_checks")
    got = migrate.new_checks({"AA001": {"level": "warn"}},
                             {"AA001": {"level": "warn"},
                              "BB001": {"level": "error", "title": "a new one"},
                              "CC001": {"level": "retired", "title": "gone"}})
    want = [{"code": "BB001", "level": "error", "title": "a new one"}]
    problems = [] if got == want else [f"new_checks returned {got}, expected {want}"]
    head = subprocess.run(["git", "-C", str(ROOT), "show", "HEAD~1:skills/audit/checks.yaml"],
                          capture_output=True, text=True)
    if head.returncode == 0:
        before = yaml.safe_load(head.stdout).get("checks") or {}
        listed = {c["code"] for c in migrate.new_checks(before, CHECKS["checks"])}
        missing = {c for c in CHECKS["checks"] if c not in before
                   and CHECKS["checks"][c].get("level") != "retired"} - listed
        if missing:
            problems.append(f"checks in the catalog and not before, not listed: {sorted(missing)}")
    return problems


_AGENTS_BODY = """---
schema: framework/agents-control-plane/v1
last_review: 2026-01-01 10:00
---

# Instructions for agents

## Non negotiable rules

1. Ours, numbered our way.

## Mandatory updates

Our table.

## Commands

```bash
## a comment in a fence, not a heading
make test
```
"""


def _write_sections(migrate, root: Path, plan: dict) -> None:
    path = root / "AGENTS.md"
    path.write_text(migrate.apply_sections(path.read_text(encoding="utf-8"), plan)[0],
                    encoding="utf-8")


@check("--adopt brings the framework's sections of AGENTS.md, and leaves what the project wrote")
def _adopt_writes_the_framework_sections():
    # A RULE THAT TRAVELS BY HAND REACHES THE PROJECTS WHOSE PEOPLE READ RELEASE NOTES. The
    # review section of 4.6.0 reached a project only if somebody pasted it, with a path written
    # over a placeholder, into the one file a session running no skill reads.
    migrate = _load(ROOT / "skills" / "audit" / "scripts" / "migrate.py", "migrate_sections")
    problems = []
    template = (ROOT / migrate.TEMPLATE_REL).read_text(encoding="utf-8")
    marked = {s["id"]: s for s in migrate.sections(template) if s["id"]}
    if set(marked) != {"review", "exchange"}:
        problems.append(f"the template marks {sorted(marked)} as the framework's, expected "
                        "the review and the exchange")
    for sid, s in marked.items():
        left = set(re.findall(r"\b[A-Z]+_[A-Z_]+\b", migrate.render(s["text"], None)))
        if left - set(migrate.PLACEHOLDERS):
            problems.append(f"section {sid!r} carries {sorted(left - set(migrate.PLACEHOLDERS))}, "
                            "which `migrate.py` does not fill")
    if marked.get("exchange", {}).get("requires") != ["exchange"]:
        problems.append("the exchange section does not require an exchange: --adopt would "
                        "write it into every project")
    # What the template says to whoever fills a marked section is a blockquote, and `render`
    # drops every one: a blockquote further down would be content, dropped in silence.
    for sid, s in marked.items():
        blocks = [b for b in s["text"].split("\n\n")[1:] if b.strip()]
        quoted = [i for i, b in enumerate(blocks) if b.lstrip().startswith(">")]
        if quoted and quoted != list(range(len(quoted))):
            problems.append(f"section {sid!r} has a blockquote after its text begins: "
                            "`migrate.py` drops every blockquote of a marked section")
    if migrate.SESSIONS_KEYS != _load(VALIDATE, "validate_sessions_keys").SESSIONS_KEYS:
        problems.append("migrate.py and validate.py disagree on the keys of "
                        ".framework/sessions.yaml")

    with tempfile.TemporaryDirectory() as tmp:
        repo = _Repo(Path(tmp) / "docs")
        root = repo.path
        repo.write("AGENTS.md", _AGENTS_BODY)
        repo.commit("day one", "2026-01-01T10:00:00+00:00")
        fw_path = migrate.framework_path(root, ROOT)[0]

        plan = {e["id"]: e for e in migrate.section_plan(root, ROOT, None)["sections"]}
        if plan.get("exchange", {}).get("action") != "skipped":
            problems.append(f"with no exchange declared the exchange section is "
                            f"{plan.get('exchange')}, expected skipped")
        review = plan.get("review", {})
        if (review.get("action"), review.get("where"), review.get("anchor")) \
                != ("add", "before", "Mandatory updates"):
            problems.append(f"the review section is planned as {plan.get('review')}, expected "
                            "added before the project's 'Mandatory updates'")

        # The report run writes nothing; the plan is what `--adopt` would do.
        r = subprocess.run([sys.executable, str(ROOT / "skills/audit/scripts/migrate.py"),
                            "--root", str(root), "--json"], capture_output=True, text=True)
        try:
            listed = {e["id"]: e["action"] for e in json.loads(r.stdout)["agents"]["sections"]}
        except (ValueError, KeyError, TypeError):
            listed = None
        if listed != {"review": "add", "exchange": "skipped"}:
            problems.append(f"the report lists {listed}, expected review add and exchange "
                            f"skipped: {(r.stdout + r.stderr)[-300:]}")
        if r.returncode != 1:
            problems.append(f"a project on this version with a section to write exits "
                            f"{r.returncode}: 0 says there is nothing to do")
        if (root / "AGENTS.md").read_text(encoding="utf-8") != _AGENTS_BODY:
            problems.append("the report run wrote AGENTS.md: only --adopt writes")

        # Uncommitted lines in AGENTS.md are somebody's: --adopt writes nothing at all.
        repo.write("AGENTS.md", _AGENTS_BODY + "\nA line nobody committed.\n")
        r = subprocess.run([sys.executable, str(ROOT / "skills/audit/scripts/migrate.py"),
                            "--root", str(root), "--adopt"], capture_output=True, text=True)
        if "not committed" not in r.stdout or "A line nobody committed." not in \
                (root / "AGENTS.md").read_text(encoding="utf-8") or "## Review" in \
                (root / "AGENTS.md").read_text(encoding="utf-8"):
            problems.append("--adopt over an AGENTS.md with uncommitted changes did not refuse "
                            f"and leave it: {r.stdout[-300:]}")
        repo.write("AGENTS.md", _AGENTS_BODY)

        repo.write(".framework/sessions.yaml", "exchange: ../_ours\n")
        plan = migrate.section_plan(root, ROOT, None)
        _write_sections(migrate, root, plan)
        text = (root / "AGENTS.md").read_text(encoding="utf-8")
        have = [s["heading"] for s in migrate.sections(text)]
        if have != ["Non negotiable rules", "Sessions and the exchange",
                    "Review before presenting, and before committing", "Mandatory updates",
                    "Commands"]:
            problems.append(f"after writing, the sections are {have}")
        for must, why in (("`../_ours`", "the declared exchange"),
                          (f"`{fw_path}/agents/critic.md`", "the framework's path"),
                          ("<!-- framework-section: review -->", "the marker"),
                          ("last_review: 2026-01-01 10:00", "the front matter, untouched"),
                          ("1. Ours, numbered our way.", "the project's own sections")):
            if must not in text:
                problems.append(f"after writing, {why} is not in AGENTS.md ({must})")
        for must_not in ("EXCHANGE_PATH", "FRAMEWORK_PATH", "\n> "):
            if must_not in text:
                problems.append(f"after writing, {must_not.strip()!r} is still in AGENTS.md: a "
                                "placeholder or an instruction to whoever fills the template")
        if migrate.section_plan(root, ROOT, None)["sections"]:
            problems.append("a second migration plans more work on sections just written")

        # The previous template's text with only the placeholders filled: nobody edited it,
        # so the new text replaces it. One word changed by the project: it is theirs.
        old = template.replace("Give it the diff as a", "Hand it the diff as a")
        unedited = text.replace("Give it the diff as a", "Hand it the diff as a")
        repo.write("AGENTS.md", unedited)
        plan = {e["id"]: e for e in migrate.section_plan(root, ROOT, old)["sections"]}
        if plan.get("review", {}).get("action") != "update":
            problems.append(f"an unedited copy of the previous text is planned as "
                            f"{plan.get('review')}, expected update")
        _write_sections(migrate, root, migrate.section_plan(root, ROOT, old))
        if (root / "AGENTS.md").read_text(encoding="utf-8") != text:
            problems.append("updating an unedited section did not give the new text back")
        repo.write("AGENTS.md", unedited.replace("the absolute path", "the full path"))
        plan = {e["id"]: e for e in migrate.section_plan(root, ROOT, old)["sections"]}
        if plan.get("review", {}).get("action") != "left":
            problems.append(f"a section the project edited is planned as "
                            f"{plan.get('review')}, expected left")

        # A hand copy under the same heading, with no marker: found, and not doubled; and
        # put in its place only when somebody asks for it by name.
        hand = _AGENTS_BODY.replace("## Mandatory updates",
                                    "## Review before presenting, and before committing\n\n"
                                    "Our own wording of the rule.\n\n## Mandatory updates")
        repo.write("AGENTS.md", hand)
        plan = {e["id"]: e for e in migrate.section_plan(root, ROOT, None)["sections"]}
        if plan.get("review", {}).get("action") != "left":
            problems.append(f"a hand copy under the same heading is planned as "
                            f"{plan.get('review')}, expected left")
        forced = migrate.section_plan(root, ROOT, None, replace={"review"})
        after, _ = migrate.apply_sections(hand, forced)
        if "Our own wording" in after or after.count("## Review before presenting") != 1:
            problems.append("--replace-section review did not put the template's text in "
                            "place of the hand copy, once")

        # The template's text with the paths the project wrote: kept on the next migration,
        # whatever machine runs it. Without its marker, it gets one.
        tpl_review = next(x for x in migrate.sections(template) if x["id"] == "review")
        own = migrate.render(tpl_review["text"], {"FRAMEWORK_PATH": ROOT.as_posix()})
        unmarked = "\n".join(line for line in own.split("\n")
                             if not migrate.SECTION_MARK.match(line.strip()))
        repo.write("AGENTS.md", _AGENTS_BODY.replace("## Mandatory updates",
                                                     unmarked + "\n\n## Mandatory updates"))
        plan = migrate.section_plan(root, ROOT, None)
        review = next((e for e in plan["sections"] if e["id"] == "review"), {})
        if review.get("action") != "update" or f"{ROOT.as_posix()}/agents/critic.md" not in \
                review.get("text", "") or "<!-- framework-section: review -->" not in \
                review.get("text", ""):
            problems.append(f"a copy with the project's own path and no marker is planned as "
                            f"{review.get('action')}: expected an update that keeps the path "
                            "and adds the marker")
        repo.write("AGENTS.md", _AGENTS_BODY.replace("## Mandatory updates",
                                                     own + "\n\n## Mandatory updates"))
        if any(e["id"] == "review" for e in migrate.section_plan(root, ROOT, None)["sections"]):
            problems.append("the template's text with the project's own path is planned for "
                            "work: a value the project wrote is not an edit")
        # A path that leads to no framework here may hold where it was written: said, and
        # kept, or two machines would rewrite it back and forth.
        repo.write("AGENTS.md", _AGENTS_BODY.replace(
            "## Mandatory updates",
            migrate.render(tpl_review["text"], {"FRAMEWORK_PATH": "../nowhere"})
            + "\n\n## Mandatory updates"))
        plan = migrate.section_plan(root, ROOT, None)
        if any(e["id"] == "review" for e in plan["sections"]):
            problems.append(f"a path leading to no framework here is planned for work: "
                            f"{[(e['id'], e['action']) for e in plan['sections']]}")
        if not any("../nowhere" in n for n in plan.get("notes", [])):
            problems.append("a path leading to no framework here is not said in the report")
        # A rule elsewhere that names the critic is not a copy of the section.
        repo.write("AGENTS.md", _AGENTS_BODY.replace(
            "1. Ours, numbered our way.",
            "1. Ours, numbered our way.\n2. Run `framework-data-ai:critic` on every plan."))
        review = next((e for e in migrate.section_plan(root, ROOT, None)["sections"]
                       if e["id"] == "review"), {})
        if review.get("action") != "add":
            problems.append(f"a rule that names the critic made the review section "
                            f"{review.get('action')}: expected add")
        # A copy that kept the template's blockquote is still to clean, filled or not.
        quoted = tpl_review["text"].replace("FRAMEWORK_PATH", ROOT.as_posix())
        repo.write("AGENTS.md", _AGENTS_BODY.replace("## Mandatory updates",
                                                     quoted + "\n\n## Mandatory updates"))
        review = next((e for e in migrate.section_plan(root, ROOT, None)["sections"]
                       if e["id"] == "review"), {})
        if review.get("action") != "update" or "\n> " in review.get("text", "\n> "):
            problems.append("a filled copy that kept the template's blockquote is not cleaned")
        # Under another heading, with no marker: not added beside it.
        renamed = unmarked.replace("## Review before presenting, and before committing",
                                   "## Revisione prima di presentare")
        repo.write("AGENTS.md", _AGENTS_BODY.replace("## Mandatory updates",
                                                     renamed + "\n\n## Mandatory updates"))
        review = next((e for e in migrate.section_plan(root, ROOT, None)["sections"]
                       if e["id"] == "review"), {})
        if review.get("action") != "left" or "Revisione" not in review.get("at", ""):
            problems.append(f"a copy under another heading is planned as "
                            f"{review.get('action')}: the rule would be in the file twice")
        after, _ = migrate.apply_sections(
            (root / "AGENTS.md").read_text(encoding="utf-8"),
            migrate.section_plan(root, ROOT, None, replace={"review"}))
        if "Revisione" in after or after.count("framework-data-ai:critic") != 1:
            problems.append("--replace-section on a copy under another heading did not put "
                            "the template's text in its place, once")

        # What `start` made by copying the body as it was: blockquote and placeholder in it.
        raw = _AGENTS_BODY.replace("## Mandatory updates",
                                   tpl_review["text"] + "\n\n## Mandatory updates")
        repo.write("AGENTS.md", raw)
        written, _ = migrate.apply_sections(raw, migrate.section_plan(root, ROOT, None))
        if "FRAMEWORK_PATH" in written or "\n> " in written or \
                f"`{fw_path}/agents/critic.md`" not in written:
            problems.append("a raw copy of the template's section, as `start` made it, was not "
                            "filled and cleaned by the migration")

        # The copy 4.6.0 asked people to make by hand: the text of that template, the path
        # written over its placeholder, the lines broken elsewhere by an editor. It is
        # unedited, so it moves to the new text -- recognised from the framework's history,
        # with no previous tree, as on a run that repeats on the same version.
        shown = subprocess.run(["git", "-C", str(ROOT), "show", "a4f49e6:templates/AGENTS.md"],
                               capture_output=True, text=True)
        was = next((x for x in migrate.sections(shown.stdout)
                    if x["heading"].startswith("Review before presenting")), None)
        if shown.returncode or was is None:
            problems.append("NOT RUN, the 4.6.0 copy: a4f49e6 is not in this history")
        else:
            copied = migrate.render(was["text"], {"FRAMEWORK_PATH": ROOT.as_posix()})
            copied = copied.replace("it returns\nfor each", "it\nreturns for each")
            repo.write("AGENTS.md", _AGENTS_BODY.replace(
                "## Mandatory updates", copied + "\n\n## Mandatory updates"))
            plan = migrate.section_plan(root, ROOT, None)
            review = next((e for e in plan["sections"] if e["id"] == "review"), {})
            if review.get("action") != "update" or f"`{ROOT.as_posix()}/agents/critic.md`" \
                    not in review.get("text", ""):
                problems.append(f"the 4.6.0 hand copy, lines rebroken, is planned as "
                                f"{review.get('action')}: expected an update keeping its path")
            # Under another heading, with no marker: an earlier text shares fewer words with
            # today's, and is still a copy of the rule, not a place to add it beside.
            elsewhere = "## Revisione prima di presentare\n" + copied.split("\n", 1)[1]
            repo.write("AGENTS.md", _AGENTS_BODY.replace(
                "## Mandatory updates", elsewhere + "\n\n## Mandatory updates"))
            review = next((e for e in migrate.section_plan(root, ROOT, None)["sections"]
                           if e["id"] == "review"), {})
            if review.get("action") != "left" or "Revisione" not in review.get("at", ""):
                problems.append(f"the 4.6.0 copy under another heading is planned as "
                                f"{review.get('action')}: the rule would be in the file twice")
            # Reworded by hand: left, and replaced only by name, with the path still read.
            reworded = copied.replace("Exempt:", "Not reviewed:")
            repo.write("AGENTS.md", _AGENTS_BODY.replace(
                "## Mandatory updates", reworded + "\n\n## Mandatory updates"))
            forced = {e["id"]: e for e in migrate.section_plan(
                root, ROOT, None, replace={"review"})["sections"]}
            if forced.get("review", {}).get("action") != "update" or \
                    f"`{ROOT.as_posix()}/agents/critic.md`" not in \
                    forced.get("review", {}).get("text", ""):
                problems.append("--replace-section on a reworded copy lost the path it carried")

        # The exchange comes from where it is declared: a copy naming another place is out of
        # date, and a place with a space in it is a place.
        repo.write(".framework/sessions.yaml", "exchange: ../the exchange\n")
        repo.write("AGENTS.md", _AGENTS_BODY)
        _write_sections(migrate, root, migrate.section_plan(root, ROOT, None))
        if any(e["id"] == "exchange" for e in migrate.section_plan(root, ROOT, None)["sections"]):
            problems.append("an exchange whose path has a space is planned again after being "
                            "written: the tool does not recognise its own text")
        repo.write(".framework/sessions.yaml", "exchange: ../moved\n")
        plan = {e["id"]: e for e in migrate.section_plan(root, ROOT, None)["sections"]}
        if plan.get("exchange", {}).get("action") != "update" or "`../moved`" not in \
                plan.get("exchange", {}).get("text", ""):
            problems.append(f"an exchange declared elsewhere after adopting is planned as "
                            f"{plan.get('exchange', {}).get('action')}: the copy would keep "
                            "the old place")

        # `exchange:` with nothing after it declares nothing; it is not the place "None".
        repo.write(".framework/sessions.yaml", "exchange:\n")
        plan = {e["id"]: e for e in migrate.section_plan(root, ROOT, None)["sections"]}
        if plan.get("exchange", {}).get("action") != "skipped":
            problems.append(f"an empty exchange is planned as "
                            f"{plan.get('exchange', {}).get('action')}: expected skipped")
        repo.write(".framework/sessions.yaml", "exchange: ../moved\n")

        # The marker above two sections: which one is the framework's is not guessed.
        text = (root / "AGENTS.md").read_text(encoding="utf-8")
        repo.write("AGENTS.md", text.replace("## Mandatory updates",
                                             "<!-- framework-section: review -->\n"
                                             "## Mandatory updates"))
        plan = {e["id"]: e for e in migrate.section_plan(root, ROOT, None)["sections"]}
        if plan.get("review", {}).get("action") != "left":
            problems.append(f"a marker above two sections is planned as "
                            f"{plan.get('review', {}).get('action')}, expected left")
        # A heading inside a `~~~` fence is code, not a section.
        repo.write("AGENTS.md", _AGENTS_BODY.replace("```bash", "~~~bash").replace("```\n", "~~~\n"))
        if [x["heading"] for x in migrate.sections((root / "AGENTS.md").read_text())][-1] \
                != "Commands":
            problems.append("a `## ` line inside a ~~~ fence was read as a heading")
        repo.write(".framework/sessions.yaml", "exchange: ../_ours\n")

        # A value the validator refuses is not a place, and is not written as one.
        repo.write(".framework/sessions.yaml", "exchange: /srv/somewhere\n")
        repo.write("AGENTS.md", _AGENTS_BODY)
        plan = {e["id"]: e for e in migrate.section_plan(root, ROOT, None)["sections"]}
        if plan.get("exchange", {}).get("action") != "skipped":
            problems.append(f"an absolute exchange is planned as {plan.get('exchange')}: "
                            "expected skipped, with the validator's reason")

        bad = subprocess.run([sys.executable, str(ROOT / "skills/audit/scripts/migrate.py"),
                              "--root", str(root), "--replace-section", "nonsense"],
                             capture_output=True, text=True)
        if "nonsense" not in bad.stdout or bad.returncode != 1:
            problems.append("--replace-section with an id the template does not mark was "
                            "ignored in silence")

        # End to end: `--adopt` against a clean committed framework writes the number and the
        # section together, which is what a project runs. The file's line endings stay CRLF.
        if subprocess.run(["git", "-C", str(ROOT), "status", "--porcelain"],
                          capture_output=True, text=True).stdout.strip():
            return problems + ["NOT RUN, the end to end half: the working tree is dirty, and "
                               "`--adopt` binds committed bytes. Commit and run again"]
        clone = Path(tmp) / "trusted"
        c = subprocess.run(["git", "clone", "-q", "--local", "--no-hardlinks", str(ROOT),
                            str(clone)], capture_output=True, text=True)
        if c.returncode:
            return problems + ["could not clone the framework for the adoption: " + c.stderr]
        repo.write(".framework/sessions.yaml", "exchange: ../_ours\n")
        (root / "AGENTS.md").write_bytes(_AGENTS_BODY.replace("\n", "\r\n").encode())
        repo.commit("an exchange", "2026-01-02T10:00:00+00:00")
        r = subprocess.run([sys.executable, str(ROOT / "skills/audit/scripts/migrate.py"),
                            "--root", str(root), "--framework", str(clone), "--adopt"],
                           capture_output=True, text=True)
        text = (root / "AGENTS.md").read_bytes().decode("utf-8")
        if text.replace("\r\n", "").count("\n"):
            problems.append("--adopt rewrote the line endings of an AGENTS.md written with CRLF")
        if r.returncode or "<!-- framework-section: review -->" not in text or \
                "<!-- framework-section: exchange" not in text:
            problems.append(f"--adopt on a clean framework did not write both sections "
                            f"(exit {r.returncode}): {r.stdout[-400:]}")
        elif "attest.py" not in r.stdout:
            problems.append("--adopt wrote AGENTS.md and did not say that a person rereads "
                            "and attests it")
    return problems


# ─────────────────────────────────────────────────────────────────────────────

@check("contribution authority, artifact resolution and code-observation contracts")
def _contracts_suite():
    r = subprocess.run([sys.executable, "-B", "-m", "unittest", "discover",
                        "-s", "tests/memory", "-v"], cwd=ROOT,
                       capture_output=True, text=True)
    if r.returncode:
        return [r.stdout + r.stderr]
    # Print the unit-suite summary, including any future expected failures rather than
    # concealing them behind selfcheck's overall result.
    for line in r.stderr.splitlines():
        if "expected failure" in line or line.startswith(("Ran ", "OK")):
            print(f"       {line}")
    return []


print()
if failures:
    print(f"{len(failures)} problem(s).")
    sys.exit(1)
print("The framework is consistent with itself.")
