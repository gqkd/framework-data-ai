#!/usr/bin/env python3
"""Check a presentation outline, and render it to a .pptx a customer can be handed.

    python render.py _meta/presentation/PRS-001-atlas-2026-10-02/outline.yaml --check
    python render.py _meta/presentation/PRS-001-atlas-2026-10-02/outline.yaml

THE OUTLINE IS THE DOCUMENT, THE .PPTX IS ITS PRINT. Every rule the `presentation` skill
states about what may reach a customer is checked here on the outline, before anything is
drawn, so that the part a person reviews is text and the part that leaves is generated from
it. The checks need nothing but PyYAML; the rendering needs python-pptx.

    lang: it                       # it | en: the language of the few fixed labels
    title: Atlas                   # the title slide, from PBR §One line
    subtitle: Riconcilia i movimenti bancari senza lavoro manuale
    customer: null                 # who it is for, or null for nobody in particular
    footer: Atlas                  # optional, on every slide but the title
    theme: {ink: "#1f2937", accent: "#2563eb"}   # optional; keys below in THEME
    sources: [PBR §One line]       # what the title slide says it from
    slides:
      - kind: bullets              # a title and at most five one-line points
        title: Che cosa fa oggi
        bullets: [Abbina da solo i movimenti alle scritture, ...]
        sources: [PBR §Current capabilities, REL-004 §What changes]
      - kind: cards                # two to six tiles, a title and a sentence each
        title: Per chi è
        cards: [{title: Operatore contabile, text: Lavora solo i movimenti che non tornano}]
        sources: [PBR §Actors]
      - kind: flow                 # the "how it works" diagram, drawn as editable shapes
        title: Come funziona
        steps: [{title: Riceve i movimenti, text: dal file della banca}, ...]
        focal: 2                   # optional, 1-based: the step the reader should remember
        caption: Dai movimenti della banca alla chiusura del mese
        sources: [ARC#current §Components]
      - kind: image                # a picture made elsewhere, when one is better
        title: Come funziona
        image: diagrams/current.png          # relative to the outline
        sources: [ARC#current §Components]
      - kind: roadmap              # exactly one: done and to do, never when
        title: A che punto siamo
        done: [Abbinamento automatico]
        todo: [Chiusura del mese guidata, {text: Riconciliazione multi-banca, planned: true}]
        sources: [PBR §Current capabilities, RMP §Increments]
        # or, for a suite, one row per product:
        # lanes: [{name: Atlas, done: [...], todo: [...]}, {name: Borea, done: [], todo: [...]}]

`sources` never reaches the .pptx, speaker notes included: notes travel with the file, and
identifiers are what a customer must not be handed. They are kept in the outline, which is
where a reviewer checks that every line came from somewhere.
"""

from __future__ import annotations

import argparse
import re
import sys
from pathlib import Path

import yaml

MAX_SLIDES = 10         # the title included
MAX_BULLETS = 5
MAX_ITEMS = 6           # per roadmap column, all lanes together
MAX_LINE = 110          # characters; a point that needs two lines is a technical point
MAX_CARD = 150          # a card's sentence, which has a tile to itself

LABELS = {
    "it": {"done": "Fatto", "todo": "Da fare", "planned": "previsto",
           "nothing": "Ancora nulla in esercizio"},
    "en": {"done": "Done", "todo": "To do", "planned": "planned",
           "nothing": "Nothing in production yet"},
}

THEME = {"ink": "#1f2937", "muted": "#6b7280", "accent": "#2563eb", "done": "#15803d",
         "paper": "#f4f6f9", "cover": "#111827", "on_cover": "#ffffff"}

TOP_KEYS = {"lang", "title", "subtitle", "customer", "footer", "theme", "sources", "slides"}
SLIDE_KEYS = {
    "bullets": {"kind", "title", "bullets", "sources"},
    "cards": {"kind", "title", "cards", "sources"},
    "flow": {"kind", "title", "steps", "focal", "caption", "sources"},
    "image": {"kind", "title", "image", "caption", "sources"},
    "roadmap": {"kind", "title", "done", "todo", "lanes", "sources"},
}

# An identifier of this framework or a path into the repository, anywhere on a slide.
INTERNAL = [
    (re.compile(r"\b[A-Z]{2,4}-\d{2,4}\b"), "an artifact identifier"),
    (re.compile(r"`"), "a backtick, which is how code reaches a slide"),
    (re.compile(r"\b[\w-]+/[\w./-]*\.(?:md|ya?ml|py|json|sql)\b|\b[\w-]+\.(?:md|ya?ml|py)\b"),
     "a repository path"),
]

# Time, on the roadmap only. Elsewhere "il report di fine mese" is a capability, and
# refusing it would teach whoever writes the outline to blur the product to get past a regex.
MONTHS = ("gennaio febbraio marzo aprile maggio giugno luglio agosto settembre ottobre "
          "novembre dicembre january february march april may june july august september "
          "october november december")
TIME = [
    (re.compile(r"\b\d{4}-\d{2}(?:-\d{2})?\b|\b\d{1,2}/\d{1,2}(?:/\d{2,4})?\b"), "a date"),
    (re.compile(r"\b(?:Q|T)[1-4]\b|\bH[12]\b"), "a quarter or a half"),
    (re.compile(r"\d\s*%"), "a percentage"),
    (re.compile(r"\b(?:settiman[ae]|mes[ei]|trimestr[ei]|giorni|anno|entro|presto|"
                r"weeks?|months?|quarters?|days|years?|soon|by the end)\b", re.I),
     "a duration or a deadline"),
    (re.compile(r"\b(?:" + "|".join(MONTHS.split()) + r")\b", re.I), "a month"),
    (re.compile(r"\b20\d{2}\b"), "a year"),
]
HEX = re.compile(r"^#[0-9a-fA-F]{6}$")


def _texts(v) -> list[str]:
    if isinstance(v, str):
        return [v]
    if isinstance(v, dict):
        return [t for k in ("name", "title", "text", "done", "todo") for t in _texts(v.get(k))]
    if isinstance(v, list):
        return [t for x in v for t in _texts(x)]
    return []


def _visible(slide: dict) -> list[str]:
    out = _texts(slide.get("title"))
    for key in ("bullets", "caption", "done", "todo", "lanes", "cards", "steps"):
        out += _texts(slide.get(key))
    return out


def _sources(where: str, v) -> list[str]:
    if not isinstance(v, list) or not v or not all(isinstance(s, str) and s.strip() for s in v):
        return [f"{where}: no `sources`. Every line on a slide comes from an artifact section, "
                "and this is where a reviewer looks it up"]
    return []


def _tiles(where: str, items, key: str, lo: int, hi: int) -> list[str]:
    if not isinstance(items, list) or not lo <= len(items) <= hi:
        return [f"{where}: `{key}` is a list of {lo} to {hi} {{title, text}}"]
    problems = []
    for x in items:
        if not (isinstance(x, dict) and set(x) <= {"title", "text"}
                and isinstance(x.get("title"), str) and x["title"].strip()
                and isinstance(x.get("text", ""), str)):
            problems.append(f"{where}: {x!r} in `{key}` is {{title, text}}")
        elif len(x.get("text", "")) > MAX_CARD:
            problems.append(f"{where}: {x['title']!r} has {len(x['text'])} characters of text; "
                            f"at most {MAX_CARD}")
    return problems


def _roadmap(where: str, s: dict) -> list[str]:
    problems = []
    if "lanes" in s:
        if "done" in s or "todo" in s:
            return [f"{where}: `lanes`, or `done` and `todo`, not both"]
        lanes = s["lanes"]
        if not isinstance(lanes, list) or not 1 <= len(lanes) <= 3:
            return [f"{where}: `lanes` is a list of one to three products"]
        for ln in lanes:
            if not (isinstance(ln, dict) and set(ln) <= {"name", "done", "todo"}
                    and isinstance(ln.get("name"), str) and ln["name"].strip()):
                problems.append(f"{where}: {ln!r} is {{name, done, todo}}")
        groups = [ln for ln in lanes if isinstance(ln, dict)]
    else:
        groups = [s]
    for col in ("done", "todo"):
        count = 0
        for g in groups:
            items = g.get(col)
            if not isinstance(items, list):
                problems.append(f"{where}: `{col}` is a list, empty if there is nothing")
                continue
            count += len(items)
            for x in items:
                ok = isinstance(x, str) or (
                    isinstance(x, dict) and isinstance(x.get("text"), str)
                    and set(x) <= {"text", "planned"}
                    and isinstance(x.get("planned", False), bool)
                    and (col == "todo" or not x.get("planned")))
                if not ok:
                    problems.append(
                        f"{where}: {x!r} in `{col}` is a sentence, or in `todo` "
                        "{text, planned: true}. There is no other label: what is "
                        "conditional does not reach a customer at all")
        if count > MAX_ITEMS * len(groups) or (len(groups) > 1 and count > MAX_ITEMS + 3):
            problems.append(f"{where}: {count} items in `{col}`; at most {MAX_ITEMS} a product, "
                            f"and {MAX_ITEMS + 3} on one slide")
    for t in _visible(s):
        for rx, what in TIME:
            if rx.search(t):
                problems.append(f"{where}: {t!r} carries {what}; the roadmap says "
                                "what is done and what is left, never when")
    return problems


def check(outline, base: Path) -> list[str]:
    """Every reason this outline cannot be handed to a customer. Empty means it can."""
    if not isinstance(outline, dict):
        return ["the outline is not a mapping"]
    problems: list[str] = []
    for k in sorted(set(outline) - TOP_KEYS):
        problems.append(f"unknown key {k!r} at the top: nothing reads it")
    if outline.get("lang", "it") not in LABELS:
        problems.append(f"lang is {outline.get('lang')!r}; it is one of {sorted(LABELS)}")
    for k in ("title", "subtitle"):
        if not isinstance(outline.get(k), str) or not outline[k].strip():
            problems.append(f"no `{k}`: the title slide is the product's name and its one line")
    theme = outline.get("theme", {})
    if not isinstance(theme, dict) or any(k not in THEME or not isinstance(v, str)
                                          or not HEX.match(v) for k, v in theme.items()):
        problems.append(f"`theme` maps some of {sorted(THEME)} to colours like #1f2937")
    problems += _sources("title slide", outline.get("sources"))

    slides = outline.get("slides")
    if not isinstance(slides, list) or not slides:
        return problems + ["no `slides`"]
    if len(slides) + 1 > MAX_SLIDES:
        problems.append(f"{len(slides) + 1} slides, the title included; at most {MAX_SLIDES}")
    kinds = [s.get("kind") if isinstance(s, dict) else None for s in slides]
    if kinds.count("roadmap") != 1:
        problems.append(f"{kinds.count('roadmap')} roadmap slides; there is exactly one")

    visible = [("title slide", t) for t in _texts([outline.get("title"),
                                                   outline.get("subtitle"),
                                                   outline.get("footer")])]
    for i, s in enumerate(slides, start=2):
        where = f"slide {i}"
        if not isinstance(s, dict) or s.get("kind") not in SLIDE_KEYS:
            problems.append(f"{where}: kind is one of {sorted(SLIDE_KEYS)}")
            continue
        where = f"slide {i} ({s['kind']})"
        for k in sorted(set(s) - SLIDE_KEYS[s["kind"]]):
            problems.append(f"{where}: unknown key {k!r}")
        if not isinstance(s.get("title"), str) or not s["title"].strip():
            problems.append(f"{where}: no `title`")
        problems += _sources(where, s.get("sources"))

        if s["kind"] == "bullets":
            b = s.get("bullets")
            if not isinstance(b, list) or not b or not all(isinstance(x, str) for x in b):
                problems.append(f"{where}: `bullets` is a list of sentences")
            elif len(b) > MAX_BULLETS:
                problems.append(f"{where}: {len(b)} bullets; at most {MAX_BULLETS}")
        elif s["kind"] == "cards":
            problems += _tiles(where, s.get("cards"), "cards", 2, 6)
        elif s["kind"] == "flow":
            problems += _tiles(where, s.get("steps"), "steps", 2, 5)
            f = s.get("focal")
            n = len(s["steps"]) if isinstance(s.get("steps"), list) else 0
            if f is not None and not (isinstance(f, int) and 1 <= f <= n):
                problems.append(f"{where}: `focal` is the number of one step, 1 to {n}")
        elif s["kind"] == "image":
            img = s.get("image")
            if not isinstance(img, str) or not (base / img).is_file():
                problems.append(f"{where}: image {img!r} is not there, relative to the outline")
            elif Path(img).suffix.lower() not in {".png", ".jpg", ".jpeg"}:
                problems.append(f"{where}: {img} is not a PNG or a JPEG, which a slide embeds")
        else:
            problems += _roadmap(where, s)
        visible += [(where, t) for t in _visible(s)]

    for where, t in visible:
        if len(t) > MAX_CARD or (len(t) > MAX_LINE and where.endswith(("(bullets)",
                                                                      "(roadmap)"))):
            problems.append(f"{where}: {t[:40]!r}... is {len(t)} characters, too long for "
                            "one line")
        for rx, what in INTERNAL:
            if rx.search(t):
                problems.append(f"{where}: {t!r} carries {what}; it goes in `sources`, "
                                "which the customer never sees")
    return problems


def render(outline: dict, base: Path, out: Path) -> None:
    from pptx import Presentation
    from pptx.dml.color import RGBColor
    from pptx.enum.dml import MSO_LINE_DASH_STYLE
    from pptx.enum.shapes import MSO_SHAPE
    from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
    from pptx.oxml.ns import qn
    from pptx.util import Emu, Inches, Pt

    th = {**THEME, **(outline.get("theme") or {})}
    C = {k: RGBColor.from_string(v[1:]) for k, v in th.items()}
    white = RGBColor(0xFF, 0xFF, 0xFF)
    labels = LABELS[outline.get("lang", "it")]

    prs = Presentation()
    prs.slide_width, prs.slide_height = Inches(13.333), Inches(7.5)
    blank = prs.slide_layouts[6]
    W, H, M = prs.slide_width, prs.slide_height, Inches(0.7)

    def plain(shape):
        # The default theme draws a shadow through `effectRef` whatever the shape says, and
        # LibreOffice and Keynote honour it; index 0 is "no effect".
        shape.shadow.inherit = False
        style = shape._element.find(qn("p:style"))
        if style is not None:
            style.find(qn("a:effectRef")).set("idx", "0")
        return shape

    def rect(slide, x, y, w, h, fill=None, line=None, width=1.0, rounded=True, dash=False):
        shp = plain(slide.shapes.add_shape(
            MSO_SHAPE.ROUNDED_RECTANGLE if rounded else MSO_SHAPE.RECTANGLE, x, y, w, h))
        if rounded:
            shp.adjustments[0] = 0.08
        if fill is None:
            shp.fill.background()
        else:
            shp.fill.solid()
            shp.fill.fore_color.rgb = fill
        if line is None:
            shp.line.fill.background()
        else:
            shp.line.color.rgb = line
            shp.line.width = Pt(width)
            if dash:
                shp.line.dash_style = MSO_LINE_DASH_STYLE.DASH
        return shp

    def text(slide, x, y, w, h, runs, size, color=None, bold=False, align=PP_ALIGN.LEFT,
             anchor=MSO_ANCHOR.TOP, frame=None):
        """`runs` is a string or a list of paragraphs, each a list of (text, overrides)."""
        tf = frame or slide.shapes.add_textbox(x, y, w, h).text_frame
        tf.word_wrap = True
        tf.vertical_anchor = anchor
        paras = [[(runs, {})]] if isinstance(runs, str) else runs
        for i, para in enumerate(paras):
            p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
            p.alignment = align
            for t, o in para:
                r = p.add_run()
                r.text = t
                r.font.size = Pt(o.get("size", size))
                r.font.bold = o.get("bold", bold)
                r.font.italic = o.get("italic", False)
                r.font.color.rgb = o.get("color", color or C["ink"])
            if "space" in (para[0][1] if para else {}):
                p.space_before = Pt(para[0][1]["space"])
        return tf

    footer = outline.get("footer")
    page = [1]

    def content(title):
        page[0] += 1
        s = prs.slides.add_slide(blank)
        text(s, M, Inches(0.45), W - 2 * M, Inches(0.8), title, 30, bold=True)
        rect(s, M, Inches(1.22), Inches(0.9), Inches(0.06), fill=C["accent"], rounded=False)
        y = H - Inches(0.5)
        if footer:
            text(s, M, y, Inches(8), Inches(0.3), footer, 10, C["muted"])
        text(s, W - M - Inches(1), y, Inches(1), Inches(0.3), str(page[0]), 10, C["muted"],
             align=PP_ALIGN.RIGHT)
        return s

    # The cover: the one slide that is not white, so the deck has a front.
    s = prs.slides.add_slide(blank)
    rect(s, 0, 0, W, H, fill=C["cover"], rounded=False)
    rect(s, M, Inches(2.55), Inches(1.2), Inches(0.08), fill=C["accent"], rounded=False)
    text(s, M, Inches(2.8), W - 2 * M, Inches(1.3), outline["title"], 46, C["on_cover"],
         bold=True)
    text(s, M, Inches(4.05), Inches(10.5), Inches(1.4), outline["subtitle"], 22,
         C["on_cover"])

    top, bottom = Inches(1.75), H - Inches(0.75)
    for sl in outline["slides"]:
        s = content(sl["title"])
        kind = sl["kind"]
        if kind == "bullets":
            text(s, M, top, W - 2 * M, bottom - top,
                 [[("•  " + b, {"space": 0 if i == 0 else 14})]
                  for i, b in enumerate(sl["bullets"])], 22)
        elif kind == "cards":
            cards = sl["cards"]
            cols = 3 if len(cards) in (3, 5, 6) else 2
            rows = (len(cards) + cols - 1) // cols
            gap = Inches(0.3)
            cw = (W - 2 * M - gap * (cols - 1)) // cols
            ch = min(Inches(2.4), (bottom - top - gap * (rows - 1)) // rows)
            for i, c in enumerate(cards):
                x = M + (i % cols) * (cw + gap)
                y = top + (i // cols) * (ch + gap)
                box = rect(s, x, y, cw, ch, fill=C["paper"])
                rect(s, x, y, Inches(0.08), ch, fill=C["accent"], rounded=False)
                tf = box.text_frame
                tf.margin_left, tf.margin_right = Inches(0.35), Inches(0.25)
                tf.margin_top = Inches(0.2)
                text(s, 0, 0, 0, 0, [[(c["title"], {"bold": True, "size": 18})],
                                     [(c.get("text", ""), {"size": 15, "color": C["ink"],
                                                           "space": 8})]],
                     15, frame=tf)
        elif kind == "flow":
            steps, focal = sl["steps"], sl.get("focal")
            n = len(steps)
            cap = Inches(0.6) if sl.get("caption") else 0

            def chip_at(x, y, i, hot):
                chip = plain(s.shapes.add_shape(MSO_SHAPE.OVAL, x, y, Inches(0.5), Inches(0.5)))
                chip.fill.solid()
                chip.fill.fore_color.rgb = white if hot else C["accent"]
                chip.line.fill.background()
                text(s, 0, 0, 0, 0, str(i + 1), 16, C["accent"] if hot else white, bold=True,
                     align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE, frame=chip.text_frame)

            def arrow_at(x, y, w, h, shape):
                a = plain(s.shapes.add_shape(shape, x, y, w, h))
                a.fill.solid()
                a.fill.fore_color.rgb = C["muted"]
                a.line.fill.background()

            if n <= 3:
                # Side by side: three boxes are wide enough for a title and a sentence.
                arrow = Inches(0.55)
                bw = (W - 2 * M - arrow * (n - 1)) // n
                bh = Inches(3.2)
                y = top + Inches(0.3)
                for i, st in enumerate(steps):
                    x = M + i * (bw + arrow)
                    hot = focal == i + 1
                    box = rect(s, x, y, bw, bh, fill=C["accent"] if hot else C["paper"])
                    chip_at(x + Inches(0.25), y + Inches(0.25), i, hot)
                    tf = box.text_frame
                    tf.margin_left = tf.margin_right = Inches(0.3)
                    tf.margin_top, tf.margin_bottom = Inches(0.95), Inches(0.2)
                    ink = white if hot else C["ink"]
                    text(s, 0, 0, 0, 0, [[(st["title"], {"bold": True, "size": 20, "color": ink})],
                                         [(st.get("text", ""), {"size": 16, "color": ink,
                                                                "space": 10})]],
                         16, frame=tf)
                    if i < n - 1:
                        arrow_at(x + bw + Inches(0.1), y + bh // 2 - Inches(0.18),
                                 arrow - Inches(0.2), Inches(0.36), MSO_SHAPE.RIGHT_ARROW)
                y_end = y + bh
            else:
                # Four or five steps side by side break words in half, so they stack: one
                # row a step, the title on the left and its sentence beside it.
                gap = Inches(0.16)
                rh = min(Inches(0.95), (bottom - top - cap - gap * (n - 1)) // n)
                tw = Inches(3.6)
                for i, st in enumerate(steps):
                    y = top + i * (rh + gap)
                    hot = focal == i + 1
                    ink = white if hot else C["ink"]
                    rect(s, M, y, W - 2 * M, rh, fill=C["accent"] if hot else C["paper"])
                    chip_at(M + Inches(0.25), y + (rh - Inches(0.5)) // 2, i, hot)
                    text(s, M + Inches(0.95), y, tw, rh, st["title"], 19, ink, bold=True,
                         anchor=MSO_ANCHOR.MIDDLE)
                    text(s, M + Inches(0.95) + tw, y, W - 2 * M - Inches(1.2) - tw, rh,
                         st.get("text", ""), 16, ink, anchor=MSO_ANCHOR.MIDDLE)
                y_end = top + n * (rh + gap) - gap
            if sl.get("caption"):
                text(s, M, y_end + Inches(0.2), W - 2 * M, Inches(0.5), sl["caption"], 17,
                     C["muted"], align=PP_ALIGN.CENTER)
        elif kind == "image":
            from PIL import Image
            with Image.open(base / sl["image"]) as im:
                iw, ih = im.size
            cap = Inches(0.7) if sl.get("caption") else 0
            bw, bh = W - 2 * M, bottom - top - cap
            scale = min(bw / iw, bh / ih)
            pw, ph = int(iw * scale), int(ih * scale)
            s.shapes.add_picture(str(base / sl["image"]), Emu((W - pw) // 2), top, pw, ph)
            if cap:
                text(s, M, top + ph + Inches(0.15), W - 2 * M, cap, sl["caption"], 17,
                     C["muted"], align=PP_ALIGN.CENTER)
        else:
            lanes = sl.get("lanes") or [{"name": None, "done": sl["done"], "todo": sl["todo"]}]
            named = lanes[0]["name"] is not None
            lw = Inches(1.6) if named else 0
            gap = Inches(0.6)
            colw = (W - 2 * M - lw - gap) // 2
            xs = (M + lw, M + lw + colw + gap)
            for c, (key, color) in enumerate((("done", C["done"]), ("todo", C["accent"]))):
                text(s, xs[c], top - Inches(0.1), colw, Inches(0.5), labels[key], 22, color,
                     bold=True)
            a = plain(s.shapes.add_shape(MSO_SHAPE.RIGHT_ARROW, xs[0] + colw + Inches(0.2),
                                         top, gap - Inches(0.4), Inches(0.3)))
            a.fill.solid()
            a.fill.fore_color.rgb = C["muted"]
            a.line.fill.background()
            y0 = top + Inches(0.55)
            slots = [max(len(ln["done"]), len(ln["todo"]), 1) for ln in lanes]
            step = min(Inches(0.7), (bottom - y0 - Inches(0.2) * (len(lanes) - 1))
                       // sum(slots))
            y = y0
            for ln, k in zip(lanes, slots):
                if named:
                    band = rect(s, M, y, lw - Inches(0.2), step * k - Inches(0.1),
                                fill=C["paper"])
                    text(s, 0, 0, 0, 0, ln["name"], 18, bold=True, align=PP_ALIGN.CENTER,
                         anchor=MSO_ANCHOR.MIDDLE, frame=band.text_frame)
                for c, (key, color) in enumerate((("done", C["done"]), ("todo", C["accent"]))):
                    items = ln[key] or ([None] if key == "done" else [])
                    for j, item in enumerate(items):
                        bx = rect(s, xs[c], y + j * step, colw, step - Inches(0.08),
                                  fill=white if item is not None else None,
                                  line=color if item is not None else None, width=1.25,
                                  dash=isinstance(item, dict) and bool(item.get("planned")))
                        tf = bx.text_frame
                        tf.margin_left = tf.margin_right = Inches(0.15)
                        tf.margin_top = tf.margin_bottom = Inches(0.02)
                        if item is None:
                            text(s, 0, 0, 0, 0, [[(labels["nothing"], {"italic": True})]], 14,
                                 C["muted"], anchor=MSO_ANCHOR.MIDDLE, frame=tf)
                            continue
                        t = item if isinstance(item, str) else item["text"]
                        # To do is in dependency order, so it is numbered: the reader sees
                        # that the third thing waits for the first two.
                        mark = "✓\t" if key == "done" else f"{j + 1}\t"
                        runs = [(mark + t, {})]
                        if isinstance(item, dict) and item.get("planned"):
                            runs.append((f"  ({labels['planned']})",
                                         {"italic": True, "size": 12, "color": C["muted"]}))
                        text(s, 0, 0, 0, 0, [runs], 14, anchor=MSO_ANCHOR.MIDDLE, frame=tf)
                        # A hanging indent, so a wrapped line starts under the text and not
                        # under the mark.
                        pPr = tf.paragraphs[0]._p.get_or_add_pPr()
                        pPr.set("marL", str(Pt(22)))
                        pPr.set("indent", str(-Pt(22)))
                y += step * k + Inches(0.2)
    prs.save(str(out))


def main(argv: list[str] | None = None) -> int:
    ap = argparse.ArgumentParser(description=__doc__.split("\n\n")[0])
    ap.add_argument("outline", type=Path, help="the outline.yaml of one presentation")
    ap.add_argument("--check", action="store_true",
                    help="check the outline and write nothing")
    ap.add_argument("--out", type=Path,
                    help="where the .pptx goes; default: beside the outline, named after "
                         "its directory")
    a = ap.parse_args(argv)
    try:
        outline = yaml.safe_load(a.outline.read_text(encoding="utf-8"))
    except (OSError, yaml.YAMLError) as e:
        print(f"cannot read {a.outline}: {e}", file=sys.stderr)
        return 2
    base = a.outline.resolve().parent
    problems = check(outline, base)
    for p in problems:
        print(p)
    if problems:
        return 1
    if a.check:
        print("ok")
        return 0
    try:
        import pptx  # noqa: F401
    except ImportError:
        print("python-pptx is not installed: pip install -r requirements.txt", file=sys.stderr)
        return 2
    out = a.out or base / f"{base.name}.pptx"
    render(outline, base, out)
    print(out)
    return 0


if __name__ == "__main__":
    sys.exit(main())
